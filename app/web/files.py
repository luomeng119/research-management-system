from __future__ import annotations

import json
from itertools import islice
from pathlib import Path

from flask import Blueprint, current_app, jsonify, request, send_file, url_for

from app.security.auth import business_required, current_identity
from app.services.files import (
    FileServiceError,
    TEXT_EXTENSIONS,
    TEXT_PREVIEW_MAX_BYTES,
    office_archive_is_preview_safe,
)


bp = Blueprint("files", __name__, url_prefix="/api/files")
OFFICE_PREVIEW_CHARACTER_BUDGET = 2 * 1024 * 1024


def _take_text(value, budget):
    text = str(value or "")
    if len(text) > budget[0]:
        raise ValueError("preview character budget exceeded")
    budget[0] -= len(text)
    return text


def business_user_required(function):
    """Retain existing imports while sharing the canonical formal-role policy."""
    return business_required(function)


def _error(error: FileServiceError):
    return jsonify({"error": {"code": error.code, "message": error.message, "requestId": getattr(request, "request_id", None)}}), error.status_code


def _service():
    return current_app.extensions["file_service"]


def _record_event(operation: str, *, result: str, error_code: str | None = None, file_id: str | None = None, opened: dict | None = None, legacy_project_id: str | None = None):
    identity = current_identity()
    project_id = legacy_project_id or (opened or {}).get("legacyProjectId")
    if project_id:
        service = _service()
        properties = {"operation": operation}
        if opened:
            properties.update(file_type=Path(opened["originalName"]).suffix.lstrip("."),
                              size_bucket=service._size_bucket(opened["sizeBytes"]))
        try:
            with service.repository.engine.begin() as connection:
                service.audit_service.record(
                    connection, event_name="file_operation_completed", user_id=identity.user_id,
                    object_type="PROJECT", object_id=project_id, result=result,
                    duration_ms=0,
                    request_id=getattr(request, "request_id", "unknown"), error_code=error_code,
                    properties=properties,
                )
            return
        except Exception as exc:
            raise FileServiceError("FILE_OPERATION_FAILED", "文件审计失败", 500) from exc
    _service().record_event(
        operation=operation,
        actor_user_id=identity.user_id,
        request_id=getattr(request, "request_id", "unknown"),
        file_id=file_id,
        result=result,
        error_code=error_code,
        file_type=Path(opened["originalName"]).suffix.lstrip(".").lower() if opened else None,
        size_bytes=opened["sizeBytes"] if opened else None,
    )


def _record_failure(operation: str, error: FileServiceError, *, file_id: str | None = None, opened: dict | None = None, legacy_project_id: str | None = None):
    try:
        _record_event(operation, result="FAILURE", error_code=error.code, file_id=file_id,
                      opened=opened, legacy_project_id=legacy_project_id)
    except FileServiceError:
        current_app.logger.warning("file failure audit unavailable")


def _object_reference():
    object_type = request.values.get("objectType", "").strip().upper()
    object_id = request.values.get("objectId", "").strip()
    if not object_type or not object_id:
        raise FileServiceError("OBJECT_REFERENCE_REQUIRED", "必须提供业务对象类型和编号")
    return object_type, object_id


@bp.get("")
@business_user_required
def list_files():
    try:
        object_type, object_id = _object_reference()
        return jsonify({"files": _service().list_for_object(object_type=object_type, object_id=object_id)})
    except FileServiceError as error:
        _record_failure("LIST", error)
        return _error(error)


@bp.post("")
@business_user_required
def upload_file():
    try:
        object_type, object_id = _object_reference()
        uploaded = request.files.get("file")
        if uploaded is None or not uploaded.filename:
            raise FileServiceError("FILE_REQUIRED", "请选择文件")
        identity = current_identity()
        result = _service().upload(
            uploaded.stream, original_name=uploaded.filename, object_type=object_type,
            object_id=object_id, actor_user_id=identity.user_id,
            request_id=getattr(request, "request_id", "unknown"),
        )
        return jsonify({key: value for key, value in result.items() if key != "storagePath"}), 201
    except FileServiceError as error:
        _record_failure("UPLOAD", error)
        return _error(error)


@bp.post("/<file_id>/versions")
@business_user_required
def add_version(file_id: str):
    try:
        object_type, object_id = _object_reference()
        uploaded = request.files.get("file")
        if uploaded is None or not uploaded.filename:
            raise FileServiceError("FILE_REQUIRED", "请选择文件")
        try:
            expected_version = int(request.form.get("expectedVersion", ""))
        except ValueError:
            raise FileServiceError("EXPECTED_VERSION_REQUIRED", "必须提供当前版本号")
        identity = current_identity()
        result = _service().add_version(
            file_id, uploaded.stream, original_name=uploaded.filename, object_type=object_type,
            object_id=object_id, expected_version=expected_version,
            actor_user_id=identity.user_id, request_id=getattr(request, "request_id", "unknown"),
        )
        return jsonify({key: value for key, value in result.items() if key != "storagePath"}), 201
    except FileServiceError as error:
        _record_failure("ADD_VERSION", error, file_id=file_id)
        return _error(error)


@bp.get("/<file_id>/versions")
@business_user_required
def list_versions(file_id: str):
    try:
        object_type, object_id = _object_reference()
        return jsonify(_service().list_versions(file_id, object_type=object_type, object_id=object_id))
    except FileServiceError as error:
        _record_failure("LIST", error, file_id=file_id)
        return _error(error)


@bp.post("/<file_id>/archive")
@business_user_required
def archive_file(file_id: str):
    try:
        object_type, object_id = _object_reference()
        identity = current_identity()
        return jsonify(_service().archive(
            file_id, object_type=object_type, object_id=object_id,
            actor_user_id=identity.user_id, request_id=getattr(request, "request_id", "unknown"),
        ))
    except FileServiceError as error:
        _record_failure("ARCHIVE", error, file_id=file_id)
        return _error(error)


def _opened(file_id: str, version_no: int):
    object_type, object_id = _object_reference()
    return _service().open_version_stream(file_id, version_no, object_type=object_type, object_id=object_id)


def _send(opened: dict, *, attachment: bool):
    response = send_file(
        opened["stream"], mimetype=opened["mediaType"], as_attachment=attachment,
        download_name=opened["originalName"], conditional=True, etag=opened["sha256"],
    )
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Content-Security-Policy"] = "default-src 'none'; sandbox"
    return response


@bp.get("/<file_id>/versions/<int:version_no>/download")
@business_user_required
def download_file(file_id: str, version_no: int):
    opened = None
    try:
        opened = _opened(file_id, version_no)
        response = _send(opened, attachment=True)
        _record_event("DOWNLOAD", result="SUCCESS", file_id=file_id, opened=opened)
        return response
    except (FileServiceError, ValueError, OSError) as raw_error:
        error = raw_error if isinstance(raw_error, FileServiceError) else FileServiceError("FILE_OPERATION_FAILED", "文件下载失败", 500)
        if opened and not opened["stream"].closed:
            opened["stream"].close()
        _record_failure("DOWNLOAD", error, file_id=file_id)
        return _error(error)


def _preview_unavailable(file_id: str, version_no: int, opened: dict | None = None, *, download_url=None):
    if opened is not None:
        if not opened["stream"].closed:
            opened["stream"].close()
        _record_event(
            "PREVIEW", result="FAILURE", error_code="PREVIEW_UNAVAILABLE",
            file_id=file_id, opened=opened,
        )
    query = {"objectType": request.args.get("objectType"), "objectId": request.args.get("objectId")}
    response, status = _error(FileServiceError("PREVIEW_UNAVAILABLE", "该类型暂不支持在线预览，请下载查看", 409))
    payload = response.get_json()
    payload["downloadUrl"] = download_url or url_for("files.download_file", file_id=file_id, version_no=version_no, **query)
    response.set_data(json.dumps(payload, ensure_ascii=False))
    response.content_type = "application/json; charset=utf-8"
    return response, status


@bp.get("/<file_id>/versions/<int:version_no>/preview")
@business_user_required
def preview_file(file_id: str, version_no: int):
    return render_file_preview(lambda: _opened(file_id, version_no), file_id=file_id, version_no=version_no)


def render_file_preview(opener, *, file_id=None, version_no=1, download_url=None, legacy_project_id=None):
    """Render an authorized stream; callers own object binding, never raw paths."""
    opened = None
    def unavailable():
        return _preview_unavailable(file_id, version_no, opened, download_url=download_url)
    try:
        opened = opener()
        extension = Path(opened["originalName"]).suffix.lower()
        if opened["sizeBytes"] > _service().preview_max_bytes:
            return unavailable()
        if extension in {".pdf", ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}:
            response = _send(opened, attachment=False)
            _record_event("PREVIEW", result="SUCCESS", file_id=file_id, opened=opened)
            return response
        if extension in TEXT_EXTENSIONS:
            if opened["sizeBytes"] > TEXT_PREVIEW_MAX_BYTES:
                return unavailable()
            content = opened["stream"].read(TEXT_PREVIEW_MAX_BYTES + 1).decode("utf-8")
            opened["stream"].close()
            response = jsonify({"success": True, "type": "text", "filename": opened["originalName"], "content": content})
            _record_event("PREVIEW", result="SUCCESS", file_id=file_id, opened=opened)
            return response
        if extension == ".docx":
            try:
                from docx import Document
                if not office_archive_is_preview_safe(opened["stream"]):
                    return unavailable()
                opened["stream"].seek(0)
                document = Document(opened["stream"])
                budget = [OFFICE_PREVIEW_CHARACTER_BUDGET]
                paragraphs = []
                for item in islice(document.paragraphs, 500):
                    if item.text.strip():
                        paragraphs.append(_take_text(item.text, budget))
                tables = []
                for table in islice(document.tables, 20):
                    table_rows = []
                    for row in islice(table.rows, 100):
                        table_rows.append([_take_text(cell.text, budget) for cell in islice(row.cells, 50)])
                    tables.append(table_rows)
                opened["stream"].close()
                response = jsonify({"success": True, "type": "word", "filename": opened["originalName"], "paragraphs": paragraphs, "tables": tables})
                _record_event("PREVIEW", result="SUCCESS", file_id=file_id, opened=opened)
                return response
            except FileServiceError:
                raise
            except Exception:
                return unavailable()
        if extension == ".xlsx":
            workbook = None
            try:
                import openpyxl
                if not office_archive_is_preview_safe(opened["stream"]):
                    return unavailable()
                opened["stream"].seek(0)
                workbook = openpyxl.load_workbook(opened["stream"], data_only=True, read_only=True)
                sheets = []
                budget = [OFFICE_PREVIEW_CHARACTER_BUDGET]
                for sheet in list(workbook.worksheets)[:5]:
                    rows = [[_take_text(value, budget) for value in row] for row in sheet.iter_rows(max_row=100, max_col=100, values_only=True)]
                    sheets.append({"name": sheet.title, "rows": rows})
                workbook.close()
                opened["stream"].close()
                response = jsonify({"success": True, "type": "excel", "filename": opened["originalName"], "sheets": sheets})
                _record_event("PREVIEW", result="SUCCESS", file_id=file_id, opened=opened)
                return response
            except FileServiceError:
                raise
            except Exception:
                return unavailable()
            finally:
                if workbook is not None:
                    workbook.close()
        return unavailable()
    except (FileServiceError, ValueError, OSError) as raw_error:
        error = raw_error if isinstance(raw_error, FileServiceError) else FileServiceError("FILE_OPERATION_FAILED", "文件预览失败", 500)
        if opened and not opened["stream"].closed:
            opened["stream"].close()
        _record_failure("PREVIEW", error, file_id=file_id, opened=opened,
                        legacy_project_id=legacy_project_id)
        if download_url and error.code == "PREVIEW_UNAVAILABLE":
            return _preview_unavailable(file_id, version_no, download_url=download_url)
        return _error(error)
