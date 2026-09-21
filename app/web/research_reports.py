"""Authenticated plain-text research report versions and stored-version exports."""
from io import BytesIO
import re
import json
import threading
import uuid
from functools import wraps

from sqlalchemy.exc import SQLAlchemyError

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from flask import Blueprint, abort, current_app, jsonify, redirect, render_template, request, send_file, url_for

from app.security.auth import business_required, current_identity
from app.services.research_reports import ResearchReportServiceError
from app.services.files import FileServiceError
from app.ai.local_selection import LocalSelectionAssistant
from app.services.sensitive_terms import replace_sensitive_terms

bp = Blueprint('research_reports', __name__)
_selection_jobs = {}
_selection_jobs_lock = threading.Lock()


def ai_features_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if current_app.config.get('AI_FEATURES_VISIBLE') is not True:
            abort(404)
        return view(*args, **kwargs)
    return wrapped


def _service():
    service = current_app.extensions.get('research_report_service')
    if service is None:
        raise ResearchReportServiceError('UNAVAILABLE', '科研报告服务尚未就绪', 503)
    return service


def _actor():
    return dict(actor_user_id=current_identity().user_id, request_id=getattr(request, 'request_id', 'unknown'))


def _draft_service():
    service = current_app.extensions.get('research_report_draft_service')
    if service is None:
        raise ResearchReportServiceError('UNAVAILABLE', '本地报告整合服务尚未就绪', 503)
    return service


def _source_context(report, selected=None):
    file_service = current_app.extensions.get('file_service')
    available = file_service.list_for_object(object_type=report['objectType'], object_id=report['objectId']) if file_service else []
    generation_available=(current_app.config.get('AI_FEATURES_VISIBLE') is True
                          and current_app.extensions.get('research_report_draft_service') is not None)
    generation_dictionary_version=None
    generation_error=None
    if generation_available:
        try:
            metadata,_=_service().redaction_context(required=True)
            generation_dictionary_version=metadata['version']
        except ResearchReportServiceError:
            generation_available=False
            generation_error='词库服务暂不可用，资料整合已暂停，仍可核对和编辑正文。'
    return dict(available_sources=available,
                selected_pairs={(item['fileId'], item['versionNo']) for item in (selected or [])},
                generation_available=generation_available,generation_dictionary_version=generation_dictionary_version,generation_error=generation_error)


def _selections():
    selections = []
    submitted = request.form.getlist('sources')
    if len(submitted) > 8:
        raise ResearchReportServiceError('INVALID_SOURCES', '请选择不超过8份资料', 422)
    for value in submitted:
        try:
            pair = json.loads(value)
            if not isinstance(pair, list) or len(pair) != 2 or not isinstance(pair[0], str) or not pair[0] or len(pair[0]) > 200 or type(pair[1]) is not int or pair[1] < 1:
                raise ValueError()
        except (ValueError, TypeError):
            raise ResearchReportServiceError('INVALID_SOURCES', '资料版本无效，请重新选择', 422)
        selections.append(dict(fileId=pair[0], versionNo=pair[1]))
    return selections


def _positive_integer(value, field):
    if not isinstance(value, str) or not re.fullmatch(r'[1-9][0-9]{0,8}', value):
        raise ResearchReportServiceError('INVALID_INTEGER', '版本或页码必须是正整数', 422, fields={field: '请输入正整数'})
    return int(value)


def _dictionary_form_payload(form, field='dictionaryVersion', target='dictionaryVersion'):
    value=form.get(field)
    if value in (None,''):
        return {}
    if not isinstance(value,str) or not re.fullmatch(r'0|[1-9][0-9]{0,8}',value):
        raise ResearchReportServiceError('VALIDATION_ERROR','词库版本号无效',422)
    return {target:int(value)}


def _log_service_error(error):
    if error.status_code >= 500:
        current_app.logger.error('research report service unavailable')


def _error(error):
    _log_service_error(error)
    return render_template('research_reports/error.html', error=error.message), error.status_code


def _configure_report_styles(document):
    for name, size, bold in [('Normal', 11, False), ('Title', 22, True), ('Heading 1', 16, True), ('Heading 2', 14, True)]:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            style.element.get_or_add_rPr().append(fonts)
        fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    normal = document.styles['Normal'].paragraph_format
    normal.line_spacing = 1.5
    normal.space_after = Pt(6)
    title_properties = document.styles['Title'].element.get_or_add_pPr()
    borders = title_properties.find(qn('w:pBdr'))
    if borders is not None:
        title_properties.remove(borders)


@bp.get('/research-reports')
@business_required
def list_page():
    object_type, object_id = request.args.get('objectType', ''), request.args.get('objectId', '')
    try:
        result = _service().list_for_object(object_type, object_id, page=_positive_integer(request.args.get('page', '1'), 'page'))
        return render_template('research_reports/list.html', result=result, object_type=object_type, object_id=object_id)
    except ResearchReportServiceError as error:
        return _error(error)


@bp.route('/research-reports/new', methods=['GET', 'POST'])
@business_required
def new_page():
    object_type, object_id = request.args.get('objectType', ''), request.args.get('objectId', '')
    form = request.form.to_dict() if request.method == 'POST' else {}
    try:
        # Validate the bound object even before rendering an empty creation form.
        _service().list_for_object(object_type, object_id, page_size=1)
        if request.method == 'POST':
            try:
                report = _service().create({key: form.get(key, '') for key in ('title', 'purpose', 'body', 'note')} | dict(objectType=object_type, objectId=object_id), **_actor())
                return redirect(url_for('research_reports.detail_page', report_id=report['id']))
            except ResearchReportServiceError as error:
                _log_service_error(error)
                return render_template('research_reports/new.html', form=form, error=error.message, object_type=object_type, object_id=object_id), error.status_code
            except SQLAlchemyError:
                current_app.logger.error('research report create failed: database unavailable')
                return render_template('research_reports/new.html', form=form, error='保存失败，服务暂不可用。输入已保留，请复制正文后重试并核对报告列表。', object_type=object_type, object_id=object_id), 503
        return render_template('research_reports/new.html', form=form, object_type=object_type, object_id=object_id)
    except ResearchReportServiceError as error:
        _log_service_error(error)
        return render_template('research_reports/new.html', form=form, error=error.message, object_type=object_type, object_id=object_id), error.status_code
    except SQLAlchemyError:
        current_app.logger.error('research report create request failed: database unavailable')
        return render_template('research_reports/new.html', form=form, error='服务暂不可用。输入已保留，请复制正文后重试。', object_type=object_type, object_id=object_id), 503


@bp.route('/research-reports/<report_id>', methods=['GET', 'POST'])
@business_required
def detail_page(report_id):
    try:
        service = _service()
        report = service.get(report_id)
        current = service.current(report_id)
        history = service.history(report_id, page=1 if request.method == 'POST' else _positive_integer(request.args.get('page', '1'), 'page'))
        form = request.form.to_dict() if request.method == 'POST' else dict(body=current['body'], note='', baseVersion=current['versionNo'])
        error_message, status = None, 200
        if request.method == 'POST':
            try:
                service.save_version(report_id, dict(baseVersion=_positive_integer(form.get('baseVersion', ''), 'baseVersion'), body=form.get('body', ''), note=form.get('note', ''), **_dictionary_form_payload(form)), **_actor())
                return redirect(url_for('research_reports.detail_page', report_id=report_id))
            except ResearchReportServiceError as error:
                _log_service_error(error)
                # Keep the submitted baseVersion and body: never silently rebase over a conflict.
                error_message, status = error.message, error.status_code
                if status == 409:
                    report, current = service.get(report_id), service.current(report_id)
            except SQLAlchemyError:
                current_app.logger.error('research report save failed: database unavailable')
                error_message, status = '保存失败，服务暂不可用。输入已保留，请复制正文后重试并核对版本。', 503
        return render_template('research_reports/detail.html', report=report, current=current, form=form, history=history, error=error_message, conflict=status == 409, **_source_context(report)), status
    except ResearchReportServiceError as error:
        return _error(error)
    except SQLAlchemyError:
        current_app.logger.error('research report request failed: database unavailable')
        return render_template('research_reports/error.html', error='服务暂不可用，请保留输入后重试。', recovery_form=request.form.to_dict() if request.method == 'POST' else None), 503


@bp.post('/research-reports/<report_id>/generate')
@ai_features_required
@business_required
def generate_page(report_id):
    form = request.form.to_dict()
    selected = []
    try:
        report = _service().get(report_id)
        current = _service().current(report_id)
        history = _service().history(report_id)
        source_context = _source_context(report)
        try:
            selected = _selections()
            source_context['selected_pairs'] = {(item['fileId'], item['versionNo']) for item in selected}
            if source_context['generation_dictionary_version'] is None:
                raise ResearchReportServiceError('REDACTION_UNAVAILABLE','当前词库不可用，未执行资料整合',503)
            target_dictionary=_dictionary_form_payload(form,'generationDictionaryVersion')
            target_dictionary.setdefault('dictionaryVersion',source_context['generation_dictionary_version'])
            if target_dictionary['dictionaryVersion']!=source_context['generation_dictionary_version']:
                raise ResearchReportServiceError('REDACTION_VERSION_CONFLICT','词库已更新，请核对本页显示的整合词库版本后重试',409)
            draft = _draft_service().generate(report_id, dict(baseVersion=_positive_integer(form.get('baseVersion', ''), 'baseVersion'), body=form.get('body', ''), selections=selected, **target_dictionary, **_dictionary_form_payload(form,'dictionaryVersion','inputDictionaryVersion')), **_actor())
            return redirect(url_for('research_reports.draft_page', report_id=report_id, draft_id=draft['id']))
        except ResearchReportServiceError as error:
            _log_service_error(error)
            return render_template('research_reports/detail.html', report=report, current=current, history=history, form=form, error=error.message, conflict=error.status_code == 409, **source_context), error.status_code
        except SQLAlchemyError:
            current_app.logger.error('research report generation failed: database unavailable')
            return render_template('research_reports/detail.html', report=report, current=current, history=history, form=form, error='整合失败，输入已保留，请稍后重试。', **source_context), 503
    except ResearchReportServiceError as error:
        _log_service_error(error)
        return render_template('research_reports/error.html', error=error.message, recovery_form=form, recovery_sources=request.form.getlist('sources')), error.status_code
    except SQLAlchemyError:
        current_app.logger.error('research report generation failed: database unavailable')
        return render_template('research_reports/error.html', error='整合失败，输入尚未保存，请复制保留后重试。', recovery_form=form, recovery_sources=request.form.getlist('sources')), 503


@bp.route('/research-reports/<report_id>/drafts/<draft_id>', methods=['GET', 'POST'])
@ai_features_required
@business_required
def draft_page(report_id, draft_id):
    form = request.form.to_dict() if request.method == 'POST' else None
    try:
        report = _service().get(report_id)
        draft = _draft_service().get(report_id, draft_id)
        if form is None:
            form = dict(body=draft['body'], baseVersion=draft['baseVersion'], note='')
        error_message, status = None, 200
        if request.method == 'POST':
            try:
                _service().save_version(report_id, dict(baseVersion=_positive_integer(form.get('baseVersion', ''), 'baseVersion'), body=form.get('body', ''), note=form.get('note', ''), draftId=draft_id, **_dictionary_form_payload(form)), **_actor())
                return redirect(url_for('research_reports.detail_page', report_id=report_id))
            except ResearchReportServiceError as error:
                _log_service_error(error)
                error_message, status = error.message, error.status_code
            except SQLAlchemyError:
                current_app.logger.error('research report draft save failed: database unavailable')
                error_message, status = '保存失败，输入已保留，请复制正文后重试并核对版本。', 503
        return render_template('research_reports/draft.html', report=report, draft=draft, form=form, error=error_message, conflict=status == 409), status
    except ResearchReportServiceError as error:
        _log_service_error(error)
        return render_template('research_reports/error.html', error=error.message, recovery_form=form, recovery_sources=request.form.getlist('sources')), error.status_code
    except SQLAlchemyError:
        current_app.logger.error('research report draft request failed: database unavailable')
        return render_template('research_reports/error.html', error='服务暂不可用，请先复制未保存内容。', recovery_form=form), 503


@bp.get('/research-reports/<report_id>/versions/<int:version_no>')
@business_required
def version_page(report_id, version_no):
    try:
        return render_template('research_reports/version.html', report=_service().get(report_id), version=_service().get_version(report_id, version_no))
    except ResearchReportServiceError as error:
        return _error(error)


@bp.get('/research-reports/<report_id>/versions/<int:version_no>/model-draft')
@ai_features_required
@business_required
def model_draft_page(report_id, version_no):
    try:
        report = _service().get(report_id)
        version = _service().get_version(report_id, version_no)
        draft_id = version.get('draftId')
        if not draft_id:
            raise ResearchReportServiceError('DRAFT_NOT_FOUND', '此版本没有关联模型原稿', 404)
        draft = _draft_service().get(report_id, draft_id)
        if str(draft.get('reportId')) != str(report_id) or str(draft.get('id')) != str(draft_id):
            raise ResearchReportServiceError('DRAFT_NOT_FOUND', '模型原稿不存在或不属于此报告', 404)
        return render_template('research_reports/model_draft.html', report=report, version=version, draft=draft)
    except ResearchReportServiceError as error:
        return _error(error)


@bp.get('/research-reports/<report_id>/versions/<int:version_no>/export.docx')
@business_required
def export_version(report_id, version_no):
    try:
        report, version = _service().get(report_id), _service().get_version(report_id, version_no)
        document = Document()
        _configure_report_styles(document)
        document.add_heading(version.get('title',report['title']), 0)
        dictionary_label=f"v{version['redactionSnapshot']['version']}" if version.get('redactionSnapshot') else '未记录'
        document.add_paragraph(f"版本 v{version['versionNo']} · {version['createdAt']} · 作者：{version.get('createdByName') or '姓名未记录'} · 词库：{dictionary_label}")
        document.add_paragraph(f"用途：{version.get('purpose',report['purpose'])}")
        document.add_paragraph(f"修改说明：{version['note'] or '未填写'}")
        # Interpret stored newline conventions only for DOCX paragraph boundaries.
        export_body = version['body'].replace('\r\n', '\n').replace('\r', '\n')
        for line in export_body.split('\n'):
            document.add_paragraph(line)
        image_sources = [source for source in version.get('sources', []) if source.get('sourceKind') == 'IMAGE']
        if image_sources:
            document.add_heading('图像资料', level=1)
            file_service = current_app.extensions.get('file_service')
            for image_number, source in enumerate(image_sources, 1):
                caption = f"图{image_number} [{source.get('label', '')}] {source.get('filename', '')}"
                if source.get('originalImageBlockedByRedaction') is True:
                    document.add_paragraph(caption + '（原图命中敏感词库，未嵌入）')
                    lines = [line.strip() for line in source.get('text', '').splitlines() if line.strip()]
                    summary = next((line for line in lines if line.startswith('摘要：')), None)
                    if summary is None:
                        summary = next((line for line in lines if not line.startswith('【图片资料')), '')
                    if summary:
                        document.add_paragraph('脱敏提取摘要：' + summary.removeprefix('摘要：'))
                    continue
                if file_service is None:
                    raise ResearchReportServiceError('IMAGE_SOURCE_UNAVAILABLE', '图像来源服务暂不可用，DOCX未导出', 503)
                try:
                    opened = file_service.open_version_stream(
                        source['fileId'], source['versionNo'],
                        object_type=source['objectType'], object_id=source['objectId'],
                    )
                    if opened.get('sha256') != source.get('sha256'):
                        raise FileServiceError('FILE_INTEGRITY_FAILED', '文件完整性校验失败', 409)
                    try:
                        document.add_picture(opened['stream'], width=Inches(6.0))
                    finally:
                        opened['stream'].close()
                except (FileServiceError, KeyError, ValueError, OSError) as error:
                    raise ResearchReportServiceError('IMAGE_SOURCE_UNAVAILABLE', '图像来源不可用，DOCX未导出，请核对附件版本', 409) from error
                document.add_paragraph(caption + '（由程序按固定版式插入）')
        if version.get('sources'):
            document.add_heading('来源清单', level=1)
            for source in version['sources']:
                document.add_paragraph(f"[{source.get('label', '')}] {source.get('filename') or source.get('originalName', '')} · v{source['versionNo']}")
                document.add_paragraph(f"文件编号：{source['fileId']} · SHA-256：{source['sha256']}")
            if current_app.config.get('AI_FEATURES_VISIBLE') is True and version.get('model'):
                document.add_paragraph(f"辅助整合模型：{version['model']}")
        output = BytesIO()
        document.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"research-report-{report['id']}-v{version['versionNo']}.docx", mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except ResearchReportServiceError as error:
        return _error(error)


def _selection_error(message, status, code):
    return jsonify(error=dict(code=code,message=message)), status


@bp.post('/research-reports/<report_id>/selection-suggestions')
@ai_features_required
@business_required
def selection_suggestions(report_id):
    payload=request.get_json(silent=True)
    if not isinstance(payload,dict):
        return _selection_error('请提供有效请求',422,'INVALID_SELECTION')
    body,prefix,selected,suffix=(payload.get(key) for key in ('body','prefix','selectedText','suffix'))
    base=payload.get('baseVersion')
    if (any(not isinstance(value,str) for value in (body,prefix,selected,suffix))
            or len(body)>200000 or len(prefix)>200000 or len(suffix)>200000
            or not selected.strip() or len(selected)>2000 or prefix+selected+suffix!=body
            or type(base) is not int or base<1):
        return _selection_error('选区已失效或超出长度限制，请重新框选',422,'INVALID_SELECTION')
    try:
        report=_service().get(report_id)
        if report['currentVersion']!=base:
            return _selection_error('报告版本已更新，请先核对最新版本',409,'VERSION_CONFLICT')
        draft_assistant=_draft_service().assistant
        reference=_draft_service().get(report_id,payload['draftId']) if payload.get('draftId') else _service().current(report_id)
        metadata,compiled=_service().redaction_context(reference.get('redactionSnapshot'),version_no=payload.get('dictionaryVersion'),required=True)
        dictionary_version=metadata['version']
        processor=lambda text: replace_sensitive_terms(text,compiled)['text']
        from app.services.local_model_runtime import configured, make_assistant
        if configured(current_app.config):
            assistant = make_assistant(LocalSelectionAssistant, current_app.config, processor=processor)
        else:
            assistant=LocalSelectionAssistant(base_url=draft_assistant.base_url,model=draft_assistant.model_version,processor=processor)
        operation_id=str(uuid.UUID(payload.get('operationId'))) if payload.get('operationId') else uuid.uuid4().hex
        key=(current_identity().user_id,report_id,operation_id)
        cancelled=threading.Event()
        with _selection_jobs_lock:
            if key in _selection_jobs or len(_selection_jobs)>=32:
                return _selection_error('建议请求正在处理中，请稍后再试',429,'SELECTION_BUSY')
            _selection_jobs[key]=cancelled
        try:
            result=assistant.suggest(selected_text=selected,prefix=prefix,suffix=suffix,cancel_check=cancelled.is_set)
            if cancelled.is_set():
                return _selection_error('建议请求已取消',409,'SELECTION_CANCELLED')
            if _service().get(report_id)['currentVersion']!=base:
                return _selection_error('生成期间报告版本已更新，建议未采用',409,'VERSION_CONFLICT')
            result['dictionaryVersion']=dictionary_version
            return jsonify(data=result)
        finally:
            with _selection_jobs_lock:
                _selection_jobs.pop(key,None)
    except ResearchReportServiceError as error:
        _log_service_error(error)
        return _selection_error(error.message,error.status_code,error.code)
    except (ValueError,TypeError,AttributeError):
        return _selection_error('选区或模型建议未通过校验，请保留原文并重新选择',422,'SELECTION_VALIDATION_FAILED')
    except InterruptedError:
        return _selection_error('建议请求已取消',409,'SELECTION_CANCELLED')
    except (TimeoutError,RuntimeError,SQLAlchemyError):
        current_app.logger.error('local selection suggestion unavailable')
        return _selection_error('本地建议暂不可用，原文保持不变',503,'SELECTION_UNAVAILABLE')


@bp.post('/research-reports/<report_id>/selection-suggestions/cancel')
@ai_features_required
@business_required
def cancel_selection_suggestions(report_id):
    payload=request.get_json(silent=True)
    try:
        operation_id=str(uuid.UUID(payload['operationId']))
    except (ValueError,TypeError,KeyError,AttributeError):
        return _selection_error('取消请求无效',422,'INVALID_SELECTION')
    key=(current_identity().user_id,report_id,operation_id)
    with _selection_jobs_lock:
        event=_selection_jobs.get(key)
        if event is not None:
            event.set()
    return jsonify(data=dict(cancelled=event is not None))
