# -*- coding: utf-8 -*-
"""
方案论证模块 - 路由
"""
import os
from pathlib import Path
from flask import Blueprint, render_template, request, jsonify, session, redirect, url_for, send_file
from flask import current_app
import json
import uuid
import shutil

# 导入模型
from app.models_argumentation import (
    init_argumentation_db,
    get_template_by_category,
    save_template,
    get_document_by_project,
    get_document_by_id,
    save_revision,
    get_document_versions,
    get_version_by_num
)
from app.models_argumentation import get_projects, get_equipment
from app.repositories.argumentation import ArgumentationConflict
from app.services.files import FileServiceError

argumentation_bp = Blueprint('argumentation', __name__)

# 板块分类
CATEGORIES = {
    '科研项目': 'research',
    '密码项目': 'crypto', 
    '安全项目': 'security'
}

# 允许的文件扩展名
ALLOWED_EXTENSIONS = {'docx', 'doc'}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def init_template_data():
    """初始化示例模板数据"""
    # 检查是否已有模板
    for cat in ['research', 'crypto', 'security']:
        existing = get_template_by_category(cat)
        if existing:
            continue
        
        # 为不同类型创建不同模板
        if cat == 'research':
            template_name = "科研项目论证方案模板"
        elif cat == 'crypto':
            template_name = "密码项目论证方案模板"
        else:
            template_name = "安全项目论证方案模板"
        
        sample_template = {
            "template_id": f"{cat}_v1",
            "name": template_name,
            "category": cat,
            "version": 1,
            "chapters": [
                {"id": "ch1", "title": "一、项目概述", "type": "text",
                 "children": [
                     {"id": "ch1_1", "title": "1.1 项目背景", "type": "text"},
                     {"id": "ch1_2", "title": "1.2 研究目标", "type": "text"}
                 ]},
                {"id": "ch2", "title": "二、需求分析", "type": "text",
                 "children": [
                     {"id": "ch2_1", "title": "2.1 功能需求", "type": "text"},
                     {"id": "ch2_2", "title": "2.2 性能需求", "type": "text"}
                 ]},
                {"id": "ch3", "title": "三、方案设计", "type": "text",
                 "children": [
                     {"id": "ch3_1", "title": "3.1 总体设计", "type": "text"},
                     {"id": "ch3_2", "title": "3.2 详细设计", "type": "text"}
                 ]},
                {"id": "ch4", "title": "四、经费预算", "type": "budget"},
                {"id": "ch5", "title": "五、进度安排", "type": "text"},
                {"id": "ch6", "title": "六、风险分析", "type": "text"},
                {"id": "ch7", "title": "七、结论", "type": "text"}
            ]
        }
        
        save_template(sample_template['template_id'], sample_template['name'], cat, '', json.dumps(sample_template, ensure_ascii=False))
def _template_data(category):
    stored = get_template_by_category(category)
    if stored and stored.get('chapter_tree'):
        return json.loads(stored['chapter_tree'])
    return {'chapters': [
        {'id': 'ch1', 'title': '一、项目概述', 'type': 'text',
         'children': [{'id': 'ch1_1', 'title': '1.1 项目背景', 'type': 'text'}]},
        {'id': 'ch2', 'title': '二、需求分析', 'type': 'text',
         'children': [{'id': 'ch2_1', 'title': '2.1 功能需求', 'type': 'text'}]},
        {'id': 'ch3', 'title': '三、方案设计', 'type': 'text',
         'children': [{'id': 'ch3_1', 'title': '3.1 总体设计', 'type': 'text'}]},
        {'id': 'ch4', 'title': '四、经费预算', 'type': 'text'},
        {'id': 'ch5', 'title': '五、进度安排', 'type': 'text'},
        {'id': 'ch6', 'title': '六、风险分析', 'type': 'text'},
        {'id': 'ch7', 'title': '七、结论', 'type': 'text'},
    ]}


def _editor_context(category, project, document, *, version=None):
    template_data = _template_data(category)
    content = json.loads(document['content']) if document and document.get('content') else {}
    template_content = {}
    for chapter in template_data.get('chapters', []):
        for item in [chapter, *chapter.get('children', [])]:
            if 'content' in item:
                template_content[item['id']] = item['content']
    equipment = get_equipment() if version is None else []
    return dict(
        category=category,
        category_name={'research': '科研项目', 'security': '安全保密项目', 'crypto': '密码应用项目'}[category],
        project=project, template=template_data, template_data=template_data,
        template_content=template_content, document=document, saved_content=content,
        equipment_list=equipment,
        categories=sorted({item['category'] for item in equipment if item.get('category')}),
        project_equipment=[],
        versions=get_document_versions(document['doc_id']) if document and version is None else [],
        back_url=url_for('argumentation.edit', category=category, project_id=project['project_id']) if version else url_for('argumentation.index', category=category),
        is_version_view=version is not None, current_version=version,
    )


@argumentation_bp.route('/<category>')
def index(category):
    """Retain the category entry point, reading only the runtime database."""
    try:
        projects = get_projects(category)
    except ValueError:
        return "项目类别不存在", 404
    template = get_template_by_category(category)
    return render_template(
        'argumentation/index.html', category=category,
        category_name={'research': '科研项目', 'crypto': '密码应用项目', 'security': '安全保密项目'}[category],
        projects=projects, template=template, template_data=_template_data(category),
    )


@argumentation_bp.route('/<category>/<project_id>')
def edit(category, project_id):
    try:
        projects = get_projects(category, project_id)
    except ValueError:
        return "项目类别不存在", 404
    if not projects:
        return "项目不存在", 404
    document = get_document_by_project(project_id, category)
    return render_template('argumentation/edit.html',
                           **_editor_context(category, projects[0], document))


@argumentation_bp.route('/save', methods=['POST'])
def save():
    """保存论证文档"""
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})
    
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({'success': False, 'message': '参数格式不正确'}), 400
    project_id = data.get('project_id')
    category = data.get('category')
    raw_content = data.get('content', {})
    if not isinstance(raw_content, dict):
        return jsonify({'success': False, 'message': '正文必须为章节对象'}), 400
    content = json.dumps(raw_content, ensure_ascii=False)
    change_note = data.get('change_note', '')
    
    if not isinstance(project_id, str) or not project_id or not isinstance(category, str):
        return jsonify({'success': False, 'message': '参数不完整'}), 400
    try:
        doc_id, version_num = save_revision(project_id, category, content, change_note, data.get('expected_version'))
    except ArgumentationConflict as exc:
        return jsonify({'success': False, 'message': str(exc)}), 409
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    except LookupError:
        return jsonify({'success': False, 'message': '项目不存在'}), 404
    
    return jsonify({
        'success': True, 
        'message': f'保存成功（版本 {version_num}）',
        'doc_id': doc_id,
        'version': version_num
    })


@argumentation_bp.route('/generate_word', methods=['POST'])
def generate_word():
    """生成Word文档 - 前端组装数据，后端直接生成"""
    try:
        data = request.json
        project_name = data.get('project_name')          # 项目名称（前端传）
        category = data.get('category')              # 类型
        template = data.get('template', [])         # 模板结构
        content = data.get('content', {})            # 内容
        
        if not project_name or not category:
            return jsonify({'success': False, 'message': '参数不完整'})
        
        # 文件名
        word_filename = f"论证方案-{uuid.uuid4().hex}.docx"
        word_dir = os.path.join(current_app.config.get('DATA_DIR', 'data'), 'documents')
        os.makedirs(word_dir, exist_ok=True)
        word_path = os.path.join(word_dir, word_filename)
        
        # 生成Word
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return jsonify({'success': False, 'message': '缺少python-docx库'})
        
        doc = Document()
        title = doc.add_heading(project_name, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 添加内容
        import re
        for chapter in template:
            doc.add_heading(chapter.get('title', ''), level=1)
            for sub in chapter.get('children', []):
                sub_id = sub.get('id', '')
                doc.add_heading(sub.get('title', ''), level=2)
                sub_content = content.get(sub_id, '')
                if sub_content:
                    text = re.sub('<[^>]+>', '', sub_content)
                    text = text.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>')
                    if text.strip():
                        doc.add_paragraph(text)
            if not chapter.get('children'):
                ch_content = content.get(chapter.get('id', ''), '')
                if ch_content:
                    text = re.sub('<[^>]+>', '', ch_content)
                    text = text.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>')
                    if text.strip():
                        doc.add_paragraph(text)
        
        doc.save(word_path)
        
        # 返回相对路径
        relative_path = f'data/documents/{word_filename}'
        
        return jsonify({
            'success': True, 
            'message':'生成word成功',
            'word_path': relative_path
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})
            

@argumentation_bp.route('/versions/<doc_id>')
def versions(doc_id):
    """查看文档版本历史"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    versions = get_document_versions(doc_id)
    document = get_document_by_id(doc_id)
    if not document:
        return "文档不存在", 404
    
    return render_template('argumentation/versions.html',
                           versions=versions,
                           document=document)


@argumentation_bp.route('/versions/<doc_id>/<int:version_num>')
def view_version(doc_id, version_num):
    """查看指定版本"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    version = get_version_by_num(doc_id, version_num)
    if not version:
        return "版本不存在", 404
    
    document = get_document_by_id(doc_id)
    if not document:
        return "文档不存在", 404
    projects = get_projects(document['category'], document['project_id'])
    if not projects:
        return "项目不存在", 404
    historical_document = {**document, 'content': version['content']}
    return render_template('argumentation/edit.html', **_editor_context(
        document['category'], projects[0], historical_document, version=version_num,
    ))


@argumentation_bp.route('/template/upload', methods=['POST'])
def upload_template():
    """上传新模板"""
    if 'user' not in session:
        return jsonify({'success': False, 'message': '权限不足'})
    
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请选择文件'})
    
    file = request.files['file']
    category = request.form.get('category')
    name = request.form.get('name', '未命名模板')

    if category not in CATEGORIES.values():
        return jsonify({'success': False, 'message': '请选择有效板块'}), 400
    
    if file.filename == '':
        return jsonify({'success': False, 'message': '请选择文件'})
    
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'message': '只支持docx格式'})

    data_root = Path(current_app.config.get('DATA_DIR', 'data'))
    upload_dir = data_root / 'templates'
    try:
        if data_root.is_symlink() or upload_dir.is_symlink():
            raise OSError('Template directory is a symbolic link')
        upload_dir.mkdir(parents=True, exist_ok=True)
        if upload_dir.resolve().parent != data_root.resolve():
            raise OSError('Template directory escapes data root')
    except OSError:
        return jsonify({'success': False, 'message': '模板存储目录不可用'}), 500

    file_service = current_app.extensions.get('file_service')
    if file_service is None:
        return jsonify({'success': False, 'message': '文件服务不可用，请稍后重试'}), 503
    try:
        with file_service.inspect_upload(file.stream, file.filename):
            file.stream.seek(0)
    except FileServiceError as exc:
        return jsonify({'success': False, 'message': exc.message}), exc.status_code
    
    file_path = upload_dir / f"{uuid.uuid4().hex}.docx"
    
    # 上传替换源文件，不清空已编辑章节，也不冒充自动解析 Word。
    chapter_tree = json.dumps(_template_data(category), ensure_ascii=False)
    
    # 保存模板
    template_id = f"{category}_v1"
    created = False
    try:
        with file_path.open('xb') as destination:
            created = True
            shutil.copyfileobj(file.stream, destination)
        save_template(template_id, name, category, str(file_path), chapter_tree)
    except Exception:
        if created:
            file_path.unlink(missing_ok=True)
        current_app.logger.exception('Template upload persistence failed')
        return jsonify({'success': False, 'message': '模板保存失败，请稍后重试'}), 500
    
    return jsonify({'success': True, 'message': '模板上传成功'})


@argumentation_bp.route('/template/edit', methods=['POST'])
def edit_template():
    """编辑模板章节结构"""
    if 'user' not in session:
        return jsonify({'success': False, 'message': '请先登录'})
    
    data = request.json
    category = data.get('category')
    chapter_tree = data.get('chapter_tree')
    
    if not category or not chapter_tree:
        return jsonify({'success': False, 'message': '参数不完整'})
    
    template_id = f"{category}_v1"
    category_name = {'research': '科研项目论证方案模板', 'crypto': '密码项目论证方案模板', 'security': '安全项目论证方案模板'}.get(category, '论证方案模板')
    
    existing = get_template_by_category(category)
    
    if existing:
        # 更新
        save_template(template_id, existing['name'], category, existing['file_path'], chapter_tree)
    else:
        # 创建新模板
        save_template(template_id, category_name, category, '', chapter_tree)
    
    return jsonify({'success': True, 'message': '模板保存成功'})


def init_argumentation_routes(app):
    """初始化路由"""
    app.register_blueprint(argumentation_bp, url_prefix='/argumentation')
    init_argumentation_db()
    init_template_data()


@argumentation_bp.route('/download')
def download_word():
    """下载Word文档"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    filepath = request.args.get('path', '')
    if not filepath:
        return "缺少文件路径", 400
    
    # Only accept the legacy reference emitted by generate_word, never an OS path.
    prefix = 'data/documents/'
    if not filepath.startswith(prefix):
        return "非法文档引用", 400
    filename = filepath[len(prefix):]
    if (not filename or any(part in filename for part in ('/', '\\', '\x00', '%'))
            or Path(filename).suffix.lower() not in {'.docx', '.doc'}):
        return "非法文档引用", 400
    data_root = Path(current_app.config['DATA_DIR']).resolve()
    document_root = data_root / 'documents'
    candidate = document_root / filename
    try:
        if (document_root.is_symlink() or candidate.is_symlink()
                or candidate.resolve().parent != document_root
                or not candidate.is_file()):
            return "文件不存在", 404
    except (OSError, ValueError):
        return "文件不存在", 404
    return send_file(candidate, as_attachment=True)
