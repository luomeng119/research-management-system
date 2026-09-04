# -*- coding: utf-8 -*-
from flask import Blueprint, current_app, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from datetime import datetime
import json
import zipfile
import openpyxl
from io import BytesIO

bp = Blueprint('experts', __name__, url_prefix='/experts')

def _resources_service():
    service = current_app.extensions.get('resources_service')
    if service is None:
        raise RuntimeError('专家库服务未就绪')
    return service


def _actor():
    return {
        'user_id': int(session.get('user_id') or 0),
        'name': session.get('name') or session.get('user') or '',
        'request_id': getattr(request, 'request_id', 'legacy-expert-request'),
    }


def _template_expert(item):
    return {
        'id': item.get('id'),
        'expert_id': item.get('expertId'),
        'name': item.get('name'),
        'unit': item.get('unit'),
        'position': item.get('position'),
        'expertise': item.get('expertise'),
        'phone': item.get('phone'),
        'id_card': item.get('idCard'),
        'bank_card': item.get('bankCard'),
        'bank_name': item.get('bankName'),
        'uploader': item.get('uploader'),
        'created_at': item.get('createdAt'),
        'updated_at': item.get('updatedAt'),
    }


def _form_payload():
    return {
        'name': request.form.get('name', ''),
        'unit': request.form.get('unit', ''),
        'position': request.form.get('position', ''),
        'expertise': request.form.get('expertise', ''),
        'bankCard': request.form.get('bank_card', ''),
        'bankName': request.form.get('bank_name', ''),
        'phone': request.form.get('phone', ''),
        'idCard': request.form.get('id_card', ''),
    }

@bp.route('/')
def index():
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    keyword = request.args.get('search', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
    except (TypeError, ValueError):
        page = 1
    per_page = 20
    partial = request.args.get('partial') == '1'

    result = _resources_service().list_experts(
        page=page, page_size=per_page, keyword=keyword
    )
    experts_page = [_template_expert(item) for item in result['items']]
    total = result['total']
    total_pages = (total + per_page - 1) // per_page if total > 0 else 1

    if partial:
        return render_template('experts/partial_table.html',
            experts=experts_page, search_keyword=keyword, page=page,
            total_pages=total_pages, total=total)

    return render_template('experts/index.html',
        experts=experts_page, search_keyword=keyword, page=page,
        total_pages=total_pages, total=total)

@bp.route('/add', methods=['GET', 'POST'])
def add():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    if request.method == 'POST':
        payload = _form_payload()
        name = payload['name'].strip()
        
        if not name:
            flash('姓名不能为空', 'error')
            return redirect(url_for('experts.add'))
        
        actor = _actor()
        _resources_service().create_expert(
            payload, uploader=actor['name'], actor_user_id=actor['user_id'],
            request_id=actor['request_id'],
        )
        
        flash('专家信息添加成功', 'success')
        return redirect(url_for('experts.index'))
    
    return render_template('experts/add.html')

@bp.route('/edit/<expert_id>', methods=['GET', 'POST'])
def edit(expert_id):
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    service = _resources_service()
    actor = _actor()
    try:
        expert = _template_expert(service.get_expert_sensitive(
            expert_id, actor_user_id=actor['user_id'],
            request_id=actor['request_id'], operation='EDIT',
        ))
    except Exception:
        expert = None
    
    if not expert:
        flash('专家不存在', 'error')
        return redirect(url_for('experts.index'))
    
    if request.method == 'POST':
        payload = _form_payload()
        name = payload['name'].strip()
        
        if not name:
            flash('姓名不能为空', 'error')
            return redirect(url_for('experts.edit', expert_id=expert_id))
        
        actor = _actor()
        service.update_expert(
            expert_id, payload, actor_user_id=actor['user_id'],
            request_id=actor['request_id'],
        )
        
        flash(f'更新成功', 'success')
        return redirect(url_for('experts.index'))
    
    return render_template('experts/edit.html', expert=expert)

@bp.route('/delete/<expert_id>', methods=['POST'])
def delete(expert_id):
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    try:
        _resources_service().delete_expert(expert_id)
    except Exception as error:
        return jsonify({'success': False, 'message': getattr(error, 'message', '删除失败')}), getattr(error, 'status_code', 400)
    
    return jsonify({'success': True, 'message': '删除成功'})

@bp.route('/export')
def export():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    actor = _actor()
    experts = [_template_expert(item) for item in _resources_service().export_experts_sensitive(
        actor_user_id=actor['user_id'], request_id=actor['request_id']
    )]
    
    output = BytesIO()
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = '专家库'
    
    ws.append(['序号', '专家编号', '姓名', '单位', '职务', '专业领域', '银行卡号', '开户行', '上传人', '创建时间', '更新时间'])
    
    for i, e in enumerate(experts, 1):
        ws.append([_xlsx_cell(value) for value in [i, e['expert_id'], e['name'], e['unit'], e['position'], e['expertise'], e['bank_card'] or '', e['bank_name'] or '', e['uploader'], e['created_at'], e['updated_at']]])
    
    workbook.save(output)
    output.seek(0)
    
    return send_file(output, download_name=f'专家库_{datetime.now().strftime("%Y%m%d")}.xlsx', as_attachment=True)

@bp.route('/export/selected', methods=['POST'])
def export_selected():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    expert_ids = request.form.getlist('expert_ids')
    if not expert_ids:
        flash('请选择要导出的专家', 'warning')
        return redirect(url_for('experts.index'))
    
    actor = _actor()
    selected_experts = [_template_expert(item) for item in _resources_service().export_experts_sensitive(
        expert_ids=expert_ids, actor_user_id=actor['user_id'],
        request_id=actor['request_id'],
    )]
    
    output = BytesIO()
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = '专家库'
    
    ws.append(['序号', '专家编号', '姓名', '单位', '职务', '专业领域', '银行卡号', '开户行', '上传人', '创建时间', '更新时间'])
    
    for i, e in enumerate(selected_experts, 1):
        ws.append([_xlsx_cell(value) for value in [i, e['expert_id'], e['name'], e['unit'], e['position'], e['expertise'], e['bank_card'] or '', e['bank_name'] or '', e['uploader'], e['created_at'], e['updated_at']]])
    
    workbook.save(output)
    output.seek(0)
    
    return send_file(output, download_name=f'专家库_选中_{datetime.now().strftime("%Y%m%d")}.xlsx', as_attachment=True)

@bp.route('/export/selected/word', methods=['POST'])
def export_selected_word():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    
    expert_ids = request.form.getlist('expert_ids')
    if not expert_ids:
        flash('请选择要导出的专家', 'warning')
        return redirect(url_for('experts.index'))
    
    actor = _actor()
    selected_experts = [_template_expert(item) for item in _resources_service().export_experts_sensitive(
        expert_ids=expert_ids, actor_user_id=actor['user_id'],
        request_id=actor['request_id'],
    )]
    
    # 尝试生成 Word 文档
    try:
        from docx import Document
        from docx.shared import Pt
        
        doc = Document()
        doc.add_heading('专家库导出', 0)
        
        table = doc.add_table(rows=1, cols=8)
        table.style = 'Table Grid'
        
        header_cells = table.rows[0].cells
        header_cells[0].text = '序号'
        header_cells[1].text = '姓名'
        header_cells[2].text = '单位'
        header_cells[3].text = '职务'
        header_cells[4].text = '专业领域'
        header_cells[5].text = '电话'
        header_cells[6].text = '银行卡号'
        header_cells[7].text = '开户行'
        
        for i, e in enumerate(selected_experts, 1):
            row = table.add_row()
            row.cells[0].text = str(i)
            row.cells[1].text = e.get('name') or ''
            row.cells[2].text = e.get('unit') or ''
            row.cells[3].text = e.get('position') or ''
            row.cells[4].text = e.get('expertise') or ''
            row.cells[5].text = e.get('phone') or ''
            row.cells[6].text = e.get('bank_card') or ''
            row.cells[7].text = e.get('bank_name') or ''
        
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        
        flash('导出Word格式成功', 'success')
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document', as_attachment=True, download_name=f'专家库_{datetime.now().strftime("%Y%m%d")}.docx')
    
    except ImportError:
        # 回退到txt
        content = '专家库导出\n'
        content += '='*80 + '\n'
        content += f"{'序号':<4} {'姓名':<10} {'单位':<20} {'职务':<12} {'专业领域':<15} {'电话':<12} {'银行卡号':<20} {'开户行':<15}\n"
        content += '-'*90 + '\n'

        for i, e in enumerate(selected_experts, 1):
            content += f"{i:<4} {e.get('name',''):<10} {e.get('unit',''):<20} {e.get('position',''):<12} {e.get('expertise',''):<15} {e.get('phone',''):<12} {e.get('bank_card',''):<20} {e.get('bank_name',''):<15}\n"

        output = BytesIO(content.encode('utf-8'))
        output.seek(0)
        flash('导出TXT格式成功（未安装python-docx）', 'success')
        return send_file(output, mimetype='text/plain;charset=utf-8', as_attachment=True, download_name=f'专家库_{datetime.now().strftime("%Y%m%d")}.txt')
    return send_file(output, mimetype='text/plain;charset=utf-8', as_attachment=True, download_name=f'专家库_{datetime.now().strftime("%Y%m%d")}.txt')


# ============= 批量导入 =============

# 导入可用的数据库字段白名单（不含 expert_id/id/created_at/updated_at/uploader 不可选但 uploader 自动填）
IMPORT_FIELDS = ['name', 'unit', 'position', 'expertise', 'phone', 'id_card', 'bank_card', 'bank_name']
EXPERT_IMPORT_MAX_BYTES = 10 * 1024 * 1024
EXPERT_IMPORT_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _xlsx_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


@bp.route('/import', methods=['GET'])
def import_page():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    return render_template('experts/import_step1.html')


@bp.route('/import/preview', methods=['POST'])
def import_preview():
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请上传文件'})
    file = request.files['file']
    if not file.filename:
        return jsonify({'success': False, 'message': '请选择文件'})
    if file.filename.endswith('.xls'):
        return jsonify({'success': False, 'message': '暂不支持 .xls 格式，请另存为 .xlsx 后重试'})
    if not file.filename.endswith('.xlsx'):
        return jsonify({'success': False, 'message': '请上传 xlsx 格式文件'})

    preview_only = request.form.get('preview_only', '') == '1'

    try:
        source_bytes = file.stream.read(EXPERT_IMPORT_MAX_BYTES + 1)
        if len(source_bytes) > EXPERT_IMPORT_MAX_BYTES:
            return jsonify({'success': False, 'message': '导入文件不能超过 10MB'}), 413
        with zipfile.ZipFile(BytesIO(source_bytes)) as archive:
            entries = archive.infolist()
            uncompressed = sum(entry.file_size for entry in entries)
            if uncompressed > EXPERT_IMPORT_MAX_UNCOMPRESSED_BYTES:
                return jsonify({'success': False, 'message': '导入文件解压后过大'}), 413
            if any(
                entry.file_size > 1024 * 1024
                and entry.file_size > max(1, entry.compress_size) * 100
                for entry in entries
            ):
                return jsonify({'success': False, 'message': '导入文件压缩比异常'}), 400
        stream = BytesIO(source_bytes)
        workbook = openpyxl.load_workbook(stream, data_only=True, read_only=True)
        ws = workbook.active
        all_rows = [list(r) for r in ws.iter_rows(values_only=True)]

        if len(all_rows) < 2:
            return jsonify({'success': False, 'message': '文件无数据行（少于 2 行）'})

        headers = [str(h).strip() if h is not None else '' for h in all_rows[0]]
        data_rows = all_rows[1:]

        if len(data_rows) > 1000:
            return jsonify({'success': False, 'message': f'超出 1000 条上限（当前 {len(data_rows)} 条），请拆分后导入'})

        if preview_only:
            preview = [[str(c) if c is not None else '' for c in r] for r in data_rows[:10]]
            return jsonify({
                'success': True,
                'headers': headers,
                'preview_rows': preview,
                'total_rows': len(data_rows),
            })

        # 完整模式：接收列映射 + 解析全部行
        try:
            column_mapping = json.loads(request.form.get('column_mapping_json', '{}'))
        except Exception:
            return jsonify({'success': False, 'message': '列映射参数错误'})

        processed = []
        for row_idx, row in enumerate(data_rows):
            record = {'_row_idx': row_idx}
            for excel_col_idx, db_field in column_mapping.items():
                if not db_field or db_field == '__ignore__':
                    continue
                if db_field not in IMPORT_FIELDS:
                    continue
                try:
                    col_idx = int(excel_col_idx)
                except (ValueError, TypeError):
                    continue
                if 0 <= col_idx < len(row):
                    val = row[col_idx]
                    record[db_field] = str(val).strip() if val is not None else ''
                else:
                    record[db_field] = ''
            processed.append(record)

        valid = [r for r in processed if r.get('name', '').strip()]
        skip = [r for r in processed if not r.get('name', '').strip()]

        statistics = {
            'total': len(processed),
            'valid': len(valid),
            'skip': len(skip),
        }

        batch = _resources_service().create_expert_import_preview(
            source_name=file.filename,
            source_bytes=source_bytes,
            rows=processed,
            owner_user_id=_actor()['user_id'],
        )
        session['expert_import_task_id'] = batch['batchId']
        session['expert_import_filename'] = file.filename
        return jsonify({'success': True, 'redirect': url_for('.import_preview_page')})

    except Exception:
        current_app.logger.warning(
            '专家导入解析失败 request_id=%s',
            getattr(request, 'request_id', 'unknown'), exc_info=True,
        )
        return jsonify({'success': False, 'message': '文件解析失败，请检查 xlsx 内容'}), 400


@bp.route('/import/preview', methods=['GET'])
def import_preview_page():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    task_id = session.get('expert_import_task_id')
    if not task_id:
        return redirect(url_for('.import_page'))
    try:
        batch = _resources_service().get_expert_import_batch(
            task_id, owner_user_id=_actor()['user_id']
        )
    except Exception:
        flash('导入会话已过期，请重新上传', 'warning')
        return redirect(url_for('.import_page'))
    return render_template(
        'experts/import_step2.html', processed=batch['rows'],
        statistics={
            'total': len(batch['rows']), 'valid': batch['validCount'],
            'skip': batch['errorCount'] + batch['duplicateCount'],
        },
        filename=batch['sourceName'],
    )


@bp.route('/import/commit', methods=['POST'])
def import_commit():
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    task_id = session.get('expert_import_task_id')
    if not task_id:
        return jsonify({'success': False, 'message': '会话已过期，请重新导入'})
    try:
        confirmed = json.loads(request.form.get('confirmed_rows', '[]'))
    except Exception:
        return jsonify({'success': False, 'message': '参数错误'}), 400
    actor = _actor()
    try:
        result = _resources_service().commit_expert_import(
            task_id, confirmed_rows=confirmed,
            owner_user_id=actor['user_id'], uploader=actor['name'],
            request_id=actor['request_id'],
        )
    except Exception as error:
        return jsonify({
            'success': False,
            'message': getattr(error, 'message', '导入未完成'),
        }), getattr(error, 'status_code', 400)

    session.pop('expert_import_task_id', None)
    session.pop('expert_import_filename', None)
    session['expert_import_results'] = {
        'success': result['created'],
        'fail': [], 'skip': result['skipCount'],
        'filename': result['sourceName'],
    }
    return jsonify({'success': True, 'redirect': url_for('.import_done_page')})


@bp.route('/import/cancel', methods=['POST'])
def import_cancel():
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'}), 401
    task_id = session.get('expert_import_task_id')
    if not task_id:
        return jsonify({'success': True})
    try:
        _resources_service().cancel_expert_import(
            task_id, owner_user_id=_actor()['user_id']
        )
    except Exception as error:
        return jsonify({
            'success': False,
            'message': getattr(error, 'message', '取消未完成'),
        }), getattr(error, 'status_code', 400)
    session.pop('expert_import_task_id', None)
    session.pop('expert_import_filename', None)
    return jsonify({'success': True, 'redirect': url_for('.import_page')})


@bp.route('/import/done', methods=['GET'])
def import_done_page():
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    results = session.get('expert_import_results')
    if not results:
        return redirect(url_for('.import_page'))
    return render_template('experts/import_done.html', results=results)
