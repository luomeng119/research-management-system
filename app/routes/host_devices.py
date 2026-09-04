# -*- coding: utf-8 -*-
from flask import Blueprint, current_app, render_template, request, redirect, url_for, session, flash, jsonify
from app.services.resources import ResourceServiceError

bp = Blueprint('host_devices', __name__, url_prefix='/equipment/hosts')


def _service():
    service = current_app.extensions.get('equipment_resources_service')
    if service is None:
        raise RuntimeError('equipment resources service is not configured')
    return service


def _check_login():
    if 'user' not in session:
        return False
    return True


@bp.route('/')
def index():
    if not _check_login():
        return redirect(url_for('auth.login'))

    category_filter = request.args.get('category', '')
    form_filter = request.args.get('form', '')
    keyword = request.args.get('keyword', '')
    page = int(request.args.get('page', 1))
    per_page = 20

    result = _service().list_host_devices(
        category=category_filter, form=form_filter, keyword=keyword,
        page=page, page_size=per_page,
    )
    all_hosts, total = result['items'], result['total']

    total_pages = (total + per_page - 1) // per_page if total > 0 else 1

    return render_template('host_devices/index.html',
                         hosts=all_hosts,
                         categories=_service().list_host_categories(),
                         category_filter=category_filter,
                         form_filter=form_filter,
                         keyword=keyword,
                         page=page,
                         total_pages=total_pages,
                         total=total)


@bp.route('/add', methods=['GET', 'POST'])
def add():
    if not _check_login():
        return redirect(url_for('auth.login'))

    service = _service()

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        model = request.form.get('model', '').strip()
        category = request.form.get('category', '').strip()
        form = request.form.get('form', '').strip()

        if not name:
            flash('名称不能为空', 'error')
            return render_template('host_devices/add.html',
                                 categories=service.list_host_categories(),
                                 selected_relations=[])

        # 处理新增类型（需要先处理，因为要用最终类型名称保存宿主设备）
        new_category = request.form.get('new_category', '').strip()
        if new_category:
            try:
                service.create_host_category(new_category)
            except ResourceServiceError as error:
                if error.code != 'RESOURCE_DUPLICATE':
                    raise
            category = new_category
        device_ids = request.form.getlist('device_id')
        quantities = request.form.getlist('quantity')
        relations = [
            {'device_id': device_id, 'quantity': qty or 1}
            for device_id, qty in zip(device_ids, quantities) if device_id
        ]
        try:
            service.create_host_device(
                {'name': name, 'model': model, 'category': category, 'form': form},
                relations=relations,
            )
        except ResourceServiceError as error:
            flash(error.message, 'error')
            return render_template(
                'host_devices/add.html', categories=service.list_host_categories(),
                selected_relations=[],
            ), error.status_code

        flash('宿主设备添加成功', 'success')
        return redirect(url_for('host_devices.index'))

    return render_template('host_devices/add.html',
                         categories=service.list_host_categories(),
                         selected_relations=[])


@bp.route('/edit/<host_id>', methods=['GET', 'POST'])
def edit(host_id):
    if not _check_login():
        return redirect(url_for('auth.login'))

    service = _service()
    host = service.get_host_device(host_id)
    if not host:
        flash('宿主设备不存在', 'error')
        return redirect(url_for('host_devices.index'))

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        model = request.form.get('model', '').strip()
        category = request.form.get('category', '').strip()
        form = request.form.get('form', '').strip()

        if not name:
            flash('名称不能为空', 'error')
            return render_template('host_devices/edit.html',
                                 host=host,
                                 categories=service.list_host_categories(),
                                 selected_relations=service.get_devices_by_host(host_id))

        device_ids = request.form.getlist('device_id')
        quantities = request.form.getlist('quantity')
        relations = [
            {'device_id': device_id, 'quantity': qty or 1}
            for device_id, qty in zip(device_ids, quantities) if device_id
        ]
        new_category = request.form.get('new_category', '').strip()
        if new_category:
            try:
                service.create_host_category(new_category)
            except ResourceServiceError as error:
                if error.code != 'RESOURCE_DUPLICATE':
                    raise
            category = new_category
        try:
            service.update_host_device(
                host_id, {'name': name, 'model': model, 'category': category, 'form': form},
                relations=relations,
            )
        except ResourceServiceError as error:
            flash(error.message, 'error')
            return render_template(
                'host_devices/edit.html', host=host,
                categories=service.list_host_categories(),
                selected_relations=service.get_devices_by_host(host_id),
            ), error.status_code

        flash('宿主设备更新成功', 'success')
        return redirect(url_for('host_devices.detail', host_id=host_id))

    selected_relations = service.get_devices_by_host(host_id)
    return render_template('host_devices/edit.html',
                         host=host,
                         categories=service.list_host_categories(),
                         selected_relations=selected_relations)


@bp.route('/detail/<host_id>')
def detail(host_id):
    if not _check_login():
        return redirect(url_for('auth.login'))

    service = _service()
    host = service.get_host_device(host_id)
    if not host:
        flash('宿主设备不存在', 'error')
        return redirect(url_for('host_devices.index'))

    related_devices = service.get_devices_by_host(host_id)
    return render_template('host_devices/detail.html',
                         host=host,
                         related_devices=related_devices)


@bp.route('/delete/<host_id>', methods=['POST'])
def delete(host_id):
    if not _check_login():
        return redirect(url_for('auth.login'))
    try:
        _service().delete_host_device(host_id)
    except ResourceServiceError as error:
        flash(error.message, 'error')
    else:
        flash('宿主设备已删除', 'success')
    return redirect(url_for('host_devices.index'))


@bp.route('/categories')
def categories():
    if not _check_login():
        return redirect(url_for('auth.login'))
    all_categories = _service().list_host_categories()

    return render_template('host_devices/categories.html',
                         categories=all_categories)


@bp.route('/categories/add', methods=['POST'])
def categories_add():
    if not _check_login():
        return jsonify({'success': False, 'message': '未登录'})
    name = request.form.get('name', '').strip()
    if not name:
        return jsonify({'success': False, 'message': '类型名称不能为空'})
    try:
        _service().create_host_category(name)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/categories/merge', methods=['POST'])
def categories_merge():
    if not _check_login():
        return jsonify({'success': False, 'message': '未登录'})
    old_name = request.form.get('old_name', '').strip()
    new_name = request.form.get('new_name', '').strip()
    if not old_name or not new_name:
        return jsonify({'success': False, 'message': '参数不完整'})
    if old_name == new_name:
        return jsonify({'success': False, 'message': '不能合并到自身'})
    try:
        _service().merge_host_category(old_name, new_name)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/categories/delete', methods=['POST'])
def categories_delete():
    if not _check_login():
        return jsonify({'success': False, 'message': '未登录'})
    name = request.form.get('name', '').strip()
    try:
        _service().delete_host_category(name)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


# ---------- API ----------

@bp.route('/api/relations/<host_id>')
def api_relations(host_id):
    """获取宿主设备关联的密码设备列表（含数量）"""
    if not _check_login():
        return jsonify({'success': False})
    devices = _service().get_devices_by_host(host_id)
    return jsonify({'success': True, 'data': devices})


@bp.route('/api/relations', methods=['POST'])
def api_relations_add():
    """批量添加关联"""
    if not _check_login():
        return jsonify({'success': False})
    data = request.get_json() or {}
    host_id = data.get('host_id')
    relations = data.get('relations', [])  # [{device_id, quantity}, ...]
    if not host_id:
        return jsonify({'success': False, 'message': 'host_id required'})
    try:
        _service().upsert_host_relations(host_id, relations)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/api/relations/<device_id>/<host_id>', methods=['PATCH'])
def api_relation_patch(device_id, host_id):
    """修改关联数量"""
    if not _check_login():
        return jsonify({'success': False})
    data = request.get_json() or {}
    quantity = data.get('quantity', 1)
    try:
        _service().update_host_relation(host_id, device_id, quantity)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/api/relations/<device_id>/<host_id>', methods=['DELETE'])
def api_relation_delete(device_id, host_id):
    """移除关联"""
    if not _check_login():
        return jsonify({'success': False})
    try:
        _service().remove_host_relation(host_id, device_id)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/api/search')
def api_search():
    """搜索宿主设备（用于关联选择弹框中过滤已选）"""
    if not _check_login():
        return jsonify({'success': False})
    keyword = request.args.get('keyword', '')
    page = int(request.args.get('page', 1))
    per_page = 20
    result = _service().list_host_devices(keyword=keyword, page=page, page_size=per_page)
    hosts, total = result['items'], result['total']
    return jsonify({'success': True, 'data': hosts, 'total': total, 'page': page})


# ---------- 导入导出 ----------

@bp.route('/export')
def export():
    """导出全部宿主设备为 Excel"""
    if not _check_login():
        return redirect(url_for('auth.login'))

    from io import BytesIO
    import openpyxl

    all_hosts = _service().export_host_devices()

    output = BytesIO()
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = '宿主设备'

    # 表头（增强：增加关联设备ID列）
    ws.append(['宿主设备编号', '名称', '型号', '类型', '形态', '关联设备数',
               '关联设备ID', '关联设备名称', '创建时间', '更新时间'])

    for h in all_hosts:
        devices = h['devices']
        device_count = len(devices)
        device_ids = '、'.join([d['device_id'] for d in devices]) if devices else ''
        device_names = '、'.join([d['name'] for d in devices]) if devices else ''
        ws.append([
            h['host_id'],
            h['name'],
            h['model'] or '',
            h['category'] or '',
            h['form'] or '',
            device_count,
            device_ids,
            device_names,
            h['created_at'] or '',
            h['updated_at'] or '',
        ])

    # 自动列宽
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    workbook.save(output)
    output.seek(0)
    from flask import send_file
    return send_file(output, download_name='宿主设备库.xlsx',
                     as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/import', methods=['GET'])
def import_page():
    """步骤1：导入上传页"""
    if not _check_login():
        return redirect(url_for('auth.login'))
    return render_template('host_devices/import_step1.html')


@bp.route('/import/preview', methods=['POST'])
def import_preview():
    """
    步骤1提交：解析文件 + 列映射 + 执行模糊匹配
    返回 JSON，前端收到后跳转到预览页

    支持两种模式：
    - preview_only=1: 只返回表头+前几行数据（用于列映射页面预览）
    - 正常模式: 完整解析+模糊匹配，存入 session，重定向到预览页
    """
    if not _check_login():
        return jsonify({'success': False, 'message': '未登录'})

    from io import BytesIO
    import openpyxl
    import json

    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请上传文件'})

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': '请选择文件'})

    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'message': '请上传 xlsx 文件'})

    preview_only = request.form.get('preview_only', '') == '1'

    try:
        stream = BytesIO(file.read())
        workbook = openpyxl.load_workbook(stream, data_only=True)
        ws = workbook.active

        # 读取所有行
        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))

        if len(all_rows) < 2:
            return jsonify({'success': False, 'message': '文件数据少于2行'})

        header = [str(h).strip() if h is not None else '' for h in all_rows[0]]
        data_rows = all_rows[1:]
        preview_rows = data_rows[:10]

        # preview_only 模式：只返回预览数据
        if preview_only:
            return jsonify({
                'success': True,
                'headers': header,
                'preview_rows': [[str(c) if c is not None else '' for c in row] for row in preview_rows],
            })

        # 正常模式：完整解析 + 模糊匹配
        try:
            column_mapping = json.loads(request.form.get('column_mapping_json', '{}'))
        except Exception:
            return jsonify({'success': False, 'message': '列映射参数错误'})

        service = _service()
        existing_categories = {item['name'] for item in service.list_host_categories()}

        related_col_idx = column_mapping.get('related_equipment')

        processed = []
        exact_count = 0
        fuzzy_count = 0
        unmatched_count = 0
        duplicate_count = 0

        for row_idx, row in enumerate(data_rows):
            if not row or not any(v for v in row if v is not None):
                continue

            name = _get_row_value(row, column_mapping, 'name')
            model = _get_row_value(row, column_mapping, 'model')
            category = _get_row_value(row, column_mapping, 'category')
            form = _get_row_value(row, column_mapping, 'form')
            host_id_raw = _get_row_value(row, column_mapping, 'host_id')

            if not name:
                continue

            # 重复检测
            duplicate_status = None
            existing_host = None
            if host_id_raw and str(host_id_raw).startswith('HD'):
                existing_host = service.get_host_device(str(host_id_raw))
                if existing_host:
                    duplicate_status = 'exists'

            # 关联设备匹配
            related_result = {'original': '', 'devices': []}
            if related_col_idx is not None and related_col_idx < len(row):
                related_text = str(row[related_col_idx] or '').strip()
                related_result = service.match_equipment_text(related_text)

                for dev in related_result['devices']:
                    if dev['status'] == 'exact':
                        exact_count += 1
                    elif dev['status'] == 'fuzzy':
                        fuzzy_count += 1
                    else:
                        unmatched_count += 1

            # 新增类型检测
            is_new_category = False
            if category:
                if category not in existing_categories:
                    is_new_category = True

            processed.append({
                'row_idx': row_idx,
                'host_id_raw': host_id_raw,
                'name': name,
                'model': model,
                'category': category,
                'form': form,
                'is_new_category': is_new_category,
                'duplicate_status': duplicate_status,
                'existing_host': dict(existing_host) if existing_host else None,
                'related_result': related_result,
            })

            if duplicate_status == 'exists':
                duplicate_count += 1

        statistics = {
            'total': len(processed),
            'exact': exact_count,
            'fuzzy': fuzzy_count,
            'unmatched': unmatched_count,
            'duplicate': duplicate_count,
        }

        # 存入 session
        session['host_import_preview'] = {
            'filename': file.filename,
            'column_mapping': column_mapping,
            'processed': processed,
            'statistics': statistics,
        }

        return jsonify({'success': True, 'redirect': url_for('host_devices.import_preview_page')})

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'message': f'解析失败: {str(e)}'})


def _get_row_value(row, column_mapping, field):
    """根据列映射获取某行某字段的值"""
    col_idx = column_mapping.get(field)
    if col_idx is None or col_idx >= len(row):
        return ''
    val = row[col_idx]
    return str(val).strip() if val is not None else ''


@bp.route('/import/preview', methods=['GET'])
def import_preview_page():
    """步骤2：匹配预览页面"""
    if not _check_login():
        return redirect(url_for('auth.login'))

    preview_data = session.get('host_import_preview')
    if not preview_data:
        flash('请先上传文件', 'error')
        return redirect(url_for('host_devices.import_page'))

    return render_template('host_devices/import_step2.html',
                           filename=preview_data['filename'],
                           processed=preview_data['processed'],
                           statistics=preview_data['statistics'])


@bp.route('/import/commit', methods=['POST'])
def import_commit():
    """步骤2确认提交：执行导入"""
    if not _check_login():
        return redirect(url_for('auth.login'))

    preview_data = session.get('host_import_preview')
    if not preview_data:
        return jsonify({'success': False, 'message': '会话过期，请重新上传'})

    import json

    import_rows = []
    skipped_relations = []

    # 遍历每行，按用户修正执行写入
    for row_data in preview_data['processed']:
        row_idx = row_data['row_idx']

        # 读取用户修正后的值
        name = request.form.get(f'name_{row_idx}', row_data['name']).strip()
        model = request.form.get(f'model_{row_idx}', row_data['model']).strip()
        category = request.form.get(f'category_{row_idx}', row_data['category']).strip()
        form = request.form.get(f'form_{row_idx}', row_data['form']).strip()
        # 重复处理选择
        dup_action = request.form.get(f'dup_action_{row_idx}', '')

        if not name:
            import_rows.append({'name': '', 'action': 'skip'})
            continue

        device_ids = []
        related_result = row_data['related_result']
        for dev_idx, dev_match in enumerate(related_result.get('devices', [])):
            # 用户选择的设备ID
            selected_id = request.form.get(f'related_select_{row_idx}_{dev_idx}', '').strip()
            if selected_id == '_skip_':
                # 用户跳过该关联
                skipped_relations.append({
                    'name': row_data['name'],
                    'device': dev_match.get('name', ''),
                })
                continue
            if not selected_id:
                # 未选择，尝试用自动匹配的
                if dev_match.get('selected'):
                    selected_id = dev_match['selected'].get('equipment_id')
                elif dev_match.get('exact'):
                    selected_id = dev_match['exact'][0].get('equipment_id')
                elif dev_match.get('fuzzy'):
                    selected_id = dev_match['fuzzy'][0].get('equipment_id')

            if selected_id:
                device_ids.append(selected_id)

        import_rows.append({
            'name': name, 'model': model, 'category': category, 'form': form,
            'existing_host_id': (
                row_data.get('existing_host', {}).get('host_id')
                if row_data.get('existing_host') else None
            ),
            'action': dup_action or 'new', 'device_ids': device_ids,
        })

    try:
        result = _service().import_host_devices(import_rows)
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code
    result['skipped_relations'] = skipped_relations

    # 清除 session
    session.pop('host_import_preview', None)

    # 跳转到结果页
    return redirect(url_for('host_devices.import_done',
                            new=result['new'],
                            overwrite=result['overwrite'],
                            skip=result['skip'],
                            relations=result['relations'],
                            skipped_relations=json.dumps(result['skipped_relations'])))


@bp.route('/import/done')
def import_done():
    """步骤3：导入结果页"""
    if not _check_login():
        return redirect(url_for('auth.login'))

    import json
    new_count = int(request.args.get('new', 0))
    overwrite_count = int(request.args.get('overwrite', 0))
    skip_count = int(request.args.get('skip', 0))
    relations_count = int(request.args.get('relations', 0))
    try:
        skipped_relations = json.loads(request.args.get('skipped_relations', '[]'))
    except Exception:
        skipped_relations = []

    return render_template('host_devices/import_done.html',
                           new_count=new_count,
                           overwrite_count=overwrite_count,
                           skip_count=skip_count,
                           relations_count=relations_count,
                           skipped_relations=skipped_relations)
