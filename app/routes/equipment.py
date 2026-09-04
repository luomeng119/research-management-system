# -*- coding: utf-8 -*-
from flask import Blueprint, render_template, request, redirect, url_for, session, flash, jsonify, send_file, current_app, Response, stream_with_context
import os
import re
import csv
import io
import json
import threading
import time

from app.utils import import_progress as ip
from app.security.auth import BUSINESS_USER, current_identity
from app.services.files import FileServiceError
from app.services.resources import ResourceServiceError

bp = Blueprint('equipment', __name__, url_prefix='/equipment')


def _equipment_resources_service():
    service = current_app.extensions.get('equipment_resources_service')
    if service is None:
        raise RuntimeError('equipment resources service is not configured')
    return service


def _file_service():
    service = current_app.extensions.get('file_service')
    if service is None:
        raise FileServiceError('FILE_SERVICE_UNAVAILABLE', '附件服务不可用', 503)
    return service


def _business_identity():
    identity = current_identity()
    if identity is None:
        raise FileServiceError('UNAUTHORIZED', '未登录', 401)
    if identity.role != BUSINESS_USER:
        raise FileServiceError('FORBIDDEN', '无权访问业务附件', 403)
    return identity


def _can_access_equipment_files():
    identity = current_identity()
    return bool(identity and identity.role == BUSINESS_USER)


def _equipment_files(equipment_id):
    if not _can_access_equipment_files():
        return []
    return _file_service().list_for_object(
        object_type='EQUIPMENT', object_id=equipment_id
    )


def _record_equipment_file_failure(operation, error, original_name=None):
    identity = current_identity()
    service = current_app.extensions.get('file_service')
    if identity is None or service is None:
        return
    try:
        service.record_event(
            operation=operation, actor_user_id=identity.user_id,
            request_id=getattr(request, 'request_id', 'equipment-file-upload'),
            file_id=None, result='FAILURE', error_code=error.code,
            file_type=os.path.splitext(original_name or '')[1].lstrip('.').lower() or None,
        )
    except FileServiceError:
        current_app.logger.warning('equipment file failure audit unavailable')


def _upload_equipment_images(equipment_id, image_files):
    identity = _business_identity()
    uploaded = 0
    for image_file in image_files:
        if not image_file or not image_file.filename:
            continue
        extension = os.path.splitext(image_file.filename)[1].lower()
        if extension not in {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp'}:
            raise FileServiceError('INVALID_FILE', '设备图片格式不支持')
        _file_service().upload(
            image_file.stream, original_name=image_file.filename,
            object_type='EQUIPMENT', object_id=equipment_id,
            actor_user_id=identity.user_id,
            request_id=getattr(request, 'request_id', 'equipment-image-upload'),
        )
        uploaded += 1
    return uploaded


def _safe_filename(filename):
    """净化文件名,防止路径穿越"""
    filename = os.path.basename(filename)  # 去掉路径,只保留文件名
    filename = re.sub(r'[^\w\s.-]', '_', filename)  # 只保留字母数字下划线短横线点和空格
    filename = filename.strip('. ')  # 去掉首尾点和空格
    return filename or 'unnamed_file'


def get_subclasses_json():
    """获取所有子类，按分类分组，用于前端联动下拉"""
    all_subclasses = _equipment_resources_service().list_subclasses()
    result = {}
    for sc in all_subclasses:
        cat = sc.get('parent_category', '')
        if cat not in result:
            result[cat] = []
        result[cat].append({'subclass_name': sc.get('subclass_name', '')})
    return result


@bp.route('/')
def index():
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    equipment_service = _equipment_resources_service()
    category_filter = request.args.get('category', '')
    form_filter = request.args.get('form', '')
    tech_status_filter = request.args.get('tech_status', '')
    keyword = request.args.get('keyword', '')
    subclass_filter = request.args.get('subclass', '')
    page = int(request.args.get('page', 1))
    per_page = 20

    result = equipment_service.list_equipment(
        page=page, page_size=per_page, category=category_filter,
        form=form_filter, tech_status=tech_status_filter,
        keyword=keyword, subclass=subclass_filter,
    )
    equipment = result['items']
    total = result['total']
    total_pages = (total + per_page - 1) // per_page if total > 0 else 1

    # 统计
    stats = {
        'total': len(equipment),
        'security': len([e for e in equipment if e.get('category') == '安全设备']),
        'crypto': len([e for e in equipment if e.get('category') == '密码设备']),
        'general': len([e for e in equipment if e.get('category') == '通用设备']),
        'total_value': sum([float(e.get('price') or 0) for e in equipment])
    }

    category_groups = {'安全设备': [], '密码设备': [], '通用设备': []}
    for e in equipment:
        cat = e.get('category') or '通用设备'
        if cat in category_groups:
            category_groups[cat].append(e)

    all_subclasses = equipment_service.list_subclasses()
    subclasses_json = {}
    for sc in all_subclasses:
        cat = sc.get('parent_category', '')
        if cat not in subclasses_json:
            subclasses_json[cat] = []
        subclasses_json[cat].append({'subclass_name': sc.get('subclass_name', '')})

    # REQ-015: ?fragment=1 时返回纯表格片段（AJAX 用）
    if request.args.get('fragment') == '1':
        return render_template('equipment/_table_fragment.html',
                             equipment=equipment,
                             category_filter=category_filter,
                             form_filter=form_filter,
                             tech_status_filter=tech_status_filter,
                             keyword=keyword,
                             page=page,
                             total_pages=total_pages,
                             total=total,
                             per_page=per_page,
                             subclass_filter=subclass_filter)

    return render_template('equipment/index.html',
                         equipment=equipment,
                         category_groups=category_groups,
                         category_filter=category_filter,
                         form_filter=form_filter,
                         tech_status_filter=tech_status_filter,
                         keyword=keyword,
                         stats=stats,
                         page=page,
                         total_pages=total_pages,
                         total=total,
                         per_page=per_page,
                         subclass_filter=subclass_filter,
                         subclasses_json=subclasses_json)

@bp.route('/stats')
def stats():
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    result = _equipment_resources_service().equipment_stats()

    return render_template('equipment/stats.html',
                         by_category=result['by_category'],
                         by_status=result['by_status'],
                         total_value=result['total_value'],
                         total=result['total'])

@bp.route('/add', methods=['GET', 'POST'])
def add():
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        model = request.form.get('model', '').strip()
        category = request.form.get('category', '').strip()
        form = request.form.get('form', '').strip()
        price = request.form.get('price', '').strip()
        tech_index = request.form.get('tech_index', '').strip()
        installation_requirements = request.form.get('installation_requirements', '').strip()
        tech_status = request.form.get('tech_status', '').strip()
        manufacturer = request.form.get('manufacturer', '').strip()
        main_purpose = request.form.get('main_purpose', '').strip()
        former_name = request.form.get('former_name', '').strip()
        resource_guarantee = request.form.get('resource_guarantee', '').strip()

        subclass = request.form.get('subclass', '').strip()
        image_files = request.files.getlist('equipment_images')
        if image_files and any(f.filename for f in image_files):
            try:
                _business_identity()
            except FileServiceError as error:
                return render_template(
                    'equipment/add.html', subclasses_json=get_subclasses_json(),
                    can_access_files=False, page_error=error.message,
                ), error.status_code
        try:
            identity = current_identity()
            created = _equipment_resources_service().create_equipment({
                'name': name, 'model': model, 'category': category, 'form': form,
                'price': price, 'techIndex': tech_index, 'techStatus': tech_status,
                'installationRequirements': installation_requirements,
                'manufacturer': manufacturer, 'mainPurpose': main_purpose,
                'formerName': former_name, 'resourceGuarantee': resource_guarantee,
                'subclass': subclass,
            }, actor_user_id=identity.user_id if identity else None,
               request_id=getattr(request, 'request_id', 'equipment-create'))
        except ResourceServiceError as error:
            flash(error.message, 'error')
            return render_template('equipment/add.html', subclasses_json=get_subclasses_json()), error.status_code
        equipment_id = created['equipment_id']

        if image_files and any(f.filename for f in image_files):
            try:
                _upload_equipment_images(equipment_id, image_files)
            except FileServiceError as error:
                first_name = next((f.filename for f in image_files if f.filename), None)
                _record_equipment_file_failure('UPLOAD', error, first_name)
                flash(f'设备已创建，但图片上传失败：{error.message}', 'warning')

        flash('设备添加成功', 'success')
        return redirect(url_for('equipment.index'))

    return render_template(
        'equipment/add.html', subclasses_json=get_subclasses_json(),
        can_access_files=_can_access_equipment_files(),
    )

@bp.route('/detail/<equipment_id>')
def detail(equipment_id):
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    equipment = _equipment_resources_service().get_equipment(equipment_id)

    if not equipment:
        flash('设备不存在', 'error')
        return redirect(url_for('equipment.index'))

    try:
        related_hosts = _equipment_resources_service().get_hosts_by_device(equipment_id)
    except ResourceServiceError:
        related_hosts = []

    file_error = None
    try:
        files = _equipment_files(equipment_id)
    except FileServiceError as error:
        _record_equipment_file_failure('LIST', error)
        files = []
        file_error = error.message
    return render_template(
        'equipment/detail.html', equipment=equipment, related_hosts=related_hosts,
        image_files=[item for item in files if item.get('mediaType', '').startswith('image/')],
        related_files=[item for item in files if not item.get('mediaType', '').startswith('image/')],
        can_access_files=_can_access_equipment_files(),
        file_error=file_error,
    )

@bp.route('/edit/<equipment_id>', methods=['GET', 'POST'])
def edit(equipment_id):
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    equipment_service = _equipment_resources_service()
    equipment = equipment_service.get_equipment(equipment_id)

    if not equipment:
        flash('设备不存在', 'error')
        return redirect(url_for('equipment.index'))

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        model = request.form.get('model', '').strip()
        category = request.form.get('category', '').strip()
        form = request.form.get('form', '').strip()
        price = request.form.get('price', '').strip()
        tech_index = request.form.get('tech_index', '').strip()
        installation_requirements = request.form.get('installation_requirements', '').strip()
        tech_status = request.form.get('tech_status', '').strip()
        manufacturer = request.form.get('manufacturer', '').strip()
        main_purpose = request.form.get('main_purpose', '').strip()
        former_name = request.form.get('former_name', '').strip()
        resource_guarantee = request.form.get('resource_guarantee', '').strip()
        subclass = request.form.get('subclass', '').strip()
        image_files = request.files.getlist('equipment_images')
        if image_files and any(f.filename for f in image_files):
            try:
                _business_identity()
            except FileServiceError as error:
                return render_template(
                    'equipment/edit.html', equipment=equipment,
                    subclasses_json=get_subclasses_json(), image_files=[],
                    can_access_files=False, page_error=error.message,
                ), error.status_code

        try:
            identity = current_identity()
            equipment_service.update_equipment(equipment_id, {
                'name': name, 'model': model, 'category': category, 'form': form,
                'price': price, 'techIndex': tech_index,
                'installationRequirements': installation_requirements,
                'techStatus': tech_status, 'manufacturer': manufacturer,
                'mainPurpose': main_purpose, 'formerName': former_name,
                'resourceGuarantee': resource_guarantee, 'subclass': subclass,
            }, actor_user_id=identity.user_id if identity else None,
               request_id=getattr(request, 'request_id', 'equipment-update'))
        except ResourceServiceError as error:
            flash(error.message, 'error')
            return render_template(
                'equipment/edit.html', equipment=equipment,
                subclasses_json=get_subclasses_json(),
            ), error.status_code

        if image_files and any(f.filename for f in image_files):
            try:
                _upload_equipment_images(equipment_id, image_files)
            except FileServiceError as error:
                first_name = next((f.filename for f in image_files if f.filename), None)
                _record_equipment_file_failure('UPLOAD', error, first_name)
                flash(f'设备信息已保存，但图片上传失败：{error.message}', 'warning')

        flash('设备更新成功', 'success')
        return redirect(url_for('equipment.detail', equipment_id=equipment_id))

    file_error = None
    try:
        files = _equipment_files(equipment_id)
    except FileServiceError as error:
        _record_equipment_file_failure('LIST', error)
        files = []
        file_error = error.message
    return render_template(
        'equipment/edit.html', equipment=equipment,
        subclasses_json=get_subclasses_json(),
        image_files=[item for item in files if item.get('mediaType', '').startswith('image/')],
        can_access_files=_can_access_equipment_files(),
        file_error=file_error,
    )

@bp.route('/delete/<equipment_id>', methods=['POST'])
def delete(equipment_id):
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    equipment_service = _equipment_resources_service()
    try:
        identity = current_identity()
        equipment_service.delete_equipment(
            equipment_id,
            actor_user_id=identity.user_id if identity else None,
            request_id=getattr(request, 'request_id', 'equipment-delete'),
        )
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code

    return jsonify({'success': True, 'message': '删除成功'})


@bp.route('/upload_file/<equipment_id>', methods=['POST'])
def upload_file(equipment_id):
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    try:
        identity = _business_identity()
    except FileServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code

    if 'related_file' not in request.files:
        error = FileServiceError('FILE_REQUIRED', '请选择文件')
        _record_equipment_file_failure('UPLOAD', error)
        return jsonify({'success': False, 'message': error.message}), error.status_code

    file = request.files['related_file']
    if file.filename == '':
        error = FileServiceError('FILE_REQUIRED', '请选择文件')
        _record_equipment_file_failure('UPLOAD', error)
        return jsonify({'success': False, 'message': error.message}), error.status_code

    try:
        result = _file_service().upload(
            file.stream, original_name=file.filename,
            object_type='EQUIPMENT', object_id=equipment_id,
            actor_user_id=identity.user_id,
            request_id=getattr(request, 'request_id', 'equipment-file-upload'),
        )
        return jsonify({
            'success': True, 'message': '上传成功',
            'file': {key: value for key, value in result.items() if key != 'storagePath'},
        })
    except FileServiceError as error:
        if error.code not in {'UNAUTHORIZED', 'FORBIDDEN'}:
            _record_equipment_file_failure('UPLOAD', error, file.filename)
        return jsonify({'success': False, 'message': error.message}), error.status_code

@bp.route('/export')
def export():
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    from io import BytesIO
    import openpyxl

    all_equipment = _equipment_resources_service().export_equipment()

    output = BytesIO()
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = '设备知识库'

    # 表头（增强：关联宿主设备ID/名称列）
    ws.append([
        '设备编号', '设备名称', '曾用名', '型号', '分类', '设备子类', '形态',
        '单价', '主要用途', '资源保障要求', '功能技术指标',
        '技术状态', '研制单位', '创建时间',
        '关联宿主设备ID', '关联宿主设备名称'
    ])

    for e in all_equipment:
        host_rels = e['hosts']
        host_ids = '、'.join([h['host_id'] for h in host_rels]) if host_rels else ''
        host_names = '、'.join([h['name'] for h in host_rels]) if host_rels else ''
        ws.append([
            e.get('equipment_id'), e.get('name'), e.get('former_name'), e.get('model'),
            e.get('category'), e.get('subclass'), e.get('form'), e.get('price'), e.get('main_purpose'),
            e.get('resource_guarantee'), e.get('tech_index'), e.get('tech_status'),
            e.get('manufacturer'), e.get('created_at'),
            host_ids, host_names
        ])

    # 自动列宽
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    workbook.save(output)
    output.seek(0)
    return send_file(output,
                     download_name='设备知识库.xlsx',
                     as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------- 三步导入流程 ----------


@bp.route('/import', methods=['GET'])
def import_page():
    """步骤1：导入上传页"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    return render_template('equipment/import_step1.html')


@bp.route('/import/preview', methods=['POST'])
def import_preview():
    """
    步骤1提交：解析文件 + 列映射 + 执行模糊匹配
    支持 preview_only=1 模式：只返回预览数据
    """
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})

    from io import BytesIO
    import openpyxl
    import json

    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请上传文件'})

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': '请选择文件'})

    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'message': '请上传 xlsx 文件'})

    preview_only = request.form.get('preview_only', '') == '1'

    try:
        source_bytes = file.read()
        stream = BytesIO(source_bytes)
        workbook = openpyxl.load_workbook(stream, data_only=True, read_only=True)
        ws = workbook.active

        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))
            if len(all_rows) >= (11 if preview_only else 1002):
                break

        if len(all_rows) < 2:
            return jsonify({'success': False, 'message': '文件数据少于2行'})

        header = [str(h).strip() if h is not None else '' for h in all_rows[0]]
        data_rows = all_rows[1:]
        preview_rows = data_rows[:10]

        if preview_only:
            return jsonify({
                'success': True,
                'headers': header,
                'preview_rows': [[str(c) if c is not None else '' for c in row] for row in preview_rows],
            })
        if len(data_rows) > 1000:
            return jsonify({'success': False, 'message': '单次导入不能超过 1000 条'}), 422

        # 完整解析模式
        try:
            column_mapping = json.loads(request.form.get('column_mapping_json', '{}'))
        except Exception:
            return jsonify({'success': False, 'message': '列映射参数错误'})

        # [REQ-012] 读取用户选定的批量分类（必填）
        default_category = request.form.get('default_category', '').strip()
        valid_categories = ('安全设备', '密码设备', '通用设备')
        if not default_category:
            return jsonify({'success': False, 'message': '请先选择批量分类'}), 400
        if default_category not in valid_categories:
            return jsonify({'success': False, 'message': f'无效的分类：{default_category}（仅支持 安全设备/密码设备/通用设备）'}), 400

        equipment_service = _equipment_resources_service()
        related_col_idx = column_mapping.get('related_host_devices')

        processed = []
        exact_count = 0
        fuzzy_count = 0
        unmatched_count = 0
        duplicate_count = 0
        new_category_count = 0

        # 固定分类（不允许动态新增）
        valid_categories = {'安全设备', '密码设备', '通用设备', '其他设备'}

        for row_idx, row in enumerate(data_rows):
            if not row or not any(v for v in row if v is not None):
                continue

            name = _eq_get_row_value(row, column_mapping, 'name')
            model = _eq_get_row_value(row, column_mapping, 'model')
            # [REQ-012] category 不再从 xlsx 读取，统一使用用户在步骤1选定的批量分类
            category = default_category
            form = _eq_get_row_value(row, column_mapping, 'form')
            price_str = _eq_get_row_value(row, column_mapping, 'price')
            tech_index = _eq_get_row_value(row, column_mapping, 'tech_index')
            tech_status = _eq_get_row_value(row, column_mapping, 'tech_status')
            manufacturer = _eq_get_row_value(row, column_mapping, 'manufacturer')
            main_purpose = _eq_get_row_value(row, column_mapping, 'main_purpose')
            former_name = _eq_get_row_value(row, column_mapping, 'former_name')
            resource_guarantee = _eq_get_row_value(row, column_mapping, 'resource_guarantee')
            installation_requirements = _eq_get_row_value(row, column_mapping, 'installation_requirements')
            equipment_id_raw = _eq_get_row_value(row, column_mapping, 'equipment_id')
            subclass = _eq_get_row_value(row, column_mapping, 'subclass')

            if not name:
                continue

            # 重复检测：按 name + model 组合
            duplicate_status = None
            existing_eq = None
            if equipment_id_raw:
                existing_eq = equipment_service.find_import_duplicate(
                    equipment_id=equipment_id_raw
                )
            else:
                existing_eq = equipment_service.find_import_duplicate(
                    name=name, model=model
                )

            if existing_eq:
                duplicate_status = 'exists'

            # 类型归一
            is_new_category = False
            if category and category not in valid_categories:
                is_new_category = True
                new_category_count += 1

            # 关联宿主设备匹配
            related_result = {'original': '', 'devices': []}
            if related_col_idx is not None and related_col_idx < len(row):
                related_text = str(row[related_col_idx] or '').strip()
                related_result = equipment_service.match_host_device_text(related_text)

                for dev in related_result['devices']:
                    if dev['status'] == 'exact':
                        exact_count += 1
                    elif dev['status'] == 'fuzzy':
                        fuzzy_count += 1
                    else:
                        unmatched_count += 1

            # 研制单位模糊匹配
            manufacturer_result = {'original': manufacturer, 'units': []}
            if manufacturer:
                # 按 、 ， ； 分隔
                parts = re.split(r'[、，；]', manufacturer)
                for part in parts:
                    part = part.strip()
                    if not part:
                        continue
                    mru = equipment_service.match_research_unit_name(part)
                    unit_entry = {
                        'original': part,
                        'matched': mru['matched'],
                        'exact': mru['exact'],
                        'method': mru['method'],
                        'candidates': mru.get('candidates', []),
                    }
                    manufacturer_result['units'].append(unit_entry)
                    if mru['exact']:
                        exact_count += 1
                    elif mru['matched']:
                        fuzzy_count += 1
                    else:
                        unmatched_count += 1

            # 设备子类模糊匹配
            subclass_result = {'original': subclass, 'matched': None, 'exact': False, 'method': 'none'}
            if subclass:
                ms = equipment_service.match_subclass_name(subclass, category)
                subclass_result = {
                    'original': subclass,
                    'matched': ms['matched'],
                    'exact': ms['exact'],
                    'method': ms['method'],
                    'candidates': ms.get('candidates', []),
                }
                if ms['exact']:
                    exact_count += 1
                elif ms['matched']:
                    fuzzy_count += 1

            # 价格字段校验
            price_error = ''
            if price_str:
                try:
                    float(price_str)
                except Exception:
                    price_error = '价格格式错误'

            processed.append({
                'row_idx': row_idx,
                'equipment_id_raw': equipment_id_raw,
                'name': name,
                'model': model,
                'category': category,
                'form': form,
                'price': price_str,
                'tech_index': tech_index,
                'tech_status': tech_status,
                'manufacturer': manufacturer,
                'main_purpose': main_purpose,
                'former_name': former_name,
                'resource_guarantee': resource_guarantee,
                'installation_requirements': installation_requirements,
                'is_new_category': is_new_category,
                'price_error': price_error,
                'duplicate_status': duplicate_status,
                'existing_equipment': ({
                    'equipment_id': str(existing_eq.get('equipment_id') or ''),
                    'name': str(existing_eq.get('name') or ''),
                    'model': str(existing_eq.get('model') or ''),
                } if existing_eq else None),
                'related_result': related_result,
                'manufacturer_result': manufacturer_result,
                'subclass_result': subclass_result,
            })

            if duplicate_status == 'exists':
                duplicate_count += 1

        statistics = {
            'total': len(processed),
            'exact': exact_count,
            'fuzzy': fuzzy_count,
            'unmatched': unmatched_count,
            'duplicate': duplicate_count,
            'new_category': new_category_count,
        }

        identity = current_identity()
        if identity is None:
            return jsonify({'success': False, 'message': '未登录'}), 401
        batch = equipment_service.create_equipment_import_preview(
            source_name=file.filename, source_bytes=source_bytes,
            rows=processed, statistics=statistics,
            owner_user_id=identity.user_id,
        )
        session['equipment_import_batch_id'] = batch['batchId']
        session.pop('equipment_import_preview', None)

        return jsonify({'success': True, 'redirect': url_for('equipment.import_preview_page')})

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'message': f'解析失败: {str(e)}'})


def _eq_get_row_value(row, column_mapping, field):
    col_idx = column_mapping.get(field)
    if col_idx is None or col_idx >= len(row):
        return ''
    val = row[col_idx]
    return str(val).strip() if val is not None else ''


@bp.route('/import/save_row', methods=['POST'])
def import_save_row():
    """AJAX: 将单行修改写回当前用户的持久化导入批次。"""
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'}), 401
    batch_id = session.get('equipment_import_batch_id')
    if not batch_id:
        return jsonify({'success': False, 'message': '会话过期'}), 400

    try:
        row_idx = int(request.form.get('row_idx', -1))
    except (ValueError, TypeError):
        return jsonify({'success': False, 'message': 'row_idx 无效'}), 400

    field = request.form.get('field', '').strip()
    value = request.form.get('value', '').strip()
    identity = current_identity()
    if identity is None:
        return jsonify({'success': False, 'message': '未登录'}), 401
    try:
        _equipment_resources_service().update_equipment_import_row(
            batch_id, owner_user_id=identity.user_id, row_idx=row_idx,
            field=field, value=value,
        )
        return jsonify({'success': True, 'row_idx': row_idx, 'field': field, 'value': value})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/import/preview', methods=['GET'])
def import_preview_page():
    """步骤2：匹配预览页面（支持分页 + 标签筛选）

    [REQ-013-fix] 修复 1000+ 条数据时 HTML 8.2MB 渲染超时问题：
    - 服务端按 page/filter 切片返回数据
    - 每页 100 行，避免一次渲染所有行
    - 翻页/筛选通过 GET URL 参数实现（无 JS 状态）
    """
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    batch_id = session.get('equipment_import_batch_id')
    if not batch_id:
        flash('请先上传文件', 'error')
        return redirect(url_for('equipment.import_page'))
    identity = current_identity()
    if identity is None:
        return redirect(url_for('auth.login'))
    try:
        batch = _equipment_resources_service().get_equipment_import_batch(
            batch_id, owner_user_id=identity.user_id
        )
    except ResourceServiceError:
        session.pop('equipment_import_batch_id', None)
        flash('导入批次不存在或已过期', 'error')
        return redirect(url_for('equipment.import_page'))
    preview_data = {
        'filename': batch['sourceName'], 'processed': batch['rows'],
        'statistics': batch['statistics'],
    }

    # 读取分页/筛选参数
    try:
        page = int(request.args.get('p', 1))
    except (ValueError, TypeError):
        page = 1
    try:
        per_page = int(request.args.get('per_page', 100))
    except (ValueError, TypeError):
        per_page = 100
    if per_page not in (50, 100, 200):
        per_page = 100
    if page < 1:
        page = 1

    filter_type = request.args.get('filter', 'all')
    if filter_type not in ('all', 'exact', 'fuzzy', 'unmatched', 'duplicate', 'newcat'):
        filter_type = 'all'

    all_processed = preview_data['processed']

    # 标签筛选
    if filter_type == 'all':
        filtered_rows = all_processed
    elif filter_type == 'duplicate':
        filtered_rows = [r for r in all_processed if r.get('duplicate_status') == 'exists']
    elif filter_type == 'newcat':
        filtered_rows = [r for r in all_processed if r.get('is_new_category')]
    else:
        # exact / fuzzy / unmatched (基于 subclass_result.method 或 related_result)
        filtered_rows = []
        for r in all_processed:
            # 优先看 subclass_result.method
            sr = r.get('subclass_result') or {}
            if sr.get('method') == filter_type:
                filtered_rows.append(r)
                continue
            # 再看 related_result.devices 中的状态
            for d in (r.get('related_result') or {}).get('devices', []):
                if d.get('status') == filter_type:
                    filtered_rows.append(r)
                    break

    # 分页切片
    total_rows = len(filtered_rows)
    total_pages = max(1, (total_rows + per_page - 1) // per_page)
    if page > total_pages:
        page = total_pages
    start = (page - 1) * per_page
    end = start + per_page
    page_rows = filtered_rows[start:end]

    return render_template('equipment/import_step2.html',
                           filename=preview_data['filename'],
                           processed=page_rows,
                           statistics=preview_data['statistics'],
                           page=page,
                           per_page=per_page,
                           total_pages=total_pages,
                           total_rows=total_rows,
                           filter_type=filter_type)


@bp.route('/import/commit', methods=['POST'])
def import_commit():
    """步骤2确认提交：异步执行导入

    [REQ-014] 性能优化 + 实时进度反馈
    - 异步后台线程跑导入（避免阻塞 HTTP）
    - 整批事务原子（成功全提交 / 失败全回滚）
    - SSE 实时推送进度
    - 单价默认 0（不是 NULL）
    """
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录', 'redirect': url_for('auth.login')}), 401
    batch_id = session.get('equipment_import_batch_id')
    if not batch_id:
        return jsonify({'success': False, 'message': '会话过期，请重新上传'}), 400

    # [REQ-014-fix] 改用 JSON body（form 受 Flask max_form_memory_size 限制，大数据量 413）
    form_data = request.get_json(silent=True) or {}

    identity = current_identity()
    if identity is None:
        return jsonify({'success': False, 'message': '未登录'}), 401
    try:
        batch = _equipment_resources_service().get_equipment_import_batch(
            batch_id, owner_user_id=identity.user_id
        )
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code
    total = len(batch.get('rows', []))

    # 创建进度任务
    task_id = ip.create_task(
        total=total, user=session.get('user', ''),
        owner_user_id=identity.user_id,
    )

    # 关键：路由内立即 pop session，避免重复提交
    session.pop('equipment_import_batch_id', None)

    # 启动后台线程跑导入
    t = threading.Thread(
        target=_do_import_thread,
        args=(
            task_id, _equipment_resources_service(), batch_id, form_data,
            identity.user_id,
            getattr(request, 'request_id', 'equipment-import'),
        ),
        daemon=True,
        name=f'import-{task_id[:8]}',
    )


    t.start()

    return jsonify({
        'success': True,
        'task_id': task_id,
        'total': total,
        'sse_url': f'/equipment/import/progress/{task_id}',
    })


@bp.route('/api/search')
def api_search():
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'}), 401
    try:
        result = _equipment_resources_service().search_equipment_candidates(
            keyword=request.args.get('keyword', ''),
            category=request.args.get('category', ''),
            page=request.args.get('page', 1), page_size=20,
        )
        return jsonify({'success': True, **result})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


def _do_import_thread(
    task_id, service, batch_id, form_data, owner_user_id, request_id
):
    """在无 Flask 请求上下文的后台线程中执行原子导入。"""
    try:
        result = service.commit_equipment_import(
            batch_id, form_data=form_data, owner_user_id=owner_user_id,
            request_id=request_id,
            cancel_check=lambda: ip.is_cancel_requested(task_id),
            begin_commit_callback=lambda: ip.begin_commit(task_id),
            progress_callback=lambda current, name: ip.update_progress(
                task_id, current, name
            ),
        )
        ip.mark_done(
            task_id, saved_new=result['saved_new'],
            saved_overwrite=result['saved_overwrite'], skipped=result['skipped'],
            relations_created=result['relations_created'],
            skipped_relations=result['skipped_relations'],
        )

        # 5 分钟后清理进度
        def _delayed_cleanup():
            time.sleep(300)
            ip.cleanup_task(task_id)
        threading.Thread(target=_delayed_cleanup, daemon=True).start()

    except ResourceServiceError as error:
        if error.code == 'IMPORT_CANCELLED':
            ip.mark_cancelled(task_id, saved_count=0)
        else:
            ip.mark_failed(task_id, error=f'{error.code}: {error.message}', failed_row=None)
    except Exception as error:
        ip.mark_failed(task_id, error=f'{type(error).__name__}: {error}', failed_row=None)


@bp.route('/import/progress/<task_id>')
def import_progress_stream(task_id):
    """SSE 实时推送导入进度

    客户端：new EventSource('/equipment/import/progress/<task_id>')
    事件格式：data: {json}\n\n
    """
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'}), 401
    identity = current_identity()
    if identity is None or ip.get_progress(task_id, identity.user_id) is None:
        return jsonify({'success': False, 'message': '任务不存在'}), 404

    def generate():
        # 最多推送 5 分钟（防止泄漏）
        max_iterations = 5 * 60 / 0.2  # 1500
        seen_status = None
        for _ in range(int(max_iterations)):
            p = ip.get_progress(task_id, identity.user_id)
            if p is None:
                # 任务不存在或已清理
                yield f"data: {json.dumps({'status': 'not_found', 'error': '任务不存在或已清理'})}\n\n"
                break
            yield f"data: {json.dumps(p, ensure_ascii=False)}\n\n"
            if p['status'] in ('done', 'failed', 'cancelled'):
                seen_status = p['status']
                break
            time.sleep(0.2)
        # SSE 标准做法：done 后多发一个 comment 让客户端知道 stream 结束
        if seen_status:
            yield f"event: end\ndata: {json.dumps({'status': seen_status})}\n\n"

    return Response(stream_with_context(generate()), mimetype='text/event-stream',
                    headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})


@bp.route('/import/cancel/<task_id>', methods=['POST'])
def import_cancel_task(task_id):
    """请求取消导入任务

    异步：标记 cancel_requested=True，后台线程每 10 行检查一次
    """
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'}), 401
    identity = current_identity()
    if identity is None:
        return jsonify({'success': False, 'message': '未登录'}), 401

    if ip.request_cancel(task_id, identity.user_id):
        return jsonify({'success': True, 'message': '已发送取消请求'})
    else:
        p = ip.get_progress(task_id, identity.user_id)
        if p is None:
            return jsonify({'success': False, 'message': '任务不存在'}), 404
        return jsonify({'success': False, 'message': f"任务状态 {p['status']}，无法取消"})


@bp.route('/import/done')
def import_done():
    """步骤3：导入结果页"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))
    import json
    new_count = int(request.args.get('new', 0))
    overwrite_count = int(request.args.get('overwrite', 0))
    skip_count = int(request.args.get('skip', 0))
    relations_count = int(request.args.get('relations', 0))
    try:
        skipped_relations = json.loads(request.args.get('skipped_relations', '[]'))
    except Exception:
        skipped_relations = []
    return render_template('equipment/import_done.html',
                           new_count=new_count,
                           overwrite_count=overwrite_count,
                           skip_count=skip_count,
                           relations_count=relations_count,
                           skipped_relations=skipped_relations)


# 旧的 CSV 导入保留（兼容）
@bp.route('/import/csv', methods=['POST'])
def import_csv():
    """旧的 CSV 导入（兼容现有调用）"""
    if 'user' not in session:
        return jsonify({'success': False, 'message': '未登录'})
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '请选择文件'})
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': '请选择文件'})
    if not file.filename.endswith('.csv'):
        return jsonify({'success': False, 'message': '请上传CSV文件'})
    try:
        content = file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(content))
        rows = []
        for row in reader:
            rows.append({
                'name': row.get('设备名称', ''), 'model': row.get('型号', ''),
                'category': row.get('分类', '通用设备'), 'form': row.get('形态', ''),
                'price': row.get('单价', ''), 'tech_index': row.get('功能技术指标', ''),
                'tech_status': row.get('技术状态', '货架产品'),
                'manufacturer': row.get('研制单位', ''),
                'main_purpose': row.get('主要用途', ''),
                'former_name': row.get('曾用名', ''),
                'resource_guarantee': row.get('资源保障要求', ''),
            })
            if len(rows) > 1000:
                return jsonify({'success': False, 'message': '单次导入不能超过 1000 条'}), 422
        count = _equipment_resources_service().import_equipment_rows(rows)
        return jsonify({'success': True, 'message': f'成功导入 {count} 条记录'})
    except UnicodeDecodeError:
        return jsonify({'success': False, 'message': 'CSV 文件必须使用 UTF-8 编码'}), 422
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code
    except Exception as e:
        return jsonify({'success': False, 'message': f'导入失败: {str(e)}'})

@bp.route('/logs/<module_name>')
def logs(module_name):
    """操作日志页面"""
    if 'user' not in session:
        return redirect(url_for('auth.login'))

    # 获取筛选参数
    operator = request.args.get('operator')
    operation_type = request.args.get('operation_type')
    file_name = request.args.get('file_name')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    logs = _equipment_resources_service().list_logs(
        module_name,
        operation=operation_type,
        operator=operator,
        file_name=file_name,
        start_date=start_date,
        end_date=end_date,
    )

    module_names = {
        'equipment': '设备知识库',
        'standards': '标准法规库',
        'templates': '科研模板'
    }

    return render_template('equipment/logs.html',
                          logs=logs,
                          module_name=module_name,
                          module_title=module_names.get(module_name, '日志'))


# ---------- 密码设备关联宿主设备 API ----------

@bp.route('/api/hosts/<device_id>', methods=['GET'])
def api_hosts_by_device(device_id):
    """获取密码设备关联的宿主设备列表"""
    if 'user' not in session:
        return jsonify({'success': False})
    try:
        hosts = _equipment_resources_service().get_hosts_by_device(device_id)
        return jsonify({'success': True, 'data': hosts})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/api/hosts', methods=['POST'])
def api_add_host_relation():
    """从密码设备侧添加宿主设备关联（无数量概念）"""
    if 'user' not in session:
        return jsonify({'success': False})
    data = request.get_json() or {}
    device_id = data.get('device_id')
    host_id = data.get('host_id')
    if not device_id or not host_id:
        return jsonify({'success': False, 'message': '缺少参数'})
    try:
        _equipment_resources_service().add_device_host_relation(device_id, host_id)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code


@bp.route('/api/hosts/<device_id>/<host_id>', methods=['DELETE'])
def api_remove_host_relation(device_id, host_id):
    """从密码设备侧移除宿主设备关联"""
    if 'user' not in session:
        return jsonify({'success': False})
    try:
        _equipment_resources_service().remove_device_host_relation(device_id, host_id)
        return jsonify({'success': True})
    except ResourceServiceError as error:
        return jsonify({'success': False, 'message': error.message}), error.status_code
