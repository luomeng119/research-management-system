"""Retained argumentation: real login, routes and migrated PostgreSQL tables."""

import json
import io
import os
from pathlib import Path
import re
import sqlite3
import uuid
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
import sqlalchemy as sa

from app import create_app
from app.security.auth import hash_password


def template_version(client, category='research'):
    if category not in {'research', 'crypto', 'security'}:
        return 0
    return client.get(f'/template/api/get/{category}').json['version']


def template_identity(client, category='research'):
    if category not in {'research', 'crypto', 'security'}:
        return ''
    return client.get(f'/template/api/get/{category}').json['template_id'] or ''


def test_template_version_cannot_target_another_template(argumentation_runtime):
    _, client, headers, engine, _ = argumentation_runtime
    first, second = f'a-{uuid.uuid4().hex}', f'b-{uuid.uuid4().hex}'
    with engine.begin() as connection:
        for template_id in [first, second]:
            connection.execute(sa.text(
                "INSERT INTO doc_templates(template_id,name,category,version,chapter_tree) "
                "VALUES (:id,:id,'security',1000,'{\"chapters\": []}'::jsonb)"
            ), {'id': template_id})
    response = client.post('/template/api/save', headers=headers, json={
        'category': 'security', 'expected_version': 1000, 'expected_template_id': first,
        'template_data': {'chapters': [{'id': 'x', 'title': '不可写入另一个模板'}]},
    })
    assert response.status_code == 409
    with engine.connect() as connection:
        assert connection.execute(sa.text(
            'SELECT version FROM doc_templates WHERE template_id=:id'
        ), {'id': second}).scalar_one() == 1000
    with engine.begin() as connection:
        connection.execute(sa.text('DELETE FROM doc_templates WHERE template_id IN (:a,:b)'),
                           {'a': first, 'b': second})


@pytest.mark.parametrize('initial_version', [0, 1])
def test_template_simultaneous_writers_have_one_winner(argumentation_runtime, initial_version):
    from app.repositories.argumentation import ArgumentationConflict, ArgumentationRepository

    _, _, _, engine, _ = argumentation_runtime
    repository = ArgumentationRepository(engine)
    template_id = f'concurrent-{uuid.uuid4().hex}'
    if initial_version:
        repository.save_template(template_id, 'initial', 'research', None,
                                 {'chapters': []}, expected_version=0)
    barrier = Barrier(2)

    def save(name):
        barrier.wait(timeout=10)
        try:
            version = repository.save_template(template_id, name, 'research', None,
                                               {'chapters': [{'id': 'a', 'title': name}]},
                                               expected_version=initial_version, expected_template_id=template_id)
            return name, version
        except ArgumentationConflict:
            return name, None

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(save, ['first', 'second']))
    winners = [(name, version) for name, version in outcomes if version is not None]
    assert len(winners) == 1
    assert winners[0][1] == initial_version + 1
    with engine.connect() as connection:
        row = connection.execute(sa.text(
            'SELECT name, version, chapter_tree FROM doc_templates WHERE template_id=:id'
        ), {'id': template_id}).one()
    assert row.name == winners[0][0]
    assert row.version == initial_version + 1
    assert row.chapter_tree['chapters'][0]['title'] == winners[0][0]
    with engine.begin() as connection:
        connection.execute(sa.text('DELETE FROM doc_templates WHERE template_id=:id'), {'id': template_id})


def test_stale_template_upload_preserves_existing_file(argumentation_runtime):
    from docx import Document

    app, client, headers, engine, _ = argumentation_runtime
    version = template_version(client)

    def upload(expected):
        content = io.BytesIO()
        Document().save(content)
        content.seek(0)
        return client.post('/argumentation/template/upload', headers=headers, data={
            'category': 'research', 'expected_version': expected, 'expected_template_id': template_identity(client),
            'file': (content, 'template.docx'),
        })

    assert upload(version).status_code == 200
    root = Path(app.config['DATA_DIR']) / 'templates'
    before = {p: p.read_bytes() for p in root.iterdir()}
    assert upload(version).status_code == 409
    assert {p: p.read_bytes() for p in root.iterdir()} == before
    current = template_version(client)
    assert client.post('/template/api/save', headers=headers, json={
        'category': 'research', 'template_data': {'chapters': []}, 'expected_version': current, 'expected_template_id': template_identity(client),
    }).status_code == 200
    with engine.connect() as connection:
        path = connection.execute(sa.text(
            "SELECT file_path FROM doc_templates WHERE template_id='research_v1'"
        )).scalar_one()
    assert Path(path) in before
    assert Path(path).read_bytes() == before[Path(path)]


@pytest.mark.parametrize('endpoint', ['/template/api/save', '/argumentation/template/edit'])
def test_template_stale_editor_is_rejected(argumentation_runtime, endpoint):
    _, client, headers, _, _ = argumentation_runtime
    version = client.get('/template/api/get/research').json.get('version', 0)
    first = {'chapters': [{'id': 'saved', 'title': '先保存的内容'}]}
    assert client.post('/template/api/save', headers=headers, json={
        'category': 'research', 'template_data': first, 'expected_version': version, 'expected_template_id': template_identity(client),
    }).status_code == 200
    response = client.post(endpoint, headers=headers, json={
        'category': 'research', 'template_data': {'chapters': []},
        'chapter_tree': json.dumps({'chapters': []}), 'expected_version': version, 'expected_template_id': template_identity(client),
    })
    assert response.status_code == 409
    assert response.json['success'] is False
    assert client.get('/template/api/get/research').json['template'] == first


def test_template_upload_rejects_symlink_directory(argumentation_runtime, tmp_path):
    from docx import Document

    app, client, headers, engine, _ = argumentation_runtime
    outside = tmp_path / "outside"
    outside.mkdir()
    templates = Path(app.config["DATA_DIR"]) / "templates"
    templates.symlink_to(outside, target_is_directory=True)
    with engine.connect() as connection:
        before = connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all()
    content = io.BytesIO()
    Document().save(content)
    content.seek(0)
    response = client.post(
        "/argumentation/template/upload",
        data={"category": "research", "expected_version": template_version(client), "expected_template_id": template_identity(client), "file": (content, "template.docx")},
        headers=headers,
    )
    assert response.status_code == 500
    assert response.json["success"] is False
    assert list(outside.iterdir()) == []
    assert templates.is_symlink()
    with engine.connect() as connection:
        assert connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all() == before


def test_template_upload_rejects_forged_docx_without_persistence(argumentation_runtime):
    app, client, headers, engine, _ = argumentation_runtime
    with engine.connect() as connection:
        before = connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all()
    response = client.post(
        "/argumentation/template/upload",
        data={"category": "research", "expected_version": template_version(client), "expected_template_id": template_identity(client), "file": (io.BytesIO(b"not a Word document"), "fake.docx")},
        headers=headers,
    )
    assert response.status_code == 415
    assert response.json["success"] is False
    assert not list((Path(app.config["DATA_DIR"]) / "templates").glob("*"))
    assert not list((Path(app.config["FILE_STORAGE_ROOT"]) / ".staging").glob("*"))
    with engine.connect() as connection:
        assert connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all() == before


def test_template_database_failure_removes_only_new_upload(argumentation_runtime):
    from docx import Document

    app, client, headers, engine, _ = argumentation_runtime
    root = Path(app.config["DATA_DIR"]) / "templates"
    root.mkdir(parents=True, exist_ok=True)
    existing = root / "retained.docx"
    existing.write_bytes(b"retained original bytes")
    content = io.BytesIO()
    Document().save(content)
    content.seek(0)

    def fail_insert(conn, cursor, statement, parameters, context, many):
        if statement.lstrip().upper().startswith(("INSERT INTO DOC_TEMPLATES", "UPDATE DOC_TEMPLATES")):
            raise RuntimeError("injected template persistence failure")

    sa.event.listen(engine, "before_cursor_execute", fail_insert)
    try:
        response = client.post(
            "/argumentation/template/upload",
            data={"category": "research", "expected_version": template_version(client), "expected_template_id": template_identity(client), "file": (content, "template.docx")},
            headers=headers,
        )
    finally:
        sa.event.remove(engine, "before_cursor_execute", fail_insert)
    assert response.status_code >= 500
    assert {p.name: p.read_bytes() for p in root.iterdir()} == {
        "retained.docx": b"retained original bytes"
    }


def test_template_upload_preserves_existing_chapters(argumentation_runtime):
    from docx import Document

    _, client, headers, _, _ = argumentation_runtime
    tree = {"name": "已有模板", "chapters": [{"id": "existing", "title": "已有研究内容"}]}
    assert client.post("/template/api/save", headers=headers, json={
        "category": "research", "template_data": tree, "expected_version": template_version(client), "expected_template_id": template_identity(client),
    }).json["success"] is True
    content = io.BytesIO()
    Document().save(content)
    content.seek(0)
    response = client.post(
        "/argumentation/template/upload",
        data={"category": "research", "expected_version": template_version(client), "expected_template_id": template_identity(client), "file": (content, "template.docx")},
        headers=headers,
    )
    assert response.status_code == 200
    assert client.get("/template/api/get/research").json["template"] == tree


@pytest.mark.parametrize("category", ["../escaped", "unknown", "", "科研项目"])
def test_template_upload_rejects_category_before_file_write(argumentation_runtime, category):
    app, client, headers, engine, _ = argumentation_runtime
    root = Path(app.config["DATA_DIR"])
    log = Path(app.config["LOG_FILE"])
    before = {p: p.read_bytes() for p in root.rglob("*") if p.is_file() and p != log}
    with engine.connect() as connection:
        templates_before = connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all()
    response = client.post(
        "/argumentation/template/upload",
        data={"category": category, "expected_version": template_version(client, category), "expected_template_id": template_identity(client, category), "file": (io.BytesIO(b"test"), "template.docx")},
        headers=headers,
    )
    assert response.status_code == 400
    assert response.json["success"] is False
    assert {p: p.read_bytes() for p in root.rglob("*") if p.is_file() and p != log} == before
    with engine.connect() as connection:
        assert connection.execute(sa.text(
            "SELECT template_id, version, file_path FROM doc_templates ORDER BY template_id"
        )).all() == templates_before


@pytest.mark.parametrize("category", ["research", "crypto", "security"])
def test_template_upload_accepts_existing_category_values(argumentation_runtime, category):
    from docx import Document

    _, client, headers, engine, _ = argumentation_runtime
    document = Document()
    document.add_heading("演练模板", level=1)
    content = io.BytesIO()
    document.save(content)
    expected = content.getvalue()
    content.seek(0)
    response = client.post(
        "/argumentation/template/upload",
        data={"category": category, "expected_version": template_version(client, category), "expected_template_id": template_identity(client, category), "file": (content, "template.docx")},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json["success"] is True
    with engine.connect() as connection:
        path = connection.execute(sa.text(
            "SELECT file_path FROM doc_templates WHERE template_id = :template_id"
        ), {"template_id": f"{category}_v1"}).scalar_one()
    assert Path(path).read_bytes() == expected


@pytest.fixture
def argumentation_runtime(tmp_path, monkeypatch):
    if not os.environ.get("TEST_DATABASE_URL"):
        pytest.skip("requires the isolated PostgreSQL contract runner")
    engine = sa.create_engine(os.environ["TEST_DATABASE_URL"])
    suffix = uuid.uuid4().hex
    username = f"argumentation_{suffix}"
    project_id = f"ARG-{suffix}"
    metadata = sa.MetaData()
    users = sa.Table("users", metadata, autoload_with=engine)
    projects = sa.Table("projects", metadata, autoload_with=engine)
    with engine.begin() as connection:
        connection.execute(
            users.insert().values(
                username=username,
                password=hash_password("TestPassword123!"),
                name="张老师",
                role="BUSINESS_USER",
                status="active",
                must_change_password=False,
                version=1,
            )
        )
        connection.execute(
            projects.insert().values(
                project_id=project_id,
                name="演练：科研方案保留验证",
                leader="张老师",
            )
        )
    app = create_app(
        {
            "TESTING": True,
            "PROPAGATE_EXCEPTIONS": False,
            "SECRET_KEY": "argumentation-test-only-secret",
            "DATABASE_ENGINE": engine,
            "DATA_DIR": str(tmp_path / "data"),
            "UPLOAD_DIR": str(tmp_path / "uploads"),
            "DOCUMENTS_DIR": str(tmp_path / "documents"),
            "SESSION_FILE_DIR": str(tmp_path / "sessions"),
            "SECURITY_AUTH_ENABLED": True,
            "CSRF_ENABLED": True,
        }
    )
    client = app.test_client()
    login = client.get("/auth/login")
    token = re.search(r'name="_csrf_token" value="([^"]+)"', login.text).group(1)
    assert (
        client.post(
            "/auth/login",
            data={
                "username": username,
                "password": "TestPassword123!",
                "_csrf_token": token,
            },
        ).status_code
        == 302
    )
    page = client.get("/users/change-password")
    token = re.search(r'name="_csrf_token" value="([^"]+)"', page.text).group(1)

    def forbidden_sqlite(*args, **kwargs):
        raise sqlite3.OperationalError("retained runtime must not access SQLite")

    monkeypatch.setattr(sqlite3, "connect", forbidden_sqlite)
    yield app, client, {"X-CSRF-Token": token}, engine, project_id
    engine.dispose()


def test_template_editor_persists_to_postgres_without_sqlite(argumentation_runtime):
    app, client, headers, engine, _ = argumentation_runtime
    tree = {"name": "演练模板", "chapters": [{"id": "ch1", "title": "研究目标"}]}
    response = client.post(
        "/template/api/save",
        json={
            "category": "research",
            "template_data": tree, "expected_version": template_version(client), "expected_template_id": template_identity(client),
        },
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json["success"] is True
    assert client.get("/template/api/get/research").json["template"] == tree
    with engine.connect() as connection:
        saved = connection.execute(
            sa.text(
                "SELECT chapter_tree FROM doc_templates WHERE template_id='research_v1'"
            )
        ).scalar_one()
    assert saved == tree
    assert client.get("/template/manage").status_code == 200
    assert client.get("/template/edit/research").status_code == 200


def test_document_save_keeps_identity_and_immutable_versions(argumentation_runtime):
    app, client, headers, engine, project_id = argumentation_runtime
    payload = {
        "project_id": project_id,
        "category": "research",
        "expected_version": 0,
        "content": {"ch1": "<p>第一版研究目标</p>"},
        "change_note": "初稿",
    }
    first = client.post("/argumentation/save", json=payload, headers=headers)
    assert first.status_code == 200
    assert first.json["success"] is True
    assert first.json["version"] == 1
    payload["content"] = {"ch1": "<p>第二版研究目标</p>"}
    payload["expected_version"] = 1
    second = client.post("/argumentation/save", json=payload, headers=headers)
    assert second.status_code == 200
    assert second.json["success"] is True
    assert second.json["version"] == 2
    assert second.json["doc_id"] == first.json["doc_id"]
    with engine.connect() as connection:
        documents = (
            connection.execute(
                sa.text(
                    "SELECT doc_id,content FROM project_documents WHERE project_id=:project"
                ),
                {"project": project_id},
            )
            .mappings()
            .all()
        )
        versions = (
            connection.execute(
                sa.text(
                    "SELECT version_num,content FROM document_versions WHERE doc_id=:doc ORDER BY version_num"
                ),
                {"doc": first.json["doc_id"]},
            )
            .mappings()
            .all()
        )
    assert len(documents) == 1
    assert json.loads(documents[0]["content"]) == {"ch1": "<p>第二版研究目标</p>"}
    assert [row["version_num"] for row in versions] == [1, 2]
    assert json.loads(versions[0]["content"]) == {"ch1": "<p>第一版研究目标</p>"}


def test_editor_and_history_pages_read_postgres(argumentation_runtime):
    _, client, headers, engine, project_id = argumentation_runtime
    saved = client.post(
        "/argumentation/save",
        json={
            "project_id": project_id,
            "category": "research",
            "content": {"ch1": "已保存目标"},
            "expected_version": 0,
        },
        headers=headers,
    ).json
    assert saved["success"]
    index = client.get("/argumentation/research")
    assert index.status_code == 200
    assert project_id in index.text
    editor = client.get(f"/argumentation/research/{project_id}")
    assert editor.status_code == 200
    assert "已保存目标" in editor.text or "\\u5df2\\u4fdd\\u5b58" in editor.text
    assert client.get(f"/argumentation/versions/{saved['doc_id']}").status_code == 200
    history = client.get(f"/argumentation/versions/{saved['doc_id']}/1")
    assert history.status_code == 200
    assert "历史版本" in history.text
    assert 'onclick="saveDocument()"' not in history.text
    assert client.get("/argumentation/versions/no-such-document").status_code == 404


def test_stale_document_save_is_rejected_without_overwriting(argumentation_runtime):
    _, client, headers, engine, project_id = argumentation_runtime
    payload = {
        "project_id": project_id,
        "category": "research",
        "expected_version": 0,
        "content": {"ch1": "甲已保存"},
    }
    first = client.post("/argumentation/save", json=payload, headers=headers)
    assert first.json["success"]
    payload["content"] = {"ch1": "乙基于旧页面的修改"}
    stale = client.post("/argumentation/save", json=payload, headers=headers)
    assert stale.status_code == 409
    with engine.connect() as connection:
        content = connection.execute(
            sa.text("SELECT content FROM project_documents WHERE project_id=:project"),
            {"project": project_id},
        ).scalar_one()
        count = connection.execute(
            sa.text("SELECT count(*) FROM document_versions WHERE doc_id=:doc"),
            {"doc": first.json["doc_id"]},
        ).scalar_one()
    assert json.loads(content) == {"ch1": "甲已保存"}
    assert count == 1


def test_default_chapters_and_unrendered_legacy_content_are_preserved(
    argumentation_runtime,
):
    _, client, headers, engine, project_id = argumentation_runtime
    with engine.begin() as connection:
        connection.execute(
            sa.text("DELETE FROM doc_templates WHERE category='research'")
        )
    editor = client.get(f"/argumentation/research/{project_id}")
    assert 'data-field="ch1_1"' in editor.text
    assert 'data-field="ch7"' in editor.text
    payload = {
        "project_id": project_id,
        "category": "research",
        "expected_version": 0,
        "content": {"ch1_1": "原始背景", "custom_legacy": "不能遗失的旧正文"},
    }
    assert client.post("/argumentation/save", json=payload, headers=headers).json[
        "success"
    ]
    payload.update(expected_version=1, content={"ch1_1": "更新背景"})
    assert client.post("/argumentation/save", json=payload, headers=headers).json[
        "success"
    ]
    with engine.connect() as connection:
        content = connection.execute(
            sa.text("SELECT content FROM project_documents WHERE project_id=:project"),
            {"project": project_id},
        ).scalar_one()
    assert json.loads(content) == {"ch1_1": "更新背景", "custom_legacy": "不能遗失的旧正文"}


@pytest.mark.parametrize(
    "tree",
    [
        {"chapters": "invalid"},
        {"chapters": ["invalid"]},
        {"chapters": [{"id": "a", "title": "x", "children": {}}]},
    ],
)
def test_invalid_template_structure_does_not_replace_saved_template(
    argumentation_runtime, tree
):
    _, client, headers, _, _ = argumentation_runtime
    before = client.get("/template/api/get/research").json
    response = client.post(
        "/template/api/save",
        json={
            "category": "research",
            "template_data": tree, "expected_version": template_version(client), "expected_template_id": template_identity(client),
        },
        headers=headers,
    )
    assert response.status_code == 400
    assert client.get("/template/api/get/research").json == before


@pytest.mark.parametrize("failure_point", ["version", "audit"])
def test_save_failure_rolls_back_body_history_and_audit(
    argumentation_runtime, failure_point
):
    app, client, headers, engine, project_id = argumentation_runtime

    def fail_at_statement(conn, cursor, statement, parameters, context, many):
        target = "document_versions" if failure_point == "version" else "audit_events"
        if statement.lstrip().upper().startswith("INSERT INTO " + target.upper()):
            raise RuntimeError("injected persistence failure")

    sa.event.listen(engine, "before_cursor_execute", fail_at_statement)
    try:
        response = client.post(
            "/argumentation/save",
            headers=headers,
            json={
                "project_id": project_id,
                "category": "research",
                "expected_version": 0,
                "content": {"ch1": "不得留下半提交"},
            },
        )
        assert response.status_code >= 500
    finally:
        sa.event.remove(engine, "before_cursor_execute", fail_at_statement)
    with engine.connect() as connection:
        assert (
            connection.execute(
                sa.text(
                    "SELECT count(*) FROM project_documents WHERE project_id=:project"
                ),
                {"project": project_id},
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                sa.text(
                    "SELECT count(*) FROM document_versions WHERE content LIKE '%不得留下半提交%'"
                )
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                sa.text("SELECT count(*) FROM audit_events WHERE object_id=:project"),
                {"project": project_id},
            ).scalar_one()
            == 0
        )


def test_simultaneous_first_save_accepts_only_one_editor(argumentation_runtime):
    _, _, _, engine, project_id = argumentation_runtime
    from app.repositories.argumentation import (
        ArgumentationRepository,
        ArgumentationConflict,
    )

    repository = ArgumentationRepository(engine)
    barrier = Barrier(2)

    def write(content):
        with engine.begin() as connection:
            barrier.wait(timeout=10)
            try:
                return repository.save_revision(
                    connection,
                    project_id=project_id,
                    category="research",
                    content=json.dumps({"ch1": content}),
                    editor="张老师",
                    change_note="并发验证",
                    expected_version=0,
                )[1]
            except ArgumentationConflict:
                return "conflict"

    with ThreadPoolExecutor(max_workers=2) as workers:
        results = list(workers.map(write, ["甲", "乙"]))
    assert results.count(1) == 1
    assert results.count("conflict") == 1
    with engine.connect() as connection:
        assert (
            connection.execute(
                sa.text(
                    "SELECT count(*) FROM project_documents WHERE project_id=:project"
                ),
                {"project": project_id},
            ).scalar_one()
            == 1
        )


def test_migrated_doc_id_and_text_version_number_are_preserved(argumentation_runtime):
    _, client, headers, engine, project_id = argumentation_runtime
    doc_id = "legacy-" + uuid.uuid4().hex
    with engine.begin() as connection:
        connection.execute(
            sa.text(
                "INSERT INTO project_documents(doc_id,project_id,category,content) "
                "VALUES(:doc,:project,'research',:content)"
            ),
            {"doc": doc_id, "project": project_id, "content": '{"ch1":"旧正文"}'},
        )
        connection.execute(
            sa.text(
                "INSERT INTO document_versions(version_id,doc_id,version_number,content) "
                "VALUES(:version,:doc,'7',:content)"
            ),
            {"version": uuid.uuid4().hex, "doc": doc_id, "content": '{"ch1":"旧正文"}'},
        )
    response = client.post(
        "/argumentation/save",
        headers=headers,
        json={
            "project_id": project_id,
            "category": "research",
            "expected_version": 7,
            "content": {"ch1": "新版正文"},
        },
    )
    assert response.status_code == 200
    assert response.json["doc_id"] == doc_id
    assert response.json["version"] == 8
    with engine.connect() as connection:
        assert (
            connection.execute(
                sa.text(
                    "SELECT version_number FROM document_versions WHERE doc_id=:doc AND version_num=8"
                ),
                {"doc": doc_id},
            ).scalar_one()
            == "8"
        )
