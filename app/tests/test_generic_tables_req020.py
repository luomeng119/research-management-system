from __future__ import annotations

import io
from pathlib import Path

import pytest
import sqlalchemy as sa
from sqlalchemy.pool import StaticPool

from app import create_app
from app.tests.test_generic_tables_postgres import _schema


def _client(tmp_path):
    engine = sa.create_engine(
        "sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    _schema(engine)
    app = create_app({
        "TESTING": True, "SECRET_KEY": "generic-test", "DATABASE_ENGINE": engine,
        "DATA_DIR": str(tmp_path), "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "SECURITY_AUTH_ENABLED": False, "CSRF_ENABLED": False,
        "AI_PROVIDER": "DISABLED", "LOG_FILE": None,
    })
    client = app.test_client()
    with client.session_transaction() as session:
        session["user"] = "张老师"
        session["user_id"] = 7
    return client


def test_req020_snapshot_current_and_history_readonly(tmp_path):
    client = _client(tmp_path)
    created = client.post("/api/generic-tables", json={"name": "科研台账"})
    assert created.status_code == 200
    table_id = created.get_json()["table_id"]
    versions = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"]
    current = versions[0]["version_id"]
    saved = client.post(
        f"/api/generic-tables/{table_id}/versions",
        json={"method": "snapshot", "source_version_id": current, "label": "初版"},
    )
    assert saved.status_code == 200
    body = saved.get_json()
    assert body["snapshot_id"] != body["new_current_id"]
    readonly = client.get(f"/tables/{table_id}/version/{body['snapshot_id']}/")
    assert readonly.status_code == 200
    assert "只读快照" in readonly.get_data(as_text=True)


def test_req020_locked_write_and_legacy_xls_are_rejected(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    current = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    snapshot = client.post(f"/api/generic-tables/{table_id}/versions", json={"method": "snapshot", "source_version_id": current}).get_json()["snapshot_id"]
    write = client.post(f"/api/generic-tables/versions/{snapshot}/rows", json={"row_data": {"name": "x"}})
    assert write.status_code == 409
    legacy = client.post(
        f"/api/generic-tables/versions/{snapshot}/rows/import",
        data={"file": (io.BytesIO(b"legacy"), "legacy.xls")},
        content_type="multipart/form-data",
    )
    assert legacy.status_code == 400
    assert ".xlsx" in legacy.get_json()["error"]


def test_list_import_invalid_upload_does_not_create_or_switch_version(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    before = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"]
    source = before[0]["version_id"]
    response = client.post(
        f"/api/generic-tables/{table_id}/versions/import",
        data={"source_version_id": source, "file": (io.BytesIO(b"not a zip"), "broken.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json()["code"] == "INVALID_XLSX"
    after = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"]
    assert after == before


def test_legacy_two_step_import_cannot_create_empty_current(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    before = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"]
    response = client.post(f"/api/generic-tables/{table_id}/versions", json={"method": "import", "source_version_id": before[0]["version_id"]})
    assert response.status_code == 409
    assert response.get_json()["code"] == "ATOMIC_IMPORT_REQUIRED"
    assert client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"] == before


def test_list_import_switches_current_only_with_successful_xlsx(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    source = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    stream = io.BytesIO()
    import openpyxl
    workbook = openpyxl.Workbook(); workbook.active.append(["名称"]); workbook.active.append(["课题甲"]); workbook.save(stream); stream.seek(0)
    response = client.post(
        f"/api/generic-tables/{table_id}/versions/import",
        data={"source_version_id": source, "file": (stream, "valid.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    body = response.get_json()
    versions = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"]
    assert body["imported"] == 1
    assert next(version for version in versions if version["is_current"])["version_id"] == body["version_id"]


def test_noncurrent_unlocked_history_has_no_editable_javascript_state(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    source = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    saved = client.post(f"/api/generic-tables/{table_id}/versions", json={"method": "snapshot", "source_version_id": source}).get_json()
    page = client.get(f"/tables/{table_id}/version/{source}/").get_data(as_text=True)
    assert "const IS_CURRENT = false" in page
    assert "const CAN_EDIT = IS_CURRENT && !IS_LOCKED" in page
    assert saved["new_current_id"] != source


def test_unexpected_snapshot_failure_does_not_leak_exception_text(tmp_path, monkeypatch):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    source = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    service = client.application.extensions["generic_tables_service"]
    monkeypatch.setattr(service, "save_snapshot", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("secret-dsn")))
    response = client.post(f"/api/generic-tables/{table_id}/versions", json={"method": "snapshot", "source_version_id": source})
    assert response.status_code == 500
    assert response.get_json()["code"] == "SNAPSHOT_FAILED"
    assert "secret-dsn" not in response.get_data(as_text=True)


def test_export_temp_file_is_removed_when_response_closes(tmp_path, monkeypatch):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    version = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    exported = tmp_path / "export.xlsx"
    exported.write_bytes(b"xlsx-bytes")
    service = client.application.extensions["generic_tables_service"]
    monkeypatch.setattr(service, "export_to_excel", lambda _version_id: str(exported))
    response = client.get(f"/api/generic-tables/versions/{version}/export", buffered=False)
    assert response.status_code == 200 and exported.exists()
    _ = response.get_data()
    response.close()
    assert not exported.exists()


def test_export_sanitizes_crlf_name_and_cleans_if_response_build_fails(tmp_path, monkeypatch):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "台\r\nInjected: yes"}).get_json()["table_id"]
    version = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    service = client.application.extensions["generic_tables_service"]
    first = tmp_path / "first.xlsx"; first.write_bytes(b"xlsx")
    monkeypatch.setattr(service, "export_to_excel", lambda _version_id: str(first))
    response = client.get(f"/api/generic-tables/versions/{version}/export", buffered=False)
    assert response.status_code == 200
    assert "\r" not in response.headers["Content-Disposition"] and "\n" not in response.headers["Content-Disposition"]
    response.close()
    assert not first.exists()

    second = tmp_path / "second.xlsx"; second.write_bytes(b"xlsx")
    monkeypatch.setattr(service, "export_to_excel", lambda _version_id: str(second))
    monkeypatch.setattr("app.routes.generic_tables.send_file", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("header rejected")))
    with pytest.raises(ValueError, match="header rejected"):
        client.get(f"/api/generic-tables/versions/{version}/export")
    assert not second.exists()


def test_export_cleans_temp_if_post_export_table_lookup_fails(tmp_path, monkeypatch):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    version = client.get(f"/api/generic-tables/{table_id}/versions").get_json()["versions"][0]["version_id"]
    exported = tmp_path / "lookup-failure.xlsx"
    exported.write_bytes(b"xlsx")
    service = client.application.extensions["generic_tables_service"]
    monkeypatch.setattr(service, "export_to_excel", lambda _version_id: str(exported))
    monkeypatch.setattr(service, "get_by_id", lambda _table_id: (_ for _ in ()).throw(RuntimeError("lookup failed")))
    with pytest.raises(RuntimeError, match="lookup failed"):
        client.get(f"/api/generic-tables/versions/{version}/export")
    assert not exported.exists()


def test_detail_get_does_not_repair_missing_current_pointer(tmp_path):
    client = _client(tmp_path)
    table_id = client.post("/api/generic-tables", json={"name": "科研台账"}).get_json()["table_id"]
    service = client.application.extensions["generic_tables_service"]
    versions_before = service.get_versions(table_id)
    with service.repository.engine.begin() as connection:
        connection.execute(service.repository.tables.update().where(
            service.repository.tables.c.table_id == table_id,
        ).values(current_version_id=None))
    response = client.get(f"/tables/{table_id}")
    assert response.status_code == 409
    assert response.get_data(as_text=True) == "表格当前版本数据异常"
    assert service.get_by_id(table_id)["current_version_id"] is None
    assert service.get_versions(table_id) == versions_before
