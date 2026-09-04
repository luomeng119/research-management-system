from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import uuid

from flask import Flask, request
import openpyxl
import pytest
import sqlalchemy as sa
from sqlalchemy.pool import StaticPool

from app.repositories.resources import EquipmentResourcesRepository, ResearchResourcesRepository
from app.services.resources import EquipmentResourcesService, ResearchResourcesService, ResourceServiceError


ROOT = Path(__file__).resolve().parents[2]


class AuditRecorder:
    def __init__(self, *, fail=False):
        self.events = []
        self.fail = fail

    def record(self, connection, **event):
        if self.fail:
            raise RuntimeError("audit unavailable")
        self.events.append(event)
        return "audit-id"


def _expert_schema(engine):
    metadata = sa.MetaData()
    sa.Table(
        "experts", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("expert_id", sa.Text, nullable=False, unique=True),
        sa.Column("name", sa.Text, nullable=False),
        sa.Column("unit", sa.Text),
        sa.Column("position", sa.Text),
        sa.Column("expertise", sa.Text),
        sa.Column("bank_card", sa.Text),
        sa.Column("bank_name", sa.Text),
        sa.Column("uploader", sa.Text),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.Column("updated_at", sa.DateTime(timezone=True)),
        sa.Column("phone", sa.Text),
        sa.Column("id_card", sa.Text),
    )
    sa.Table(
        "expert_groups", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("group_id", sa.Text, nullable=False, unique=True),
        sa.Column("meeting_name", sa.Text, nullable=False),
        sa.Column("creator", sa.Text, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
    )
    sa.Table(
        "expert_group_members", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("group_id", sa.Text, nullable=False),
        sa.Column("expert_id", sa.Text, nullable=False),
        sa.Column("selected_by", sa.Text, nullable=False),
        sa.Column("selected_at", sa.DateTime(timezone=True), nullable=False),
        sa.UniqueConstraint("group_id", "expert_id"),
    )
    sa.Table(
        "expert_import_batches", metadata,
        sa.Column("id", sa.String(36), primary_key=True),
        sa.Column("owner_user_id", sa.Integer, nullable=False),
        sa.Column("source_name", sa.Text, nullable=False),
        sa.Column("source_sha256", sa.String(64), nullable=False),
        sa.Column("status", sa.Text, nullable=False),
        sa.Column("rows", sa.JSON, nullable=False),
        sa.Column("valid_count", sa.Integer, nullable=False),
        sa.Column("error_count", sa.Integer, nullable=False),
        sa.Column("duplicate_count", sa.Integer, nullable=False),
        sa.Column("result", sa.JSON),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("committed_at", sa.DateTime(timezone=True)),
        sa.Column("version", sa.Integer, nullable=False),
    )
    sa.Table(
        "equipment", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("equipment_id", sa.Text, nullable=False, unique=True),
        sa.Column("name", sa.Text, nullable=False),
        sa.Column("model", sa.Text), sa.Column("category", sa.Text),
        sa.Column("subclass", sa.Text), sa.Column("form", sa.Text),
        sa.Column("price", sa.Numeric(18, 2)), sa.Column("tech_index", sa.Text),
        sa.Column("tech_status", sa.Text), sa.Column("installation_requirements", sa.Text),
        sa.Column("manufacturer", sa.Text), sa.Column("equipment_image", sa.Text),
        sa.Column("related_files", sa.JSON), sa.Column("main_purpose", sa.Text),
        sa.Column("former_name", sa.Text), sa.Column("resource_guarantee", sa.Text),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.Column("updated_at", sa.DateTime(timezone=True)),
    )
    sa.Table(
        "knowledge_subclasses", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("parent_category", sa.Text, nullable=False),
        sa.Column("subclass_name", sa.Text, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.UniqueConstraint("parent_category", "subclass_name"),
    )
    sa.Table(
        "research_units", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("unit_id", sa.Text, nullable=False, unique=True),
        sa.Column("name", sa.Text, nullable=False), sa.Column("alias", sa.Text),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.Column("updated_at", sa.DateTime(timezone=True)),
    )
    sa.Table(
        "equipment_groups", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("group_id", sa.Text, nullable=False, unique=True),
        sa.Column("project_name", sa.Text, nullable=False),
        sa.Column("project_id", sa.Text, unique=True),
        sa.Column("creator", sa.Text, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
    )
    sa.Table(
        "equipment_group_members", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("group_id", sa.Text, nullable=False),
        sa.Column("equipment_id", sa.Text, nullable=False),
        sa.Column("quantity", sa.Integer, nullable=False),
        sa.Column("selected_by", sa.Text, nullable=False),
        sa.Column("selected_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("location", sa.Text),
        sa.UniqueConstraint("group_id", "equipment_id"),
    )
    sa.Table(
        "host_devices", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("host_id", sa.Text, nullable=False, unique=True),
        sa.Column("name", sa.Text, nullable=False), sa.Column("model", sa.Text),
        sa.Column("category", sa.Text), sa.Column("form", sa.Text),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.Column("updated_at", sa.DateTime(timezone=True)),
    )
    sa.Table(
        "host_device_categories", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("name", sa.Text, nullable=False, unique=True),
        sa.Column("created_at", sa.DateTime(timezone=True)),
    )
    sa.Table(
        "device_host_relations", metadata,
        sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
        sa.Column("device_id", sa.Text, nullable=False),
        sa.Column("host_id", sa.Text, nullable=False),
        sa.Column("quantity", sa.Integer, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True)),
        sa.Column("updated_at", sa.DateTime(timezone=True)),
        sa.UniqueConstraint("device_id", "host_id"),
    )
    metadata.create_all(engine)
    return metadata


@pytest.fixture
def expert_engine():
    engine = sa.create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    metadata = _expert_schema(engine)
    now = datetime.now(timezone.utc)
    with engine.begin() as connection:
        connection.execute(metadata.tables["experts"].insert(), [
            {
                "expert_id": "EXP-001", "name": "张老师", "unit": "第一研究室",
                "position": "研究员", "expertise": "材料", "phone": "13812345678",
                "id_card": "110101199001011234", "bank_card": "6222021234567890",
                "bank_name": "本地银行", "uploader": "tester", "created_at": now, "updated_at": now,
            },
            {
                "expert_id": "EXP-002", "name": "李老师", "unit": "第二研究室",
                "position": "高级工程师", "expertise": "通信", "phone": "13987654321",
                "id_card": "110101198802022345", "bank_card": "6222029876543210",
                "bank_name": "本地银行", "uploader": "tester", "created_at": now, "updated_at": now,
            },
        ])
    yield engine
    engine.dispose()


@pytest.fixture
def expert_service(expert_engine):
    audit = AuditRecorder()
    return ResearchResourcesService(ResearchResourcesRepository(expert_engine), audit), audit


@pytest.fixture
def equipment_service(expert_engine):
    return EquipmentResourcesService(EquipmentResourcesRepository(expert_engine))


@pytest.fixture
def equipment_routes(expert_engine, equipment_service):
    from app.routes.equipment import bp as equipment_bp
    from app.routes.research_units import bp as units_bp

    app = Flask(
        __name__, template_folder=str(ROOT / "app/templates"),
        static_folder=str(ROOT / "app/static"),
    )
    app.config.update(TESTING=True, SECRET_KEY="test-secret")
    app.jinja_env.globals["csrf_token"] = lambda: "test-csrf"
    app.extensions["database_engine"] = expert_engine
    app.extensions["equipment_resources_service"] = equipment_service
    app.add_url_rule("/test/login", endpoint="auth.login", view_func=lambda: "")
    app.add_url_rule("/test/logout", endpoint="auth.logout", view_func=lambda: "")
    app.add_url_rule("/test/users", endpoint="users.index", view_func=lambda: "")
    app.add_url_rule("/test/password", endpoint="users.change_password", view_func=lambda: "")
    app.register_blueprint(equipment_bp)
    app.register_blueprint(units_bp)
    client = app.test_client()
    with client.session_transaction() as active_session:
        active_session.update(
            user_id=7, user="zhang", name="张老师",
            role="BUSINESS_USER", account_version=1,
        )
    return client


def test_equipment_service_uses_bounded_filters_and_preserves_fields(equipment_service):
    first = equipment_service.create_equipment({
        "name": "密码机", "model": "M-1", "category": "密码设备",
        "subclass": "加密设备", "form": "机架式", "price": "1200.50",
        "techStatus": "在研", "manufacturer": "第一研究室",
        "mainPurpose": "试验", "resourceGuarantee": "实验室",
    })
    equipment_service.create_equipment({"name": "交换机", "category": "通用设备"})

    page = equipment_service.list_equipment(
        page=1, page_size=1, category="密码设备", keyword="密码"
    )
    assert page["total"] == 1
    assert page["items"][0]["equipment_id"] == first["equipment_id"]
    assert str(page["items"][0]["price"]) == "1200.50"
    updated = equipment_service.update_equipment(
        first["equipment_id"], {"name": "密码设备", "techStatus": "定型"}
    )
    assert updated["name"] == "密码设备"
    assert updated["tech_status"] == "定型"


def test_equipment_stats_merge_null_category_into_general(equipment_service):
    equipment_service.create_equipment({"name": "未分类设备", "price": "10"})
    equipment_service.create_equipment({
        "name": "通用设备", "category": "通用设备", "price": "20"
    })
    stats = equipment_service.equipment_stats()
    assert stats["by_category"]["通用设备"] == {"count": 2, "value": 30.0}
    assert stats["total"] == 2


def test_equipment_dictionaries_are_database_backed(equipment_service):
    unit = equipment_service.create_research_unit("第一研究室", "一室")
    subclass = equipment_service.create_subclass("密码设备", "加密设备")

    assert equipment_service.list_research_units()[0]["unit_id"] == unit["unit_id"]
    assert equipment_service.list_subclasses("密码设备")[0]["id"] == subclass["id"]
    with pytest.raises(ResourceServiceError) as caught:
        equipment_service.create_research_unit("第一研究室", "重复")
    assert caught.value.code == "RESOURCE_DUPLICATE"


def test_equipment_groups_preserve_quantity_and_location(equipment_service):
    equipment = equipment_service.create_equipment({"name": "密码机"})
    group = equipment_service.get_or_create_project_group(
        "P-001", "课题一", "张老师"
    )
    same = equipment_service.get_or_create_project_group(
        "P-001", "课题一", "李老师"
    )
    equipment_service.add_group_member(
        group["group_id"], equipment["equipment_id"],
        quantity=2, location="一号实验室", selected_by="张老师",
    )
    detail = equipment_service.get_equipment_group(group["group_id"])
    assert same["group_id"] == group["group_id"]
    assert detail["members"][0]["quantity"] == 2
    assert detail["members"][0]["location"] == "一号实验室"


def test_host_device_relations_are_replaced_atomically(equipment_service):
    first = equipment_service.create_equipment({"name": "密码机A"})
    second = equipment_service.create_equipment({"name": "密码机B"})
    host = equipment_service.create_host_device(
        {"name": "服务器", "category": "计算存储"},
        relations=[
            {"device_id": first["equipment_id"], "quantity": 1},
            {"device_id": second["equipment_id"], "quantity": 3},
        ],
    )
    equipment_service.replace_host_relations(host["host_id"], [
        {"device_id": second["equipment_id"], "quantity": 2},
    ])
    related = equipment_service.get_devices_by_host(host["host_id"])
    assert [(row["device_id"], row["quantity"]) for row in related] == [
        (second["equipment_id"], 2)
    ]

    original = equipment_service.repository.replace_host_relations
    def fail_after_delete(connection, host_id, relations, now):
        connection.execute(
            equipment_service.repository.device_host_relations.delete().where(
                equipment_service.repository.device_host_relations.c.host_id == host_id
            )
        )
        raise RuntimeError("injected relation failure")

    equipment_service.repository.replace_host_relations = fail_after_delete
    try:
        with pytest.raises(RuntimeError, match="injected relation failure"):
            equipment_service.replace_host_relations(host["host_id"], [
                {"device_id": first["equipment_id"], "quantity": 1},
            ])
    finally:
        equipment_service.repository.replace_host_relations = original
    assert equipment_service.get_devices_by_host(host["host_id"])[0]["device_id"] == second["equipment_id"]


def test_relation_quantity_rejects_fraction_and_malformed_rows(equipment_service):
    with pytest.raises(ResourceServiceError) as fraction:
        equipment_service.create_host_device(
            {"name": "宿主"}, relations=[{"device_id": "EQP-X", "quantity": 1.9}]
        )
    assert fraction.value.code == "VALIDATION_ERROR"
    with pytest.raises(ResourceServiceError) as malformed:
        equipment_service.create_host_device({"name": "宿主"}, relations=["EQP-X"])
    assert malformed.value.code == "VALIDATION_ERROR"


def test_equipment_and_unit_routes_use_postgres_service(equipment_routes, equipment_service):
    equipment_service.create_equipment({"name": "密码机", "category": "密码设备"})
    listing = equipment_routes.get("/equipment/?category=密码设备&fragment=1")
    assert listing.status_code == 200
    assert "密码机" in listing.get_data(as_text=True)

    created = equipment_routes.post(
        "/research-units/api", json={"name": "第二研究室", "alias": "二室"}
    )
    assert created.status_code == 201
    units = equipment_routes.get("/research-units/api/list").get_json()["units"]
    assert [item["name"] for item in units] == ["第二研究室"]

    invalid = equipment_routes.post("/equipment/add", data={
        "name": "", "price": "not-a-number",
    })
    assert invalid.status_code == 422
    assert "设备单价格式无效" in invalid.get_data(as_text=True)


@pytest.fixture
def expert_routes(expert_engine, expert_service):
    from app.routes.expert_groups import bp as groups_bp
    from app.routes.experts import bp as experts_bp

    service, audit = expert_service
    app = Flask(
        __name__, template_folder=str(ROOT / "app/templates"),
        static_folder=str(ROOT / "app/static"),
    )
    app.config.update(TESTING=True, SECRET_KEY="test-secret")
    app.jinja_env.globals["csrf_token"] = lambda: "test-csrf"
    app.extensions["database_engine"] = expert_engine
    app.extensions["resources_service"] = service
    app.add_url_rule("/test/users", endpoint="users.index", view_func=lambda: "")
    app.add_url_rule(
        "/test/change-password", endpoint="users.change_password", view_func=lambda: ""
    )
    app.add_url_rule("/test/logout", endpoint="auth.logout", view_func=lambda: "")
    app.add_url_rule("/test/login", endpoint="auth.login", view_func=lambda: "")
    app.register_blueprint(experts_bp)
    app.register_blueprint(groups_bp)

    @app.before_request
    def request_id():
        request.request_id = "req-route"

    client = app.test_client()
    with client.session_transaction() as active_session:
        active_session.update(
            user_id=7, user="zhang", name="张老师", role="BUSINESS_USER"
        )
    return client, service, audit


def test_expert_list_uses_server_pagination_and_masks_sensitive_fields(expert_service):
    service, _audit = expert_service

    result = service.list_experts(page=1, page_size=1, keyword="老师")

    assert result["total"] == 2
    assert result["page"] == 1
    assert result["pageSize"] == 1
    assert [item["name"] for item in result["items"]] == ["张老师"]
    assert result["items"][0]["phone"] == "138****5678"
    assert result["items"][0]["idCard"] == "110***********1234"
    assert result["items"][0]["bankCard"] == "6222********7890"


def test_sensitive_expert_view_returns_full_values_and_writes_redacted_audit(expert_service):
    service, audit = expert_service

    result = service.get_expert_sensitive(
        "EXP-001", actor_user_id=7, request_id="req-view"
    )

    assert result["phone"] == "13812345678"
    assert result["idCard"] == "110101199001011234"
    assert result["bankCard"] == "6222021234567890"
    assert audit.events == [{
        "event_name": "expert_sensitive_accessed",
        "user_id": 7,
        "object_type": "EXPERT",
        "object_id": "EXP-001",
        "result": "SUCCESS",
        "request_id": "req-view",
        "duration_ms": 0,
        "properties": {"operation": "VIEW", "record_count": 1},
    }]


def test_missing_sensitive_expert_does_not_write_success_audit(expert_service):
    service, audit = expert_service

    with pytest.raises(ResourceServiceError) as caught:
        service.get_expert_sensitive(
            "EXP-404", actor_user_id=7, request_id="req-missing"
        )

    assert caught.value.code == "EXPERT_NOT_FOUND"
    assert audit.events == []


def test_expert_create_and_update_preserve_the_original_business_fields(expert_service):
    service, _audit = expert_service

    created = service.create_expert(
        {
            "name": "王老师", "unit": "第三研究室", "position": "研究员",
            "expertise": "软件", "phone": "13711112222", "idCard": "110101198703033456",
            "bankCard": "6222021111222233", "bankName": "本地银行",
        },
        uploader="operator", actor_user_id=7, request_id="req-create",
    )
    updated = service.update_expert(
        created["expertId"],
        {
            "name": "王老师", "unit": "联合实验室", "position": "研究员",
            "expertise": "软件工程", "phone": "13711112222",
            "idCard": "110101198703033456", "bankCard": "6222021111222233",
            "bankName": "本地银行",
        },
        actor_user_id=7, request_id="req-update",
    )

    assert created["expertId"].startswith("EXP")
    assert updated["unit"] == "联合实验室"
    assert updated["expertise"] == "软件工程"
    assert updated["phone"] == "137****2222"


def test_expert_group_is_a_reusable_group_not_a_meeting_workflow(expert_service):
    service, _audit = expert_service

    group = service.create_expert_group(
        {"groupName": "通信与材料专家组"}, creator="operator"
    )
    service.add_expert_group_member(
        group["groupId"], "EXP-001", selected_by="operator"
    )
    detail = service.get_expert_group(group["groupId"])

    assert detail["groupName"] == "通信与材料专家组"
    assert "meetingName" not in detail
    assert [member["name"] for member in detail["members"]] == ["张老师"]
    assert detail["members"][0]["phone"] == "138****5678"
    with pytest.raises(ResourceServiceError) as caught:
        service.add_expert_group_member(
            group["groupId"], "EXP-001", selected_by="operator"
        )
    assert caught.value.code == "GROUP_MEMBER_EXISTS"


def test_expert_groups_support_paged_list_remove_and_delete(expert_service):
    service, _audit = expert_service
    group = service.create_expert_group(
        {"groupName": "通信专家组"}, creator="operator"
    )
    service.add_expert_group_member(
        group["groupId"], "EXP-001", selected_by="operator"
    )

    page = service.list_expert_groups(page=1, page_size=10)
    assert page["total"] == 1
    assert page["items"][0]["groupName"] == "通信专家组"
    assert page["items"][0]["memberCount"] == 1

    service.remove_expert_group_member(group["groupId"], "EXP-001")
    assert service.get_expert_group(group["groupId"])["members"] == []
    service.delete_expert_group(group["groupId"])
    assert service.list_expert_groups(page=1, page_size=10)["total"] == 0


def test_expert_group_sensitive_export_is_explicitly_audited(expert_service):
    service, audit = expert_service
    group = service.create_expert_group(
        {"groupName": "材料专家组"}, creator="operator"
    )
    service.add_expert_group_member(
        group["groupId"], "EXP-001", selected_by="operator"
    )

    exported = service.export_expert_group_sensitive(
        group["groupId"], actor_user_id=7, request_id="req-group-export"
    )

    assert exported["members"][0]["phone"] == "13812345678"
    assert audit.events[-1]["properties"] == {
        "operation": "EXPORT", "record_count": 1,
    }
    assert "name" not in audit.events[-1]["properties"]


def test_referenced_expert_cannot_be_deleted(expert_service):
    service, _audit = expert_service
    group = service.create_expert_group(
        {"groupName": "科研专家组"}, creator="operator"
    )
    service.add_expert_group_member(
        group["groupId"], "EXP-002", selected_by="operator"
    )

    with pytest.raises(ResourceServiceError) as caught:
        service.delete_expert("EXP-002")

    assert caught.value.code == "EXPERT_IN_USE"


def test_expert_routes_use_postgres_service_and_mask_list(expert_routes):
    client, service, audit = expert_routes

    listing = client.get("/experts/")
    assert listing.status_code == 200
    html = listing.get_data(as_text=True)
    assert "138****5678" in html
    assert "13812345678" not in html
    assert "6222021234567890" not in html

    detail = client.get("/experts/edit/EXP-001")
    assert detail.status_code == 200
    assert "13812345678" in detail.get_data(as_text=True)
    assert audit.events[-1]["event_name"] == "expert_sensitive_accessed"

    created = client.post("/experts/add", data={
        "name": "王老师", "unit": "第三研究室", "expertise": "软件",
    })
    assert created.status_code == 302
    assert service.list_experts(keyword="王老师")["total"] == 1

    expert_id = service.list_experts(keyword="王老师")["items"][0]["expertId"]
    deleted = client.post(f"/experts/delete/{expert_id}")
    assert deleted.status_code == 200
    assert deleted.get_json()["success"] is True


def test_expert_import_route_session_holds_only_batch_reference(expert_routes):
    client, service, _audit = expert_routes
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["姓名", "单位"])
    sheet.append(["王老师", "第三研究室"])
    payload = BytesIO()
    workbook.save(payload)

    preview = client.post(
        "/experts/import/preview",
        data={
            "file": (BytesIO(payload.getvalue()), "experts.xlsx"),
            "column_mapping_json": '{"0":"name","1":"unit"}',
        },
        content_type="multipart/form-data",
    )
    assert preview.status_code == 200
    with client.session_transaction() as active_session:
        batch_id = active_session["expert_import_task_id"]
        assert isinstance(batch_id, str)
        assert "processed" not in active_session

    page = client.get("/experts/import/preview")
    assert page.status_code == 200
    assert "王老师" in page.get_data(as_text=True)

    cancelled = client.post("/experts/import/cancel")
    assert cancelled.status_code == 200
    assert service.get_expert_import_batch(batch_id, owner_user_id=7)["status"] == "CANCELLED"
    assert service.list_experts(keyword="王老师")["total"] == 0


def test_import_rejects_duplicate_row_selection(expert_service):
    service, _audit = expert_service
    batch = service.create_expert_import_preview(
        source_name="duplicate.xlsx", source_bytes=b"test",
        rows=[{"_row_idx": 9, "name": "王老师", "unit": "第三研究室"}],
        owner_user_id=7,
    )

    with pytest.raises(ResourceServiceError) as caught:
        service.commit_expert_import(
            batch["batchId"],
            confirmed_rows=[
                {"_row_idx": 9, "name": "王老师"},
                {"_row_idx": 9, "name": "王老师"},
            ],
            owner_user_id=7, uploader="张老师", request_id="req-duplicate",
        )
    assert caught.value.code == "VALIDATION_ERROR"


def test_import_rechecks_edited_rows_against_existing_experts(expert_service):
    service, _audit = expert_service
    batch = service.create_expert_import_preview(
        source_name="edited-after-preview.xlsx", source_bytes=b"test",
        rows=[{
            "_row_idx": 5, "name": "王老师", "unit": "第三研究室",
            "phone": "13700000000",
        }],
        owner_user_id=7,
    )

    with pytest.raises(ResourceServiceError) as caught:
        service.commit_expert_import(
            batch["batchId"],
            confirmed_rows=[{
                "_row_idx": 5, "name": "张老师", "unit": "第一研究室",
                "phone": "13812345678",
            }],
            owner_user_id=7, uploader="李老师", request_id="req-edited-duplicate",
        )

    assert caught.value.code == "EXPERT_DUPLICATE"
    assert service.list_experts(page=1, page_size=20)["total"] == 2
    assert service.get_expert_import_batch(
        batch["batchId"], owner_user_id=7
    )["status"] == "PREVIEW"


def test_legacy_null_identity_is_normalized_for_create_and_import(expert_service):
    service, _audit = expert_service
    now = datetime.now(timezone.utc)
    with service.repository.engine.begin() as connection:
        connection.execute(service.repository.experts.insert().values(
            expert_id="EXP-LEGACY-NULL", name="王老师", unit=None, phone=None,
            position=None, expertise=None, bank_card=None, bank_name=None,
            id_card=None, uploader="legacy", created_at=now, updated_at=now,
        ))

    with pytest.raises(ResourceServiceError) as caught:
        service.create_expert(
            {"name": "王老师", "unit": "", "phone": ""},
            uploader="张老师", actor_user_id=7, request_id="req-null-create",
        )
    assert caught.value.code == "EXPERT_DUPLICATE"

    preview = service.create_expert_import_preview(
        source_name="legacy-null.xlsx", source_bytes=b"test",
        rows=[{"_row_idx": 0, "name": "王老师", "unit": "", "phone": ""}],
        owner_user_id=7,
    )
    assert preview["duplicateCount"] == 1
    assert preview["rows"][0]["_duplicate"] is True


def test_import_rejects_non_object_confirmed_row(expert_service):
    service, _audit = expert_service
    batch = service.create_expert_import_preview(
        source_name="invalid-row.xlsx", source_bytes=b"test",
        rows=[{"_row_idx": 1, "name": "王老师"}], owner_user_id=7,
    )

    with pytest.raises(ResourceServiceError) as caught:
        service.commit_expert_import(
            batch["batchId"], confirmed_rows=[1],
            owner_user_id=7, uploader="李老师", request_id="req-invalid-row",
        )

    assert caught.value.code == "VALIDATION_ERROR"


def test_expert_export_neutralizes_formulas_and_group_search_uses_text_content(expert_routes):
    client, service, _audit = expert_routes
    service.create_expert(
        {"name": '=HYPERLINK("https://example.invalid","x")', "unit": "=1+1"},
        uploader="张老师", actor_user_id=7, request_id="req-formula",
    )

    exported = client.get("/experts/export")
    workbook = openpyxl.load_workbook(BytesIO(exported.data), data_only=False)
    formula_row = next(
        row for row in workbook.active.iter_rows(min_row=2)
        if "HYPERLINK" in str(row[2].value)
    )
    assert formula_row[2].data_type == "s"
    assert formula_row[2].value.startswith("'=")
    assert formula_row[3].data_type == "s"

    template = (ROOT / "app/templates/experts/group_edit.html").read_text(encoding="utf-8")
    assert "cell.textContent = e[field] || '';" in template
    assert "<td>${e.name" not in template


def test_expert_group_routes_preserve_grouping_without_meeting_scope(expert_routes):
    client, service, audit = expert_routes

    created = client.post(
        "/experts/groups/new", data={"group_name": "材料与制造专家组"}
    )
    assert created.status_code == 302
    group_id = service.list_expert_groups()["items"][0]["groupId"]

    added = client.post(
        f"/experts/groups/edit/{group_id}",
        data={"action": "add", "expert_id": "EXP-001"},
    )
    assert added.status_code == 302
    page = client.get(f"/experts/groups/edit/{group_id}")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "张老师" in html
    assert "138****5678" in html
    assert "会议流程" not in html

    exported = client.get(f"/experts/groups/export/{group_id}")
    assert exported.status_code == 200
    assert audit.events[-1]["event_name"] == "expert_sensitive_accessed"


@pytest.mark.skipif(
    not os.environ.get("T10_TEST_DATABASE_URL"),
    reason="requires isolated PostgreSQL",
)
def test_postgresql_expert_runtime_contract():
    engine = sa.create_engine(os.environ["T10_TEST_DATABASE_URL"])
    repository = ResearchResourcesRepository(engine)
    service = ResearchResourcesService(repository, AuditRecorder())
    suffix = uuid.uuid4().hex[:10]
    username = f"t10_{suffix}"
    keyword = f"PG专家{suffix}"
    users = sa.Table("users", sa.MetaData(), autoload_with=engine)

    with engine.begin() as connection:
        owner_id = connection.execute(users.insert().values(
            username=username, password="test-only", role="BUSINESS_USER",
            name="王老师", status="active", directory_permissions={},
            must_change_password=False, version=1,
        ).returning(users.c.id)).scalar_one()
    try:
        created = service.create_expert(
            {"name": keyword, "unit": "第三研究室", "phone": "13612345678"},
            uploader="王老师", actor_user_id=owner_id, request_id="req-pg-create",
        )
        assert service.list_experts(keyword=keyword)["items"][0]["phone"] == "136****5678"

        batch = service.create_expert_import_preview(
            source_name="pg.xlsx", source_bytes=b"pg-test",
            rows=[{"_row_idx": 0, "name": f"{keyword}导入", "unit": "第三研究室"}],
            owner_user_id=owner_id,
        )
        committed = service.commit_expert_import(
            batch["batchId"], confirmed_rows=[{"_row_idx": 0, "name": f"{keyword}导入"}],
            owner_user_id=owner_id, uploader="王老师", request_id="req-pg-import",
        )
        assert committed["successCount"] == 1
        assert service.list_experts(keyword=keyword)["total"] == 2
    finally:
        with engine.begin() as connection:
            connection.execute(repository.expert_import_batches.delete().where(
                repository.expert_import_batches.c.owner_user_id == owner_id
            ))
            connection.execute(repository.experts.delete().where(
                repository.experts.c.name.like(f"{keyword}%")
            ))
            connection.execute(users.delete().where(users.c.id == owner_id))
        engine.dispose()


@pytest.mark.skipif(
    not os.environ.get("T10_TEST_DATABASE_URL"),
    reason="requires isolated PostgreSQL",
)
def test_postgresql_equipment_resource_runtime_contract():
    engine = sa.create_engine(os.environ["T10_TEST_DATABASE_URL"])
    repository = EquipmentResourcesRepository(engine)
    service = EquipmentResourcesService(repository)
    suffix = uuid.uuid4().hex[:10]
    equipment_name = f"PG设备{suffix}"
    unit_name = f"PG研制单位{suffix}"
    try:
        equipment = service.create_equipment({
            "name": equipment_name, "category": "密码设备", "price": "99.50"
        })
        unit = service.create_research_unit(unit_name, f"PG别名{suffix}")
        subclass = service.create_subclass("密码设备", f"PG子类{suffix}")
        group = service.get_or_create_project_group(
            f"PG-PROJECT-{suffix}", f"PG课题{suffix}", "张老师"
        )
        service.add_group_member(
            group["group_id"], equipment["equipment_id"], quantity=2,
            location="PG实验室", selected_by="张老师",
        )
        host = service.create_host_device(
            {"name": f"PG宿主{suffix}", "category": "计算存储"},
            relations=[{"device_id": equipment["equipment_id"], "quantity": 1}],
        )
        page = service.list_equipment(page=1, page_size=1, keyword=equipment_name)
        assert page["total"] == 1
        assert page["items"][0]["equipment_id"] == equipment["equipment_id"]
        assert any(
            row["unit_id"] == unit["unit_id"] for row in service.list_research_units()
        )
        assert any(
            row["id"] == subclass["id"] for row in service.list_subclasses("密码设备")
        )
        assert service.get_equipment_group(group["group_id"])["members"][0]["location"] == "PG实验室"
        assert service.get_devices_by_host(host["host_id"])[0]["quantity"] == 1
    finally:
        cleanup = sa.create_engine(os.environ.get(
            "MIGRATION_DATABASE_URL", os.environ["T10_TEST_DATABASE_URL"]
        ))
        with cleanup.begin() as connection:
            connection.execute(repository.device_host_relations.delete().where(
                repository.device_host_relations.c.host_id == host["host_id"]
            ))
            connection.execute(repository.host_devices.delete().where(
                repository.host_devices.c.host_id == host["host_id"]
            ))
            connection.execute(repository.equipment_group_members.delete().where(
                repository.equipment_group_members.c.group_id == group["group_id"]
            ))
            connection.execute(repository.equipment_groups.delete().where(
                repository.equipment_groups.c.group_id == group["group_id"]
            ))
            connection.execute(repository.knowledge_subclasses.delete().where(
                repository.knowledge_subclasses.c.subclass_name == f"PG子类{suffix}"
            ))
            connection.execute(repository.research_units.delete().where(
                repository.research_units.c.name == unit_name
            ))
            connection.execute(repository.equipment.delete().where(
                repository.equipment.c.name == equipment_name
            ))
        cleanup.dispose()
        engine.dispose()


@pytest.mark.skipif(
    not os.environ.get("T10_TEST_DATABASE_URL"),
    reason="requires isolated PostgreSQL",
)
def test_postgresql_expert_runtime_contract_reference_library_uses_runtime_metadata(tmp_path):
    """The T10 runner must exercise retained files outside doc_templates."""
    from app.repositories.files import FilesRepository
    from app.repositories.audit import AuditRepository
    from app.repositories.reference_library import ReferenceLibraryRepository
    from app.services.files import FileService
    from app.services.audit import AuditService
    from app.services.reference_library import ReferenceLibraryService

    engine = sa.create_engine(os.environ["T10_TEST_DATABASE_URL"])
    suffix = uuid.uuid4().hex[:10]
    users = sa.Table("users", sa.MetaData(), autoload_with=engine)
    doc_templates = sa.Table("doc_templates", sa.MetaData(), autoload_with=engine)
    with engine.begin() as connection:
        owner_id = connection.execute(users.insert().values(
            username=f"t10_reference_{suffix}", password="test-only", role="BUSINESS_USER",
            name="参考库测试", status="active", directory_permissions={}, must_change_password=False,
            version=1,
        ).returning(users.c.id)).scalar_one()
        before_schemas = connection.execute(sa.select(sa.func.count()).select_from(doc_templates)).scalar_one()
    try:
        audit = AuditService(AuditRepository(engine), app_version="test-v1")
        service = ReferenceLibraryService(
            ReferenceLibraryRepository(engine), audit,
            FileService(FilesRepository(engine), audit, storage_root=tmp_path / "reference-files", max_bytes=1024 * 1024, preview_max_bytes=1024),
        )
        standard = service.upload_standard(BytesIO(b"%PDF-1.4\npostgres"), "postgres.pdf", f"PG标准{suffix}", "国家标准", owner_id, f"req-pg-standard-{suffix}")
        template = service.upload_template(BytesIO(b"%PDF-1.4\npostgres-template"), "postgres-template.pdf", "财务模板", f"PG模板{suffix}.pdf", actor_user_id=owner_id, request_id=f"req-pg-template-{suffix}")
        assert service.open_standard_download(standard["docId"])["stream"].read().startswith(b"%PDF-")
        assert service.resolve_template_path(f"财务模板/PG模板{suffix}.pdf")["templateId"] == template["templateId"]
        service.create_folder(f"PG目录{suffix}", "财务模板", actor_user_id=owner_id, request_id=f"req-pg-folder-{suffix}")
        nested = service.upload_template(BytesIO(b"%PDF-1.4\npostgres-nested"), "nested.pdf", f"财务模板/PG目录{suffix}", f"嵌套{suffix}.pdf", actor_user_id=owner_id, request_id=f"req-pg-nested-{suffix}")
        service.archive_folder(f"财务模板/PG目录{suffix}", actor_user_id=owner_id, request_id=f"req-pg-archive-{suffix}")
        with pytest.raises(Exception) as archived_write:
            service.file_service.add_version(nested["file"]["fileId"], BytesIO(b"%PDF-1.4\nnew"), original_name="nested.pdf", object_type="TEMPLATE", object_id=nested["templateId"], expected_version=1, actor_user_id=owner_id, request_id=f"req-pg-write-{suffix}")
        assert getattr(archived_write.value, "code", None) == "OBJECT_READ_ONLY"
        assert any(log["operation_code"] == "ARCHIVE_FOLDER" for log in service.list_logs("templates", operation="ARCHIVE_FOLDER"))
        with engine.connect() as connection:
            assert connection.execute(sa.select(sa.func.count()).select_from(doc_templates)).scalar_one() == before_schemas
    finally:
        # The runtime role deliberately has INSERT/SELECT-only access to audit_events.
        # Use the runner's migration owner for teardown; production code never gains
        # DELETE access merely to make this test convenient.
        cleanup_engine = sa.create_engine(os.environ.get("MIGRATION_DATABASE_URL", os.environ["T10_TEST_DATABASE_URL"]))
        with cleanup_engine.begin() as connection:
            metadata = sa.MetaData()
            links = sa.Table("object_files", metadata, autoload_with=cleanup_engine)
            versions = sa.Table("stored_file_versions", metadata, autoload_with=cleanup_engine)
            files = sa.Table("stored_files", metadata, autoload_with=cleanup_engine)
            items = sa.Table("reference_template_items", metadata, autoload_with=cleanup_engine)
            folders = sa.Table("reference_template_folders", metadata, autoload_with=cleanup_engine)
            standards = sa.Table("standards", metadata, autoload_with=cleanup_engine)
            audit_events = sa.Table("audit_events", metadata, autoload_with=cleanup_engine)
            connection.execute(audit_events.delete().where(audit_events.c.actor_user_id == owner_id))
            connection.execute(links.delete().where(links.c.created_by == owner_id))
            connection.execute(versions.delete().where(versions.c.created_by == owner_id))
            connection.execute(files.delete().where(files.c.created_by == owner_id))
            connection.execute(items.delete().where(items.c.created_by == owner_id))
            connection.execute(folders.delete().where(folders.c.created_by == owner_id))
            connection.execute(standards.delete().where(standards.c.uploader == str(owner_id)))
            connection.execute(users.delete().where(users.c.id == owner_id))
        cleanup_engine.dispose()
        engine.dispose()


def test_expert_import_preview_and_commit_are_database_backed_and_atomic(expert_service):
    service, audit = expert_service
    preview = service.create_expert_import_preview(
        source_name="专家导入.xlsx",
        source_bytes=b"workbook",
        rows=[
            {"_row_idx": 0, "name": "王老师", "unit": "第三研究室", "expertise": "软件"},
            {"_row_idx": 1, "name": "陈老师", "unit": "第四研究室", "expertise": "测试"},
        ],
        owner_user_id=7,
    )

    assert preview["batchId"]
    assert preview["validCount"] == 2
    committed = service.commit_expert_import(
        preview["batchId"],
        confirmed_rows=preview["rows"],
        owner_user_id=7,
        uploader="operator",
        request_id="req-import",
    )

    assert committed["status"] == "COMMITTED"
    assert committed["successCount"] == 2
    listed = service.list_experts(page=1, page_size=20)
    assert listed["total"] == 4
    assert audit.events[-1]["event_name"] == "import_batch_completed"
    assert audit.events[-1]["properties"] == {
        "module": "experts", "valid_count": 2,
        "error_count": 0, "duplicate_count": 0,
    }


def test_expert_import_rolls_back_rows_and_batch_status_when_audit_fails(expert_engine):
    service = ResearchResourcesService(
        ResearchResourcesRepository(expert_engine), AuditRecorder(fail=True)
    )
    preview = service.create_expert_import_preview(
        source_name="专家导入.xlsx",
        source_bytes=b"workbook",
        rows=[{"_row_idx": 0, "name": "王老师", "unit": "第三研究室"}],
        owner_user_id=7,
    )

    with pytest.raises(RuntimeError, match="audit unavailable"):
        service.commit_expert_import(
            preview["batchId"], confirmed_rows=preview["rows"],
            owner_user_id=7, uploader="operator", request_id="req-import",
        )

    assert service.list_experts(page=1, page_size=20)["total"] == 2
    assert service.get_expert_import_batch(
        preview["batchId"], owner_user_id=7
    )["status"] == "PREVIEW"


def test_cancelled_expert_import_cannot_be_committed(expert_service):
    service, _audit = expert_service
    preview = service.create_expert_import_preview(
        source_name="专家导入.xlsx", source_bytes=b"workbook",
        rows=[{"_row_idx": 0, "name": "王老师"}], owner_user_id=7,
    )

    cancelled = service.cancel_expert_import(
        preview["batchId"], owner_user_id=7
    )

    assert cancelled["status"] == "CANCELLED"
    with pytest.raises(ResourceServiceError) as caught:
        service.commit_expert_import(
            preview["batchId"], confirmed_rows=preview["rows"],
            owner_user_id=7, uploader="operator", request_id="req-import",
        )
    assert caught.value.code == "IMPORT_BATCH_CLOSED"


def test_sensitive_export_is_audited_once_without_sensitive_metadata(expert_service):
    service, audit = expert_service

    rows = service.export_experts_sensitive(
        expert_ids=["EXP-001", "EXP-002"], actor_user_id=7,
        request_id="req-export",
    )

    assert [row["phone"] for row in rows] == ["13812345678", "13987654321"]
    assert audit.events == [{
        "event_name": "expert_sensitive_accessed",
        "user_id": 7,
        "object_type": "EXPERT",
        "object_id": None,
        "result": "SUCCESS",
        "request_id": "req-export",
        "duration_ms": 0,
        "properties": {"operation": "EXPORT", "record_count": 2},
    }]
