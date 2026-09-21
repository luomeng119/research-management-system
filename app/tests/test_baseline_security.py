import builtins
import logging
import sqlite3

import pytest
import sqlalchemy as sa

from app import create_app


def app_config(tmp_path, **overrides):
    config = {
        "TESTING": True,
        "SECRET_KEY": "test-only-secret",
        "DATA_DIR": str(tmp_path / "data"),
        "UPLOAD_DIR": str(tmp_path / "uploads"),
        "DOCUMENTS_DIR": str(tmp_path / "documents"),
        "SESSION_FILE_DIR": str(tmp_path / "sessions"),
    }
    config.update(overrides)
    return config


@pytest.fixture()
def app_factory(tmp_path, monkeypatch):
    def database_access_during_startup(*args, **kwargs):
        raise AssertionError("application startup must not access a database")

    real_import = builtins.__import__

    def import_without_optional_ocr(name, *args, **kwargs):
        if name == "rapidocr_onnxruntime":
            raise ImportError("optional OCR dependency is not installed")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(sqlite3, "connect", database_access_during_startup)
    monkeypatch.setattr(builtins, "__import__", import_without_optional_ocr)

    def factory(**overrides):
        return create_app(app_config(tmp_path, **overrides))

    return factory


def test_factory_starts_without_database_access_when_optional_ocr_is_missing(app_factory):
    app = app_factory()
    assert app.config["TESTING"] is True


def test_factory_rejects_production_without_flask_secret_key(tmp_path, monkeypatch):
    monkeypatch.delenv("FLASK_SECRET_KEY", raising=False)

    with pytest.raises(RuntimeError, match="FLASK_SECRET_KEY"):
        create_app(app_config(tmp_path, TESTING=False, SECRET_KEY=None))


def test_factory_uses_explicit_stable_flask_secret_key_in_production(tmp_path, monkeypatch):
    deployment_key = "stable-deployment-secret-key"
    monkeypatch.setenv("FLASK_SECRET_KEY", deployment_key)
    engine = sa.create_engine("sqlite+pysqlite:///:memory:")

    config = app_config(
        tmp_path,
        TESTING=False,
        SECRET_KEY=None,
        DATABASE_ENGINE=engine,
        SECURITY_AUTH_ENABLED=False,
    )
    first = create_app(config)
    second = create_app(config)

    assert first.config["SECRET_KEY"] == deployment_key
    assert second.config["SECRET_KEY"] == deployment_key


def test_factory_does_not_create_a_default_user(app_factory, monkeypatch):
    from app.models import UserModel

    created_users = []
    monkeypatch.setattr(UserModel, "__init__", lambda self: None)
    monkeypatch.setattr(UserModel, "get_by_username", lambda self, username: None)
    monkeypatch.setattr(
        UserModel,
        "add",
        lambda self, username, *args, **kwargs: created_users.append(username),
    )

    app_factory()

    assert created_users == []


def test_only_login_intro_and_health_are_public(app_factory):
    app = app_factory()
    client = app.test_client()

    login_response = client.get("/auth/login")
    intro_response = client.get("/system-intro")
    health_response = client.get("/healthz")

    assert login_response.status_code == 200
    assert intro_response.status_code == 200
    assert 'data-intro-section="overview"' in intro_response.text
    assert 'data-intro-section="workflow"' in intro_response.text
    assert 'data-intro-section="features"' in intro_response.text
    assert 'data-intro-section="data-flow"' in intro_response.text
    assert 'data-intro-section="local-ai"' not in intro_response.text
    assert 'data-intro-section="architecture"' in intro_response.text
    assert 'data-intro-section="delivery"' in intro_response.text
    assert 'href="/auth/login"' in intro_response.text
    assert "http://127.0.0.1:8893/auth/login" in intro_response.text
    assert "https://" not in intro_response.text
    assert health_response.status_code == 200
    assert health_response.get_json() == {"status": "ok"}

    with client.session_transaction() as authenticated_session:
        authenticated_session["user"] = "admin"
    authenticated_intro = client.get("/system-intro")
    assert authenticated_intro.status_code == 200
    assert 'data-intro-section="overview"' in authenticated_intro.text
    assert 'class="app-shell"' not in authenticated_intro.text
    assert 'href="/">返回系统</a>' in authenticated_intro.text
    with client.session_transaction() as authenticated_session:
        authenticated_session.clear()

    protected_urls = (
        "/auth/register",
        "/projects/",
        "/security_projects/",
        "/crypto_projects/",
        "/experts/",
        "/utils/",
        "/expense/",
        "/api/tree",
        "/uploads/private.txt",
    )
    for url in protected_urls:
        response = client.get(url)
        assert response.status_code == 302, url
        assert response.headers["Location"].endswith("/auth/login"), url


def test_test_and_bare_upload_routes_are_not_registered(app_factory):
    app = app_factory()
    rules = {rule.rule for rule in app.url_map.iter_rules()}

    assert "/test/document_preview" not in rules
    assert "/test/tree" not in rules
    assert "/uploads/<path:filename>" not in rules


@pytest.mark.parametrize("path_kind", ["absolute", "relative", "encoded", "symlink"])
def test_argumentation_download_cannot_escape_document_directory(app_factory, tmp_path, path_kind):
    app = app_factory()
    private = tmp_path / "private.txt"
    private.write_text("not-a-business-document", encoding="utf-8")
    document_dir = tmp_path / "data" / "documents"
    document_dir.mkdir(parents=True)
    link = document_dir / "linked.docx"
    link.symlink_to(private)
    path = {
        "absolute": str(private),
        "relative": "../private.txt",
        "encoded": "%2e%2e%2fprivate.txt",
        "symlink": "data/documents/linked.docx",
    }[path_kind]
    client = app.test_client()
    with client.session_transaction() as session:
        session["user"] = "zhanglaoshi"
    response = client.get("/argumentation/download", query_string={"path": path})
    assert response.status_code in (400, 404)
    assert b"not-a-business-document" not in response.data
    assert str(tmp_path).encode() not in response.data


def test_argumentation_download_retains_generated_document(app_factory, tmp_path):
    app = app_factory()
    document_dir = tmp_path / "data" / "documents"
    document_dir.mkdir(parents=True)
    document = document_dir / "科研方案.docx"
    document.write_bytes(b"generated-document-test-fixture")
    client = app.test_client()
    with client.session_transaction() as session:
        session["user"] = "zhanglaoshi"
    response = client.get("/argumentation/download", query_string={"path": "data/documents/科研方案.docx"})
    assert response.status_code == 200
    assert response.data == b"generated-document-test-fixture"


def test_argumentation_generate_download_uses_safe_web_reference(app_factory, tmp_path):
    import io
    from docx import Document

    app = app_factory()
    client = app.test_client()
    with client.session_transaction() as session:
        session["user"] = "zhanglaoshi"
    response = client.post("/argumentation/generate_word", json={
        "project_name": "../科研方案", "category": "research",
        "template": [{"id": "background", "title": "研究背景"}],
        "content": {"background": "已有实验材料"},
    })
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["success"] is True
    reference = payload["word_path"]
    assert reference.startswith("data/documents/")
    assert ".." not in reference
    assert "\\" not in reference
    assert reference.count("/") == 2
    download = client.get("/argumentation/download", query_string={"path": reference})
    assert download.status_code == 200
    content = Document(io.BytesIO(download.data))
    assert content.paragraphs[0].text == "../科研方案"
    assert content.paragraphs[-1].text == "已有实验材料"
    assert len(list((tmp_path / "data" / "documents").glob("*.docx"))) == 1


def test_critical_existing_business_urls_remain_registered(app_factory):
    app = app_factory()
    rules = {rule.rule for rule in app.url_map.iter_rules()}

    assert {
        "/projects/",
        "/security_projects/",
        "/crypto_projects/",
        "/experts/",
        "/utils/",
        "/expense/",
        "/tables",
        "/argumentation/<category>",
    } <= rules


def test_login_does_not_write_credentials_to_stdout(app_factory, monkeypatch, capsys):
    app = app_factory()
    class FakeUserModel:
        def get_by_username(self, username):
            return {
                "username": username,
                "name": "测试用户",
                "role": "用户",
                "status": "active",
            }

        def verify(self, username, password):
            return False

    from app.routes import auth

    monkeypatch.setattr(auth, "UserModel", FakeUserModel)
    secret = "credential-that-must-not-be-logged"

    response = app.test_client().post(
        "/auth/login",
        data={"username": "tester", "password": secret},
    )

    assert response.status_code == 200
    captured = capsys.readouterr()
    assert secret not in captured.out
    assert secret not in captured.err


def test_successful_login_does_not_log_sensitive_data(
    app_factory, monkeypatch, capsys, caplog
):
    app = app_factory()
    caplog.set_level(logging.DEBUG)

    class FakeUserModel:
        def get_by_username(self, username):
            return {
                "username": username,
                "name": "sensitive-display-name",
                "role": "用户",
                "status": "active",
            }

        def verify(self, username, password):
            return True

    from app.routes import auth

    monkeypatch.setattr(auth, "UserModel", FakeUserModel)
    username = "sensitive-success-user"
    password = "sensitive-success-password"

    response = app.test_client().post(
        "/auth/login",
        data={"username": username, "password": password},
    )

    assert response.status_code == 302
    captured = capsys.readouterr()
    logged_text = caplog.text
    for sensitive_value in (username, password, "sensitive-display-name"):
        assert sensitive_value not in captured.out
        assert sensitive_value not in captured.err
        assert sensitive_value not in logged_text
