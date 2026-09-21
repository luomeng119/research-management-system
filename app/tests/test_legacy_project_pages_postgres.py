"""Retained project pages against PostgreSQL, real authentication and services."""

import json
import hashlib
import io
import os
from pathlib import Path
import re
import uuid
import zipfile
import time
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event, get_ident
import subprocess
from html.parser import HTMLParser

import pytest
import sqlalchemy as sa

from app import create_app
from app.security.auth import hash_password
from app.services.files import FileServiceError


@pytest.mark.parametrize("path,expected", [
    (".history/input_v1.txt", True),
    (".history/input_versions.json", True),
    (".history/unrelated_v1.txt", False),
    (".history/input_v1.pdf", False),
])
def test_deleted_root_history_path_mapping(path, expected):
    from app.routes._project_bridge import is_deleted_project_path
    assert is_deleted_project_path(path, {"input.txt"}) is expected


@pytest.mark.parametrize("category,prefix", [
    ("GENERAL_RESEARCH", "projects"),
    ("SECURITY_CONFIDENTIALITY", "security_projects"),
    ("CRYPTO_APPLICATION", "crypto_projects"),
])
@pytest.mark.parametrize("operation", ["detail", "upload", "path_versions", "legacy_versions", "legacy_concurrent", "legacy_bad_json", "legacy_gap", "legacy_orphan", "legacy_rollback", "legacy_whitespace", "legacy_other_extension", "legacy_delete", "legacy_preview", "legacy_office", "folder_upload", "folder_delete", "folder_restore_failure", "file_rename", "file_archive", "file_delete"])
def test_retained_project_reads_and_uploads(tmp_path, category, prefix, operation, monkeypatch):
    if not os.environ.get("T02_ISOLATED_POSTGRES_ROOT"):
        pytest.skip("requires owned PostgreSQL harness")
    expected_data = Path(os.environ["T02_ISOLATED_POSTGRES_ROOT"]).resolve() / "data"
    owned_pid = (expected_data / "postmaster.pid").read_text().splitlines()
    expected_port = int(os.environ["T02_ISOLATED_POSTGRES_PORT"])
    expected_database = os.environ["T02_ISOLATED_POSTGRES_DATABASE"]
    assert Path(owned_pid[1]).resolve() == expected_data
    assert int(owned_pid[3]) == expected_port
    os.kill(int(owned_pid[0]), 0)
    url = sa.engine.make_url(os.environ["TEST_DATABASE_URL"])
    assert url.host in {"localhost", "127.0.0.1", "::1"}
    assert url.port == expected_port and url.database == expected_database
    engine = sa.create_engine(url)
    username = "pages_" + uuid.uuid4().hex
    try:
        with engine.connect() as connection:
            database, port, address = connection.execute(sa.text(
                "SELECT current_database(), inet_server_port(), host(inet_server_addr())"
            )).one()
            assert database == expected_database and port == expected_port
            assert address in {"127.0.0.1", "::1"}
        users = sa.Table("users", sa.MetaData(), autoload_with=engine)
        with engine.begin() as connection:
            user_id = connection.execute(users.insert().values(
                username=username, password=hash_password("TestPassword123!"),
                name="王老师", role="BUSINESS_USER", status="active",
                must_change_password=False, version=1,
            ).returning(users.c.id)).scalar_one()
        app = create_app({
            "TESTING": True, "SECRET_KEY": "retained-pages-test-only",
            "DATABASE_ENGINE": engine,
            "DATA_DIR": str(tmp_path / "data"),
            "UPLOAD_DIR": str(tmp_path / "uploads"),
            "DOCUMENTS_DIR": str(tmp_path / "documents"),
            "SESSION_FILE_DIR": str(tmp_path / "sessions"),
            "SECURITY_AUTH_ENABLED": True, "CSRF_ENABLED": True,
        })
        project = app.extensions["project_service"].create_standalone(
            category, {"name": "演练-三类旧项目详情", "leader": "王老师"},
            actor_user_id=user_id,
        )
        client = app.test_client()
        login = client.get("/auth/login")
        token = re.search(r'name="_csrf_token" value="([^"]+)"', login.text)[1]
        assert client.post("/auth/login", data={
            "username": username, "password": "TestPassword123!", "_csrf_token": token,
        }).status_code == 302
        response = client.get(f"/{prefix}/detail/{project['businessId']}")
        assert response.status_code == 200
        assert "演练-三类旧项目详情" in response.text
        assert project["businessId"] in response.text
        folders = json.loads(re.search(r"var folderData = (.*);", response.text)[1])
        assert [folder["name"] for folder in folders] == [
            "任务输入文件", "研究成果文件", "院内审查文件", "机关审查文件", "成果上报文件",
        ]
        stored = app.extensions["project_service"].get(project["id"])
        assert stored["category"] == category
        assert stored["businessId"] == project["businessId"]
        if operation == "legacy_office":
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            fixtures = Path(__file__).parent / "fixtures/legacy"
            for filename in ("office-sample.doc", "office-source.xls", "office-source.ppt"):
                data = (fixtures / filename).read_bytes()
                # Same-length text edits preserve CFB offsets and Office record
                # lengths while making the two actual version payloads distinct.
                before, after = ((b"Equipment A", b"Equipment B") if filename.endswith(".xls")
                                 else ("Legacy".encode("utf-16-le"), "Review".encode("utf-16-le")))
                assert data.count(before) == 1
                versions = {1: data, 2: data.replace(before, after)}
                assert versions[1] != versions[2]
                file_id = None
                for version in (1, 2):
                    upload = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                        "folder": "任务输入文件", "file": (io.BytesIO(versions[version]), filename),
                    }, headers={"X-CSRF-Token": csrf})
                    assert upload.status_code == 200 and upload.json["success"] is True
                    assert upload.json["versionNo"] == version
                    if file_id is None:
                        file_id = upload.json["fileId"]
                    assert upload.json["fileId"] == file_id
                path = "任务输入文件/" + filename
                download = client.get(f"/{prefix}/download/{project['businessId']}/{path}", follow_redirects=True)
                assert download.status_code == 200 and download.data == versions[2]
                assert "attachment" in download.headers["Content-Disposition"]
                assert download.headers["X-Content-Type-Options"] == "nosniff"
                for version in (1, 2):
                    previous = client.get(f"/api/files/{file_id}/versions/{version}/download",
                                          query_string={"objectType": "PROJECT", "objectId": project["businessId"]})
                    assert previous.status_code == 200 and previous.data == versions[version]
                preview = client.get("/preview/project", query_string={
                    "projectId": project["businessId"], "category": category, "filepath": path,
                }, follow_redirects=True)
                assert preview.status_code == 409
                assert preview.json["error"]["code"] == "PREVIEW_UNAVAILABLE"
                assert client.get(preview.json["downloadUrl"]).data == versions[2]
                assert app.test_client().get(preview.json["downloadUrl"]).status_code == 401
            folder_upload = client.post(f"/{prefix}/upload_folder/{project['businessId']}", data={
                "files": [(io.BytesIO((fixtures / filename).read_bytes()), "旧Office/" + filename)
                          for filename in ("office-sample.doc", "office-source.xls", "office-source.ppt")],
            }, headers={"X-CSRF-Token": csrf})
            assert folder_upload.status_code == 200 and folder_upload.json["success"] is True
            assert folder_upload.json["uploadedCount"] == 3
            for filename in ("office-sample.doc", "office-source.xls", "office-source.ppt"):
                downloaded = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/旧Office/{filename}",
                                        follow_redirects=True)
                assert downloaded.status_code == 200 and downloaded.data == (fixtures / filename).read_bytes()
            return
        if operation == "legacy_preview":
            audits = sa.Table("audit_events", sa.MetaData(), autoload_with=engine)
            directory = tmp_path / "uploads" / project["businessId"] / "任务输入文件"
            directory.mkdir(parents=True, exist_ok=True)
            (directory / "bad-text.txt").write_bytes(b"\xff\xfe invalid UTF-8")
            reference = {"projectId": project["businessId"], "category": category,
                         "filepath": "任务输入文件/bad-text.txt"}
            anonymous = app.test_client().get("/preview/project", query_string=reference)
            assert anonymous.status_code == 302
            assert anonymous.headers["Location"].endswith("/auth/login")
            upload_root = app.config["UPLOAD_DIR"]
            root_alias = tmp_path / "upload-alias"
            root_alias.symlink_to(upload_root, target_is_directory=True)
            service = app.extensions["file_service"]
            maximum = service.preview_max_bytes
            for case, status, error_code in (("decode", 500, "FILE_OPERATION_FAILED"),
                                             ("root_link", 404, "FILE_NOT_FOUND"),
                                             ("oversize", 409, "PREVIEW_UNAVAILABLE")):
                with engine.connect() as connection:
                    before_ids = set(connection.execute(sa.select(audits.c.id)).scalars())
                try:
                    if case == "root_link":
                        app.config["UPLOAD_DIR"] = str(root_alias)
                    elif case == "oversize":
                        service.preview_max_bytes = 1
                    failed = client.get("/preview/project", query_string=reference)
                finally:
                    app.config["UPLOAD_DIR"] = upload_root
                    service.preview_max_bytes = maximum
                assert failed.status_code == status
                with engine.connect() as connection:
                    events = [row for row in connection.execute(sa.select(audits)).mappings()
                              if row["id"] not in before_ids and row["action"] == "file_operation_completed"]
                assert len(events) == 1
                assert (events[0]["object_type"], events[0]["object_id"]) == ("PROJECT", project["businessId"])
                assert events[0]["result"] == "FAILURE" and events[0]["metadata"]["error_code"] == error_code
                assert str(events[0]["actor_user_id"]) == str(user_id)
                if case in {"root_link", "oversize"}:
                    assert "size_bucket" not in events[0]["metadata"]
                    assert "file_type" not in events[0]["metadata"]
            assert (directory / "bad-text.txt").read_bytes() == b"\xff\xfe invalid UTF-8"
            directory = tmp_path / "uploads" / project["businessId"] / "任务输入文件"
            directory.mkdir(parents=True, exist_ok=True)
            original = directory / "旧材料.txt"
            payload = "张老师整理的已有研究依据\n<script>不执行</script>".encode()
            original.write_bytes(payload)
            before = original.stat()
            # Generic filesystem references remain prohibited; the replacement is project-bound.
            assert client.get("/preview/file", query_string={
                "path": str(original), "type": "project",
            }).status_code == 400
            preview = client.get("/preview/project", query_string={
                "projectId": project["businessId"], "category": category,
                "filepath": "任务输入文件/旧材料.txt",
            })
            assert preview.status_code == 200, preview.text
            assert preview.json["success"] is True
            assert preview.json["type"] == "text"
            assert preview.json["content"] == payload.decode()
            assert original.read_bytes() == payload
            assert original.stat().st_mtime_ns == before.st_mtime_ns
            assert app.extensions["file_service"].list_project_paths(project["businessId"]) == []
            reference = {"projectId": project["businessId"], "category": category,
                         "filepath": "任务输入文件/旧材料.txt"}
            for bad in ("../旧材料.txt", "/etc/passwd", "任务输入文件\\旧材料.txt",
                        "任务输入文件/.history/旧材料.txt"):
                assert client.get("/preview/project", query_string={**reference, "filepath": bad}).status_code == 400
            assert client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/%2e%2e/旧材料.txt"}).status_code == 404
            other = next(value for value in ("GENERAL_RESEARCH", "SECURITY_CONFIDENTIALITY", "CRYPTO_APPLICATION") if value != category)
            assert client.get("/preview/project", query_string={**reference, "category": other}).status_code == 404
            assert client.get("/preview/project", query_string={**reference, "projectId": "missing-project"}).status_code == 404
            link = directory / "linked.txt"
            link.symlink_to(original)
            assert client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/linked.txt"}).status_code == 404
            assert client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件"}).status_code == 404
            assert original.read_bytes() == payload
            from docx import Document
            from openpyxl import Workbook
            from PIL import Image
            document = Document()
            document.add_paragraph("已有论证材料")
            document.save(directory / "材料.docx")
            workbook = Workbook()
            workbook.active.append(["设备", "数量"])
            workbook.active.append(["设备甲", 2])
            workbook.save(directory / "材料.xlsx")
            workbook.close()
            Image.new("RGB", (2, 2), "white").save(directory / "图片.png")
            (directory / "100%进度.txt").write_text("50%", encoding="utf-8")
            (directory / "损坏.docx").write_bytes(b"not an Office archive")
            # Build a valid, empty one-page PDF with an independently defined object graph.
            pdf = bytearray(b"%PDF-1.4\n")
            offsets = [0]
            for number, body in enumerate((b"<< /Type /Catalog /Pages 2 0 R >>",
                                            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                                            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] >>"), 1):
                offsets.append(len(pdf))
                pdf.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
            xref = len(pdf)
            pdf.extend(b"xref\n0 4\n0000000000 65535 f \n")
            for offset in offsets[1:]:
                pdf.extend(f"{offset:010d} 00000 n \n".encode())
            pdf.extend(f"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
            (directory / "材料.pdf").write_bytes(pdf)
            paths = [path for path in directory.iterdir() if path.is_file() and not path.is_symlink()]
            snapshot = {path.name: (path.read_bytes(), path.stat().st_mtime_ns) for path in paths}
            for filename, expected in (("材料.docx", "word"), ("材料.xlsx", "excel"), ("100%进度.txt", "text")):
                result = client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/" + filename})
                assert result.status_code == 200, result.text
                assert result.json["type"] == expected
                if expected == "word":
                    assert result.json["paragraphs"] == ["已有论证材料"]
                elif expected == "excel":
                    rows = result.json["sheets"][0]["rows"]
                    assert [row[:2] for row in rows[:2]] == [["设备", "数量"], ["设备甲", "2"]]
                    assert all(not cell for row in rows[:2] for cell in row[2:])
                    assert all(not cell for row in rows[2:] for cell in row)
                else:
                    assert result.json["content"] == "50%"
            for filename, media in (("图片.png", "image/png"), ("材料.pdf", "application/pdf")):
                result = client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/" + filename})
                assert result.status_code == 200
                assert result.mimetype == media
                assert result.data == snapshot[filename][0]
                assert result.headers["X-Content-Type-Options"] == "nosniff"
                assert "sandbox" in result.headers["Content-Security-Policy"]
            result = client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/损坏.docx"})
            assert result.status_code == 409
            assert result.json["error"]["code"] == "PREVIEW_UNAVAILABLE"
            downloaded = client.get(result.json["downloadUrl"])
            assert downloaded.status_code == 200 and downloaded.data == b"not an Office archive"
            assert {path.name: (path.read_bytes(), path.stat().st_mtime_ns) for path in paths} == snapshot
            # Simulate a migrated controlled version while the read-only source still exists.
            files_service = app.extensions["file_service"]
            controlled = files_service.upload_project_path(
                io.BytesIO(b"controlled latest"), original_name="旧材料.txt", folder="任务输入文件",
                project_id=project["businessId"], actor_user_id=user_id, request_id="preview-priority-fixture",
            )
            redirected = client.get("/preview/project", query_string=reference)
            assert redirected.status_code == 302
            assert f"/api/files/{controlled['fileId']}/versions/1/preview?" in redirected.location
            current = client.get(redirected.location)
            assert current.status_code == 200 and current.json["content"] == "controlled latest"
            files_service.delete_project_path(project["businessId"], "任务输入文件/旧材料.txt",
                                              expected_file_id=controlled["fileId"], actor_user_id=user_id,
                                              request_id="preview-tombstone-fixture")
            assert client.get("/preview/project", query_string=reference).status_code == 404
            assert original.read_bytes() == payload
            # Intermediate links are refused even when pointing at an in-root directory.
            (directory / "linked-dir").symlink_to(directory, target_is_directory=True)
            assert client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/linked-dir/材料.pdf"}).status_code == 404
            os.mkfifo(directory / "pipe.txt")
            assert client.get("/preview/project", query_string={**reference, "filepath": "任务输入文件/pipe.txt"}).status_code == 404
            return
        if operation == "legacy_delete":
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            directory = tmp_path / "uploads" / project["businessId"] / "任务输入文件"
            directory.mkdir(parents=True, exist_ok=True)
            original = directory / "old.txt"
            original_bytes = b"legacy-current-must-remain-recoverable"
            original.write_bytes(original_bytes)
            history = directory / ".history"
            history.mkdir()
            version = history / "old_v1.txt"
            version.write_bytes(b"legacy-version-one")
            metadata = history / "old_versions.json"
            records = {"old.txt": [{"version": 1, "original_name": "old.txt"}],
                       "old.pdf": [{"version": 1, "original_name": "old.pdf"}]}
            metadata.write_text(json.dumps(records), encoding="utf-8")
            other_history = history / "old_v1.pdf"
            other_history.write_bytes(b"other-extension-untouched")
            expected = {str(path.relative_to(directory.parent)): path.read_bytes()
                        for path in (original, version, metadata)}
            sibling = directory / "keep.txt"
            sibling.write_bytes(b"untouched")
            real_rename = os.rename
            changed_bytes = b"concurrently-changed-current"
            def replace_before_move(source, target):
                if Path(source) == original:
                    original.write_bytes(changed_bytes)
                return real_rename(source, target)
            with monkeypatch.context() as faults:
                faults.setattr(os, "rename", replace_before_move)
                conflict = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                       data={"file_path": "任务输入文件/old.txt"},
                                       headers={"X-CSRF-Token": csrf})
                assert conflict.status_code == 409
            assert original.read_bytes() == changed_bytes
            assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json"))
            original.write_bytes(original_bytes)
            audit = app.extensions["file_service"].audit_service
            real_read_bytes = Path.read_bytes
            concurrent_records = {**records, "another.pdf": []}
            def change_metadata_after_parse(path, *args, **kwargs):
                value = real_read_bytes(path, *args, **kwargs)
                if path == metadata:
                    metadata.write_text(json.dumps(concurrent_records), encoding="utf-8")
                return value
            with monkeypatch.context() as faults:
                faults.setattr(Path, "read_bytes", change_metadata_after_parse)
                conflict = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                       data={"file_path": "任务输入文件/old.txt"},
                                       headers={"X-CSRF-Token": csrf})
                assert conflict.status_code == 409
            assert json.loads(metadata.read_text()) == concurrent_records
            assert original.read_bytes() == original_bytes
            metadata.write_bytes(expected["任务输入文件/.history/old_versions.json"])
            real_record = audit.record
            def fail_delete_audit(*args, **kwargs):
                if kwargs.get("properties", {}).get("operation") == "DELETE_LEGACY_FILE":
                    raise RuntimeError("injected legacy deletion audit failure")
                return real_record(*args, **kwargs)
            with monkeypatch.context() as faults:
                faults.setattr(audit, "record", fail_delete_audit)
                rejected = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                       data={"file_path": "任务输入文件/old.txt"},
                                       headers={"X-CSRF-Token": csrf})
                assert rejected.status_code == 500
            for relative, content in expected.items():
                assert (directory.parent / relative).read_bytes() == content
            assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json"))
            for fail_at in (1, 2, 3):
                moves = []
                def fail_move(source, target):
                    if Path(source) in (original, version, metadata):
                        moves.append(Path(source))
                        if len(moves) == fail_at:
                            raise OSError("injected move failure")
                    return real_rename(source, target)
                with monkeypatch.context() as faults:
                    faults.setattr(os, "rename", fail_move)
                    failed = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                         data={"file_path": "任务输入文件/old.txt"},
                                         headers={"X-CSRF-Token": csrf})
                assert failed.status_code == 500
                assert len(moves) == fail_at
                for relative, content in expected.items():
                    assert (directory.parent / relative).read_bytes() == content
                assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json"))
            commit_attempts = []
            def mark_commit(connection, **kwargs):
                result = real_record(connection, **kwargs)
                if kwargs.get("properties", {}).get("operation") == "DELETE_LEGACY_FILE":
                    connection.info["fail_legacy_delete_commit"] = True
                return result
            def fail_commit(connection):
                if connection.info.pop("fail_legacy_delete_commit", False):
                    commit_attempts.append(True)
                    raise RuntimeError("injected failure before DBAPI commit")
            with monkeypatch.context() as faults:
                faults.setattr(audit, "record", mark_commit)
                sa.event.listen(engine, "commit", fail_commit)
                try:
                    failed = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                         data={"file_path": "任务输入文件/old.txt"},
                                         headers={"X-CSRF-Token": csrf})
                finally:
                    sa.event.remove(engine, "commit", fail_commit)
            assert failed.status_code == 500 and commit_attempts == [True]
            for relative, content in expected.items():
                assert (directory.parent / relative).read_bytes() == content
            assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json"))
            deleted = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                  data={"file_path": "任务输入文件/old.txt"},
                                  headers={"X-CSRF-Token": csrf})
            assert deleted.status_code == 200 and deleted.json["success"] is True
            assert not original.exists()
            assert sibling.read_bytes() == b"untouched"
            manifests = list((tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json"))
            assert len(manifests) == 1
            manifest = json.loads(manifests[0].read_text())
            assert manifest["kind"] == "LEGACY_PROJECT_FILE_DELETE"
            assert manifest["projectId"] == project["businessId"]
            assert manifest["category"] == category
            assert manifest["filepath"] == "任务输入文件/old.txt"
            assert {item["sourceRelativePath"] for item in manifest["files"]} == set(expected)
            for item in manifest["files"]:
                content = expected[item["sourceRelativePath"]]
                assert item["sha256"] == hashlib.sha256(content).hexdigest()
                assert item["sizeBytes"] == len(content)
                assert (manifests[0].parent / item["payloadRelativePath"]).read_bytes() == content
            assert not version.exists()
            assert json.loads(metadata.read_text()) == {"old.pdf": records["old.pdf"]}
            assert other_history.read_bytes() == b"other-extension-untouched"
            uploaded = client.post(f"/{prefix}/upload/{project['businessId']}",
                                   data={"folder": "任务输入文件", "file": (io.BytesIO(b"new-v1"), "old.txt")},
                                   headers={"X-CSRF-Token": csrf})
            assert uploaded.status_code == 200
            paths = app.extensions["file_service"].list_project_paths(project["businessId"])
            assert next(row for row in paths if row["path"] == "任务输入文件/old.txt")["versionNo"] == 1
            recovery_source = directory / "recovery.txt"
            recovery_source.write_bytes(b"recover-me")
            def fail_compensation(source, target):
                if Path(target) == recovery_source:
                    raise OSError("injected compensation failure")
                return real_rename(source, target)
            with monkeypatch.context() as faults:
                faults.setattr(audit, "record", fail_delete_audit)
                faults.setattr(os, "rename", fail_compensation)
                failed = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                     data={"file_path": "任务输入文件/recovery.txt"},
                                     headers={"X-CSRF-Token": csrf})
            assert failed.status_code == 500
            assert failed.json["code"] == "FILE_RECOVERY_REQUIRED"
            assert str(tmp_path) not in failed.text
            assert not recovery_source.exists()
            recovery_records = [(path, json.loads(path.read_text())) for path in
                                (tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json")]
            recovery_path, recovery_record = next((path, record) for path, record in recovery_records
                                                  if record["filepath"] == "任务输入文件/recovery.txt")
            assert recovery_record["projectId"] == project["businessId"]
            assert len(recovery_record["files"]) == 1
            retained = recovery_path.parent / recovery_record["files"][0]["payloadRelativePath"]
            assert retained.read_bytes() == b"recover-me"
            assert recovery_record["files"][0]["sha256"] == hashlib.sha256(b"recover-me").hexdigest()
            for with_history in (False, True):
                name = "single.txt" if with_history else "plain.txt"
                source = directory / name
                source.write_bytes(b"original-" + name.encode())
                originals = [source]
                if with_history:
                    prior = history / "single_v1.txt"
                    prior.write_bytes(b"single-prior")
                    record_file = history / "single_versions.json"
                    record_file.write_text(json.dumps({name: [{"version": 1, "original_name": name}]}), encoding="utf-8")
                    originals.extend([prior, record_file])
                contents = {path.relative_to(directory.parent).as_posix(): path.read_bytes() for path in originals}
                removed = client.post(f"/{prefix}/delete_file/{project['businessId']}",
                                      data={"file_path": "任务输入文件/" + name},
                                      headers={"X-CSRF-Token": csrf})
                assert removed.status_code == 200 and removed.json["success"] is True
                private = tmp_path / "data" / "legacy-folder-archive" / removed.json["recoveryId"]
                record = json.loads((private / "recovery.json").read_text())
                assert record["projectId"] == project["businessId"] and record["category"] == category
                assert record["filepath"] == "任务输入文件/" + name
                assert {item["sourceRelativePath"] for item in record["files"]} == set(contents)
                for item in record["files"]:
                    content = contents[item["sourceRelativePath"]]
                    assert (private / item["payloadRelativePath"]).read_bytes() == content
                    assert item["sha256"] == hashlib.sha256(content).hexdigest()
                    assert item["sizeBytes"] == len(content)
                assert all(not path.exists() for path in originals)
                detail = client.get(f"/{prefix}/detail/{project['businessId']}")
                assert detail.status_code == 200
                tree = json.loads(re.search(r"var folderData = (.*);", detail.text)[1])
                assert name not in json.dumps(tree)
                uploaded = client.post(f"/{prefix}/upload/{project['businessId']}",
                                       data={"folder": "任务输入文件", "file": (io.BytesIO(b"fresh-version-one"), name)},
                                       headers={"X-CSRF-Token": csrf})
                assert uploaded.status_code == 200
                paths = app.extensions["file_service"].list_project_paths(project["businessId"])
                entry = next(row for row in paths if row["path"] == "任务输入文件/" + name)
                assert entry["versionNo"] == 1
                download = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/{name}", follow_redirects=True)
                assert download.status_code == 200 and download.data == b"fresh-version-one"
                for item in record["files"]:
                    assert (private / item["payloadRelativePath"]).read_bytes() == contents[item["sourceRelativePath"]]
            return
        if operation == "detail":
            csrf_page = client.get("/users/change-password")
            token = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            folder_data = {"parent_folder": "任务输入文件", "folder_name": "归属检查", "_csrf_token": token}
            uploads = tmp_path / "uploads"
            before = sorted(str(path.relative_to(uploads)) for path in uploads.rglob("*"))
            wrong_routes = [(route, project["businessId"]) for route in
                            ("projects", "security_projects", "crypto_projects") if route != prefix]
            for route, business_id in [*wrong_routes, (prefix, "MISSING-PROJECT")]:
                rejected = client.post(f"/{route}/create_folder/{business_id}", data=folder_data)
                assert rejected.status_code == 404
                assert rejected.json["success"] is False
                assert sorted(str(path.relative_to(uploads)) for path in uploads.rglob("*")) == before
            for parent, name in [("../../outside", "new"), ("任务输入文件", "../../../outside")]:
                rejected = client.post(f"/{prefix}/create_folder/{project['businessId']}", data={
                    **folder_data, "parent_folder": parent, "folder_name": name,
                })
                assert rejected.status_code == 400 and rejected.json["success"] is False
                assert sorted(str(path.relative_to(uploads)) for path in uploads.rglob("*")) == before
                assert not (tmp_path / "outside").exists()
            created = client.post(f"/{prefix}/create_folder/{project['businessId']}", data=folder_data)
            assert created.status_code == 200 and created.json["success"] is True
            assert (uploads / project["businessId"] / "任务输入文件" / "归属检查").is_dir()
        if operation in {"folder_delete", "folder_restore_failure"}:
            delete_code = response.text.split("// 删除文件夹按钮事件", 1)[1].split("// 删除文件按钮事件", 1)[0]
            node_result = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const result = {}; const button = {dataset:{folder:'任务输入文件/资料包'},
  addEventListener:(event, fn)=>{button.click=fn;}};
const context = {document:{querySelectorAll:()=>[button]}, projectUpdateUrl:input.prefix,
  confirm:()=>true, FormData:class{append(key,value){result[key]=value;}},
  fetch:(url)=>{result.url=url;return Promise.resolve({json:()=>Promise.resolve({success:true})});},
  alert:()=>{},location:{reload:()=>{}}};
vm.createContext(context); vm.runInContext(input.code, context);
button.click.call(button,{stopPropagation(){}});
setImmediate(()=>process.stdout.write(JSON.stringify(result)));
"""], input=json.dumps({"code": delete_code, "prefix": "/" + prefix}), text=True,
                                         capture_output=True, check=True, timeout=10)
            assert json.loads(node_result.stdout) == {
                "url": f"/{prefix}/delete_folder/{project['businessId']}",
                "folder_path": "任务输入文件/资料包",
            }
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            legacy_folder = tmp_path / "uploads" / project["businessId"] / "任务输入文件" / "资料包"
            legacy_history = legacy_folder / ".history"
            legacy_history.mkdir(parents=True)
            (legacy_folder / "old.doc").write_bytes(b"retained legacy bytes, not parsed as Word")
            (legacy_history / "old_v1.doc").write_bytes(b"retained prior bytes")
            for relative, payload in [
                ("资料包/甲/input.txt", b"nested first"),
                ("资料包/乙/input.txt", b"nested second"),
                ("资料包副本/input.txt", b"unrelated sibling"),
            ]:
                uploaded = client.post(f"/{prefix}/upload_folder/{project['businessId']}", data={
                    "folder": "任务输入文件", "files": [(io.BytesIO(payload), relative)],
                }, headers={"X-CSRF-Token": csrf})
                assert uploaded.status_code == 200 and uploaded.json["success"] is True
            service = app.extensions["file_service"]
            paths_before = service.list_project_paths(project["businessId"])
            shared = next(row for row in paths_before if row["path"] == "任务输入文件/资料包/甲/input.txt")
            other_project = app.extensions["project_service"].create_standalone(
                category, {"name": "演练-目录共享文件保留", "leader": "李老师"}, actor_user_id=user_id,
            )
            with engine.begin() as connection:
                service.repository.link_object(
                    connection, link_id=uuid.uuid4(), object_type="PROJECT", object_id=other_project["businessId"],
                    file_id=shared["fileId"], actor_user_id=user_id, purpose="PROJECT_TREE:任务输入文件/shared.txt",
                )
            real_audit = service.audit_service.record
            real_rename = os.rename
            def fail_folder_audit(connection, **kwargs):
                if kwargs.get("properties", {}).get("operation") == "DELETE_FOLDER":
                    raise RuntimeError("injected folder audit failure")
                return real_audit(connection, **kwargs)
            def fail_restore(source, destination):
                if Path(source).name == "payload":
                    raise OSError("injected folder restore failure")
                return real_rename(source, destination)
            if operation == "folder_restore_failure":
                with monkeypatch.context() as fault:
                    fault.setattr(service.audit_service, "record", fail_folder_audit)
                    fault.setattr(os, "rename", fail_restore)
                    failed = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                        "folder_path": "任务输入文件/资料包",
                    }, headers={"X-CSRF-Token": csrf})
                assert failed.status_code == 500 and failed.json["code"] == "FILE_RECOVERY_REQUIRED"
                assert service.list_project_paths(project["businessId"]) == paths_before
                recovered = list((tmp_path / "data" / "legacy-folder-archive").glob("*/payload"))
                assert len(recovered) == 1 and not legacy_folder.exists()
                assert (recovered[0] / "old.doc").read_bytes() == b"retained legacy bytes, not parsed as Word"
                assert (recovered[0] / ".history" / "old_v1.doc").read_bytes() == b"retained prior bytes"
                assert (recovered[0].parent / "recovery.json").is_file()
                assert str(tmp_path) not in failed.text
                return
            with monkeypatch.context() as fault:
                fault.setattr(service.audit_service, "record", fail_folder_audit)
                failed = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                    "folder_path": "任务输入文件/资料包",
                }, headers={"X-CSRF-Token": csrf})
            assert failed.status_code == 500
            assert service.list_project_paths(project["businessId"]) == paths_before
            assert (legacy_folder / "old.doc").read_bytes() == b"retained legacy bytes, not parsed as Word"
            assert (legacy_history / "old_v1.doc").read_bytes() == b"retained prior bytes"
            assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/payload"))
            commit_attempts = []
            def mark_delete_commit(connection, **kwargs):
                result = real_audit(connection, **kwargs)
                if kwargs.get("properties", {}).get("operation") == "DELETE_FOLDER":
                    connection.info["fail_folder_commit"] = True
                return result
            def fail_delete_commit(connection):
                if connection.info.pop("fail_folder_commit", False):
                    commit_attempts.append(True)
                    raise RuntimeError("injected failure before DBAPI commit")
            with monkeypatch.context() as fault:
                fault.setattr(service.audit_service, "record", mark_delete_commit)
                sa.event.listen(engine, "commit", fail_delete_commit)
                try:
                    failed = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                        "folder_path": "任务输入文件/资料包",
                    }, headers={"X-CSRF-Token": csrf})
                finally:
                    sa.event.remove(engine, "commit", fail_delete_commit)
            assert commit_attempts == [True]
            assert failed.status_code == 500
            assert service.list_project_paths(project["businessId"]) == paths_before
            assert (legacy_folder / "old.doc").read_bytes() == b"retained legacy bytes, not parsed as Word"
            assert (legacy_history / "old_v1.doc").read_bytes() == b"retained prior bytes"
            assert not list((tmp_path / "data" / "legacy-folder-archive").glob("*/payload"))
            deleted = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                "folder_path": "任务输入文件/资料包",
            }, headers={"X-CSRF-Token": csrf})
            assert deleted.status_code == 200 and deleted.json["success"] is True
            assert not legacy_folder.exists()
            recovered = list((tmp_path / "data" / "legacy-folder-archive").glob("*/payload"))
            assert len(recovered) == 1
            assert (recovered[0] / "old.doc").read_bytes() == b"retained legacy bytes, not parsed as Word"
            assert (recovered[0] / ".history" / "old_v1.doc").read_bytes() == b"retained prior bytes"
            assert json.loads((recovered[0].parent / "recovery.json").read_text()) == {
                "projectId": project["businessId"], "folder": "任务输入文件/资料包",
            }
            for relative in ("资料包/甲/input.txt", "资料包/乙/input.txt"):
                response = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/{relative}", follow_redirects=True)
                assert response.status_code == 404
            sibling = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/资料包副本/input.txt", follow_redirects=True)
            assert sibling.status_code == 200 and sibling.data == b"unrelated sibling"
            shared_download = client.get(f"/api/files/{shared['fileId']}/versions/1/download",
                                         query_string={"objectType": "PROJECT", "objectId": other_project["businessId"]})
            assert shared_download.status_code == 200 and shared_download.data == b"nested first"
            assert [row["fileId"] for row in service.list_project_paths(other_project["businessId"])] == [shared["fileId"]]
            renewed = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                "folder": "任务输入文件/资料包/甲", "file": (io.BytesIO(b"new identity"), "input.txt"),
            }, headers={"X-CSRF-Token": csrf})
            assert renewed.status_code == 200 and renewed.json["versionNo"] == 1
            assert renewed.json["fileId"] != shared["fileId"]
            # Remove the newly recreated folder before checking the original
            # deletion view; its archive must not replace the earlier recovery.
            assert client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                "folder_path": "任务输入文件/资料包",
            }, headers={"X-CSRF-Token": csrf}).json["success"] is True
            paths = app.extensions["file_service"].list_project_paths(project["businessId"])
            assert [item["path"] for item in paths] == ["任务输入文件/资料包副本/input.txt"]
            refreshed = client.get(f"/{prefix}/detail/{project['businessId']}")
            tree = json.loads(re.search(r"var folderData = (.*);", refreshed.text)[1])
            assert {node["name"] for node in tree[0]["children"]} == {"资料包副本"}
            remaining_paths = service.list_project_paths(project["businessId"])
            project_root = tmp_path / "uploads" / project["businessId"]
            for folder_name in ("任务输入文件/空目录", "任务输入文件/仅旧目录", "研究成果文件"):
                physical = project_root / folder_name
                physical.mkdir(parents=True, exist_ok=True)
                if folder_name.endswith("仅旧目录"):
                    (physical / "legacy.txt").write_bytes(b"legacy-only folder")
                result = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                    "folder_path": folder_name,
                }, headers={"X-CSRF-Token": csrf})
                assert result.status_code == 200 and result.json["success"] is True
                assert result.json["deletedFileCount"] == 0
                assert not physical.exists()
                archives = [path for path in (tmp_path / "data" / "legacy-folder-archive").glob("*/recovery.json")
                            if json.loads(path.read_text())["folder"] == folder_name]
                assert len(archives) == 1
                if folder_name.endswith("仅旧目录"):
                    assert (archives[0].parent / "payload" / "legacy.txt").read_bytes() == b"legacy-only folder"
            refreshed = client.get(f"/{prefix}/detail/{project['businessId']}")
            tree = json.loads(re.search(r"var folderData = (.*);", refreshed.text)[1])
            assert {node["name"] for node in tree[0]["children"]} == {"资料包副本"}
            fixed = next(node for node in tree if node["name"] == "研究成果文件")
            assert fixed["files"] == [] and fixed["children"] == []
            assert service.list_project_paths(project["businessId"]) == remaining_paths
            for invalid in ("", "/", "../outside", "任务输入文件/../资料包副本", "任务输入文件//资料包副本", "任务输入文件\\资料包副本", "任务输入文件/.history"):
                response = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                    "folder_path": invalid,
                }, headers={"X-CSRF-Token": csrf})
                assert response.status_code == 400
            outside = tmp_path / "outside-folder"
            outside.mkdir()
            (outside / "keep.txt").write_bytes(b"outside must remain")
            linked = project_root / "任务输入文件" / "外部链接"
            linked.symlink_to(outside, target_is_directory=True)
            response = client.post(f"/{prefix}/delete_folder/{project['businessId']}", data={
                "folder_path": "任务输入文件/外部链接",
            }, headers={"X-CSRF-Token": csrf})
            assert response.status_code == 400 and linked.is_symlink()
            assert (outside / "keep.txt").read_bytes() == b"outside must remain"
            assert service.list_project_paths(project["businessId"]) == remaining_paths
            return
        if operation == "folder_upload":
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            response = client.post(f"/{prefix}/upload_folder/{project['businessId']}", data={
                "files": [(io.BytesIO(b"direction a"), "资料包/方向甲/input.txt"),
                          (io.BytesIO(b"direction b"), "资料包/方向乙/input.txt")],
            }, headers={"X-CSRF-Token": csrf})
            assert response.status_code == 200 and response.json["success"] is True
            service = app.extensions["file_service"]
            paths = service.list_project_paths(project["businessId"])
            assert {item["path"] for item in paths} == {
                "任务输入文件/资料包/方向甲/input.txt", "任务输入文件/资料包/方向乙/input.txt",
            }
            for path, content in [("任务输入文件/资料包/方向甲/input.txt", b"direction a"),
                                  ("任务输入文件/资料包/方向乙/input.txt", b"direction b")]:
                downloaded = client.get(f"/{prefix}/download/{project['businessId']}/{path}", follow_redirects=True)
                assert downloaded.status_code == 200 and downloaded.data == content
            refreshed = client.get(f"/{prefix}/detail/{project['businessId']}")
            tree = json.loads(re.search(r"var folderData = (.*);", refreshed.text)[1])
            package = next(node for node in tree[0]["children"] if node["name"] == "资料包")
            assert {node["name"] for node in package["children"]} == {"方向甲", "方向乙"}
            mixed = client.post(f"/{prefix}/upload_folder/{project['businessId']}", data={
                "files": [(io.BytesIO(b"updated"), "资料包/方向甲/input.txt"),
                          (io.BytesIO(b"invalid"), "../outside.txt"),
                          (io.BytesIO(b"new valid"), "资料包/方向丙/input.txt")],
            }, headers={"X-CSRF-Token": csrf})
            assert mixed.status_code == 207 and mixed.json["success"] is False
            assert mixed.json["uploadedCount"] == 2 and mixed.json["failedCount"] == 1
            assert mixed.json["errors"][0]["path"] == "../outside.txt"
            paths = service.list_project_paths(project["businessId"])
            assert len(paths) == 3
            assert next(item for item in paths if item["path"] == "任务输入文件/资料包/方向甲/input.txt")["versionNo"] == 2
            assert not (tmp_path / "outside.txt").exists()
            submit_code = re.search(
                r"document.getElementById\('uploadForm'\).onsubmit = function\(e\).*?\n};",
                refreshed.text, re.S,
            )[0]
            submitted = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const form = {}; const sent = {}; const state = {reload: false};
const elements = {uploadForm: form, folderInput: {files: [
  {name:'input.txt', webkitRelativePath:'资料包/方向甲/input.txt'},
  {name:'input.txt', webkitRelativePath:'资料包/方向乙/input.txt'}]}, uploadFolder: {value:'任务输入文件'}};
const context = {document:{getElementById:id=>elements[id]}, projectUpdateUrl:input.prefix,
  FormData:class {constructor(){this.entries=[];} append(...args){this.entries.push(args);}},
  fetch:(url, options)=>{sent.url=url; sent.entries=options.body.entries;
    return Promise.resolve({json:()=>Promise.resolve({success:false,uploadedCount:1,message:'一项成功、一项失败'})});},
  alert:message=>{state.message=message;}, location:{reload:()=>{state.reload=true;}}};
vm.createContext(context); vm.runInContext(input.code, context);
form.onsubmit.call(form,{preventDefault(){}});
setImmediate(()=>process.stdout.write(JSON.stringify({sent,state})));
"""], input=json.dumps({"code": submit_code, "prefix": "/" + prefix}), text=True,
                                      capture_output=True, check=True)
            browser_contract = json.loads(submitted.stdout)
            assert browser_contract["sent"]["url"] == f"/{prefix}/upload_folder/{project['businessId']}"
            assert [entry[2] for entry in browser_contract["sent"]["entries"] if entry[0] == "files"] == [
                "资料包/方向甲/input.txt", "资料包/方向乙/input.txt",
            ]
            assert browser_contract["state"] == {"reload": True, "message": "一项成功、一项失败"}
            real_upload = service.upload_project_path
            for failure_type in (OSError, sa.exc.SQLAlchemyError):
                def faulty_upload(stream, *, original_name, **kwargs):
                    if original_name == "fault.txt":
                        raise failure_type("private diagnostic sentinel")
                    return real_upload(stream, original_name=original_name, **kwargs)
                # Inject a lower-level failure; successful neighbors still use real PG and storage.
                with monkeypatch.context() as fault:
                    fault.setattr(service, "upload_project_path", faulty_upload)
                    result = client.post(f"/{prefix}/upload_folder/{project['businessId']}", data={
                        "files": [(io.BytesIO(b"before"), "failure-test/before.txt"),
                                  (io.BytesIO(b"fault"), "failure-test/fault.txt"),
                                  (io.BytesIO(b"after"), "failure-test/after.txt")],
                    }, headers={"X-CSRF-Token": csrf})
                assert result.status_code == 207
                assert result.json["uploadedCount"] == 2 and result.json["failedCount"] == 1
                assert result.json["errors"][0]["code"] == "FILE_OPERATION_FAILED"
                assert "private diagnostic sentinel" not in result.text
        if operation in {"legacy_versions", "legacy_concurrent", "legacy_bad_json", "legacy_gap", "legacy_orphan", "legacy_rollback", "legacy_whitespace", "legacy_other_extension"}:
            # Original uploader moves each predecessor into .history, leaving
            # the latest physical file unnumbered. A new upload must adopt all
            # three predecessors, not reset the controlled identity to v1.
            folder = tmp_path / "uploads" / project["businessId"] / "任务输入文件"
            history = folder / ".history"
            history.mkdir(parents=True)
            originals = {
                history / "input_v1.txt": b"legacy first",
                history / "input_v2.txt": b"legacy second",
                folder / "input.txt": b"legacy current",
                history / "input_versions.json": json.dumps({"input.txt": [
                    {"version": 1, "upload_time": "2026-08-01 10:00:00",
                     "uploader": "张老师", "original_name": "input.txt"},
                    {"version": 2, "upload_time": "2026-08-02 10:00:00",
                     "uploader": "李老师", "original_name": "input.txt"},
                ]}, ensure_ascii=False).encode("utf-8"),
            }
            if operation == "legacy_other_extension":
                # Same stem is not the same file. Its metadata must not block
                # a genuinely new extension (PDF bytes are never adopted here).
                originals = {
                    path.with_suffix(".pdf") if path.suffix == ".txt" else path:
                    content.replace(b"input.txt", b"input.pdf") if path.suffix == ".json" else content
                    for path, content in originals.items()
                }
            if operation == "legacy_bad_json":
                originals[history / "input_versions.json"] = b"{broken"
            elif operation == "legacy_gap":
                del originals[history / "input_v1.txt"]
            elif operation == "legacy_orphan":
                del originals[folder / "input.txt"]
            for path, content in originals.items():
                path.write_bytes(content)
            original_mtimes = {path: path.stat().st_mtime_ns for path in originals}
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            if operation in {"legacy_bad_json", "legacy_gap", "legacy_orphan"}:
                service = app.extensions["file_service"]
                with engine.connect() as connection:
                    counts_before = [connection.scalar(sa.text(f"SELECT count(*) FROM {table}"))
                                     for table in ("stored_files", "stored_file_versions", "object_files")]
                failed = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                    "folder": "任务输入文件", "file": (io.BytesIO(b"must not replace"), "input.txt"),
                }, headers={"X-CSRF-Token": csrf})
                assert failed.status_code == 409 and failed.json["code"] == "LEGACY_FILE_CONFLICT"
                with engine.connect() as connection:
                    assert [connection.scalar(sa.text(f"SELECT count(*) FROM {table}"))
                            for table in ("stored_files", "stored_file_versions", "object_files")] == counts_before
                assert not [path for path in service.storage_root.rglob("*") if path.is_file()]
                assert {path for path in folder.rglob("*") if path.is_file()} == set(originals)
                for path, content in originals.items():
                    assert path.read_bytes() == content
                    assert path.stat().st_mtime_ns == original_mtimes[path]
                return
            if operation == "legacy_concurrent":
                from app.routes._project_bridge import legacy_project_sources
                service = app.extensions["file_service"]
                ready, release = Event(), Event()
                first_pid = []
                marker = "adoption-second-" + uuid.uuid4().hex
                real_audit = service.audit_service.record
                def pause_adoption(connection, **kwargs):
                    result = real_audit(connection, **kwargs)
                    if kwargs.get("properties", {}).get("operation") == "ADOPT_LEGACY":
                        first_pid.append(connection.scalar(sa.text("select pg_backend_pid()")))
                        ready.set()
                        assert release.wait(10), "test must release adopting transaction"
                    return result
                def write_legacy(payload, *, marked=False):
                    worker = get_ident()
                    def mark_transaction(connection):
                        if get_ident() == worker:
                            connection.execute(sa.text("select set_config('application_name', :name, true)"),
                                               {"name": marker})
                    if marked:
                        sa.event.listen(engine, "begin", mark_transaction)
                    try:
                        with app.app_context():
                            return service.upload_project_path(
                                io.BytesIO(payload), original_name="input.txt", folder="任务输入文件",
                                project_id=project["businessId"], actor_user_id=user_id,
                                request_id="legacy-concurrent",
                                legacy_sources=lambda name: legacy_project_sources(project["businessId"], "任务输入文件", name),
                            )
                    finally:
                        if marked:
                            sa.event.remove(engine, "begin", mark_transaction)
                with monkeypatch.context() as scheduling:
                    scheduling.setattr(service.audit_service, "record", pause_adoption)
                    with ThreadPoolExecutor(max_workers=2) as pool:
                        first = pool.submit(write_legacy, b"concurrent fourth")
                        try:
                            assert ready.wait(5), "first writer must reach real audit before commit"
                            second = pool.submit(write_legacy, b"concurrent fifth", marked=True)
                            deadline = time.monotonic() + 5
                            blocked = False
                            with engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
                                while time.monotonic() < deadline and not second.done():
                                    blocked = observer.scalar(sa.text(
                                        "select exists (select 1 from pg_stat_activity where application_name = :name "
                                        "and :pid = any(pg_blocking_pids(pid)))"
                                    ), {"name": marker, "pid": first_pid[0]})
                                    if blocked:
                                        break
                                    time.sleep(0.01)
                            assert blocked, "second adoption must wait for the first actual transaction"
                        finally:
                            release.set()
                        results = [first.result(timeout=5), second.result(timeout=5)]
                assert [result["versionNo"] for result in results] == [4, 5]
                assert results[0]["fileId"] == results[1]["fileId"]
                assert len(first_pid) == 1
                assert len(service.list_project_paths(project["businessId"])) == 1
                for version, expected in enumerate([
                    b"legacy first", b"legacy second", b"legacy current", b"concurrent fourth", b"concurrent fifth",
                ], start=1):
                    download = client.get(f"/api/files/{results[0]['fileId']}/versions/{version}/download",
                                          query_string={"objectType": "PROJECT", "objectId": project["businessId"]})
                    assert download.status_code == 200 and download.data == expected
                for path, content in originals.items():
                    assert path.read_bytes() == content
                    assert path.stat().st_mtime_ns == original_mtimes[path]
                return
            if operation == "legacy_rollback":
                service = app.extensions["file_service"]
                before_files = {path for path in service.storage_root.rglob("*") if path.is_file()}
                def database_counts():
                    with engine.connect() as connection:
                        return [connection.scalar(sa.text(f"SELECT count(*) FROM {table}"))
                                for table in ("stored_files", "stored_file_versions", "object_files")]
                before_counts = database_counts()
                real_replace = os.replace
                real_audit = service.audit_service.record
                for fail_at in (2, 3, 4, "audit"):
                    moves = []
                    def failing_replace(source, target):
                        if Path(source).parent == service.storage_root / ".staging":
                            moves.append(source)
                            if len(moves) == fail_at:
                                raise OSError("injected adoption move failure")
                        return real_replace(source, target)
                    def failing_audit(connection, **kwargs):
                        if fail_at == "audit" and kwargs.get("properties", {}).get("operation") == "ADOPT_LEGACY":
                            raise RuntimeError("injected adoption audit failure")
                        return real_audit(connection, **kwargs)
                    with monkeypatch.context() as fault:
                        fault.setattr(os, "replace", failing_replace)
                        fault.setattr(service.audit_service, "record", failing_audit)
                        failed = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                            "folder": "任务输入文件",
                            "file": (io.BytesIO(b"new fourth"), "input.txt"),
                        }, headers={"X-CSRF-Token": csrf})
                    assert failed.status_code == 500
                    assert database_counts() == before_counts
                    assert {path for path in service.storage_root.rglob("*") if path.is_file()} == before_files
                    for path, content in originals.items():
                        assert path.read_bytes() == content
                        assert path.stat().st_mtime_ns == original_mtimes[path]
            uploaded = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                "folder": "任务输入文件",
                "file": (io.BytesIO(b"new fourth"), " input.txt " if operation == "legacy_whitespace" else "input.txt"),
            }, headers={"X-CSRF-Token": csrf})
            assert uploaded.status_code == 200
            if operation == "legacy_other_extension":
                assert uploaded.json["versionNo"] == 1
                downloaded = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/input.txt", follow_redirects=True)
                assert downloaded.status_code == 200 and downloaded.data == b"new fourth"
                for path, content in originals.items():
                    assert path.read_bytes() == content
                    assert path.stat().st_mtime_ns == original_mtimes[path]
                return
            assert uploaded.json["versionNo"] == 4
            file_id = uploaded.json["fileId"]
            for version, content in enumerate([
                b"legacy first", b"legacy second", b"legacy current", b"new fourth",
            ], start=1):
                downloaded = client.get(
                    f"/api/files/{file_id}/versions/{version}/download",
                    query_string={"objectType": "PROJECT", "objectId": project["businessId"]},
                )
                assert downloaded.status_code == 200
                assert downloaded.data == content
            for path, content in originals.items():
                assert path.read_bytes() == content
                assert path.stat().st_mtime_ns == original_mtimes[path]
        if operation == "path_versions":
            service = app.extensions["file_service"]
            barrier = Barrier(2)
            def write_path(payload):
                barrier.wait(timeout=10)
                return service.upload_project_path(
                    io.BytesIO(payload), original_name="input.txt", folder="任务输入文件",
                    project_id=project["businessId"], actor_user_id=user_id,
                    request_id="project-path-concurrent",
                )
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(write_path, [b"first upload", b"second upload"]))
            assert results[0]["fileId"] == results[1]["fileId"]
            assert sorted(result["versionNo"] for result in results) == [1, 2]
            assert {(service.storage_root / result["storagePath"]).read_bytes()
                    for result in results} == {b"first upload", b"second upload"}
            assert len(service.list_for_object(
                object_type="PROJECT", object_id=project["businessId"],
            )) == 1
        if operation in {"upload", "file_rename", "file_archive", "file_delete"}:
            csrf_page = client.get("/users/change-password")
            csrf = re.search(r'name="_csrf_token" value="([^"]+)"', csrf_page.text)[1]
            uploaded = client.post(f"/{prefix}/upload/{project['businessId']}",
                                   data={"folder": "任务输入文件",
                                         "file": (io.BytesIO(b"research input"), "input.txt")},
                                   headers={"X-CSRF-Token": csrf})
            assert uploaded.status_code == 200
            assert uploaded.json["success"] is True
            files = app.extensions["file_service"].list_for_object(
                object_type="PROJECT", object_id=project["businessId"],
            )
            assert len(files) == 1
            assert files[0]["originalName"] == "input.txt"
            assert files[0]["versionNo"] == 1
            refreshed = client.get(f"/{prefix}/detail/{project['businessId']}")
            tree = json.loads(re.search(r"var folderData = (.*);", refreshed.text)[1])
            assert tree[0]["files"][0]["path"] == "任务输入文件/input.txt"
            downloaded = client.get(
                f"/{prefix}/download/{project['businessId']}/任务输入文件/input.txt",
                follow_redirects=True,
            )
            assert downloaded.status_code == 200
            assert downloaded.data == b"research input"
            render_code = re.search(r"function buildNode\(node\) \{.*?\n\}\n", refreshed.text, re.S)[0]
            escape_code = re.search(r"function escapeEq\(value\) \{.*?\n\}\n", refreshed.text, re.S)[0]
            rendered = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const context = {projectUpdateUrl: input.prefix};
vm.createContext(context); vm.runInContext(input.code, context);
process.stdout.write(context.buildNode(input.tree));
"""], input=json.dumps({"code": escape_code + render_code, "prefix": "/" + prefix,
                         "tree": {**tree[0], "name": '<svg onload="bad()">',
                                  "files": [{**tree[0]["files"][0], "name": '<svg onload="bad()">.txt'}]}}),
                                      text=True, capture_output=True, check=True)
            class Links(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.hrefs = []
                    self.tags = []
                    self.preview_calls = []
                    self.rename_buttons = []
                    self.delete_buttons = []
                def handle_starttag(self, tag, attrs):
                    self.tags.append(tag)
                    if "btn-rename" in dict(attrs).get("class", "").split():
                        self.rename_buttons.append(dict(attrs))
                    if "btn-del" in dict(attrs).get("class", "").split():
                        self.delete_buttons.append(dict(attrs))
                    if tag == "a":
                        self.hrefs.append(dict(attrs).get("href", ""))
                        if dict(attrs).get("onclick"):
                            self.preview_calls.append(dict(attrs)["onclick"])
            links = Links()
            links.feed(rendered.stdout)
            assert "svg" not in links.tags
            assert any(href.startswith(f"/{prefix}/download/") for href in links.hrefs)
            invoked = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const callback = (...args) => process.stdout.write(JSON.stringify(args));
vm.runInNewContext(fs.readFileSync(0, 'utf8'), {openControlledPreview: callback});
"""], input=links.preview_calls[0], text=True, capture_output=True, check=True)
            assert json.loads(invoked.stdout)[:5] == [
                files[0]["fileId"], 1, "PROJECT", project["businessId"], '<svg onload="bad()">.txt',
            ]
            assert json.loads(invoked.stdout)[5] in links.hrefs
            preview = client.get("/preview/file", query_string={
                "fileId": files[0]["fileId"], "versionNo": 1,
                "objectType": "PROJECT", "objectId": project["businessId"],
            }, follow_redirects=True)
            assert preview.status_code == 200
            assert preview.json["type"] == "text" and preview.json["content"] == "research input"
            second = client.post(f"/{prefix}/upload/{project['businessId']}",
                                 data={"folder": "任务输入文件",
                                       "file": (io.BytesIO(b"second revision"), "input.txt")},
                                 headers={"X-CSRF-Token": csrf})
            assert second.status_code == 200 and second.json["versionNo"] == 2
            assert second.json["fileId"] == files[0]["fileId"]
            latest = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/input.txt",
                                follow_redirects=True)
            assert latest.data == b"second revision"
            original = client.get(f"/api/files/{files[0]['fileId']}/versions/1/download",
                                  query_string={"objectType": "PROJECT", "objectId": project["businessId"]})
            assert original.status_code == 200 and original.data == b"research input"
            wrong_prefix = "crypto_projects" if prefix == "projects" else "projects"
            rejected = client.post(f"/{wrong_prefix}/upload/{project['businessId']}",
                                   data={"folder": "任务输入文件",
                                         "file": (io.BytesIO(b"wrong category"), "input.txt")},
                                   headers={"X-CSRF-Token": csrf})
            assert rejected.status_code == 404
            unsafe = client.post(f"/{prefix}/upload/{project['businessId']}",
                                 data={"folder": "../outside",
                                       "file": (io.BytesIO(b"invalid"), "input.txt")},
                                 headers={"X-CSRF-Token": csrf})
            assert unsafe.status_code == 400
            assert app.extensions["file_service"].list_for_object(
                object_type="PROJECT", object_id=project["businessId"],
            )[0]["versionNo"] == 2
            if operation == "file_delete":
                service = app.extensions["file_service"]
                headless = app.extensions["project_service"].create_standalone(
                    category, {"name": "直接导出演练", "leader": "李老师"},
                    actor_user_id=user_id,
                )
                service.upload_project_path(
                    io.BytesIO(b"controlled-only"), original_name="input.txt",
                    folder="任务输入文件", project_id=headless["businessId"],
                    actor_user_id=user_id, request_id="headless-archive",
                )
                assert not (tmp_path / "uploads" / headless["businessId"]).exists()
                direct_zip = client.post(
                    f"/{prefix}/archive/{headless['businessId']}",
                    headers={"X-CSRF-Token": csrf},
                )
                assert direct_zip.status_code == 200
                with zipfile.ZipFile(io.BytesIO(direct_zip.data)) as packed:
                    assert packed.namelist() == ["任务输入文件/input.txt"]
                    assert packed.read("任务输入文件/input.txt") == b"controlled-only"
                if prefix != "projects":
                    preview = client.get(f"/{prefix}/preview/{project['businessId']}/任务输入文件/input.txt")
                    assert preview.status_code == 200 and preview.data == b"second revision"
                    assert not preview.headers.get("Content-Disposition", "").startswith("attachment")
                    assert preview.headers["X-Content-Type-Options"] == "nosniff"
                    assert "sandbox" in preview.headers["Content-Security-Policy"]
                delete_code = re.search(r"document.querySelectorAll\('\.btn-del'\).*?\n    \}\);",
                                        refreshed.text, re.S)[0]
                submitted = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const button = {dataset: {path: input.button['data-path'], fileId: input.button['data-file-id']},
  addEventListener(event, handler) {this.click = handler;}};
vm.runInNewContext(input.code, {document: {querySelectorAll: () => [button]}, FormData,
  projectUpdateUrl: input.prefix, confirm: () => true, alert() {},
  fetch(url, options) {process.stdout.write(JSON.stringify({url, data: Object.fromEntries(options.body)}));
    return Promise.resolve({json: () => ({success: false})});}});
button.click({stopPropagation() {}});
"""], input=json.dumps({"code": delete_code, "button": links.delete_buttons[0], "prefix": "/" + prefix}),
                                           text=True, capture_output=True, check=True)
                payload = json.loads(submitted.stdout)
                assert payload["data"].get("expectedFileId") == files[0]["fileId"]
                assert payload["url"] == f"/{prefix}/delete_file/{project['businessId']}"
                legacy_path = tmp_path / "uploads" / project["businessId"] / "任务输入文件" / "input.txt"
                legacy_path.parent.mkdir(parents=True, exist_ok=True)
                legacy_path.write_bytes(b"retained physical predecessor")
                legacy_path.with_name("unrelated.txt").write_bytes(b"retained unrelated")
                history_dir = legacy_path.parent / ".history"
                history_dir.mkdir()
                (history_dir / "input_v1.txt").write_bytes(b"deleted predecessor history")
                (history_dir / "input_versions.json").write_text("{}")
                (history_dir / "unrelated_v1.txt").write_bytes(b"unrelated history retained")
                deleted = client.post(f"/{prefix}/delete_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/input.txt", "expectedFileId": files[0]["fileId"],
                }, headers={"X-CSRF-Token": csrf})
                assert deleted.status_code == 200 and deleted.json["success"] is True
                service = app.extensions["file_service"]
                assert service.list_project_paths(project["businessId"]) == []
                assert legacy_path.read_bytes() == b"retained physical predecessor"
                refreshed = client.get(f"/{prefix}/detail/{project['businessId']}")
                tree = json.loads(re.search(r"var folderData = (.*);", refreshed.text)[1])
                assert [entry["name"] for entry in tree[0]["files"]] == ["unrelated.txt"]
                packed = client.post(f"/{prefix}/archive/{project['businessId']}", headers={"X-CSRF-Token": csrf})
                assert packed.status_code == 200
                invalid_selection = client.post(f"/{prefix}/archive/{project['businessId']}", data="{broken",
                                                content_type="application/json", headers={"X-CSRF-Token": csrf})
                assert invalid_selection.status_code == 400
                def failed_walk(path, **kwargs):
                    if kwargs.get("onerror"):
                        kwargs["onerror"](PermissionError("injected directory scan failure"))
                    return iter(())
                with monkeypatch.context() as failure:
                    failure.setattr("app.routes._project_bridge.os.walk", failed_walk)
                    incomplete = client.post(f"/{prefix}/archive/{project['businessId']}", headers={"X-CSRF-Token": csrf})
                    assert incomplete.status_code == 500
                project_directory = legacy_path.parent.parent
                saved_directory = project_directory.with_name(project_directory.name + "-test-saved")
                foreign = tmp_path / "uploads" / "FOREIGN-PROJECT"
                foreign.mkdir()
                (foreign / "private.txt").write_bytes(b"must not export another project")
                project_directory.rename(saved_directory)
                try:
                    project_directory.symlink_to(foreign, target_is_directory=True)
                    crossed = client.post(f"/{prefix}/archive/{project['businessId']}", headers={"X-CSRF-Token": csrf})
                    assert crossed.status_code == 400
                finally:
                    if project_directory.is_symlink():
                        project_directory.unlink()
                    saved_directory.rename(project_directory)
                with zipfile.ZipFile(io.BytesIO(packed.data)) as archive:
                    assert "任务输入文件/input.txt" not in archive.namelist()
                    assert "任务输入文件/.history/input_v1.txt" not in archive.namelist()
                    assert "任务输入文件/.history/input_versions.json" not in archive.namelist()
                    assert archive.read("任务输入文件/.history/unrelated_v1.txt") == b"unrelated history retained"
                    assert archive.read("任务输入文件/unrelated.txt") == b"retained unrelated"
                assert client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/.history/input_v1.txt").status_code == 404
                assert client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/input.txt").status_code == 404
                if prefix != "projects":
                    assert client.get(f"/{prefix}/preview/{project['businessId']}/任务输入文件/input.txt").status_code == 404
                resurrect = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/input.txt", "new_name": "resurrected.txt",
                }, headers={"X-CSRF-Token": csrf})
                assert resurrect.status_code == 409 and resurrect.json["success"] is False
                assert legacy_path.read_bytes() == b"retained physical predecessor"
                for alias in ("任务输入文件/./input.txt", "任务输入文件/../任务输入文件/input.txt"):
                    assert client.get(f"/{prefix}/download/{project['businessId']}/{alias}").status_code == 400
                with engine.connect() as connection:
                    for number, expected in [(1, b"research input"), (2, b"second revision")]:
                        old = service.repository.get_version(connection, file_id=files[0]["fileId"], version_no=number)
                        assert (service.storage_root / old["storage_path"]).read_bytes() == expected
                reuploaded = client.post(f"/{prefix}/upload/{project['businessId']}", data={
                    "folder": "任务输入文件", "file": (io.BytesIO(b"replacement after delete"), "input.txt"),
                }, headers={"X-CSRF-Token": csrf})
                assert reuploaded.status_code == 200 and reuploaded.json["success"] is True
                assert reuploaded.json["fileId"] != files[0]["fileId"] and reuploaded.json["versionNo"] == 1
                stale = client.post(f"/{prefix}/delete_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/input.txt", "expectedFileId": files[0]["fileId"],
                }, headers={"X-CSRF-Token": csrf})
                assert stale.status_code == 409
                replacement = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/input.txt",
                                         follow_redirects=True)
                assert replacement.status_code == 200 and replacement.data == b"replacement after delete"
                packed = client.post(f"/{prefix}/archive/{project['businessId']}",
                                     json={"paths": ["任务输入文件"]}, headers={"X-CSRF-Token": csrf})
                assert packed.status_code == 200
                with zipfile.ZipFile(io.BytesIO(packed.data)) as archive:
                    assert archive.read("任务输入文件/input.txt") == b"replacement after delete"
                    assert archive.read("任务输入文件/unrelated.txt") == b"retained unrelated"
                if prefix != "projects":
                    current_preview = client.get(f"/{prefix}/preview/{project['businessId']}/任务输入文件/input.txt")
                    assert current_preview.status_code == 200 and current_preview.data == b"replacement after delete"
                    opened_streams = []
                    original_open = service.open_version_stream
                    def track_open(*args, **kwargs):
                        opened = original_open(*args, **kwargs)
                        opened_streams.append(opened["stream"])
                        return opened
                    def fail_preview_audit(**kwargs):
                        raise FileServiceError("FILE_OPERATION_FAILED", "injected preview audit failure", 500)
                    with monkeypatch.context() as failure:
                        failure.setattr(service, "open_version_stream", track_open)
                        failure.setattr(service, "record_event", fail_preview_audit)
                        failed_preview = client.get(f"/{prefix}/preview/{project['businessId']}/任务输入文件/input.txt")
                        assert failed_preview.status_code == 500
                    assert len(opened_streams) == 1 and opened_streams[0].closed
                other = app.extensions["project_service"].create_standalone(
                    category, {"name": "演练-共享删除保留", "leader": "李老师"}, actor_user_id=user_id,
                )
                replacement_id = reuploaded.json["fileId"]
                with engine.begin() as connection:
                    service.repository.link_object(connection, link_id=uuid.uuid4(), object_type="PROJECT",
                                                   object_id=other["businessId"], file_id=replacement_id,
                                                   actor_user_id=user_id, purpose="PROJECT_TREE:任务输入文件/input.txt")
                def fail_delete_audit(*args, **kwargs):
                    raise RuntimeError("injected delete audit failure")
                with monkeypatch.context() as failure:
                    failure.setattr(service.audit_service, "record", fail_delete_audit)
                    rejected = client.post(f"/{prefix}/delete_file/{project['businessId']}", data={
                        "file_path": "任务输入文件/input.txt", "expectedFileId": replacement_id,
                    }, headers={"X-CSRF-Token": csrf})
                    assert rejected.status_code == 500
                assert service.list_project_paths(project["businessId"])[0]["fileId"] == replacement_id
                for business_id in (project["businessId"], other["businessId"]):
                    deleted = client.post(f"/{prefix}/delete_file/{business_id}", data={
                        "file_path": "任务输入文件/input.txt", "expectedFileId": replacement_id,
                    }, headers={"X-CSRF-Token": csrf})
                    assert deleted.status_code == 200 and deleted.json["success"] is True
                    assert service.list_project_paths(business_id) == []
                    assert client.get(f"/api/files/{replacement_id}/versions/1/download",
                                      query_string={"objectType": "PROJECT", "objectId": business_id}).status_code == 404
                    if business_id == project["businessId"]:
                        retained = client.get(f"/api/files/{replacement_id}/versions/1/download",
                                              query_string={"objectType": "PROJECT", "objectId": other["businessId"]})
                        assert retained.status_code == 200 and retained.data == b"replacement after delete"
                assert legacy_path.read_bytes() == b"retained physical predecessor"
                with engine.connect() as connection:
                    table = service.repository.files
                    row = connection.execute(sa.select(table).where(table.c.id == uuid.UUID(replacement_id))).mappings().one()
                    assert row["status"] == "ARCHIVED" and row["version"] == 1
            if operation == "file_archive":
                service = app.extensions["file_service"]
                def reject_archive_audit(*args, **kwargs):
                    raise RuntimeError("injected archive audit failure")
                with monkeypatch.context() as failure:
                    failure.setattr(service.audit_service, "record", reject_archive_audit)
                    with pytest.raises(FileServiceError) as failed:
                        service.archive(files[0]["fileId"], object_type="PROJECT", object_id=project["businessId"],
                                        actor_user_id=user_id, request_id="archive-audit-failure")
                    assert failed.value.code == "FILE_OPERATION_FAILED"
                active = service.list_project_paths(project["businessId"])
                assert [(row["fileId"], row["versionNo"]) for row in active] == [(files[0]["fileId"], 2)]
                for attempt in range(2):
                    service.archive(files[0]["fileId"], object_type="PROJECT", object_id=project["businessId"],
                                    actor_user_id=user_id, request_id=f"archive-version-{attempt}")
                    with engine.connect() as connection:
                        row = service.repository.get_linked_file(
                            connection, file_id=files[0]["fileId"], object_type="PROJECT",
                            object_id=project["businessId"],
                        )
                        assert row["status"] == "ARCHIVED" and row["version"] == 2
                        for number, expected_bytes in [(1, b"research input"), (2, b"second revision")]:
                            version = service.repository.get_version(
                                connection, file_id=files[0]["fileId"], version_no=number,
                            )
                            assert (service.storage_root / version["storage_path"]).read_bytes() == expected_bytes
                        assert service.repository.get_version(
                            connection, file_id=files[0]["fileId"], version_no=3,
                        ) is None
                assert service.list_project_paths(project["businessId"]) == []
            if operation == "file_rename":
                assert links.rename_buttons[0]["data-file-id"] == files[0]["fileId"]
                submit_code = re.search(
                    r"document.getElementById\('renameForm'\).onsubmit = function\(e\).*?\n    };",
                    refreshed.text, re.S,
                )[0]
                submitted = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const fields = {renameForm: {}, renameFilePath: {value: ''},
  renameFileId: {value: ''}, newFileName: {value: ''}, renameModal: {}};
const button = {dataset: {path: input.button['data-path'], name: input.button['data-name'],
  fileId: input.button['data-file-id']}, addEventListener: function(event, handler) {
    this.click = handler;
  }};
const context = {document: {getElementById: id => fields[id],
  querySelectorAll: () => [button]}, FormData,
  bootstrap: {Modal: class {show() {}}},
  projectUpdateUrl: input.prefix,
  fetch: (url, options) => {
    process.stdout.write(JSON.stringify({url, data: Object.fromEntries(options.body)}));
    return Promise.resolve({json: () => ({success: false})});
  }, alert: () => {}};
vm.createContext(context); vm.runInContext(input.clickCode, context);
button.click({stopPropagation() {}});
fields.newFileName.value = 'renamed.txt';
vm.runInContext(input.code, context);
fields.renameForm.onsubmit({preventDefault() {}});
"""], input=json.dumps({"code": submit_code, "prefix": "/" + prefix,
                         "button": links.rename_buttons[0],
                         "clickCode": re.search(
                             r"document.querySelectorAll\('\.btn-rename'\).*?\n    \}\);",
                             refreshed.text, re.S,
                         )[0]}),
                                           text=True, capture_output=True, check=True)
                request_data = json.loads(submitted.stdout)
                assert request_data["data"].get("expectedFileId") == files[0]["fileId"]
                assert request_data["url"] == f"/{prefix}/rename_file/{project['businessId']}"
                renamed = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/input.txt", "new_name": "renamed.txt",
                    "expectedFileId": files[0]["fileId"],
                }, headers={"X-CSRF-Token": csrf})
                assert renamed.status_code == 200 and renamed.json["success"] is True
                path_rows = app.extensions["file_service"].list_project_paths(project["businessId"])
                assert [(row["fileId"], row["path"], row["versionNo"]) for row in path_rows] == [
                    (files[0]["fileId"], "任务输入文件/renamed.txt", 2),
                ]
                downloaded = client.get(f"/{prefix}/download/{project['businessId']}/任务输入文件/renamed.txt",
                                        follow_redirects=True)
                assert downloaded.status_code == 200 and downloaded.data == b"second revision"
                original = client.get(f"/api/files/{files[0]['fileId']}/versions/1/download",
                                      query_string={"objectType": "PROJECT", "objectId": project["businessId"]})
                assert original.status_code == 200 and original.data == b"research input"
                legacy_target = tmp_path / "uploads" / project["businessId"] / "任务输入文件" / "occupied.txt"
                legacy_target.parent.mkdir(parents=True, exist_ok=True)
                legacy_target.write_bytes(b"retained legacy file")
                for name, status in [("occupied.txt", 409), ("renamed.pdf", 415), ("../escape.txt", 400)]:
                    rejected = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                        "file_path": "任务输入文件/renamed.txt", "new_name": name,
                        "expectedFileId": files[0]["fileId"],
                    }, headers={"X-CSRF-Token": csrf})
                    assert rejected.status_code == status and rejected.json["success"] is False
                assert legacy_target.read_bytes() == b"retained legacy file"
                service = app.extensions["file_service"]
                def fail_audit(*args, **kwargs):
                    raise RuntimeError("injected rename audit failure")
                with monkeypatch.context() as failure:
                    failure.setattr(service.audit_service, "record", fail_audit)
                    failed = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                        "file_path": "任务输入文件/renamed.txt", "new_name": "must-not-persist.txt",
                        "expectedFileId": files[0]["fileId"],
                    }, headers={"X-CSRF-Token": csrf})
                assert failed.status_code == 500
                assert [(row["path"], row["versionNo"]) for row in service.list_project_paths(project["businessId"])] == [
                    ("任务输入文件/renamed.txt", 2),
                ]
                reverse_conflict = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/occupied.txt", "new_name": "renamed.txt",
                }, headers={"X-CSRF-Token": csrf})
                assert reverse_conflict.status_code == 409 and reverse_conflict.json["success"] is False
                assert legacy_target.read_bytes() == b"retained legacy file"
                for source, target in [
                    ("任务输入文件/occupied.txt", "./renamed.txt"),
                    ("任务输入文件/occupied.txt", "../任务输入文件/renamed.txt"),
                    ("任务输入文件/./occupied.txt", "renamed.txt"),
                    ("任务输入文件//occupied.txt", "renamed.txt"),
                ]:
                    aliased = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                        "file_path": source, "new_name": target,
                    }, headers={"X-CSRF-Token": csrf})
                    assert aliased.status_code == 400 and aliased.json["success"] is False
                    assert legacy_target.read_bytes() == b"retained legacy file"
                legacy_renamed = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/occupied.txt", "new_name": "保留旧格式.doc",
                }, headers={"X-CSRF-Token": csrf})
                assert legacy_renamed.status_code == 200 and legacy_renamed.json["success"] is True
                assert not legacy_target.exists()
                assert legacy_target.with_name("保留旧格式.doc").read_bytes() == b"retained legacy file"
                for stale_id in (str(uuid.uuid4()), ""):
                    stale = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                        "file_path": "任务输入文件/renamed.txt", "new_name": "stale-must-not-change.txt",
                        "expectedFileId": stale_id,
                    }, headers={"X-CSRF-Token": csrf})
                    assert stale.status_code == 409
                assert service.list_project_paths(project["businessId"])[0]["path"] == "任务输入文件/renamed.txt"
                other = app.extensions["project_service"].create_standalone(
                    category, {"name": "演练-共享附件保留", "leader": "李老师"}, actor_user_id=user_id,
                )
                rename_marker = "rename-test-" + uuid.uuid4().hex
                def rename_concurrently():
                    worker_id = get_ident()
                    def mark_transaction(connection):
                        if get_ident() == worker_id:
                            connection.execute(sa.text("select set_config('application_name', :name, true)"),
                                               {"name": rename_marker})
                    sa.event.listen(engine, "begin", mark_transaction)
                    try:
                        return service.rename_project_path(
                            project["businessId"], "任务输入文件/renamed.txt", "concurrent-must-not-change.txt",
                            expected_file_id=files[0]["fileId"], actor_user_id=user_id,
                            request_id="shared-link-concurrency",
                        )
                    finally:
                        sa.event.remove(engine, "begin", mark_transaction)
                with ThreadPoolExecutor(max_workers=1) as executor:
                    with engine.begin() as connection:
                        blocker_pid = connection.execute(sa.text("select pg_backend_pid()")).scalar_one()
                        service.repository.link_object(
                            connection, link_id=uuid.uuid4(), object_type="PROJECT",
                            object_id=other["businessId"], file_id=files[0]["fileId"],
                            actor_user_id=user_id, purpose="PROJECT_TREE:任务输入文件/renamed.txt",
                        )
                        pending = executor.submit(rename_concurrently)
                        deadline = time.monotonic() + 5
                        blocked = False
                        with engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
                            while time.monotonic() < deadline and not pending.done():
                                blocked = observer.execute(sa.text(
                                    "select exists (select 1 from pg_stat_activity "
                                    "where application_name = :name and :pid = any(pg_blocking_pids(pid)))"
                                ), {"pid": blocker_pid, "name": rename_marker}).scalar_one()
                                if blocked:
                                    break
                                time.sleep(0.01)
                        assert blocked, "rename must wait for the uncommitted shared-link transaction"
                    with pytest.raises(FileServiceError) as conflict:
                        pending.result(timeout=5)
                    assert conflict.value.code == "FILE_SHARED" and conflict.value.status_code == 409
                shared = client.post(f"/{prefix}/rename_file/{project['businessId']}", data={
                    "file_path": "任务输入文件/renamed.txt", "new_name": "shared-must-not-change.txt",
                    "expectedFileId": files[0]["fileId"],
                }, headers={"X-CSRF-Token": csrf})
                assert shared.status_code == 409 and shared.json["code"] == "FILE_SHARED"
                with pytest.raises(FileServiceError) as shared_archive:
                    service.archive(files[0]["fileId"], object_type="PROJECT", object_id=project["businessId"],
                                    actor_user_id=user_id, request_id="shared-archive-must-not-change")
                assert shared_archive.value.code == "FILE_SHARED" and shared_archive.value.status_code == 409
                for business_id in (project["businessId"], other["businessId"]):
                    retained = service.list_project_paths(business_id)
                    assert [(row["path"], row["name"], row["versionNo"]) for row in retained] == [
                        ("任务输入文件/renamed.txt", "renamed.txt", 2),
                    ]
                    content = client.get(f"/api/files/{files[0]['fileId']}/versions/2/download",
                                         query_string={"objectType": "PROJECT", "objectId": business_id})
                    assert content.status_code == 200 and content.data == b"second revision"
                reverse_file = service.upload_project_path(
                    io.BytesIO(b"reverse ordering bytes"), original_name="before.txt", folder="任务输入文件",
                    project_id=project["businessId"], actor_user_id=user_id, request_id="reverse-upload",
                )
                rename_ready, release_rename = Event(), Event()
                rename_pid = []
                original_audit = service.audit_service.record
                def pause_after_real_audit(connection, **kwargs):
                    result = original_audit(connection, **kwargs)
                    if kwargs["request_id"] == "reverse-rename":
                        rename_pid.append(connection.execute(sa.text("select pg_backend_pid()")).scalar_one())
                        rename_ready.set()
                        assert release_rename.wait(10), "test must release the rename transaction"
                    return result
                link_marker = "link-test-" + uuid.uuid4().hex
                def insert_after_rename():
                    with engine.begin() as connection:
                        connection.execute(sa.text("select set_config('application_name', :name, true)"),
                                           {"name": link_marker})
                        service.repository.link_object(
                            connection, link_id=uuid.uuid4(), object_type="PROJECT",
                            object_id=other["businessId"], file_id=reverse_file["fileId"],
                            actor_user_id=user_id, purpose="PROJECT_TREE:任务输入文件/after.txt",
                        )
                with monkeypatch.context() as scheduling:
                    scheduling.setattr(service.audit_service, "record", pause_after_real_audit)
                    with ThreadPoolExecutor(max_workers=2) as executor:
                        renaming = executor.submit(
                            service.rename_project_path, project["businessId"], "任务输入文件/before.txt",
                            "after.txt", expected_file_id=reverse_file["fileId"],
                            actor_user_id=user_id, request_id="reverse-rename",
                        )
                        try:
                            assert rename_ready.wait(5), "rename must reach its real audit before commit"
                            linking = executor.submit(insert_after_rename)
                            blocked = False
                            deadline = time.monotonic() + 5
                            with engine.connect().execution_options(isolation_level="AUTOCOMMIT") as observer:
                                while time.monotonic() < deadline and not linking.done():
                                    blocked = observer.execute(sa.text(
                                        "select exists (select 1 from pg_stat_activity where "
                                        "application_name = :name and :pid = any(pg_blocking_pids(pid)))"
                                    ), {"name": link_marker, "pid": rename_pid[0]}).scalar_one()
                                    if blocked:
                                        break
                                    time.sleep(0.01)
                            assert blocked, "new link must wait for the rename transaction"
                        finally:
                            release_rename.set()
                        assert renaming.result(timeout=5)["path"] == "任务输入文件/after.txt"
                        linking.result(timeout=5)
                for business_id in (project["businessId"], other["businessId"]):
                    retained = [row for row in service.list_project_paths(business_id)
                                if row["fileId"] == reverse_file["fileId"]]
                    assert [(row["path"], row["name"], row["versionNo"]) for row in retained] == [
                        ("任务输入文件/after.txt", "after.txt", 1),
                    ]
                    content = client.get(f"/api/files/{reverse_file['fileId']}/versions/1/download",
                                         query_string={"objectType": "PROJECT", "objectId": business_id})
                    assert content.status_code == 200 and content.data == b"reverse ordering bytes"
    finally:
        engine.dispose()


def test_single_file_selection_never_selects_ancestor_directories():
    template = Path(__file__).parents[1] / "templates/projects/detail.html"
    source = template.read_text(encoding="utf-8")
    function = re.search(r"function toggleFolder\([^)]*\) \{.*?\n\}", source, re.S)[0]
    result = subprocess.run(["node", "-e", """
const fs = require('fs'); const vm = require('vm');
const folder = {checked: false, indeterminate: false};
const file = {checked: true};
const node = {querySelector: () => folder, querySelectorAll: () => [file],
  parentElement: {closest: () => null}};
file.closest = () => node;
const context = {}; vm.createContext(context);
vm.runInContext(fs.readFileSync(0, 'utf8'), context);
context.toggleFolder(file);
process.stdout.write(JSON.stringify(folder));
"""], input=function, text=True, capture_output=True, check=True)
    assert json.loads(result.stdout) == {"checked": False, "indeterminate": True}


def test_unchecking_subfolder_removes_ancestor_from_export_selection():
    source = (Path(__file__).parents[1] / "templates/projects/detail.html").read_text()
    functions = "\n".join(re.search(r"function " + name + r"\([^)]*\) \{.*?\n\}", source, re.S)[0]
                          for name in ("toggleKids", "toggleFolder"))
    result = subprocess.run(["node", "-e", """
const fs=require('fs'), vm=require('vm');
const rootCb={checked:false,value:'root'}, childCb={checked:false,value:'root/excluded'};
const keep={checked:false,value:'root/keep.txt'}, omit={checked:false,value:'root/excluded/omit.txt'};
const root={querySelector:()=>rootCb, querySelectorAll:s=>s==='.children'?[]:
  s==='.file-cb'?[keep,omit]:[rootCb,childCb,keep,omit], parentElement:{closest:()=>null}};
const child={querySelector:()=>childCb, querySelectorAll:s=>s==='.children'?[]:
  s==='.file-cb'?[omit]:[childCb,omit], parentElement:{closest:()=>root}};
rootCb.closest=()=>root; childCb.closest=()=>child;
const context={}; vm.createContext(context); vm.runInContext(fs.readFileSync(0,'utf8'),context);
rootCb.checked=true; context.toggleKids(rootCb);
childCb.checked=false; context.toggleKids(childCb);
const selected=[rootCb,childCb,keep,omit].filter(c=>c.checked).map(c=>c.value);
const partial=rootCb.indeterminate;
rootCb.checked=false; context.toggleKids(rootCb);
process.stdout.write(JSON.stringify({selected,partial,resetPartial:rootCb.indeterminate}));
"""], input=functions, text=True, capture_output=True, check=True)
    assert json.loads(result.stdout) == {"selected": ["root/keep.txt"], "partial": True, "resetPartial": False}
