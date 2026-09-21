from io import BytesIO
from types import SimpleNamespace

import pytest
from flask import Flask
from jinja2 import ChoiceLoader, DictLoader
from docx import Document

from app.security.csrf import csrf_token, protect_request


@pytest.fixture
def web():
    from app.web.research_reports import bp
    from app.services.research_reports import ResearchReportServiceError
    app = Flask('app', template_folder='templates')
    app.config.update(TESTING=True, SECRET_KEY='test-only', CSRF_ENABLED=True,
                      AI_FEATURES_VISIBLE=True)
    app.jinja_loader = ChoiceLoader([DictLoader({'base.html': '{% block content %}{% endblock %}'}), app.jinja_loader])
    app.jinja_env.globals['csrf_token'] = csrf_token
    app.before_request(protect_request)
    app.add_url_rule('/login', endpoint='auth.login', view_func=lambda: 'login')
    app.register_blueprint(bp)
    app.extensions['sensitive_term_set_service']=SimpleNamespace(current=lambda:dict(version=0,rules=[]))
    report = dict(id='r1', objectType='PROPOSAL', objectId='TP-1', title='研究报告', purpose='自填用途', currentVersion=2)
    versions = {n: dict(versionNo=n, body=f'第{n}版（原文）\n<script>alert(1)</script>', note=f'修改{n}', createdBy=7, createdByName='测试作者', createdAt='2026-09-10T10:00:00') for n in [1, 2]}
    calls = []
    def save(report_id, payload, **kwargs):
        calls.append(payload)
        raise ResearchReportServiceError('VERSION_CONFLICT', '版本已更新，请核对后再保存', 409)
    app.extensions['research_report_service'] = SimpleNamespace(
        get=lambda rid: report, current=lambda rid: versions[2],
        get_version=lambda rid, n: versions[n],
        history=lambda rid, **kw: dict(items=list(versions.values()), page=1, pageSize=20, total=2),
        list_for_object=lambda *args, **kw: dict(items=[report], page=1, pageSize=20, total=1),
        save_version=save, create=lambda payload, **kw: calls.append(payload) or report,
    )
    def redaction_context(snapshot=None, *, version_no=None, required=False):
        from app.services.sensitive_terms import compile_rules
        terms=app.extensions.get('sensitive_term_set_service')
        if terms is None:
            raise ResearchReportServiceError('REDACTION_UNAVAILABLE','词库不可用',503)
        vocabulary=terms.current() if snapshot is None and version_no is None else terms.get_version(snapshot['version'] if snapshot else version_no)
        return dict(version=vocabulary['version']),compile_rules([{k:rule[k] for k in ('id','source','replacement')} for rule in vocabulary['rules'] if rule['enabled']])
    app.extensions['research_report_service'].redaction_context=redaction_context
    client = app.test_client()
    with client.session_transaction() as session:
        session.update(user_id=7, user='alice', role='BUSINESS_USER', account_version=1, csrf_token='token')
    return client, calls


def test_conflict_preserves_unsaved_input_and_base(web):
    client, calls = web
    result = client.post('/research-reports/r1', data={'_csrf_token':'token', 'baseVersion':'1', 'body':'未保存 <b>正文</b>', 'note':'我的说明'})
    assert result.status_code == 409
    text = result.get_data(as_text=True)
    assert '未保存 &lt;b&gt;正文&lt;/b&gt;' in text
    assert '我的说明' in text and 'value="1"' in text
    assert calls[0]['baseVersion'] == 1


def test_history_is_readonly_and_escaped(web):
    result = web[0].get('/research-reports/r1/versions/1')
    text = result.get_data(as_text=True)
    assert result.status_code == 200
    assert '&lt;script&gt;' in text and '<script>alert' not in text
    assert '<textarea' not in text and '保存新版本' not in text


def test_docx_uses_requested_stored_version(web):
    result = web[0].get('/research-reports/r1/versions/1/export.docx')
    assert result.status_code == 200
    paragraphs = [p.text for p in Document(BytesIO(result.data)).paragraphs]
    assert '第1版（原文）' in paragraphs
    assert '第2版（原文）' not in paragraphs
    assert '<script>alert(1)</script>' in paragraphs
    assert any('v1' in p for p in paragraphs)
    assert any('测试作者' in p for p in paragraphs)
    document=Document(BytesIO(result.data))
    assert str(document.styles['Title'].font.color.rgb)=='000000'
    assert document.styles['Title'].element.get_or_add_pPr().find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr') is None


def test_docx_embeds_clean_image_and_omits_sensitive_original(web):
    import hashlib
    from PIL import Image
    output=BytesIO(); Image.new('RGB',(80,60),'white').save(output,format='PNG'); data=output.getvalue()
    digest=hashlib.sha256(data).hexdigest()
    service=web[0].application.extensions['research_report_service']
    version=service.get_version('r1',1)
    common=dict(versionNo=1,objectType='PROPOSAL',objectId='TP-1',sourceKind='IMAGE',
                mediaType='image/png',image=dict(width=80,height=60,mediaType='image/png'))
    version['sources']=[
        dict(common,sourceId='S1',label='S1',fileId='clean',filename='clean.png',sha256=digest,
             text='无敏感内容',originalImageBlockedByRedaction=False),
        dict(common,sourceId='S2',label='S2',fileId='blocked',filename='PRJ-B.png',sha256='b'*64,
             text='※：PRJ-B',originalImageBlockedByRedaction=True),
    ]
    calls=[]
    def open_version(file_id,version_no,**owner):
        calls.append((file_id,version_no,owner))
        assert file_id=='clean'
        return dict(stream=BytesIO(data),sha256=digest)
    web[0].application.extensions['file_service']=SimpleNamespace(open_version_stream=open_version)
    result=web[0].get('/research-reports/r1/versions/1/export.docx')
    assert result.status_code==200 and calls==[('clean',1,dict(object_type='PROPOSAL',object_id='TP-1'))]
    document=Document(BytesIO(result.data)); text='\n'.join(p.text for p in document.paragraphs)
    assert len(document.inline_shapes)==1
    assert 'clean.png（由程序按固定版式插入）' in text
    assert 'PRJ-B.png（原图命中敏感词库，未嵌入' in text and '脱敏提取摘要：※：PRJ-B' in text


def test_writes_require_csrf_and_session(web):
    client, calls = web
    assert client.post('/research-reports/r1', data={'body':'x'}).status_code == 403
    assert not calls
    with client.session_transaction() as session:
        session.clear()
    assert client.get('/research-reports/r1').status_code == 302


def test_create_carries_binding_and_arbitrary_purpose(web):
    client, calls = web
    result = client.post('/research-reports/new?objectType=PROPOSAL&objectId=TP-1', data={'_csrf_token':'token', 'title':'报告', 'purpose':'用于阶段讨论', 'body':'资料正文'})
    assert result.status_code == 302
    assert calls[-1]['objectId'] == 'TP-1'
    assert calls[-1]['purpose'] == '用于阶段讨论'


# Exercise the HTTP form boundary with real isolated SQLite persistence.
from app.tests.test_research_reports import runtime


def test_real_service_form_versions_and_export(web, runtime):
    client, _ = web
    client.application.extensions['research_report_service'] = runtime[0]
    with client.session_transaction() as session:
        session['user_id'] = 1
    response = client.post('/research-reports/new?objectType=PROPOSAL&objectId=TP-TEST', data={'_csrf_token':'token', 'title':'真实服务边界测试', 'purpose':'测试用途', 'body':'第一版正文'})
    assert response.status_code == 302
    location = response.headers['Location']
    response = client.post(location, data={'_csrf_token':'token', 'baseVersion':'1', 'body':'第二版正文', 'note':'补充材料'})
    assert response.status_code == 302
    assert '第二版正文' in client.get(location).get_data(as_text=True)
    assert '第一版正文' in client.get(location + '/versions/1').get_data(as_text=True)
    result = client.get(location + '/versions/2/export.docx')
    assert result.status_code == 200
    assert '第二版正文' in [p.text for p in Document(BytesIO(result.data)).paragraphs]
    assert client.get('/research-reports?objectType=PROPOSAL&objectId=TP-TEST&page=1').status_code == 200
    response = client.post(location, data={'_csrf_token':'token', 'baseVersion':'1', 'body':'冲突未保存', 'note':'冲突'})
    assert response.status_code == 409 and '冲突未保存' in response.get_data(as_text=True)
    assert runtime[0].current(location.rsplit('/', 1)[1])['body'] == '第二版正文'


@pytest.mark.parametrize('value', ['0', '-1', '1.1', 'true', '1e1', ''])
def test_invalid_form_base_preserves_body(web, value):
    client, calls = web
    response = client.post('/research-reports/r1', data={'_csrf_token':'token', 'baseVersion':value, 'body':'输入保留'})
    assert response.status_code == 422
    assert '输入保留' in response.get_data(as_text=True)
    assert not calls


@pytest.mark.parametrize('status,version,expected', [('active',1,200), ('disabled',1,302), ('active',2,302)])
def test_actual_auth_revalidates_account(web, status, version, expected):
    client, _ = web
    client.application.config['SECURITY_AUTH_ENABLED'] = True
    client.application.extensions['users_repository'] = SimpleNamespace(get_by_id=lambda uid: dict(id=7, status=status, role='BUSINESS_USER', version=version))
    assert client.get('/research-reports/r1').status_code == expected


def test_database_failure_retains_body_without_false_success(web):
    from sqlalchemy.exc import OperationalError
    client, _ = web
    def fail(*args, **kwargs):
        raise OperationalError('statement', {}, Exception('unavailable'))
    client.application.extensions['research_report_service'].save_version = fail
    response = client.post('/research-reports/r1', data={'_csrf_token':'token','baseVersion':'2','body':'数据库故障时保留','note':'待保存'})
    assert response.status_code == 503
    text = response.get_data(as_text=True)
    assert '数据库故障时保留' in text and '待保存' in text
    assert 'statement' not in text


def test_real_audit_failure_rolls_back_and_preserves_form(web, runtime):
    client, _ = web
    service = runtime[0]
    client.application.extensions['research_report_service'] = service
    with client.session_transaction() as session:
        session['user_id'] = 1
    report = service.create(dict(objectType='PROPOSAL',objectId='TP-TEST',title='故障回滚测试',body='原稿'),actor_user_id=1,request_id='test')
    def fail(*args, **kwargs):
        raise RuntimeError('audit storage unavailable')
    service.audit_service.record = fail
    response = client.post('/research-reports/' + report['id'], data={'_csrf_token':'token','baseVersion':'1','body':'审计故障时保留','note':'未保存'})
    assert response.status_code == 503
    assert '审计故障时保留' in response.get_data(as_text=True)
    assert service.current(report['id'])['body'] == '原稿'
    assert service.history(report['id'])['total'] == 1


@pytest.mark.parametrize('failed_method', ['create', 'list_for_object'])
def test_create_database_failure_retains_all_input(web, failed_method):
    from sqlalchemy.exc import OperationalError
    client, _ = web
    def fail(*args, **kwargs):
        raise OperationalError('private statement', {}, Exception('unavailable'))
    setattr(client.application.extensions['research_report_service'], failed_method, fail)
    response = client.post('/research-reports/new?objectType=PROPOSAL&objectId=TP-1', data={'_csrf_token':'token', 'title':'故障标题','purpose':'故障用途','body':'故障正文','note':'故障说明'})
    assert response.status_code == 503
    text = response.get_data(as_text=True)
    assert all(v in text for v in ['故障标题','故障用途','故障正文','故障说明'])
    assert 'private statement' not in text


def test_real_form_crlf_exports_logical_paragraphs_without_extra_breaks(web, runtime):
    client, _ = web
    service = runtime[0]
    client.application.extensions['research_report_service'] = service
    with client.session_transaction() as session:
        session['user_id'] = 1
    original = '第一段\r\n第二段\r\n\r\n第四段\r末段\r\n'
    response = client.post('/research-reports/new?objectType=PROPOSAL&objectId=TP-TEST', data={'_csrf_token':'token','title':'换行测试','purpose':'隔离测试','body':original})
    assert response.status_code == 302
    location = response.headers['Location']
    report_id = location.rsplit('/', 1)[1]
    assert service.get_version(report_id, 1)['body'] == original
    result = client.get(location + '/versions/1/export.docx')
    assert result.status_code == 200
    paragraphs = [p.text for p in Document(BytesIO(result.data)).paragraphs]
    assert paragraphs[4:] == ['第一段', '第二段', '', '第四段', '末段', '']
    assert service.get_version(report_id, 1)['body'] == original


@pytest.fixture
def draft_web(web):
    client, calls = web
    source = dict(fileId='f1', versionNo=3, originalName='来源材料.txt', sha256='a'*64)
    draft = dict(id='d1', body='待确认 <script>正文</script>', baseVersion=2, sources=[source], model='local-test-model')
    generated = []
    client.application.extensions['file_service'] = SimpleNamespace(list_for_object=lambda **kw: [source])
    client.application.extensions['research_report_draft_service'] = SimpleNamespace(
        get=lambda rid, did: draft,
        generate=lambda rid, payload, **kw: generated.append(payload) or draft,
    )
    return client, calls, generated, draft


def test_generate_posts_unsaved_body_and_fixed_source_then_separate_draft(draft_web):
    client, calls, generated, draft = draft_web
    result = client.post('/research-reports/r1/generate', data={'_csrf_token':'token','baseVersion':'2','body':'尚未保存的修订基础','note':'我的说明','sources':'["f1",3]'})
    assert result.status_code == 302
    assert result.headers['Location'].endswith('/research-reports/r1/drafts/d1')
    assert generated == [dict(baseVersion=2, body='尚未保存的修订基础', selections=[dict(fileId='f1',versionNo=3)],dictionaryVersion=0)]
    assert not calls
    text = client.get(result.headers['Location']).get_data(as_text=True)
    assert '待确认草稿' in text and '&lt;script&gt;' in text
    assert '来源材料.txt' in text and 'local-test-model' in text


def test_generate_failure_preserves_input_selection_and_note(draft_web):
    from app.services.research_reports import ResearchReportServiceError
    client, calls, generated, draft = draft_web
    def fail(*args, **kwargs):
        raise ResearchReportServiceError('MODEL_UNAVAILABLE','本地模型暂不可用',503)
    client.application.extensions['research_report_draft_service'].generate = fail
    result = client.post('/research-reports/r1/generate', data={'_csrf_token':'token','baseVersion':'2','body':'保留草稿内容','note':'保留说明','sources':'["f1",3]'})
    assert result.status_code == 503
    text = result.get_data(as_text=True)
    assert '保留草稿内容' in text and '保留说明' in text and 'checked' in text
    assert not calls


def test_draft_save_passes_draft_id_and_conflict_keeps_candidate(draft_web):
    client, calls, _, _ = draft_web
    result = client.post('/research-reports/r1/drafts/d1', data={'_csrf_token':'token','baseVersion':'2','body':'人已修订正文','note':'确认前修改'})
    assert result.status_code == 409
    assert calls[-1] == dict(baseVersion=2,body='人已修订正文',note='确认前修改',draftId='d1')
    assert '人已修订正文' in result.get_data(as_text=True)


def test_generate_requires_csrf(draft_web):
    client, _, generated, _ = draft_web
    assert client.post('/research-reports/r1/generate',data={'body':'x'}).status_code == 403
    assert not generated


from app.tests.test_research_report_drafts import drafts


def test_real_draft_service_form_accept_and_source_export(web, drafts):
    import json
    import uuid
    client, _ = web
    draft_service, reports, _ = drafts
    client.application.extensions['research_report_service'] = reports
    client.application.extensions['research_report_draft_service'] = draft_service
    source_id = str(uuid.uuid4())
    client.application.extensions['file_service'] = SimpleNamespace(list_for_object=lambda **kw: [dict(fileId=source_id,versionNo=1,originalName='资料.txt')])
    with client.session_transaction() as session:
        session['user_id'] = 1
    report = reports.create(dict(objectType='PROPOSAL',objectId='TP-TEST',title='整合测试',body='已保存正文'),actor_user_id=1,request_id='test')
    location = '/research-reports/' + report['id']
    result = client.post(location + '/generate',data={'_csrf_token':'token','baseVersion':'1','body':'未保存输入','sources':json.dumps([source_id,1])})
    assert result.status_code == 302
    draft_url = result.headers['Location']
    assert reports.current(report['id'])['body'] == '已保存正文'
    assert draft_service.assistant.inputs['current_body'] == '未保存输入'
    assert '来源原文（不可变）' in client.get(draft_url).get_data(as_text=True)
    result = client.post(draft_url,data={'_csrf_token':'token','baseVersion':'1','body':'人工核对的整合正文[S1]','note':'已核对来源'})
    assert result.status_code == 302
    saved = reports.current(report['id'])
    assert saved['versionNo'] == 2 and saved['sources'][0]['fileId'] == source_id
    text = client.get(location + '/versions/2').get_data(as_text=True)
    assert '来源原文（不可变）' in text and saved['sources'][0]['sha256'] in text
    result = client.get(location + '/versions/2/export.docx')
    paragraphs = [p.text for p in Document(BytesIO(result.data)).paragraphs]
    assert '来源清单' in paragraphs
    assert any(source_id in p for p in paragraphs)
    assert any(saved['sources'][0]['sha256'] in p for p in paragraphs)
    assert reports.get_version(report['id'],1)['body'] == '已保存正文'


def test_generate_database_failure_keeps_selected_source(draft_web):
    from sqlalchemy.exc import OperationalError
    client, _, generated, _ = draft_web
    def fail(*args, **kwargs):
        raise OperationalError('private sql', {}, Exception('unavailable'))
    client.application.extensions['research_report_draft_service'].generate = fail
    response = client.post('/research-reports/r1/generate',data={'_csrf_token':'token','baseVersion':'2','body':'数据库故障正文','note':'故障说明','sources':'["f1",3]'})
    assert response.status_code == 503
    text = response.get_data(as_text=True)
    assert '数据库故障正文' in text and '故障说明' in text and 'checked' in text
    assert 'private sql' not in text
    assert not generated


@pytest.mark.parametrize('selection', ['["f1",true]', '["f1",0]', '{"fileId":"f1"}'])
def test_invalid_generation_source_never_calls_model(draft_web, selection):
    client, _, generated, _ = draft_web
    response = client.post('/research-reports/r1/generate',data={'_csrf_token':'token','baseVersion':'2','body':'保留输入','sources':selection})
    assert response.status_code == 422 and '保留输入' in response.get_data(as_text=True)
    assert not generated


def test_selection_api_uses_text_boundaries_with_emoji(web, monkeypatch):
    from app.web import research_reports
    client, calls = web
    received=[]
    client.application.extensions['research_report_draft_service']=SimpleNamespace(assistant=SimpleNamespace(base_url='http://127.0.0.1:18081',model_version='local-test'))
    class Selection:
        def __init__(self, **kwargs):
            pass
        def suggest(self, **kwargs):
            received.append(kwargs)
            return dict(suggestions=[dict(text='新措辞',reason='简洁')],processorStatus='PENDING_INTEGRATION')
    monkeypatch.setattr(research_reports, 'LocalSelectionAssistant', Selection)
    response=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=2,body='😀前文原措辞尾文',prefix='😀前文',selectedText='原措辞',suffix='尾文'),headers={'X-CSRF-Token':'token'})
    assert response.status_code==200
    assert received[0]['selected_text']=='原措辞' and received[0]['prefix']=='😀前文'
    assert not calls


def test_selection_api_rejects_mismatched_boundary_and_stale_version(web):
    client, _ = web
    headers={'X-CSRF-Token':'token'}
    bad=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=2,body='原文',prefix='',selectedText='别文',suffix=''),headers=headers)
    assert bad.status_code==422
    stale=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=1,body='原文',prefix='',selectedText='原文',suffix=''),headers=headers)
    assert stale.status_code==409


def test_selection_cancel_reaches_inflight_transport_for_same_identity(web, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import threading
    import time
    from app.web import research_reports
    client, calls=web
    app=client.application
    app.extensions['research_report_draft_service']=SimpleNamespace(assistant=SimpleNamespace(base_url='http://127.0.0.1:18081',model_version='local-test'))
    started=threading.Event()
    class Selection:
        def __init__(self, **kwargs): pass
        def suggest(self, **kwargs):
            started.set()
            end=time.monotonic()+3
            while time.monotonic()<end:
                if kwargs['cancel_check'](): raise InterruptedError('cancelled')
                time.sleep(0.01)
            raise AssertionError('cancel did not reach transport')
    monkeypatch.setattr(research_reports,'LocalSelectionAssistant',Selection)
    cancel_client=app.test_client()
    with cancel_client.session_transaction() as session:
        session.update(user_id=7,user='alice',role='BUSINESS_USER',account_version=1,csrf_token='token')
    operation='16a6d65e-6268-4642-a3da-a865c46471c0'
    payload=dict(operationId=operation,baseVersion=2,body='原文',prefix='',selectedText='原文',suffix='')
    with ThreadPoolExecutor(max_workers=1) as executor:
        future=executor.submit(client.post,'/research-reports/r1/selection-suggestions',json=payload,headers={'X-CSRF-Token':'token'})
        assert started.wait(2)
        response=cancel_client.post('/research-reports/r1/selection-suggestions/cancel',json={'operationId':operation},headers={'X-CSRF-Token':'token'})
        assert response.json['data']['cancelled'] is True
        assert future.result(timeout=2).status_code==409
    assert not research_reports._selection_jobs and not calls


def test_selection_checks_remote_version_again_after_model(web, monkeypatch):
    from app.web import research_reports
    client, calls=web
    service=client.application.extensions['research_report_service']
    report=service.get('r1')
    client.application.extensions['research_report_draft_service']=SimpleNamespace(assistant=SimpleNamespace(base_url='http://127.0.0.1:18081',model_version='local-test'))
    class Selection:
        def __init__(self,**kwargs):pass
        def suggest(self,**kwargs):
            report['currentVersion']=3
            return dict(suggestions=[dict(text='原文',reason='不变')])
    monkeypatch.setattr(research_reports,'LocalSelectionAssistant',Selection)
    result=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=2,body='原文',prefix='',selectedText='原文',suffix=''),headers={'X-CSRF-Token':'token'})
    assert result.status_code==409 and not calls


def test_selection_dictionary_snapshot_once_no_rules_in_business_response(web,monkeypatch):
    from app.web import research_reports
    client,_=web
    reads=[];processed=[]
    snapshot=dict(version=7,rules=[dict(id='term1',source='虚构敏感名',replacement='单位甲',enabled=True),dict(id='term2',source='禁用词',replacement='单位乙',enabled=False)])
    client.application.extensions['sensitive_term_set_service']=SimpleNamespace(current=lambda:reads.append(1) or snapshot)
    client.application.extensions['research_report_draft_service']=SimpleNamespace(assistant=SimpleNamespace(base_url='http://127.0.0.1:18081',model_version='local-test'))
    class Selection:
        def __init__(self,**kwargs):self.processor=kwargs['processor']
        def suggest(self,**kwargs):
            processed.append(self.processor(kwargs['selected_text']))
            return dict(suggestions=[dict(text='单位甲表达',reason='简洁')])
    monkeypatch.setattr(research_reports,'LocalSelectionAssistant',Selection)
    result=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=2,body='虚构敏感名表达',prefix='',selectedText='虚构敏感名表达',suffix=''),headers={'X-CSRF-Token':'token'})
    assert result.status_code==200 and reads==[1] and processed==['单位甲表达']
    assert result.json['data']['dictionaryVersion']==7
    assert '虚构敏感名' not in result.get_data(as_text=True) and 'rules' not in result.get_data(as_text=True)


def test_selection_fails_closed_without_dictionary_service(web,monkeypatch):
    from app.web import research_reports
    client,_=web
    client.application.extensions.pop('sensitive_term_set_service')
    client.application.extensions['research_report_draft_service']=SimpleNamespace(assistant=SimpleNamespace(base_url='http://127.0.0.1:18081',model_version='local-test'))
    calls=[]
    monkeypatch.setattr(research_reports,'LocalSelectionAssistant',lambda **kw:calls.append(kw))
    result=client.post('/research-reports/r1/selection-suggestions',json=dict(baseVersion=2,body='原文',prefix='',selectedText='原文',suffix=''),headers={'X-CSRF-Token':'token'})
    assert result.status_code==503 and not calls


from app.tests.test_report_redaction_snapshot import redacted


def test_selection_uses_fixed_vocabulary_and_save_carries_same_version(web,redacted,monkeypatch):
    from app.web import research_reports
    client,_=web
    draft_service,reports,_,vocabulary=redacted
    client.application.extensions['research_report_service']=reports
    client.application.extensions['research_report_draft_service']=draft_service
    draft_service.assistant.base_url='http://127.0.0.1:18081'
    draft_service.assistant.model_version='local-test'
    with client.session_transaction() as session:session['user_id']=1
    report=reports.create(dict(objectType='PROPOSAL',objectId='TP-TEST',title='保密单位报告',purpose='保密单位用途',body='保密单位表达'),actor_user_id=1,request_id='test')
    vocabulary.advance()
    class Selection:
        def __init__(self,**kwargs):self.processor=kwargs['processor']
        def suggest(self,**kwargs):return dict(suggestions=[dict(text=self.processor('保密单位表述'),reason='措辞')])
    monkeypatch.setattr(research_reports,'LocalSelectionAssistant',Selection)
    location='/research-reports/'+report['id']
    result=client.post(location+'/selection-suggestions',json=dict(baseVersion=1,body='保密单位表达',prefix='',selectedText='保密单位表达',suffix=''),headers={'X-CSRF-Token':'token'})
    assert result.status_code==200 and result.json['data']['dictionaryVersion']==1
    assert result.json['data']['suggestions'][0]['text']=='单位甲表述'
    result=client.post(location,data={'_csrf_token':'token','baseVersion':'1','dictionaryVersion':'1','body':'单位甲表述','note':'确认'})
    assert result.status_code==302
    assert reports.current(report['id'])['redactionSnapshot']['version']==1
    result=client.get(location+'/versions/2/export.docx')
    paragraphs=[p.text for p in Document(BytesIO(result.data)).paragraphs]
    assert paragraphs[0]=='单位甲报告'
    assert '保密单位' not in '\n'.join(paragraphs) and '单位乙' not in '\n'.join(paragraphs)


def test_new_generation_distinguishes_target_and_input_vocabulary(draft_web):
    client,_,generated,_=draft_web
    client.application.extensions['sensitive_term_set_service']=SimpleNamespace(current=lambda:dict(version=8,rules=[]))
    response=client.post('/research-reports/r1/generate',data={'_csrf_token':'token','baseVersion':'2','dictionaryVersion':'7','generationDictionaryVersion':'8','body':'旧代称编辑框','sources':'["f1",3]'})
    assert response.status_code==302
    assert generated[0]['dictionaryVersion']==8 and generated[0]['inputDictionaryVersion']==7


def test_generation_rejects_outdated_displayed_target_vocabulary(draft_web):
    client,_,generated,_=draft_web
    client.application.extensions['sensitive_term_set_service']=SimpleNamespace(current=lambda:dict(version=8,rules=[]))
    response=client.post('/research-reports/r1/generate',data={'_csrf_token':'token','baseVersion':'2','generationDictionaryVersion':'7','body':'保留未保存内容','sources':'["f1",3]'})
    assert response.status_code==409 and not generated
    assert '保留未保存内容' in response.get_data(as_text=True)


def test_author_html_and_docx_follow_frozen_dictionary(web, redacted):
    import sqlalchemy as sa
    client, _ = web
    _, reports, engine, vocabulary = redacted
    client.application.extensions['research_report_service'] = reports
    with engine.begin() as connection:
        connection.execute(sa.text("UPDATE users SET name='保密单位作者' WHERE id=1"))
    report = reports.create(dict(objectType='PROPOSAL',objectId='TP-TEST',title='测试报告',body='正文'),actor_user_id=1,request_id='author-test')
    vocabulary.advance()
    url = '/research-reports/' + report['id'] + '/versions/1'
    html = client.get(url).get_data(as_text=True)
    assert '单位甲作者' in html and '保密单位作者' not in html and '单位乙作者' not in html
    result = client.get(url + '/export.docx')
    text = '\n'.join(p.text for p in Document(BytesIO(result.data)).paragraphs)
    assert '单位甲作者' in text and '保密单位作者' not in text and '单位乙作者' not in text


def test_saved_model_draft_route_is_bound_and_readonly(draft_web, monkeypatch):
    from app.web import research_reports
    client, calls, _, draft = draft_web
    service = client.application.extensions['research_report_service']
    service.get_version = lambda rid, number: dict(versionNo=number, draftId='d1')
    draft['reportId'] = 'r1'
    monkeypatch.setattr(research_reports, 'render_template', lambda name, **context: dict(template=name, **context))
    url = '/research-reports/r1/versions/2/model-draft'
    result = client.get(url)
    assert result.status_code == 200
    assert result.json['draft']['id'] == 'd1'
    assert result.json['template'] == 'research_reports/model_draft.html'
    assert client.post(url, data={'_csrf_token':'token'}).status_code == 405
    draft['reportId'] = 'another-report'
    assert client.get(url).status_code == 404
    service.get_version = lambda rid, number: dict(versionNo=number, draftId=None)
    assert client.get(url).status_code == 404
    assert not calls


def test_model_draft_history_link_and_readonly_original_template(web):
    client, _ = web
    service = client.application.extensions['research_report_service']
    saved = service.get_version('r1', 1)
    saved['draftId'] = 'd1'
    html = client.get('/research-reports/r1/versions/1').get_data(as_text=True)
    assert '/research-reports/r1/versions/1/model-draft' in html
    from flask import render_template
    with client.application.test_request_context():
        html = render_template('research_reports/model_draft.html', report=service.get('r1'), version=saved, draft=dict(id='d1', body='模型原稿 <script>危险</script>\n独立正文', model='local-test', createdAt='2026-09-11', sources=[]))
    assert '模型原稿 &lt;script&gt;危险&lt;/script&gt;' in html
    assert '<textarea' not in html and '<form' not in html
    assert '独立正文' in html and '/research-reports/r1/versions/1' in html
