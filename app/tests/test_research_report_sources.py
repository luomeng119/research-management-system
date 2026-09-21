import hashlib
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image

from app.services.research_reports import ResearchReportServiceError
from app.services.research_report_sources import extract_report_source


def opened(name, data):
    return {'originalName': name, 'sizeBytes': len(data), 'stream': BytesIO(data)}


def png_bytes(size=(64, 48)):
    output = BytesIO(); Image.new('RGB', size, 'white').save(output, format='PNG')
    return output.getvalue()


class Vision:
    def generate(self, **kwargs):
        self.inputs = kwargs
        return dict(documentType='科研图表',summary='试验摘要',visibleText=['青岚-07'],
                    facts=['样本24个'],chartFindings=['阶段3得分89'],uncertainties=[],
                    model='Qwen3.5-9B',promptVersion='vision-v1',usage={'inputTokens':1,'outputTokens':2})


def test_png_uses_local_vision_and_closes_original_stream():
    data=png_bytes(); source=opened('记录.png',data); vision=Vision()
    result=extract_report_source(source,vision_assistant=vision)
    assert source['stream'].closed and result['sourceKind']=='IMAGE'
    assert result['image']=={'width':64,'height':48,'mediaType':'image/png'}
    assert '青岚-07' in result['text'] and '阶段3得分89' in result['text']
    assert result['visionModel']=='Qwen3.5-9B' and result['visionPromptVersion']=='vision-v1'
    assert vision.inputs['image_bytes']==data and vision.inputs['media_type']=='image/png'


def test_image_requires_vision_and_rejects_excessive_pixel_dimensions(monkeypatch):
    source=opened('记录.png',png_bytes())
    with pytest.raises(ResearchReportServiceError) as error:
        extract_report_source(source)
    assert error.value.status_code==422 and source['stream'].closed
    import app.services.research_report_sources as module
    monkeypatch.setattr(module,'MAX_IMAGE_PIXELS',10)
    source=opened('记录.png',png_bytes())
    with pytest.raises(ResearchReportServiceError) as error:
        extract_report_source(source,vision_assistant=Vision())
    assert error.value.status_code==422 and source['stream'].closed


@pytest.mark.parametrize('suffix', ['txt', 'md', 'csv'])
def test_utf8_preserves_text_and_closes(suffix):
    value = '科研：ＡＢＣ ① m²\r\n原文\n'
    source = opened('来源.' + suffix, value.encode())
    result = extract_report_source(source)
    assert result['text'] == value
    assert result['textSha256'] == hashlib.sha256(value.encode()).hexdigest()
    assert result['segments'] == [{'locator': '全文', 'text': value}]
    assert source['stream'].closed


@pytest.mark.parametrize('name,data', [('x.exe', b'a'), ('x.txt', b'\xff'), ('x.txt', b'\x00'), ('x.txt', b' '), ('x.pdf', b'bad'), ('x.docx', b'bad'), ('x.txt', b'a' * 12001)])
def test_failures_are_safe_422_and_close(name, data):
    source = opened(name, data)
    with pytest.raises(ResearchReportServiceError) as error:
        extract_report_source(source)
    assert error.value.status_code == 422
    assert source['stream'].closed
    assert '/Users/' not in error.value.message


def test_docx_preserves_paragraph_table_order():
    document = Document()
    document.add_paragraph('首段 ＡＢＣ m²')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '单位甲'
    table.cell(0, 1).text = '数据①'
    document.add_paragraph('尾段')
    data = BytesIO(); document.save(data)
    source = opened('来源.docx', data.getvalue())
    result = extract_report_source(source)
    assert result['text'] == '首段 ＡＢＣ m²\n单位甲\t数据①\n尾段'
    assert [row['locator'] for row in result['segments']] == ['段落1', '表格1/行1', '段落2']
    assert source['stream'].closed


def test_docx_expansion_limit():
    data = BytesIO()
    with ZipFile(data, 'w') as archive:
        archive.writestr('word/document.xml', b'a' * (20 * 1024 * 1024 + 1))
    source = opened('x.docx', data.getvalue())
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(source)
    assert source['stream'].closed


def pdf_pages(texts):
    # Minimal actual PDF with Helvetica text, no optional fixture dependency.
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'']
    kids = []
    for text in texts:
        page_id = len(objects) + 1; kids.append(f'{page_id} 0 R')
        content = f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode()
        objects += [f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> /Contents {page_id+1} 0 R >>'.encode(), b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream']
    objects[1] = f'<< /Type /Pages /Count {len(texts)} /Kids [{" ".join(kids)}] >>'.encode()
    data = b'%PDF-1.4\n'; offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += f'{index} 0 obj\n'.encode() + obj + b'\nendobj\n'
    start = len(data)
    data += f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode()
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    data += f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
    return data


def test_pdf_per_page():
    source = opened('x.pdf', pdf_pages(['first', 'second']))
    result = extract_report_source(source)
    assert [s['locator'] for s in result['segments']] == ['第1页', '第2页']
    assert 'first' in result['segments'][0]['text'] and 'second' in result['segments'][1]['text']
    assert source['stream'].closed


@pytest.mark.parametrize('texts', [['text', ''], ['text'] * 21])
def test_pdf_never_silently_skips_page(texts):
    source = opened('x.pdf', pdf_pages(texts))
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(source)
    assert source['stream'].closed


def test_input_size_limit_closes_without_reading():
    source = opened('x.txt', b'ok')
    source['sizeBytes'] = 10 * 1024 * 1024 + 1
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(source)
    assert source['stream'].closed


def test_bom_and_non_normalized_characters_preserved():
    value = '\ufeffＡ e\u0301 ①\r\n'
    assert extract_report_source(opened('x.txt', value.encode()))['text'] == value


def test_docx_tracked_changes_are_not_silently_lost():
    from docx.oxml import OxmlElement
    document = Document()
    paragraph = document.add_paragraph('可见文字')
    paragraph._p.append(OxmlElement('w:ins'))
    data = BytesIO(); document.save(data)
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(opened('x.docx', data.getvalue()))


def test_close_failure_stays_safe():
    class FailedClose(BytesIO):
        def close(self):
            super().close()
            raise OSError('/secret/local/path')
    source = {'originalName': 'x.txt', 'sizeBytes': 2, 'stream': FailedClose(b'ok')}
    with pytest.raises(ResearchReportServiceError) as error:
        extract_report_source(source)
    assert error.value.status_code == 422
    assert '/secret/' not in error.value.message


@pytest.mark.parametrize('part_name', ['word/footnotes.xml', 'word/endnotes.xml', 'word/header1.xml', 'word/footer1.xml', 'word/media/image1.png'])
def test_docx_unread_parts_rejected(part_name):
    document = Document(); document.add_paragraph('正文仍然存在')
    original = BytesIO(); document.save(original)
    data = BytesIO()
    with ZipFile(original) as source, ZipFile(data, 'w') as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr(part_name, b'<unsupported/>')
    source = opened('x.docx', data.getvalue())
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(source)
    assert source['stream'].closed


def test_docx_inline_picture_rejected_even_with_body_text():
    from docx.oxml import OxmlElement
    document = Document()
    paragraph = document.add_paragraph('正文')
    paragraph._p.append(OxmlElement('w:drawing'))
    data = BytesIO(); document.save(data)
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(opened('x.docx', data.getvalue()))


def test_pdf_image_with_text_footer_rejected(monkeypatch):
    from pdfminer.layout import LTImage, LTPage, LTTextBoxHorizontal, LTTextLineHorizontal, LTAnno
    from pdfminer.pdftypes import PDFStream
    import app.services.research_report_sources as module
    page = LTPage(1, (0, 0, 600, 800))
    footer = LTTextBoxHorizontal(); line = LTTextLineHorizontal(0.1)
    line._objs.append(LTAnno('page footer')); footer.add(line); page.add(footer)
    page.add(LTImage('scanned-body', PDFStream({}, b''), (0, 0, 600, 700)))
    monkeypatch.setattr(module, 'extract_pages', lambda *a, **k: iter([page]))
    source = opened('x.pdf', b'placeholder')
    with pytest.raises(ResearchReportServiceError):
        extract_report_source(source)
    assert source['stream'].closed


def test_docx_hyperlink_text_is_preserved():
    from docx.oxml import OxmlElement
    document = Document(); paragraph = document.add_paragraph('正文：')
    hyperlink = OxmlElement('w:hyperlink'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = '来源名称Ａ'; run.append(text); hyperlink.append(run); paragraph._p.append(hyperlink)
    data = BytesIO(); document.save(data)
    result = extract_report_source(opened('x.docx', data.getvalue()))
    assert result['text'] == '正文：来源名称Ａ'
