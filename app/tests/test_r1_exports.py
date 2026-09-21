import json
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import re
import subprocess
from types import SimpleNamespace

from docx import Document
from flask import Flask, render_template, session
import openpyxl
import pytest

from app.security.csrf import csrf_token, protect_request
from app.tests.test_legacy_modules import expert_engine, expert_service, expert_routes


ROOT = Path(__file__).resolve().parents[2]


def browser_result(html, script):
    result = subprocess.run(
        ['node', '-e', '''
const {chromium} = require('playwright');
(async () => {
  const input = JSON.parse(require('fs').readFileSync(0, 'utf8'));
  const browser = await chromium.launch({headless: true});
  try {
    const page = await browser.newPage();
    await page.route('**/*', route => route.abort());
    await page.setContent(input.html.replace(/<script\\b[^>]*>[\\s\\S]*?<\\/script>/gi, ''));
    console.log(JSON.stringify(await page.evaluate('(' + input.script + ')()')));
  } finally { await browser.close(); }
})().catch(error => {console.error(error); process.exit(1);});
'''],
        input=json.dumps({'html': html, 'script': script}),
        text=True, capture_output=True, cwd=ROOT, check=True,
    )
    return json.loads(result.stdout)


@pytest.mark.parametrize('format', ['word', 'excel'])
def test_selected_export_form_passes_csrf_and_returns_document(expert_routes, format):
    client, _service, _audit = expert_routes
    client.application.config['CSRF_ENABLED'] = True
    client.application.jinja_env.globals['csrf_token'] = csrf_token
    client.application.before_request(protect_request)
    html = client.get('/experts/').get_data(as_text=True)
    with client.session_transaction() as active_session:
        token = active_session['csrf_token']
    function = re.search(r'function exportSelected\(format\) \{[\s\S]*?\n\}', html).group()
    submitted = browser_result(html, '''() => {
window.APP_CSRF_TOKEN = %s;
HTMLFormElement.prototype.submit = function () {
  window.submitted = {path: this.getAttribute('action'), data: Array.from(new FormData(this).entries())};
};
%s
window.exportSelected = exportSelected;
var button = Array.from(document.querySelectorAll('button')).find(button => button.textContent.includes(%s));
if (!button || !button.getClientRects().length) throw new Error('Visible export button missing');
window.alert = message => {window.emptyMessage = message;};
button.click();
if (window.submitted || window.emptyMessage !== '请先选择要导出的专家') throw new Error('Empty selection must not export');
document.querySelector('.expert-cb').checked = true;
button.click();
return window.submitted;
}''' % (json.dumps(token), function, json.dumps('导出Word' if format == 'word' else '导出Excel')))
    from werkzeug.datastructures import MultiDict
    response = client.post(submitted['path'], data=MultiDict(submitted['data']))
    assert response.status_code == 200
    if format == 'word':
        assert Document(BytesIO(response.data)).tables[0].rows[1].cells[1].text == '张老师'
    else:
        assert openpyxl.load_workbook(BytesIO(response.data)).active['C2'].value == '张老师'
    assert client.post(submitted['path'], data={'expert_ids': 'EXP-001'}).status_code == 403


@pytest.mark.parametrize('endpoint', ['group', 'all', 'selected'])
def test_excel_exports_preserve_timezone_dates(endpoint):
    from app.routes.experts import bp
    from app.routes.expert_groups import bp as groups_bp
    app = Flask(__name__)
    app.config.update(TESTING=True, SECRET_KEY='r1-test')
    app.register_blueprint(bp)
    app.register_blueprint(groups_bp)
    date = datetime(2026, 9, 9, 12, 34, 56, tzinfo=timezone(timedelta(hours=8)))
    expert = dict(expertId='EXP-001', name='专家', unit='单位', position='研究员',
                  expertise='材料', phone='123', bankCard='456', bankName='银行',
                  uploader='测试', createdAt=date, updatedAt=date,
                  selectedBy='测试', selectedAt=date)
    app.extensions['resources_service'] = SimpleNamespace(
        export_experts_sensitive=lambda **kwargs: [expert],
        export_expert_group_sensitive=lambda *args, **kwargs: dict(
            groupId='group-1', groupName='测试组', creator='测试', createdAt=date,
            members=[expert],
        ),
    )
    client = app.test_client()
    with client.session_transaction() as active_session:
        active_session.update(user='tester', user_id=7)
    if endpoint == 'group':
        response = client.get('/experts/groups/export/group-1')
        cells = ['B3', 'J6']
    elif endpoint == 'all':
        response = client.get('/experts/export')
        cells = ['J2', 'K2']
    else:
        response = client.post('/experts/export/selected', data={'expert_ids': 'EXP-001'})
        cells = ['J2', 'K2']
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(BytesIO(response.data)).active
    assert [sheet[cell].value for cell in cells] == [date.isoformat(), date.isoformat()]


def test_generated_word_contains_each_top_level_chapter(tmp_path):
    from app.routes.argumentation.routes import argumentation_bp, _template_data
    app = Flask(__name__, template_folder=str(ROOT / 'app/templates'))
    app.config.update(TESTING=True, SECRET_KEY='r1-test', DATA_DIR=str(tmp_path))
    app.register_blueprint(argumentation_bp, url_prefix='/argumentation')
    for endpoint in ['users.index', 'users.change_password', 'auth.logout']:
        app.add_url_rule('/test/' + endpoint, endpoint=endpoint, view_func=lambda: '')
    app.jinja_env.globals['csrf_token'] = lambda: 'test-token'
    template = _template_data('research', None)
    with app.test_request_context('/'):
        session['user'] = 'tester'
        html = render_template(
            'argumentation/edit.html', template=template, template_data=template,
            project={'name': '测试方案', 'project_id': 'KY-TEST'}, category='research',
            category_name='科研项目', saved_content={}, template_content={},
            equipment_list=[], categories=[], project_equipment=[], versions=[],
            document=None, template_version=0, template_id=None, back_url='/',
            is_version_view=False,
        )
    function = html.split('function generateWord() {', 1)[1].split('// ==================== 章节编辑功能', 1)[0]
    payload = browser_result(html, '''() => {
window.confirm = () => true;
var editors = {};
document.querySelectorAll('.editor').forEach(el => {
  editors[el.dataset.field] = {root: {innerHTML: '<p>正文-' + el.dataset.field + '</p>'}};
});
window.fetch = (url, options) => {window.payload = JSON.parse(options.body); return new Promise(() => {});};
function generateWord() {%s
generateWord();
return window.payload;
}''' % function)
    response = app.test_client().post('/argumentation/generate_word', json=payload)
    assert response.json['success'] is True
    path = tmp_path / 'documents' / response.json['word_path'].rsplit('/', 1)[1]
    paragraphs = [paragraph.text for paragraph in Document(path).paragraphs]
    for identifier in ['ch1_1', 'ch2_1', 'ch3_1', 'ch4', 'ch5', 'ch6', 'ch7']:
        assert '正文-' + identifier in paragraphs
