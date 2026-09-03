from __future__ import annotations

from datetime import date, datetime, timezone
from concurrent.futures import ThreadPoolExecutor
import math
import os
from pathlib import Path
import re
import zipfile

import openpyxl
import pytest
import sqlalchemy as sa
from sqlalchemy.pool import StaticPool

from app.repositories.generic_tables import GenericTablesRepository
from app.services.generic_tables import MAX_ROWS, GenericTablesError, GenericTablesService


def _schema(engine):
    metadata = sa.MetaData()
    tables = sa.Table(
        "generic_tables", metadata,
        sa.Column("id", sa.Integer, primary_key=True),
        sa.Column("table_id", sa.Text, nullable=False, unique=True),
        sa.Column("name", sa.Text, nullable=False),
        sa.Column("description", sa.Text),
        sa.Column("creator", sa.Text, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("current_version_id", sa.Text),
    )
    versions = sa.Table(
        "generic_table_versions", metadata,
        sa.Column("id", sa.Integer, primary_key=True),
        sa.Column("version_id", sa.Text, nullable=False, unique=True),
        sa.Column("table_id", sa.Text, sa.ForeignKey(tables.c.table_id), nullable=False),
        sa.Column("version_number", sa.Integer, nullable=False),
        sa.Column("version_label", sa.Text),
        sa.Column("create_method", sa.Text, nullable=False),
        sa.Column("source_version_id", sa.Text),
        sa.Column("row_count", sa.Integer, nullable=False, default=0),
        sa.Column("page_size", sa.Integer, nullable=False, default=20),
        sa.Column("creator", sa.Text, nullable=False),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("note", sa.Text),
        sa.Column("is_locked", sa.Boolean, nullable=False, default=False),
        sa.UniqueConstraint("table_id", "version_number"),
    )
    sa.ForeignKeyConstraint([tables.c.current_version_id], [versions.c.version_id], use_alter=True)
    sa.Table(
        "generic_table_columns", metadata,
        sa.Column("id", sa.Integer, primary_key=True),
        sa.Column("version_id", sa.Text, sa.ForeignKey(versions.c.version_id), nullable=False),
        sa.Column("col_key", sa.Text, nullable=False),
        sa.Column("col_name", sa.Text, nullable=False),
        sa.Column("col_type", sa.Text, nullable=False),
        sa.Column("col_index", sa.Integer, nullable=False),
        sa.Column("col_width", sa.Integer, nullable=False, default=120),
        sa.Column("col_align", sa.Text, nullable=False, default="left"),
        sa.Column("col_summary", sa.Text, nullable=False, default=""),
        sa.Column("col_options", sa.JSON),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.UniqueConstraint("version_id", "col_key"),
    )
    sa.Table(
        "generic_table_data", metadata,
        sa.Column("id", sa.Integer, primary_key=True),
        sa.Column("version_id", sa.Text, sa.ForeignKey(versions.c.version_id), nullable=False),
        sa.Column("row_key", sa.Text, nullable=False),
        sa.Column("row_index", sa.Integer, nullable=False),
        sa.Column("row_data", sa.JSON, nullable=False),
        sa.Column("row_color", sa.Text, nullable=False, default=""),
        sa.Column("created_at", sa.DateTime(timezone=True), nullable=False),
        sa.Column("updated_at", sa.DateTime(timezone=True), nullable=False),
        sa.UniqueConstraint("version_id", "row_key"),
    )
    metadata.create_all(engine)
    return metadata


@pytest.fixture
def service():
    engine = sa.create_engine(
        "sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )
    _schema(engine)
    return GenericTablesService(GenericTablesRepository(engine))


def test_table_crud_and_current_write_guards(service):
    table_id = service.create("样本台账", "描述", "张老师")
    table = service.get_by_id_with_current(table_id)
    current = table["current_version_id"]
    service.upsert_column(current, "name", "名称", "text", col_options=["甲", "乙"])
    service.upsert_row(current, "r1", {"name": "甲"})
    snapshot, new_current = service.save_snapshot(table_id, current, "初版", "", "张老师")
    assert snapshot != new_current
    assert service.get_version_by_id(snapshot)["is_locked"] is True
    assert service.get_by_id(table_id)["current_version_id"] == new_current
    with pytest.raises(GenericTablesError, match="当前编辑版本"):
        service.upsert_row(current, "r2", {"name": "历史"})
    with pytest.raises(GenericTablesError, match="当前版本"):
        service.delete_version(new_current)
    assert service.delete_version(snapshot)
    service.update(table_id, name="样本台账2")
    assert service.get_by_id(table_id)["name"] == "样本台账2"


def test_cross_table_compare_set_current_and_rollback_are_rejected(service):
    first = service.create("甲表", "", "张老师")
    second = service.create("乙表", "", "李老师")
    first_version = service.get_by_id(first)["current_version_id"]
    second_version = service.get_by_id(second)["current_version_id"]
    with pytest.raises(GenericTablesError, match="不属于"):
        service.set_current_version(first, second_version)
    with pytest.raises(GenericTablesError, match="同一表格"):
        service.compare_versions(first_version, second_version)
    with pytest.raises(GenericTablesError, match="不属于"):
        service.rollback_to(first, second_version, "王老师")


def test_snapshot_failure_rolls_back_all_versions(service, monkeypatch):
    table_id = service.create("甲表", "", "张老师")
    source = service.get_by_id(table_id)["current_version_id"]
    before = service.get_versions(table_id)

    def fail(*_args, **_kwargs):
        raise RuntimeError("injected copy failure")

    monkeypatch.setattr(service.repository, "copy_rows", fail)
    with pytest.raises(RuntimeError, match="injected"):
        service.save_snapshot(table_id, source, "失败快照", "", "张老师")
    assert service.get_versions(table_id) == before
    assert service.get_by_id(table_id)["current_version_id"] == source


def test_new_version_import_failure_rolls_back_current_and_version(service, tmp_path, monkeypatch):
    table_id = service.create("甲表", "", "张老师")
    source = service.get_by_id(table_id)["current_version_id"]
    path = tmp_path / "valid.xlsx"
    workbook = openpyxl.Workbook(); workbook.active.append(["名称"]); workbook.active.append(["甲"]); workbook.save(path)
    before = service.get_versions(table_id)
    monkeypatch.setattr(service, "_import_parsed", lambda *_args: (_ for _ in ()).throw(RuntimeError("injected")))
    with pytest.raises(RuntimeError, match="injected"):
        service.import_as_new_version(table_id, source, path, "张老师")
    assert service.get_by_id(table_id)["current_version_id"] == source
    assert service.get_versions(table_id) == before


def test_json_validation_and_bounded_rows(service):
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    for bad in (["not-object"], {"n": math.nan}, {"d": date.today()}, {"x": [[[[[["deep"]]]]]]}):
        with pytest.raises(GenericTablesError):
            service.upsert_row(version, "bad", bad)
    service.upsert_row(version, "ok", {"n": 1, "nested": [True, None, {"x": "y"}]})
    assert service.get_rows(version)[0]["row_data"]["nested"][2] == {"x": "y"}


def test_xlsx_replace_append_formula_defence_and_cleanup(service, tmp_path):
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    source = tmp_path / "rows.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["名称", "日期"])
    ws.append(["=2+2", datetime(2026, 9, 3, 10, 30)])
    wb.save(source)
    assert service.import_rows_from_excel(version, source, "张老师", "replace") == 1
    assert service.import_rows_from_excel(version, source, "张老师", "append") == 1
    assert service.get_version_by_id(version)["row_count"] == 2
    out = Path(service.export_to_excel(version))
    try:
        exported = openpyxl.load_workbook(out, data_only=False).active
        assert exported.cell(2, 1).value == "'=2+2"
        assert exported.cell(2, 2).value == "2026-09-03T10:30:00"
    finally:
        out.unlink(missing_ok=True)
    legacy = tmp_path / "rows.xls"
    legacy.write_bytes(b"not-xlsx")
    with pytest.raises(GenericTablesError, match=".xlsx"):
        service.import_rows_from_excel(version, legacy, "张老师")


def test_export_escapes_formula_headers_and_serializes_nested_json(service):
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    service.upsert_column(version, "meta", "=DANGEROUS", "text")
    service.upsert_row(version, "r1", {"meta": {"level": 1}})
    out = Path(service.export_to_excel(version))
    try:
        sheet = openpyxl.load_workbook(out, data_only=False).active
        assert sheet.cell(1, 1).value == "'=DANGEROUS"
        assert sheet.cell(2, 1).value == '{"level": 1}'
    finally:
        out.unlink(missing_ok=True)


def test_import_preserves_horizontal_header_and_vertical_data_merges(service, tmp_path):
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    source = tmp_path / "merged.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["C1"] = "主题", "内容"
    sheet.merge_cells("A1:B1")
    sheet["A2"], sheet["C2"], sheet["C3"] = "课题甲", "第一项", "第二项"
    sheet.merge_cells("A2:A3")
    workbook.save(source)
    assert service.import_rows_from_excel(version, source, "张老师") == 2
    columns = service.get_columns(version)
    assert [column["col_name"] for column in columns] == ["主题", "内容"]
    rows = service.get_rows(version)
    assert [row["row_data"]["主题"] for row in rows] == ["课题甲", "课题甲"]


def test_header_only_import_still_rejects_noncurrent_unlocked_version(service, tmp_path):
    table_id = service.create("甲表", "", "张老师")
    old_current = service.get_by_id(table_id)["current_version_id"]
    service.save_snapshot(table_id, old_current, "快照", "", "张老师")
    source = tmp_path / "header.xlsx"
    workbook = openpyxl.Workbook(); workbook.active.append(["名称"]); workbook.save(source)
    with pytest.raises(GenericTablesError, match="当前编辑版本"):
        service.import_rows_from_excel(old_current, source, "张老师")


def test_compare_detects_row_color_only_change(service):
    table_id = service.create("甲表", "", "张老师")
    first = service.get_by_id(table_id)["current_version_id"]
    service.upsert_row(first, "r1", {"name": "甲"})
    snapshot, current = service.save_snapshot(table_id, first, "快照", "", "张老师")
    service.update_row_color(current, "r1", "red")
    diff = service.compare_versions(snapshot, current)
    assert diff["row_diff"][0]["status"] == "modified"
    assert diff["row_diff"][0]["old_color"] == ""
    assert diff["row_diff"][0]["new_color"] == "red"


def test_xlsx_zip_bomb_is_rejected_before_openpyxl(service, tmp_path):
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 1_000_000)
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    with pytest.raises(GenericTablesError, match="压缩比"):
        service.import_rows_from_excel(version, path, "张老师")


def test_oversized_read_is_explicit_not_truncated(service):
    table_id = service.create("甲表", "", "张老师")
    version = service.get_by_id(table_id)["current_version_id"]
    now = datetime.now(timezone.utc)
    with service.repository.engine.begin() as connection:
        connection.execute(service.repository.data.insert(), [
            {"version_id": version, "row_key": f"r{i}", "row_index": i, "row_data": {"i": i}, "row_color": "", "created_at": now, "updated_at": now}
            for i in range(MAX_ROWS + 1)
        ])
    with pytest.raises(GenericTablesError, match="超过"):
        service.get_rows(version)
    page, total = service.get_rows_page(version, page=2, page_size=100)
    assert len(page) == 100 and total == MAX_ROWS + 1


def test_source_has_no_runtime_sqlite_or_schema_ddl():
    source = Path("app/models_generic_tables.py").read_text(encoding="utf-8")
    combined = source + Path("app/repositories/generic_tables.py").read_text(encoding="utf-8")
    for forbidden in ("sqlite3", "DB_PATH", "PRAGMA", "ALTER TABLE", "CREATE TABLE", "init_db"):
        assert forbidden not in combined


def test_templates_do_not_interpolate_user_values_into_html_sinks():
    templates = Path("app/templates/generic_tables")
    source = "\n".join(path.read_text(encoding="utf-8") for path in templates.glob("*.html"))
    assert not re.search(r"onclick=[^>]*\{\{", source)
    assert not re.search(r"innerHTML[^\n]*(?:\$\{(?:v|row|item)\.|oldVal|newVal)", source)


@pytest.mark.skipif(not os.environ.get("T10_TEST_DATABASE_URL"), reason="requires isolated PostgreSQL")
def test_postgresql_concurrent_snapshot_has_one_winner_and_unique_versions():
    engine = sa.create_engine(os.environ["T10_TEST_DATABASE_URL"])
    service = GenericTablesService(GenericTablesRepository(engine))
    table_id = service.create("并发表", "", "张老师")
    source = service.get_by_id(table_id)["current_version_id"]

    def save(label):
        try:
            return "ok", service.save_snapshot(table_id, source, label, "", "张老师")
        except GenericTablesError as error:
            return error.code, None

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, ("快照A", "快照B")))
    assert [status for status, _ in results].count("ok") == 1
    assert [status for status, _ in results].count("VERSION_READ_ONLY") == 1
    versions = service.get_versions(table_id)
    assert len({row["version_number"] for row in versions}) == len(versions)
    assert service.get_by_id(table_id)["current_version_id"] in {row["version_id"] for row in versions if not row["is_locked"]}
    service.delete(table_id)
