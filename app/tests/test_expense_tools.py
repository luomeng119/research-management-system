from __future__ import annotations

from decimal import Decimal
import os
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import io
from contextlib import contextmanager
from datetime import datetime, timezone

import pytest
import sqlalchemy as sa


def test_expense_service_rejects_float_amounts_before_repository_access():
    from app.services.expenses import ExpenseService, ExpenseValidationError

    class NeverCalledRepository:
        engine = None

        def __getattr__(self, name):
            raise AssertionError(f"repository must not be called: {name}")

    service = ExpenseService(NeverCalledRepository())
    with pytest.raises(ExpenseValidationError, match="金额"):
        service.create_invoice(amount=0.1)


def test_runtime_facade_contains_no_sqlite_or_schema_ddl():
    source = (Path(__file__).parents[1] / "expense_db.py").read_text(encoding="utf-8").lower()
    assert "sqlite3" not in source
    assert "create table" not in source
    assert "pragma" not in source
    assert "db_path" not in source


def test_expense_runtime_has_bounded_queries_and_no_legacy_preview_paths():
    app_dir = Path(__file__).parents[1]
    service_source = (app_dir / "services" / "expenses.py").read_text(encoding="utf-8")
    repository_source = (app_dir / "repositories" / "expenses.py").read_text(encoding="utf-8")
    template_source = (app_dir / "templates" / "expense" / "records.html").read_text(encoding="utf-8")
    assert "MAX_PAGE_SIZE" in service_source
    assert ".limit(limit).offset(offset)" in repository_source
    assert "/app/uploads/" not in template_source
    assert "file_path" not in template_source
    assert "controlledFile" in template_source


def test_expense_full_record_reads_use_bounded_batches_not_limit_none():
    app_dir = Path(__file__).parents[1]
    service_source = (app_dir / "services" / "expenses.py").read_text(encoding="utf-8")
    repository_source = (app_dir / "repositories" / "expenses.py").read_text(encoding="utf-8")
    facade_source = (app_dir / "expense_db.py").read_text(encoding="utf-8")
    routes_source = "\n".join(
        (app_dir / "routes" / name).read_text(encoding="utf-8")
        for name in ("expense.py", "documents.py")
    )
    assert "limit=None" not in service_source
    assert "list_all_invoices" not in service_source
    assert "list_all_payments" not in service_source
    assert "all_rows=True" not in routes_source
    assert "all_rows" not in facade_source
    assert "if limit is not None" not in repository_source


def test_records_page_preserves_match_warning_and_renders_it_with_text_content(tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-warning-page", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions-warning"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state["user"] = "teacher"
    page = client.get("/expense/records")
    source = page.get_data(as_text=True)
    assert page.status_code == 200
    assert "match_warning: data.match_warning" in source
    assert "match_warning: r.match_warning" in source
    assert "warningNode.textContent" in source
    assert "warningNode.innerHTML" not in source


def test_real_expense_pages_expose_bounded_pagination_and_mixed_upload_warnings(tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-pagination-page", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions-pagination"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state["user"] = "teacher"

    records_source = client.get("/expense/records").get_data(as_text=True)
    documents_source = client.get("/expense/documents/").get_data(as_text=True)

    assert "recordsNextOffset" in records_source
    assert "invoice_next_offset" in records_source
    assert "payment_next_offset" in records_source
    assert "appendMatchWarnings(document.getElementById('uploadErrorText'), results)" in records_source
    assert "本页合计" in records_source
    assert "renderRecordsPager();\n        return;" in records_source
    assert "reimbursementsNextOffset" in documents_source
    assert "loadReimbursements(reimbursementsNextOffset)" in documents_source
    assert "selectReimbursement(parseInt(pid));" in documents_source
    assert "setInterval" not in documents_source


def test_document_template_avoids_dynamic_inline_document_handlers():
    source = (Path(__file__).parents[1] / "templates" / "expense" / "documents.html").read_text(encoding="utf-8")
    assert 'onclick="deleteDoc(' not in source
    assert "data-doc-id" in source
    assert "escapeHtml(prefillValue)" in source


@pytest.fixture(scope="module")
def expense_service():
    url = os.environ.get("T10_TEST_DATABASE_URL")
    if not url:
        pytest.skip("expense PostgreSQL tests require T10_TEST_DATABASE_URL")
    from app.repositories.expenses import ExpensesRepository
    from app.services.expenses import ExpenseService

    engine = sa.create_engine(url, pool_size=5, max_overflow=5)
    tables = ["object_files", "stored_file_versions", "stored_files", "expense_invoice_item", "expense_invoice", "expense_payment", "expense_reimbursement"]
    with engine.begin() as connection:
        for table in tables:
            connection.execute(sa.text(f"DELETE FROM {table}"))
    service = ExpenseService(ExpensesRepository(engine))
    yield service
    with engine.begin() as connection:
        for table in tables:
            connection.execute(sa.text(f"DELETE FROM {table}"))
    engine.dispose()


def test_postgresql_crud_uses_decimal_dates_items_and_protects_control_fields(expense_service):
    from app.services.expenses import ExpenseValidationError

    rid, number = expense_service.create_reimbursement(title="设备采购", approver="张老师")
    iid = expense_service.create_invoice(
        reimbursement_id=rid, invoice_no="INV-001", date="2026-09-03",
        amount="12.30", tax_amount="0.30", price_ex_tax="12.00",
        items=[{"name": "试剂", "quantity": "2.0000", "unit_price": "6.0000", "amount": "12.00"}],
    )
    pid = expense_service.create_payment(reimbursement_id=rid, amount="12.30", pay_date="2026-09-03")

    assert number.startswith("REI202")
    assert expense_service.get_invoice(iid)["items"][0]["name"] == "试剂"
    assert expense_service.recalculate_total(rid) == 12.3
    assert expense_service.get_payment(pid)["file_path"] is None
    with pytest.raises(ExpenseValidationError, match="专用操作"):
        expense_service.update_invoice(iid, reimbursement_id=None)
    with pytest.raises(ExpenseValidationError, match="专用操作"):
        expense_service.update_reimbursement(rid, status="已确认")


def test_postgresql_confirm_uses_all_rows_beyond_first_page(expense_service):
    from app.services.expenses import ExpenseValidationError

    rid, _ = expense_service.create_reimbursement(title="501 条全集确认")
    now = datetime.now(timezone.utc)
    invoices = [{
        "reimbursement_id": rid, "invoice_no": f"BULK-{index}", "amount": Decimal("1.00"),
        "tax_amount": Decimal("0.00"), "price_ex_tax": Decimal("0.00"), "status": "已匹配",
        "matched_payment_ids": [], "created_at": now,
    } for index in range(501)]
    payments = [{
        "reimbursement_id": rid, "payment_no": f"BULK-PAY-{index}",
        "amount": Decimal("1.00") if index < 500 else Decimal("2.00"),
        "status": "已匹配", "matched_invoice_ids": [], "created_at": now,
    } for index in range(501)]
    with expense_service.repository.engine.begin() as connection:
        connection.execute(expense_service.repository.invoices.insert(), invoices)
        connection.execute(expense_service.repository.payments.insert(), payments)
    with pytest.raises(ExpenseValidationError, match="金额必须相等"):
        expense_service.confirm(rid)
    assert expense_service.get_reimbursement(rid)["status"] == "草稿"


def test_postgresql_document_record_collection_reads_501_rows_in_batches(expense_service):
    rid, _ = expense_service.create_reimbursement(title="501 条文档记录")
    now = datetime.now(timezone.utc)
    with expense_service.repository.engine.begin() as connection:
        connection.execute(expense_service.repository.invoices.insert(), [{
            "reimbursement_id": rid, "invoice_no": f"DOC-{index}", "amount": Decimal("1.00"),
            "tax_amount": Decimal("0.00"), "price_ex_tax": Decimal("0.00"), "status": "已匹配",
            "matched_payment_ids": [], "created_at": now,
        } for index in range(501)])
        connection.execute(expense_service.repository.payments.insert(), [{
            "reimbursement_id": rid, "payment_no": f"DOC-PAY-{index}", "amount": Decimal("1.00"),
            "status": "已匹配", "matched_invoice_ids": [], "created_at": now,
        } for index in range(501)])
    invoices, payments = expense_service.collect_reimbursement_children(rid)
    assert len(invoices) == 501
    assert len(payments) == 501


def test_document_record_collection_rejects_explicit_v1_limit(expense_service, monkeypatch):
    from app.services.expenses import ExpenseError

    monkeypatch.setattr(
        expense_service.repository, "count_invoices",
        lambda *_args, **_kwargs: expense_service.MAX_REIMBURSEMENT_CHILDREN + 1,
    )
    monkeypatch.setattr(expense_service.repository, "count_payments", lambda *_args, **_kwargs: 0)
    monkeypatch.setattr(
        expense_service.repository, "list_invoices",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("must reject before reading rows")),
    )
    with pytest.raises(ExpenseError, match="上限"):
        expense_service.collect_reimbursement_children(1)


def test_confirmed_document_source_rows_are_immutable(expense_service):
    from app.services.expenses import ExpenseError

    rid, _ = expense_service.create_reimbursement(title="文档来源快照")
    iid = expense_service.create_invoice(reimbursement_id=rid, amount="18.20")
    pid = expense_service.create_payment(reimbursement_id=rid, amount="18.20", pay_date="2026-09-03")
    expense_service.confirm(rid)
    with pytest.raises(ExpenseError, match="草稿"):
        expense_service.update_invoice(iid, amount="19.20")
    with pytest.raises(ExpenseError, match="草稿"):
        expense_service.update_payment(pid, amount="19.20")


def test_cross_reimbursement_detach_is_rejected_without_partial_write(expense_service):
    from app.services.expenses import ExpenseError

    first, _ = expense_service.create_reimbursement(title="第一项")
    second, _ = expense_service.create_reimbursement(title="第二项")
    iid = expense_service.create_invoice(reimbursement_id=first, amount="9.90")
    with pytest.raises(ExpenseError, match="不属于"):
        expense_service.detach_invoice(second, iid)
    assert expense_service.get_invoice(iid)["reimbursement_id"] == first


def test_manual_match_is_atomic_and_records_both_sides(expense_service):
    iid = expense_service.create_invoice(amount="88.00")
    pid = expense_service.create_payment(amount="88.00", pay_date="2026-09-03")
    rid = expense_service.manual_match([iid], [pid], title="手工整理", approver="李老师")

    invoice = expense_service.get_invoice(iid)
    payment = expense_service.get_payment(pid)
    assert invoice["reimbursement_id"] == payment["reimbursement_id"] == rid
    assert invoice["matched_payment_ids"] == [pid]
    assert payment["matched_invoice_ids"] == [iid]


def test_manual_match_cannot_overwrite_an_existing_pair(expense_service):
    from app.services.expenses import ExpenseError

    first_invoice = expense_service.create_invoice(amount="14.00")
    first_payment = expense_service.create_payment(amount="14.00", pay_date="2026-09-03")
    rid = expense_service.manual_match([first_invoice], [first_payment], title="已匹配")
    second_payment = expense_service.create_payment(amount="14.00", pay_date="2026-09-03")
    with pytest.raises(ExpenseError, match="状态不可匹配"):
        expense_service.manual_match([first_invoice], [second_payment], rid=rid)
    assert expense_service.get_payment(first_payment)["matched_invoice_ids"] == [first_invoice]
    assert expense_service.get_payment(second_payment)["reimbursement_id"] is None


def test_list_api_exposes_pagination_total_and_next(expense_service, tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-page-test", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED", "EXPENSE_SERVICE": expense_service,
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": 1})
    response = client.get("/expense/api/invoices?limit=2&offset=0")
    payload = response.get_json()
    assert response.status_code == 200
    assert len(payload["invoices"]) <= 2
    assert isinstance(payload["total"], int)
    assert payload["next_offset"] == (2 if payload["total"] > 2 else None)


def test_concurrent_number_generation_has_no_duplicates(expense_service):
    def create(index):
        return expense_service.create_reimbursement(title=f"并发 {index}")[1]

    with ThreadPoolExecutor(max_workers=2) as pool:
        numbers = list(pool.map(create, range(2)))
    assert len(set(numbers)) == 2


def test_concurrent_payment_number_generation_has_no_duplicates(expense_service):
    def create(_index):
        pid = expense_service.create_payment(amount="1.00", pay_date="2026-09-03")
        return expense_service.get_payment(pid)["payment_no"]

    with ThreadPoolExecutor(max_workers=2) as pool:
        numbers = list(pool.map(create, range(2)))
    assert len(set(numbers)) == 2


def test_auto_match_rechecks_under_lock_and_does_not_reuse_records(expense_service):
    iid = expense_service.create_invoice(amount="66.60")
    pid = expense_service.create_payment(amount="66.60", pay_date="2026-09-03")

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: expense_service.auto_match(approver="王老师"), range(2)))

    assert sum(result["total_payments_matched"] for result in results) == 1
    assert expense_service.get_invoice(iid)["matched_payment_ids"] == [pid]
    assert expense_service.get_payment(pid)["matched_invoice_ids"] == [iid]


def test_documents_json_updates_are_serialized_by_row_lock(expense_service):
    rid, _ = expense_service.create_reimbursement(title="单据并发")

    def add(index):
        return expense_service.add_document(rid, {"doc_type": "差旅费报销凭证", "fields": {"序号": str(index)}})

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(add, range(2)))
    assert len(expense_service.get_documents(rid)) == 2


def test_first_attachment_failure_rolls_back_business_row(expense_service):
    from app.services.expenses import ExpenseService

    class Writer:
        def __init__(self, connection): self.connection = connection
        def execute(self, statement, parameters=None): self.connection.execute(statement, parameters or {})

    class FailingAtomicFiles:
        def __init__(self, engine): self.engine = engine
        def upload_new_object(self, stream, **kwargs):
            with self.engine.begin() as connection:
                kwargs["create_metadata"](Writer(connection))
                raise RuntimeError("injected file failure")

    service = ExpenseService(expense_service.repository, file_service=FailingAtomicFiles(expense_service.repository.engine))
    with service.repository.engine.connect() as connection:
        before = connection.scalar(sa.select(sa.func.count()).select_from(service.repository.invoices))
    with pytest.raises(RuntimeError, match="injected"):
        service.create_invoice_with_upload(
            stream=None, original_name="proof.png", actor_user_id=1, request_id="rollback",
            amount="1.00", date="2026-09-03",
        )
    with service.repository.engine.connect() as connection:
        after = connection.scalar(sa.select(sa.func.count()).select_from(service.repository.invoices))
    assert after == before


def test_generated_document_batch_failure_does_not_change_status(expense_service):
    class FailingFiles:
        def upload_expense_documents_atomic(self, *_args, **_kwargs):
            raise RuntimeError("injected generated-file failure")

    from app.services.expenses import ExpenseService
    service = ExpenseService(expense_service.repository, file_service=FailingFiles())
    rid, _ = service.create_reimbursement(title="文档原子性")
    with service.repository.engine.begin() as connection:
        service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
    with pytest.raises(RuntimeError, match="generated-file"):
        service.persist_generated_documents(
            rid, [
                ("settlement.docx", io.BytesIO(b"PKfake-one")),
                ("approval.docx", io.BytesIO(b"PKfake-two")),
            ], actor_user_id=1, request_id="doc-fail"
        )
    assert service.get_reimbursement(rid)["status"] == "已确认"


def _valid_docx_bytes(text):
    from docx import Document
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_generated_document_batch_persists_two_real_files_and_state_atomically(expense_service, tmp_path):
    from docx import Document
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','张老师') RETURNING id"
        ), {"username": f"doc-success-{os.getpid()}"})
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "success", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="真实双文档")
    with engine.begin() as connection:
        service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
    try:
        stored = service.persist_generated_documents(
            rid,
            [("结算单.docx", io.BytesIO(_valid_docx_bytes("SETTLEMENT"))),
             ("审批单.docx", io.BytesIO(_valid_docx_bytes("APPROVAL")))],
            actor_user_id=user_id, request_id="doc-success",
        )
        assert len(stored) == 2
        assert service.get_reimbursement(rid)["status"] == "已生成文档"
        linked = files.list_for_object(object_type="EXPENSE", object_id=str(rid))
        assert len(linked) == 2
        with engine.connect() as connection:
            assert connection.scalar(sa.text(
                "SELECT count(*) FROM stored_file_versions WHERE file_id IN (SELECT file_id FROM object_files WHERE object_type='EXPENSE' AND object_id=:rid)"
            ), {"rid": str(rid)}) == 2
        for item in stored:
            opened = files.open_version_stream(item["fileId"], 1, object_type="EXPENSE", object_id=str(rid))
            try:
                Document(opened["stream"])
            finally:
                opened["stream"].close()
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM object_files WHERE object_type='EXPENSE' AND object_id=:rid"), {"rid": str(rid)})
            connection.execute(sa.text("DELETE FROM stored_file_versions WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:uid)"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM stored_files WHERE created_by=:uid"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})


def test_second_generated_document_move_failure_leaves_no_db_or_disk_residue(expense_service, tmp_path, monkeypatch):
    import app.services.files as files_module
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','李老师') RETURNING id"
        ), {"username": f"doc-failure-{os.getpid()}"})
    root = tmp_path / "failure"
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=root, max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="双文档回滚")
    with engine.begin() as connection:
        service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
    real_replace = files_module.os.replace
    moves = 0

    def fail_second(source, destination):
        nonlocal moves
        moves += 1
        if moves == 2:
            raise OSError("injected second move failure")
        return real_replace(source, destination)

    monkeypatch.setattr(files_module.os, "replace", fail_second)
    try:
        with pytest.raises(OSError, match="second move"):
            service.persist_generated_documents(
                rid,
                [("结算单.docx", io.BytesIO(_valid_docx_bytes("ONE"))),
                 ("审批单.docx", io.BytesIO(_valid_docx_bytes("TWO")))],
                actor_user_id=user_id, request_id="doc-failure",
            )
        assert service.get_reimbursement(rid)["status"] == "已确认"
        with engine.connect() as connection:
            assert connection.scalar(sa.text("SELECT count(*) FROM object_files WHERE object_type='EXPENSE' AND object_id=:rid"), {"rid": str(rid)}) == 0
            assert connection.scalar(sa.text("SELECT count(*) FROM stored_files WHERE created_by=:uid"), {"uid": user_id}) == 0
        assert not [path for path in root.rglob('*') if path.is_file()]
    finally:
        monkeypatch.setattr(files_module.os, "replace", real_replace)
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})


def test_finance_attachment_writes_reject_confirmed_parent_without_orphans(expense_service, tmp_path):
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService, FileServiceError

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','王老师') RETURNING id"
        ), {"username": f"finance-policy-{os.getpid()}"})
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "policy", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="已确认不可写")
    with engine.begin() as connection:
        service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
        before = connection.scalar(sa.select(sa.func.count()).select_from(service.repository.invoices))
    try:
        with pytest.raises(FileServiceError, match="终态"):
            service.create_invoice_with_upload(
                io.BytesIO(b"\x89PNG\r\n\x1a\nblocked"), "blocked.png", reimbursement_id=rid,
                actor_user_id=user_id, request_id="blocked-child", amount="1.00",
            )
        with pytest.raises(FileServiceError, match="终态"):
            files.upload(io.BytesIO(b"\x89PNG\r\n\x1a\nblocked"), original_name="blocked.png", object_type="EXPENSE", object_id=str(rid), actor_user_id=user_id, request_id="blocked-parent")
        with engine.connect() as connection:
            assert connection.scalar(sa.select(sa.func.count()).select_from(service.repository.invoices)) == before
            assert connection.scalar(sa.text("SELECT count(*) FROM stored_files WHERE created_by=:uid"), {"uid": user_id}) == 0
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})


def test_deleting_draft_expense_archives_its_controlled_files(expense_service, tmp_path):
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','王老师') RETURNING id"
        ), {"username": f"expense-delete-{os.getpid()}"})
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "delete", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="删除附件归档")
    try:
        stored = files.upload(
            io.BytesIO(b"\x89PNG\r\n\x1a\nexpense"), original_name="expense.png",
            object_type="EXPENSE", object_id=str(rid), actor_user_id=user_id, request_id="expense-delete",
        )
        service.delete_reimbursement(rid)
        assert service.get_reimbursement(rid) is None
        with engine.connect() as connection:
            assert connection.scalar(sa.text("SELECT count(*) FROM object_files WHERE object_type='EXPENSE' AND object_id=:rid"), {"rid": str(rid)}) == 0
            assert connection.scalar(sa.text("SELECT status FROM stored_files WHERE id=:fid"), {"fid": stored["fileId"]}) == "ARCHIVED"
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM stored_file_versions WHERE file_id=:fid"), {"fid": locals().get("stored", {}).get("fileId")})
            connection.execute(sa.text("DELETE FROM stored_files WHERE id=:fid"), {"fid": locals().get("stored", {}).get("fileId")})
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})


def test_cross_detach_operations_complete_without_deadlock(expense_service):
    from app.services.expenses import ExpenseError

    iid = expense_service.create_invoice(amount="31.00")
    pid = expense_service.create_payment(amount="31.00", pay_date="2026-09-03")
    rid = expense_service.manual_match([iid], [pid], title="锁顺序")

    def detach(kind):
        try:
            if kind == "invoice": expense_service.detach_invoice(rid, iid)
            else: expense_service.detach_payment(rid, pid)
            return "ok"
        except ExpenseError as exc:
            return exc.code

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(detach, ("invoice", "payment")))
    assert "ok" in results
    assert expense_service.get_invoice(iid)["reimbursement_id"] is None
    assert expense_service.get_payment(pid)["reimbursement_id"] is None


def test_manual_match_failure_injection_rolls_back_all_relations(expense_service, monkeypatch):
    iid = expense_service.create_invoice(amount="21.00")
    pid = expense_service.create_payment(amount="21.00", pay_date="2026-09-03")
    original = expense_service.repository.update_payment

    def fail_after_invoice(*_args, **_kwargs):
        raise RuntimeError("injected relation failure")

    monkeypatch.setattr(expense_service.repository, "update_payment", fail_after_invoice)
    with pytest.raises(RuntimeError, match="injected relation failure"):
        expense_service.manual_match([iid], [pid], title="rollback")
    monkeypatch.setattr(expense_service.repository, "update_payment", original)
    assert expense_service.get_invoice(iid)["reimbursement_id"] is None
    assert expense_service.get_payment(pid)["reimbursement_id"] is None


def test_first_expense_attachment_uses_controlled_file_service_atomically(expense_service, tmp_path):
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event):
            return event

    engine = expense_service.repository.engine
    username = f"expense-file-{os.getpid()}"
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','张老师') RETURNING id"
        ), {"username": username})
    files = FileService(
        FilesRepository(engine), AuditRecorder(),
        storage_root=tmp_path / "files", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024,
    )
    service = ExpenseService(expense_service.repository, file_service=files)
    try:
        iid, stored = service.create_invoice_with_upload(
            io.BytesIO(b"\x89PNG\r\n\x1a\ncontrolled"), "invoice.png",
            actor_user_id=user_id, request_id="expense-upload", amount="7.50", date="2026-09-03",
        )
        assert stored["storagePath"].startswith("2026/")
        invoice = service.get_invoice(iid)
        assert invoice["file_path"] is None
        listed = files.list_for_object(object_type="INVOICE", object_id=str(iid))
        assert len(listed) == 1
        file_id = listed[0]["fileId"]
        assert service.get_invoice(iid)["files"][0]["fileId"] == file_id
        opened = files.open_version_stream(file_id, 1, object_type="INVOICE", object_id=str(iid))
        try:
            assert opened["stream"].read().startswith(b"\x89PNG")
        finally:
            opened["stream"].close()
        service.delete_invoice(iid)
        assert service.get_invoice(iid) is None
        with engine.connect() as connection:
            assert connection.scalar(sa.text("SELECT count(*) FROM object_files WHERE object_type='INVOICE' AND object_id=:id"), {"id": str(iid)}) == 0
            assert connection.scalar(sa.text("SELECT status FROM stored_files WHERE id=:id"), {"id": file_id}) == "ARCHIVED"
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM object_files WHERE object_type='INVOICE' AND object_id=:id"), {"id": str(locals().get("iid", -1))})
            connection.execute(sa.text("DELETE FROM stored_file_versions WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:id)"), {"id": user_id})
            connection.execute(sa.text("DELETE FROM stored_files WHERE created_by=:id"), {"id": user_id})
            connection.execute(sa.text("DELETE FROM expense_invoice WHERE id=:id"), {"id": locals().get("iid", -1)})
            connection.execute(sa.text("DELETE FROM users WHERE id=:id"), {"id": user_id})


def test_attached_invoice_validates_draft_parent_and_recalculates_total_in_file_transaction(
    expense_service, tmp_path, request
):
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseError, ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event):
            return event

    engine = expense_service.repository.engine
    username = f"expense-parent-{os.getpid()}"
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','张老师') RETURNING id"
        ), {"username": username})
    state = {"rid": None}

    def cleanup():
        with engine.begin() as connection:
            connection.execute(sa.text(
                "DELETE FROM object_files WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:uid)"
            ), {"uid": user_id})
            connection.execute(sa.text(
                "DELETE FROM stored_file_versions WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:uid)"
            ), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM stored_files WHERE created_by=:uid"), {"uid": user_id})
            if state["rid"] is not None:
                connection.execute(sa.text("DELETE FROM expense_invoice_item WHERE invoice_id IN (SELECT id FROM expense_invoice WHERE reimbursement_id=:rid)"), {"rid": state["rid"]})
                connection.execute(sa.text("DELETE FROM expense_invoice WHERE reimbursement_id=:rid"), {"rid": state["rid"]})
                connection.execute(sa.text("DELETE FROM expense_payment WHERE reimbursement_id=:rid"), {"rid": state["rid"]})
                connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": state["rid"]})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})

    request.addfinalizer(cleanup)
    files = FileService(
        FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "parent-files",
        max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024,
    )
    service = ExpenseService(expense_service.repository, file_service=files)
    rid, _ = service.create_reimbursement(title="附件发票合同")
    state["rid"] = rid

    iid, _ = service.create_invoice_with_upload(
        io.BytesIO(b"invoice contract"), "invoice.txt", actor_user_id=user_id,
        request_id="invoice-parent", reimbursement_id=rid, amount="1280.00",
        date="2026-09-04",
    )
    assert service.get_invoice(iid)["reimbursement_id"] == rid
    assert service.get_reimbursement(rid)["total_amount"] == 1280.0

    with engine.begin() as connection:
        expense_service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
    with pytest.raises(ExpenseError) as confirmed:
        service.create_invoice_with_upload(
            io.BytesIO(b"second invoice"), "second.txt", actor_user_id=user_id,
            request_id="invoice-confirmed", reimbursement_id=rid, amount="1.00",
        )
    assert confirmed.value.code == "INVALID_STATUS"
    with pytest.raises(ExpenseError) as missing:
        service.create_invoice_with_upload(
            io.BytesIO(b"missing parent"), "missing.txt", actor_user_id=user_id,
            request_id="invoice-missing", reimbursement_id=999999999, amount="1.00",
        )
    assert missing.value.code == "NOT_FOUND"


@pytest.mark.parametrize("doc_type", ["差旅费报销凭证", "因公出差审批单", "伙食补助费申报表"])
def test_three_shipped_docx_templates_generate_real_word_files(doc_type):
    from docx import Document
    from app.document_engine import DocumentFiller

    generated = DocumentFiller().generate(doc_type, {}, {"session": {}, "invoices": [], "payments": []})
    document = Document(io.BytesIO(generated))
    assert generated.startswith(b"PK")
    assert document.paragraphs or document.tables


def test_json_only_purchase_template_and_template_traversal_fail_clearly():
    from app.document_engine import DocumentFiller

    filler = DocumentFiller()
    with pytest.raises(FileNotFoundError, match="docx"):
        filler.generate("科研物资采购申请单", {}, {})
    with pytest.raises(FileNotFoundError, match="模板不存在"):
        filler.generate("../科研物资采购申请单", {}, {})


def test_multiple_documents_are_really_present_in_merged_word():
    from docx import Document
    from flask import Flask
    from app.routes.documents import _merge_docx_and_images

    parts = []
    for token in ("FIRST-UNIQUE-CONTENT", "SECOND-UNIQUE-CONTENT"):
        document = Document()
        document.add_paragraph(token)
        output = io.BytesIO()
        document.save(output)
        parts.append((token, output.getvalue()))
    app = Flask(__name__)
    with app.app_context():
        merged = Document(io.BytesIO(_merge_docx_and_images(parts, [], [])))
    text = "\n".join(paragraph.text for paragraph in merged.paragraphs)
    assert "FIRST-UNIQUE-CONTENT" in text
    assert "SECOND-UNIQUE-CONTENT" in text


def test_merge_fails_if_any_selected_docx_is_invalid():
    from flask import Flask
    from app.routes.documents import _merge_docx_and_images

    with Flask(__name__).app_context(), pytest.raises(Exception):
        _merge_docx_and_images([("broken", b"not-a-docx")], [], [])


def test_currency_helpers_keep_decimal_exactness():
    from app.document_engine import _cn_number, _parse_amount

    assert _parse_amount("0.29") == Decimal("0.29")
    assert _cn_number(Decimal("0.29")) == "贰角玖分"


def test_legacy_page_urls_redirect_instead_of_missing_template_500(tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-route-test",
        "DATA_DIR": str(tmp_path), "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "SECURITY_AUTH_ENABLED": False, "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state["user"] = "teacher"
    for path in ("/expense/upload", "/expense/payments", "/expense/pending", "/expense/fill"):
        response = client.get(path)
        assert response.status_code == 302
        assert response.headers["Location"].endswith("/expense/records")


def test_expense_api_and_page_url_contract_is_still_registered(tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-url-test",
        "DATA_DIR": str(tmp_path), "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "SECURITY_AUTH_ENABLED": False, "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
    })
    rules = {rule.rule for rule in app.url_map.iter_rules()}
    assert {
        "/expense/", "/expense/records", "/expense/upload", "/expense/payments",
        "/expense/pending", "/expense/fill", "/expense/approvals",
        "/expense/documents/", "/expense/documents/new/<int:rid>",
        "/expense/documents/templates", "/expense/documents/print/<int:rid>",
        "/expense/api/reimbursements", "/expense/api/invoices", "/expense/api/payments",
        "/expense/api/manual_match", "/expense/api/match",
    } <= rules


def test_upload_survives_ocr_outage_and_internal_errors_do_not_leak(tmp_path, monkeypatch):
    from app import create_app
    import app.routes.expense as expense_routes

    class InspectionFiles:
        @contextmanager
        def inspect_upload(self, stream, original_name):
            path = tmp_path / original_name
            path.write_bytes(stream.read())
            try:
                yield path
            finally:
                path.unlink(missing_ok=True)

    class FakeExpenseService:
        def find_duplicate_invoice(self, *_args): return None
        def create_invoice_with_upload(self, stream, original_name, **kwargs):
            assert stream.read().startswith(b"\x89PNG")
            return 17, {"fileId": "00000000-0000-0000-0000-000000000017", "versionNo": 1, "originalName": original_name}
        def list_reimbursements(self, *_args, **_kwargs):
            raise RuntimeError("/secret/internal/database/path")

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-upload-test",
        "DATA_DIR": str(tmp_path), "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "SECURITY_AUTH_ENABLED": False, "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
        "EXPENSE_SERVICE": FakeExpenseService(), "FILE_SERVICE": InspectionFiles(),
    })
    monkeypatch.setattr(expense_routes, "recognize_file", lambda _path: (_ for _ in ()).throw(RuntimeError("OCR offline")))
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": 1, "name": "张老师"})
    response = client.post(
        "/expense/api/upload", data={"type": "invoice", "file": (io.BytesIO(b"\x89PNG\r\n\x1a\nmanual"), "invoice.png")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert response.get_json()["manual_required"] is True

    leaked = client.get("/expense/api/reimbursements")
    assert leaked.status_code == 500
    assert "/secret/" not in leaked.get_data(as_text=True)


def test_auto_upload_is_validated_before_ocr_and_routes_by_filename(tmp_path, monkeypatch):
    from app import create_app
    import app.routes.expense as expense_routes
    from app.services.files import FileServiceError

    calls = []

    class InspectionFiles:
        @contextmanager
        def inspect_upload(self, stream, original_name):
            calls.append(("validate", original_name))
            if original_name.endswith(".exe"):
                raise FileServiceError("UNSUPPORTED_MEDIA_TYPE", "不支持该文件类型", 415)
            path = tmp_path / original_name
            path.write_bytes(stream.read())
            try:
                yield path
            finally:
                path.unlink(missing_ok=True)

    class AutoService:
        def find_duplicate_payment(self, *_args): return None
        def create_payment_with_upload(self, stream, original_name, **kwargs):
            calls.append(("create-payment", original_name))
            return 9, {"originalName": original_name}
        def auto_match(self, **_kwargs):
            raise RuntimeError("matching temporarily unavailable")

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-auto-test", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED", "EXPENSE_SERVICE": AutoService(),
        "FILE_SERVICE": InspectionFiles(),
    })
    monkeypatch.setattr(expense_routes, "recognize_payment", lambda _path: {"amount": "8.00", "pay_date": "2026-09-03"})
    monkeypatch.setattr(expense_routes, "recognize_file", lambda _path: (_ for _ in ()).throw(AssertionError("wrong OCR path")))
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": 1, "name": "张老师"})
    rejected = client.post("/expense/api/upload", data={"type": "auto", "file": (io.BytesIO(b"MZ"), "evil.exe")}, content_type="multipart/form-data")
    assert rejected.status_code == 415
    assert not any(kind.startswith("ocr") for kind, _ in calls)
    accepted = client.post("/expense/api/upload", data={"type": "auto", "file": (io.BytesIO(b"image"), "微信支付凭证.png")}, content_type="multipart/form-data")
    payload = accepted.get_json()
    assert accepted.status_code == 200
    assert payload["success"] is True and payload["type"] == "payment"
    assert "match_warning" in payload


def test_reimbursement_upload_compatibility_url_is_registered(tmp_path):
    from app import create_app
    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-route-alias", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
    })
    assert "/expense/api/reimbursements/<int:rid>/upload" in {rule.rule for rule in app.url_map.iter_rules()}


def test_postgresql_route_contract_supports_crud_edit_and_confirm(expense_service, tmp_path):
    from app import create_app

    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-pg-route-test",
        "DATA_DIR": str(tmp_path), "SESSION_FILE_DIR": str(tmp_path / "sessions"),
        "SECURITY_AUTH_ENABLED": False, "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED",
        "EXPENSE_SERVICE": expense_service,
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": 1, "name": "张老师"})

    created = client.post("/expense/api/reimbursements", json={"title": "路由验证"})
    assert created.status_code == 200
    payload = created.get_json()
    assert payload["success"] is True
    rid = payload["id"]
    expense_service.create_invoice(reimbursement_id=rid, amount="31.20", invoice_no="ROUTE-INV")
    expense_service.create_invoice(reimbursement_id=rid, amount="0.00", invoice_no="ROUTE-INV-2")
    iid = expense_service.create_invoice(reimbursement_id=rid, amount="0.00", invoice_no="ROUTE-INV-3")
    expense_service.create_payment(reimbursement_id=rid, amount="31.20", pay_date="2026-09-03")

    edited = client.put(f"/expense/api/invoices/{iid}", json={"amount": 0.0, "seller": "<img src=x onerror=alert(1)>"})
    assert edited.status_code == 200
    detail = client.get(f"/expense/api/reimbursements/{rid}?limit=2&offset=0").get_json()
    assert detail["success"] is True
    assert len(detail["invoices"]) == 2
    assert detail["invoice_total"] == 3
    assert detail["invoice_next_offset"] == 2
    edited_invoice = next(row for row in detail["invoices"] if row["id"] == iid)
    assert edited_invoice["seller"] == "<img src=x onerror=alert(1)>"
    assert edited_invoice["file_path"] is None
    assert edited_invoice["ocr_text"] == ""

    second_detail = client.get(
        f"/expense/api/reimbursements/{rid}?limit=2&invoice_offset=2&payment_offset=0"
    ).get_json()
    assert len(second_detail["invoices"]) == 1
    assert second_detail["invoice_next_offset"] is None
    assert len(second_detail["payments"]) == 1

    documents_page = client.get("/expense/documents/api/reimbursements?limit=1&offset=0").get_json()
    assert documents_page["success"] is True
    assert len(documents_page["reimbursements"]) == 1
    assert isinstance(documents_page["total"], int)
    assert documents_page["next_offset"] == (1 if documents_page["total"] > 1 else None)

    for index in range(3):
        expense_service.add_document(rid, {"doc_type": f"单据-{index}", "fields": {}})
    document_items = client.get(
        f"/expense/documents/api/reimbursements/{rid}/documents?limit=2&offset=0"
    ).get_json()
    assert len(document_items["documents"]) == 2
    assert document_items["total"] == 3
    assert document_items["next_offset"] == 2

    confirmed = client.post(f"/expense/api/reimbursements/{rid}/confirm")
    assert confirmed.status_code == 200
    assert expense_service.get_reimbursement(rid)["status"] == "已确认"
    rejected = client.put(f"/expense/api/invoices/{iid}", json={"amount": 32.0})
    assert rejected.status_code == 409
    assert rejected.get_json()["code"] == "INVALID_STATUS"


def test_postgresql_reimbursement_auto_upload_route_attaches_real_file(expense_service, tmp_path, monkeypatch):
    from app import create_app
    import app.routes.expense as expense_routes
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','张老师') RETURNING id"
        ), {"username": f"route-upload-{os.getpid()}"})
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "route-files", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="页面上传")
    monkeypatch.setattr(expense_routes, "recognize_file", lambda _path: {"fields": {"invoice_no": "AUTO-ROUTE", "amount": "9.80", "date": "2026-09-03"}, "text": "ignored"})
    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-route-upload", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions-upload"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED", "DATABASE_ENGINE": engine,
        "EXPENSE_SERVICE": service, "FILE_SERVICE": files,
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": user_id, "name": "张老师"})
    try:
        response = client.post(
            f"/expense/api/reimbursements/{rid}/upload",
            data={"type": "auto", "file": (io.BytesIO(b"\x89PNG\r\n\x1a\ninvoice"), "scan.png")},
            content_type="multipart/form-data",
        )
        payload = response.get_json()
        assert response.status_code == 200
        assert payload["success"] is True and payload["type"] == "invoice"
        invoice = service.get_invoice(payload["record_id"])
        assert invoice["reimbursement_id"] == rid
        assert len(invoice["files"]) == 1
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM object_files WHERE object_type='INVOICE' AND object_id IN (SELECT id::text FROM expense_invoice WHERE reimbursement_id=:rid)"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM stored_file_versions WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:uid)"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM stored_files WHERE created_by=:uid"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM expense_invoice WHERE reimbursement_id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})


def test_postgresql_generate_route_returns_two_downloadable_docx_files(expense_service, tmp_path):
    from docx import Document
    from app import create_app
    from app.repositories.files import FilesRepository
    from app.services.expenses import ExpenseService
    from app.services.files import FileService

    class AuditRecorder:
        def record(self, connection, **event): return event

    engine = expense_service.repository.engine
    with engine.begin() as connection:
        user_id = connection.scalar(sa.text(
            "INSERT INTO users (username,password,role,name) VALUES (:username,'x','user','李老师') RETURNING id"
        ), {"username": f"route-docs-{os.getpid()}"})
    files = FileService(FilesRepository(engine), AuditRecorder(), storage_root=tmp_path / "route-docs", max_bytes=1024 * 1024, preview_max_bytes=1024 * 1024)
    service = ExpenseService(expense_service.repository, audit_service=AuditRecorder(), file_service=files)
    rid, _ = service.create_reimbursement(title="生成两份文档")
    with engine.begin() as connection:
        service.repository.update_reimbursement(connection, rid, {"status": "已确认"})
    app = create_app({
        "TESTING": True, "SECRET_KEY": "expense-route-docs", "DATA_DIR": str(tmp_path),
        "SESSION_FILE_DIR": str(tmp_path / "sessions-docs"), "SECURITY_AUTH_ENABLED": False,
        "CSRF_ENABLED": False, "AI_PROVIDER": "DISABLED", "DATABASE_ENGINE": engine,
        "EXPENSE_SERVICE": service, "FILE_SERVICE": files,
    })
    client = app.test_client()
    with client.session_transaction() as state:
        state.update({"user": "teacher", "user_id": user_id, "name": "李老师"})
    try:
        response = client.post(f"/expense/api/reimbursements/{rid}/generate_docs")
        payload = response.get_json()
        assert response.status_code == 200
        assert payload["success"] is True and payload["settlement"] is True
        assert payload["settlement_path"] != payload["approval_path"]
        assert len(payload["documents"]) == 2
        for key in ("settlement_path", "approval_path"):
            downloaded = client.get(payload[key])
            assert downloaded.status_code == 200
            Document(io.BytesIO(downloaded.data))
        assert service.get_reimbursement(rid)["status"] == "已生成文档"
    finally:
        with engine.begin() as connection:
            connection.execute(sa.text("DELETE FROM object_files WHERE object_type='EXPENSE' AND object_id=:rid"), {"rid": str(rid)})
            connection.execute(sa.text("DELETE FROM stored_file_versions WHERE file_id IN (SELECT id FROM stored_files WHERE created_by=:uid)"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM stored_files WHERE created_by=:uid"), {"uid": user_id})
            connection.execute(sa.text("DELETE FROM expense_reimbursement WHERE id=:rid"), {"rid": rid})
            connection.execute(sa.text("DELETE FROM users WHERE id=:uid"), {"uid": user_id})
