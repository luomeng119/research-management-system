from __future__ import annotations

import io
import os
from queue import Queue
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event

import pytest
import sqlalchemy as sa
from flask import Flask
from sqlalchemy.pool import StaticPool

from app.repositories.files import FilesRepository
from app.services.files import FileService, FileServiceError


def _schema(engine):
    metadata = sa.MetaData()
    def common():
        return (
            sa.Column("id", sa.String(36), primary_key=True),
            sa.Column("created_by", sa.Integer),
            sa.Column("updated_by", sa.Integer),
            sa.Column("version", sa.Integer, nullable=False, server_default="1"),
        )
    sa.Table(
        "stored_files", metadata, *common(),
        sa.Column("business_id", sa.Text, nullable=False, unique=True),
        sa.Column("original_name", sa.Text, nullable=False),
        sa.Column("media_type", sa.Text, nullable=False),
        sa.Column("status", sa.Text, nullable=False, server_default="ACTIVE"),
    )
    sa.Table(
        "stored_file_versions", metadata, *common(),
        sa.Column("file_id", sa.String(36), nullable=False),
        sa.Column("version_no", sa.Integer, nullable=False),
        sa.Column("storage_path", sa.Text, nullable=False),
        sa.Column("sha256", sa.String(64), nullable=False),
        sa.Column("size_bytes", sa.BigInteger, nullable=False),
        sa.Column("media_type", sa.Text, nullable=False),
        sa.UniqueConstraint("file_id", "version_no"),
    )
    sa.Table(
        "object_files", metadata, *common(),
        sa.Column("object_type", sa.Text, nullable=False),
        sa.Column("object_id", sa.Text, nullable=False),
        sa.Column("file_id", sa.String(36), nullable=False),
        sa.Column("purpose", sa.Text),
        sa.UniqueConstraint("object_type", "object_id", "file_id"),
    )
    sa.Table("expense_reimbursement", metadata, sa.Column("id", sa.Integer, primary_key=True))
    sa.Table(
        "standards", metadata,
        sa.Column("doc_id", sa.Text, primary_key=True),
        sa.Column("name", sa.Text, nullable=False),
        sa.Column("status", sa.Text, nullable=False, server_default="ACTIVE"),
    )
    sa.Table(
        "reference_template_items", metadata,
        sa.Column("template_id", sa.Text, primary_key=True),
        sa.Column("display_name", sa.Text, nullable=False),
        sa.Column("status", sa.Text, nullable=False, server_default="ACTIVE"),
    )
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(metadata.tables["expense_reimbursement"].insert().values(id=7))


class Audit:
    def __init__(self, fail=False):
        self.fail = fail
        self.events = []

    def record(self, connection, **event):
        if self.fail:
            raise RuntimeError("audit unavailable")
        self.events.append(event)
        return "audit-id"


@pytest.fixture()
def engine():
    value = sa.create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    _schema(value)
    return value


@pytest.fixture()
def service(tmp_path, engine):
    return FileService(
        FilesRepository(engine),
        Audit(),
        storage_root=tmp_path / "files",
        max_bytes=1024,
        preview_max_bytes=512,
    )


def _pdf(payload=b"document"):
    return io.BytesIO(b"%PDF-1.7\n" + payload)


def _rows(engine, table_name):
    table = sa.Table(table_name, sa.MetaData(), autoload_with=engine)
    with engine.connect() as connection:
        return [dict(row) for row in connection.execute(sa.select(table)).mappings()]


def test_project_path_upload_keeps_versions_and_distinguishes_folders(service, engine):
    metadata = sa.MetaData()
    projects = sa.Table("project_registry", metadata,
                        sa.Column("business_id", sa.Text, primary_key=True))
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(projects.insert().values(business_id="PRJ-001"))
    def upload(folder, payload):
        return service.upload_project_path(
            _pdf(payload), original_name="report.pdf", folder=folder,
            project_id="PRJ-001", actor_user_id=1, request_id="req-path",
        )
    first = upload("任务输入文件", b"first")
    second = upload("任务输入文件", b"second")
    other = upload("研究成果文件", b"output")
    assert first["fileId"] == second["fileId"] != other["fileId"]
    assert [first["versionNo"], second["versionNo"], other["versionNo"]] == [1, 2, 1]
    assert (service.storage_root / first["storagePath"]).read_bytes() == b"%PDF-1.7\nfirst"
    assert (service.storage_root / second["storagePath"]).read_bytes() == b"%PDF-1.7\nsecond"
    assert len(_rows(engine, "stored_files")) == 2
    assert len(_rows(engine, "stored_file_versions")) == 3
    assert {row["purpose"] for row in _rows(engine, "object_files")} == {
        "PROJECT_TREE:任务输入文件/report.pdf", "PROJECT_TREE:研究成果文件/report.pdf",
    }
    before = {path: path.read_bytes() for path in service.storage_root.rglob("*") if path.is_file()}
    service.audit_service.fail = True
    with pytest.raises(FileServiceError) as error:
        upload("任务输入文件", b"must roll back")
    assert error.value.code == "FILE_OPERATION_FAILED"
    assert len(_rows(engine, "stored_file_versions")) == 3
    assert {path: path.read_bytes() for path in service.storage_root.rglob("*") if path.is_file()} == before
    assert next(row for row in _rows(engine, "stored_files") if row["id"] == first["fileId"])["version"] == 2


@pytest.mark.parametrize("folder,name", [
    ("../outside", "report.pdf"), ("/absolute", "report.pdf"),
    ("a//b", "report.pdf"), ("a/./b", "report.pdf"),
    ("a/.history", "report.pdf"), ("a\\b", "report.pdf"),
    ("a\x00b", "report.pdf"), ("a", "../report.pdf"),
])
def test_project_path_rejects_ambiguous_names_before_storage(service, engine, folder, name):
    metadata = sa.MetaData()
    projects = sa.Table("project_registry", metadata,
                        sa.Column("business_id", sa.Text, primary_key=True))
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(projects.insert().values(business_id="PRJ-001"))
    with pytest.raises(FileServiceError) as error:
        service.upload_project_path(
            _pdf(), original_name=name, folder=folder, project_id="PRJ-001",
            actor_user_id=1, request_id="invalid-path",
        )
    assert error.value.code == "INVALID_PROJECT_PATH"
    assert _rows(engine, "stored_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_creates_uuid_path_hash_version_link_and_low_sensitivity_audit(service, engine):
    result = service.upload(
        _pdf(), original_name="研究报告.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-test",
    )

    assert result["versionNo"] == 1
    assert result["storagePath"].endswith(".pdf")
    assert "研究报告" not in result["storagePath"]
    stored_path = service.storage_root / result["storagePath"]
    assert stored_path.read_bytes().startswith(b"%PDF-")
    assert len(result["sha256"]) == 64
    assert len(_rows(engine, "stored_files")) == 1
    assert len(_rows(engine, "stored_file_versions")) == 1
    assert len(_rows(engine, "object_files")) == 1
    event = service.audit_service.events[-1]
    assert event["event_name"] == "file_operation_completed"
    assert set(event["properties"]) == {"operation", "file_type", "size_bucket"}
    assert "研究报告.pdf" not in repr(event)


def test_upload_new_object_creates_standard_metadata_and_first_file_atomically(service, engine):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_standard(connection):
        assert connection.in_transaction()
        connection.execute(
            standards.insert().values(doc_id="STD-2026-001", name="测试标准", status="ACTIVE")
        )

    result = service.upload_new_object(
        _pdf(b"standard"), original_name="测试标准.pdf", object_type="STANDARD",
        object_id="STD-2026-001", create_metadata=create_standard,
        actor_user_id=1, request_id="req-standard-create",
    )

    assert _rows(engine, "standards") == [
        {"doc_id": "STD-2026-001", "name": "测试标准", "status": "ACTIVE"}
    ]
    assert len(_rows(engine, "stored_files")) == 1
    assert len(_rows(engine, "stored_file_versions")) == 1
    assert len(_rows(engine, "object_files")) == 1
    assert (service.storage_root / result["storagePath"]).read_bytes().endswith(b"standard")
    assert service.audit_service.events[-1]["properties"]["operation"] == "UPLOAD"


def test_metadata_writer_supports_bounded_read_and_update_without_connection_escape(
    service, engine
):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_and_normalize(writer):
        writer.execute(
            standards.insert().values(
                doc_id="STD-2026-BOUNDED", name="待校验标准", status="ACTIVE"
            )
        )
        found = writer.one(
            sa.select(standards).where(standards.c.doc_id == "STD-2026-BOUNDED")
        )
        assert found["name"] == "待校验标准"
        assert writer.update(
            standards.update()
            .where(standards.c.doc_id == "STD-2026-BOUNDED")
            .values(name="已校验标准")
        ) == 1

    service.upload_new_object(
        _pdf(), original_name="已校验标准.pdf", object_type="STANDARD",
        object_id="STD-2026-BOUNDED", create_metadata=create_and_normalize,
        actor_user_id=1, request_id="req-bounded-metadata",
    )

    assert _rows(engine, "standards") == [{
        "doc_id": "STD-2026-BOUNDED", "name": "已校验标准", "status": "ACTIVE"
    }]


@pytest.mark.parametrize("operation", ["commit", "rollback", "begin", "begin_nested"])
def test_upload_new_object_rejects_callback_transaction_control_without_residue(
    service, engine, operation
):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)
    rejected = False

    def create_then_attempt_transaction_control(writer):
        nonlocal rejected
        writer.execute(
            standards.insert().values(
                doc_id=f"STD-2026-{operation}", name="事务控制探针", status="ACTIVE"
            )
        )
        try:
            getattr(writer, operation)()
        except RuntimeError as exc:
            if "metadata callback" not in str(exc):
                raise AssertionError("callback transaction control was not explicitly rejected") from exc
            rejected = True
            raise
        raise AssertionError("callback transaction control was not explicitly rejected")

    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="事务控制探针.pdf", object_type="STANDARD",
            object_id=f"STD-2026-{operation}", create_metadata=create_then_attempt_transaction_control,
            actor_user_id=1, request_id=f"req-standard-{operation}",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert rejected is True
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


@pytest.mark.parametrize("escape", ["result_connection", "text_commit"])
def test_upload_new_object_writer_rejects_result_and_sql_transaction_escapes(
    service, engine, escape
):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)
    rejected = False

    def create_then_attempt_writer_escape(writer):
        nonlocal rejected
        result = writer.execute(
            standards.insert().values(
                doc_id=f"STD-2026-{escape}", name="写入器逃逸探针", status="ACTIVE"
            )
        )
        if escape == "result_connection":
            try:
                result.connection.commit()
            except AttributeError:
                rejected = True
                raise RuntimeError("safe metadata writer result has no connection")
            raise AssertionError("metadata writer result exposed a connection")
        try:
            writer.execute(sa.text("COMMIT"))
        except RuntimeError as exc:
            if "Insert" not in str(exc):
                raise AssertionError("metadata SQL transaction control was not explicitly rejected") from exc
            rejected = True
            raise
        raise AssertionError("metadata SQL transaction control was not explicitly rejected")

    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="写入器逃逸探针.pdf", object_type="STANDARD",
            object_id=f"STD-2026-{escape}", create_metadata=create_then_attempt_writer_escape,
            actor_user_id=1, request_id=f"req-standard-{escape}",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert rejected is True
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_new_object_creator_failure_leaves_no_metadata_or_file_residue(service, engine):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_then_fail(connection):
        connection.execute(
            standards.insert().values(doc_id="STD-2026-FAIL", name="失败标准", status="ACTIVE")
        )
        raise RuntimeError("metadata unavailable")

    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="失败标准.pdf", object_type="STANDARD",
            object_id="STD-2026-FAIL", create_metadata=create_then_fail,
            actor_user_id=1, request_id="req-standard-creator-failure",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


@pytest.mark.parametrize(
    "repository_method",
    ["object_allows_file_write", "create_file", "create_version", "link_object"],
)
def test_upload_new_object_repository_failure_leaves_no_metadata_or_file_residue(
    service, engine, monkeypatch, repository_method
):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_standard(writer):
        writer.execute(
            standards.insert().values(
                doc_id=f"STD-2026-{repository_method}", name="失败标准", status="ACTIVE"
            )
        )

    def fail_repository_write(*_args, **_kwargs):
        raise RuntimeError(f"injected {repository_method} failure")

    monkeypatch.setattr(service.repository, repository_method, fail_repository_write)
    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="失败标准.pdf", object_type="STANDARD",
            object_id=f"STD-2026-{repository_method}", create_metadata=create_standard,
            actor_user_id=1, request_id=f"req-standard-{repository_method}-failure",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert service.audit_service.events == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_new_object_rechecks_the_new_object_write_lock_and_rolls_back(service, engine):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_archived_standard(connection):
        connection.execute(
            standards.insert().values(doc_id="STD-2026-ARCHIVED", name="只读标准", status="ARCHIVED")
        )

    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="只读标准.pdf", object_type="STANDARD",
            object_id="STD-2026-ARCHIVED", create_metadata=create_archived_standard,
            actor_user_id=1, request_id="req-standard-read-only",
        )

    assert error.value.code == "OBJECT_READ_ONLY"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_new_object_audit_failure_rolls_back_business_metadata_and_file(tmp_path, engine):
    service = FileService(
        FilesRepository(engine), Audit(fail=True), storage_root=tmp_path / "files",
        max_bytes=1024, preview_max_bytes=512,
    )
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_standard(connection):
        connection.execute(
            standards.insert().values(doc_id="STD-2026-AUDIT", name="审计失败标准", status="ACTIVE")
        )

    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="审计失败标准.pdf", object_type="STANDARD",
            object_id="STD-2026-AUDIT", create_metadata=create_standard,
            actor_user_id=1, request_id="req-standard-audit-failure",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_new_object_final_move_failure_rolls_back_business_metadata_and_staging(
    service, engine, monkeypatch
):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_standard(connection):
        connection.execute(
            standards.insert().values(doc_id="STD-2026-MOVE", name="移动失败标准", status="ACTIVE")
        )

    def fail_replace(*_args):
        raise OSError("storage unavailable")

    monkeypatch.setattr("app.services.files.os.replace", fail_replace)
    with pytest.raises(FileServiceError) as error:
        service.upload_new_object(
            _pdf(), original_name="移动失败标准.pdf", object_type="STANDARD",
            object_id="STD-2026-MOVE", create_metadata=create_standard,
            actor_user_id=1, request_id="req-standard-move-failure",
        )

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_upload_new_object_commit_failure_removes_business_metadata_and_final_file(service, engine):
    standards = sa.Table("standards", sa.MetaData(), autoload_with=engine)

    def create_standard(connection):
        connection.execute(
            standards.insert().values(doc_id="STD-2026-COMMIT", name="提交失败标准", status="ACTIVE")
        )

    def fail_commit(connection):
        if connection.in_transaction():
            raise RuntimeError("commit unavailable")

    sa.event.listen(engine, "commit", fail_commit)
    try:
        with pytest.raises(FileServiceError) as error:
            service.upload_new_object(
                _pdf(), original_name="提交失败标准.pdf", object_type="STANDARD",
                object_id="STD-2026-COMMIT", create_metadata=create_standard,
                actor_user_id=1, request_id="req-standard-commit-failure",
            )
    finally:
        sa.event.remove(engine, "commit", fail_commit)

    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "standards") == []
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert _rows(engine, "object_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


@pytest.mark.parametrize(
    ("object_type", "table_name", "id_column", "name_column", "object_id"),
    [
        ("STANDARD", "standards", "doc_id", "name", "STD-2026-READ-ONLY"),
        ("TEMPLATE", "reference_template_items", "template_id", "display_name", "TPL-2026-READ-ONLY"),
    ],
)
def test_archived_retained_resource_rejects_upload_version_and_archive_mutations(
    service, engine, object_type, table_name, id_column, name_column, object_id
):
    resource_table = sa.Table(table_name, sa.MetaData(), autoload_with=engine)
    with engine.begin() as connection:
        connection.execute(
            resource_table.insert().values(**{id_column: object_id, name_column: "保留资源", "status": "ACTIVE"})
        )
    uploaded = service.upload(
        _pdf(b"v1"), original_name="保留资源.pdf", object_type=object_type, object_id=object_id,
        actor_user_id=1, request_id=f"req-{object_type}-active",
    )
    with engine.begin() as connection:
        connection.execute(
            resource_table.update().where(resource_table.c[id_column] == object_id).values(status="ARCHIVED")
        )

    with pytest.raises(FileServiceError) as upload_error:
        service.upload(
            _pdf(b"another"), original_name="保留资源.pdf",
            object_type=object_type, object_id=object_id,
            actor_user_id=1, request_id=f"req-{object_type}-upload",
        )
    with pytest.raises(FileServiceError) as version_error:
        service.add_version(
            uploaded["fileId"], _pdf(b"v2"), original_name="保留资源.pdf",
            object_type=object_type, object_id=object_id, expected_version=1,
            actor_user_id=1, request_id=f"req-{object_type}-version",
        )
    with pytest.raises(FileServiceError) as archive_error:
        service.archive(
            uploaded["fileId"], object_type=object_type, object_id=object_id,
            actor_user_id=1, request_id=f"req-{object_type}-archive",
        )

    assert upload_error.value.code == "OBJECT_READ_ONLY"
    assert version_error.value.code == "OBJECT_READ_ONLY"
    assert archive_error.value.code == "OBJECT_READ_ONLY"
    assert len(_rows(engine, "stored_file_versions")) == 1


def test_same_original_name_never_overwrites(service):
    first = service.upload(
        _pdf(b"one"), original_name="报告.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-1",
    )
    second = service.upload(
        _pdf(b"two"), original_name="报告.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-2",
    )

    assert first["fileId"] != second["fileId"]
    assert first["storagePath"] != second["storagePath"]
    assert (service.storage_root / first["storagePath"]).read_bytes().endswith(b"one")
    assert (service.storage_root / second["storagePath"]).read_bytes().endswith(b"two")


@pytest.mark.parametrize(
    ("name", "content", "code"),
    [
        ("../escape.pdf", b"%PDF-1.7", "INVALID_FILENAME"),
        ("folder\\escape.pdf", b"%PDF-1.7", "INVALID_FILENAME"),
        ("bad\nname.pdf", b"%PDF-1.7", "INVALID_FILENAME"),
        ("report.exe.pdf", b"%PDF-1.7", "INVALID_FILENAME"),
        ("fake.pdf", b"not a pdf", "UNSUPPORTED_MEDIA_TYPE"),
        ("empty.pdf", b"", "UNSUPPORTED_MEDIA_TYPE"),
        ("empty.txt", b"", "UNSUPPORTED_MEDIA_TYPE"),
        ("nul.txt", b"hello\x00world", "UNSUPPORTED_MEDIA_TYPE"),
        ("too-large.pdf", b"%PDF-" + b"x" * 2048, "FILE_TOO_LARGE"),
    ],
)
def test_invalid_uploads_are_rejected_without_database_or_disk_residue(
    service, engine, name, content, code
):
    with pytest.raises(FileServiceError) as error:
        service.upload(
            io.BytesIO(content), original_name=name, object_type="EXPENSE", object_id="7",
            actor_user_id=1, request_id="req-bad",
        )

    assert error.value.code == code
    assert _rows(engine, "stored_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_unknown_object_is_rejected_before_staging(service, engine):
    with pytest.raises(FileServiceError) as error:
        service.upload(
            _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="999",
            actor_user_id=1, request_id="req-missing",
        )
    assert error.value.code == "OBJECT_NOT_FOUND"
    assert _rows(engine, "stored_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_audit_failure_rolls_back_database_and_removes_final_file(tmp_path, engine):
    service = FileService(
        FilesRepository(engine), Audit(fail=True), storage_root=tmp_path / "files",
        max_bytes=1024, preview_max_bytes=512,
    )
    with pytest.raises(FileServiceError) as error:
        service.upload(
            _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
            actor_user_id=1, request_id="req-audit",
        )
    assert error.value.code == "FILE_OPERATION_FAILED"
    assert _rows(engine, "stored_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_new_version_preserves_v1_detects_conflict_and_archive_preserves_history(service, engine):
    first = service.upload(
        _pdf(b"v1"), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-v1",
    )
    assert service.count_versions(
        first["fileId"], object_type="EXPENSE", object_id="7"
    ) == 1
    v1_bytes = (service.storage_root / first["storagePath"]).read_bytes()
    second = service.add_version(
        first["fileId"], _pdf(b"v2"), original_name="report.pdf",
        object_type="EXPENSE", object_id="7", expected_version=1,
        actor_user_id=1, request_id="req-v2",
    )

    assert second["versionNo"] == 2
    assert service.count_versions(
        first["fileId"], object_type="EXPENSE", object_id="7"
    ) == 2
    assert second["storagePath"] != first["storagePath"]
    assert (service.storage_root / first["storagePath"]).read_bytes() == v1_bytes
    with pytest.raises(FileServiceError) as conflict:
        service.add_version(
            first["fileId"], _pdf(b"stale"), original_name="report.pdf",
            object_type="EXPENSE", object_id="7", expected_version=1,
            actor_user_id=1, request_id="req-stale",
        )
    assert conflict.value.code == "VERSION_CONFLICT"

    service.archive(
        first["fileId"], object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-archive",
    )
    assert len(_rows(engine, "stored_file_versions")) == 2
    archived_file = _rows(engine, "stored_files")[0]
    assert archived_file["status"] == "ARCHIVED"
    assert archived_file["version"] == 2
    assert all((service.storage_root / row["storage_path"]).exists() for row in _rows(engine, "stored_file_versions"))
    with pytest.raises(FileServiceError) as archived:
        service.open_version(first["fileId"], 2, object_type="EXPENSE", object_id="7")
    assert archived.value.code == "FILE_ARCHIVED"


def test_new_version_must_keep_the_original_file_type(service):
    uploaded = service.upload(
        _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-original",
    )
    with pytest.raises(FileServiceError) as error:
        service.add_version(
            uploaded["fileId"], io.BytesIO(b"plain text"), original_name="report.txt",
            object_type="EXPENSE", object_id="7", expected_version=1,
            actor_user_id=1, request_id="req-wrong-type",
        )
    assert error.value.code == "FILE_TYPE_MISMATCH"
    assert not any(path.suffix == ".txt" for path in service.storage_root.rglob("*"))


@pytest.mark.parametrize("storage_path", ["../../outside.pdf", "/tmp/outside.pdf", "%252e%252e/outside.pdf"])
def test_database_path_tampering_is_rejected(service, engine, storage_path):
    uploaded = service.upload(
        _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-path",
    )
    versions = sa.Table("stored_file_versions", sa.MetaData(), autoload_with=engine)
    with engine.begin() as connection:
        connection.execute(
            versions.update().where(versions.c.file_id == uploaded["fileId"]).values(storage_path=storage_path)
        )
    with pytest.raises(FileServiceError) as error:
        service.open_version(uploaded["fileId"], 1, object_type="EXPENSE", object_id="7")
    assert error.value.code == "FILE_NOT_FOUND"


def test_deleted_business_object_makes_existing_link_inaccessible(service, engine):
    uploaded = service.upload(
        _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-deleted-object",
    )
    expenses = sa.Table("expense_reimbursement", sa.MetaData(), autoload_with=engine)
    with engine.begin() as connection:
        connection.execute(expenses.delete().where(expenses.c.id == 7))
    with pytest.raises(FileServiceError) as error:
        service.open_version(uploaded["fileId"], 1, object_type="EXPENSE", object_id="7")
    assert error.value.code == "FILE_NOT_FOUND"


def test_symlinked_storage_file_is_rejected(service, tmp_path, engine):
    uploaded = service.upload(
        _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-link",
    )
    stored = service.storage_root / uploaded["storagePath"]
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"%PDF-outside")
    stored.unlink()
    stored.symlink_to(outside)
    with pytest.raises(FileServiceError) as error:
        service.open_version(uploaded["fileId"], 1, object_type="EXPENSE", object_id="7")
    assert error.value.code == "FILE_NOT_FOUND"


def test_symlinked_ancestor_outside_storage_root_name_remains_readable(tmp_path, engine):
    real_parent = tmp_path / "real-parent"
    real_parent.mkdir()
    alias_parent = tmp_path / "alias-parent"
    alias_parent.symlink_to(real_parent, target_is_directory=True)
    aliased = FileService(
        FilesRepository(engine), Audit(), storage_root=alias_parent / "files",
        max_bytes=1024, preview_max_bytes=512,
    )
    uploaded = aliased.upload(
        _pdf(), original_name="report.pdf", object_type="EXPENSE", object_id="7",
        actor_user_id=1, request_id="req-ancestor-link",
    )
    opened = aliased.open_version_stream(
        uploaded["fileId"], 1, object_type="EXPENSE", object_id="7"
    )
    try:
        assert opened["stream"].read().startswith(b"%PDF-")
    finally:
        opened["stream"].close()


@pytest.fixture()
def web_app(service):
    from app.routes.preview import bp as preview_bp
    from app.web.files import bp as files_bp

    app = Flask(__name__)
    app.config.update(TESTING=True, SECRET_KEY="file-test", SECURITY_AUTH_ENABLED=False)
    app.extensions["file_service"] = service
    app.register_blueprint(files_bp)
    app.register_blueprint(preview_bp)
    return app


def _business_client(app, role="BUSINESS_USER"):
    client = app.test_client()
    with client.session_transaction() as active_session:
        active_session.update(
            user_id=1,
            user="alice",
            name="Alice",
            role=role,
            account_version=1,
        )
    return client


def test_file_api_requires_auth_and_round_trips_upload_download(web_app):
    anonymous = web_app.test_client()
    assert anonymous.get("/api/files?objectType=EXPENSE&objectId=7").status_code == 401

    client = _business_client(web_app)
    uploaded = client.post(
        "/api/files",
        data={
            "objectType": "EXPENSE",
            "objectId": "7",
            "file": (_pdf(b"api"), "report.pdf"),
        },
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 201
    payload = uploaded.get_json()
    assert "storagePath" not in payload

    listing = client.get("/api/files?objectType=EXPENSE&objectId=7")
    assert listing.status_code == 200
    assert listing.get_json()["files"][0]["fileId"] == payload["fileId"]

    downloaded = client.get(
        f"/api/files/{payload['fileId']}/versions/1/download?objectType=EXPENSE&objectId=7"
    )
    assert downloaded.status_code == 200
    assert downloaded.data.endswith(b"api")
    assert downloaded.headers["X-Content-Type-Options"] == "nosniff"
    assert "attachment" in downloaded.headers["Content-Disposition"]


@pytest.mark.parametrize("fixture,name,media_type", [
    ("office-sample.doc", "原始材料.doc", "application/msword"),
    ("office-source.xls", "资源登记.xls", "application/vnd.ms-excel"),
    ("office-source.ppt", "研究汇报.ppt", "application/vnd.ms-powerpoint"),
])
def test_legacy_office_attachment_round_trips_without_online_conversion(web_app, service, fixture, name, media_type, monkeypatch):
    from pathlib import Path

    # Local textutil DOC / LibreOffice MS Excel 97 XLS conversions from the
    # adjacent text/CSV sources; actual Office files, not signature-only fakes.
    payload = (Path(__file__).parent / "fixtures/legacy" / fixture).read_bytes()
    service.max_bytes = 1_000_000
    service.preview_max_bytes = 1_000_000
    # A host without Office MIME registrations must still serve stable types.
    monkeypatch.setattr("app.services.files.mimetypes.guess_type", lambda *_: (None, None))
    client = _business_client(web_app)
    uploaded = client.post(
        "/api/files",
        data={"objectType": "EXPENSE", "objectId": "7",
              "file": (io.BytesIO(payload), name)},
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 201, uploaded.get_json()
    file_id = uploaded.get_json()["fileId"]
    preview = client.get(
        f"/api/files/{file_id}/versions/1/preview?objectType=EXPENSE&objectId=7"
    )
    assert preview.status_code == 409
    assert preview.get_json()["error"]["code"] == "PREVIEW_UNAVAILABLE"
    download_url = preview.get_json()["downloadUrl"]
    downloaded = client.get(download_url)
    assert downloaded.status_code == 200
    assert downloaded.data == payload
    assert downloaded.mimetype == media_type
    assert "attachment" in downloaded.headers["Content-Disposition"]
    assert downloaded.headers["X-Content-Type-Options"] == "nosniff"
    assert web_app.test_client().get(download_url).status_code == 401


@pytest.mark.parametrize("kind", ["plain", "signature_only", "truncated", "wrong_office_type", "missing_table"])
def test_binary_word_rejects_invalid_or_mislabeled_content(web_app, service, engine, kind):
    from pathlib import Path

    fixture_dir = Path(__file__).parent / "fixtures/legacy"
    payload = {
        "plain": b"not a Word file",
        "signature_only": bytes.fromhex("d0cf11e0a1b11ae1"),
        "truncated": (fixture_dir / "office-sample.doc").read_bytes()[:512],
        "wrong_office_type": (fixture_dir / "office-source.xls").read_bytes(),
        "missing_table": (fixture_dir / "office-sample.doc").read_bytes().replace(
            "1Table".encode("utf-16-le"), "NoTabl".encode("utf-16-le")
        ),
    }[kind]
    service.max_bytes = 100_000
    response = _business_client(web_app).post(
        "/api/files",
        data={"objectType": "EXPENSE", "objectId": "7",
              "file": (io.BytesIO(payload), "错误材料.doc")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 415
    assert response.get_json()["error"]["code"] == "UNSUPPORTED_MEDIA_TYPE"
    assert _rows(engine, "stored_files") == []
    assert _rows(engine, "stored_file_versions") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


@pytest.mark.parametrize("fixture,stream_name", [
    ("office-sample.doc", "WordDocument"), ("office-source.xls", "Workbook"),
    ("office-source.ppt", "PowerPoint Document"),
])
@pytest.mark.parametrize("defect", ["oversized_stream", "broken_chain"])
def test_binary_office_rejects_invalid_required_stream(web_app, service, engine, fixture, stream_name, defect):
    import struct
    from pathlib import Path

    payload = bytearray((Path(__file__).parent / "fixtures/legacy" / fixture).read_bytes())
    entry = payload.find((stream_name + "\0").encode("utf-16-le"))
    assert entry >= 0
    if defect == "oversized_stream":
        struct.pack_into("<Q", payload, entry + 120, 100_000_000)
    else:
        struct.pack_into("<I", payload, entry + 116, 0xFFFFFFFE)
    service.max_bytes = 1_000_000
    response = _business_client(web_app).post(
        "/api/files",
        data={"objectType": "EXPENSE", "objectId": "7",
              "file": (io.BytesIO(payload), fixture)},
        content_type="multipart/form-data",
    )
    assert response.status_code == 415
    assert response.get_json()["error"]["code"] == "UNSUPPORTED_MEDIA_TYPE"
    assert _rows(engine, "stored_files") == []
    assert not any(path.is_file() for path in service.storage_root.rglob("*"))


def test_maintenance_role_can_access_business_files(web_app):
    maintainer = _business_client(web_app, role="SYSTEM_MAINTAINER")
    assert maintainer.get("/api/files?objectType=EXPENSE&objectId=7").status_code == 200
    assert maintainer.post(
        "/api/files",
        data={
            "objectType": "EXPENSE", "objectId": "7",
            "file": (_pdf(), "report.pdf"),
        },
        content_type="multipart/form-data",
    ).status_code == 201


def test_preview_fallback_and_legacy_path_parameter_rejection(web_app, service):
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as value:
        value.writestr("readme.txt", "hello")
    archive.seek(0)
    client = _business_client(web_app)
    uploaded = client.post(
        "/api/files",
        data={
            "objectType": "EXPENSE",
            "objectId": "7",
            "file": (archive, "materials.zip"),
        },
        content_type="multipart/form-data",
    ).get_json()

    previewed = client.get(
        f"/api/files/{uploaded['fileId']}/versions/1/preview?objectType=EXPENSE&objectId=7"
    )
    assert previewed.status_code == 409
    assert previewed.get_json()["error"]["code"] == "PREVIEW_UNAVAILABLE"
    assert "/download?" in previewed.get_json()["downloadUrl"]
    assert service.audit_service.events[-1]["result"] == "FAILURE"
    assert service.audit_service.events[-1]["properties"]["operation"] == "PREVIEW"

    legacy = client.get("/preview/file?path=/etc/passwd&type=other")
    assert legacy.status_code == 400
    assert legacy.get_json()["error"]["code"] == "CONTROLLED_FILE_REFERENCE_REQUIRED"


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_office_preview_audit_failure_is_not_reported_as_unsupported(web_app, service, monkeypatch, extension):
    service.max_bytes = 100_000
    service.preview_max_bytes = 100_000
    data = io.BytesIO()
    if extension == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("研究依据")
        document.save(data)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["研究依据"])
        workbook.save(data)
        workbook.close()
    data.seek(0)
    client = _business_client(web_app)
    uploaded = client.post("/api/files", data={
        "objectType": "EXPENSE", "objectId": "7", "file": (data, "材料." + extension),
    }, content_type="multipart/form-data").get_json()
    record = service.audit_service.record
    def fail_success(connection, **values):
        if values.get("properties", {}).get("operation") == "PREVIEW" and values["result"] == "SUCCESS":
            raise RuntimeError("injected audit failure")
        return record(connection, **values)
    monkeypatch.setattr(service.audit_service, "record", fail_success)
    response = client.get(f"/api/files/{uploaded['fileId']}/versions/1/preview?objectType=EXPENSE&objectId=7")
    assert response.status_code == 500
    assert response.json["error"]["code"] == "FILE_OPERATION_FAILED"


def test_excel_preview_closes_workbook_when_text_budget_fails(web_app, service, monkeypatch):
    service.max_bytes = 100_000
    service.preview_max_bytes = 100_000
    import openpyxl
    import app.web.files as preview_module
    data = io.BytesIO()
    workbook = openpyxl.Workbook()
    workbook.active.append(["研究依据"])
    workbook.save(data)
    workbook.close()
    data.seek(0)
    client = _business_client(web_app)
    uploaded = client.post("/api/files", data={
        "objectType": "EXPENSE", "objectId": "7", "file": (data, "材料.xlsx"),
    }, content_type="multipart/form-data").get_json()
    load = openpyxl.load_workbook
    opened = []
    def observe(*args, **kwargs):
        result = load(*args, **kwargs)
        opened.append(result)
        return result
    def exhausted(*args):
        raise ValueError("preview character budget exceeded")
    monkeypatch.setattr(openpyxl, "load_workbook", observe)
    monkeypatch.setattr(preview_module, "_take_text", exhausted)
    response = client.get(f"/api/files/{uploaded['fileId']}/versions/1/preview?objectType=EXPENSE&objectId=7")
    assert response.status_code == 409
    assert len(opened) == 1
    assert opened[0]._archive.fp is None


def test_invalid_file_id_and_object_mismatch_are_controlled_and_audited(web_app, service):
    client = _business_client(web_app)
    invalid = client.get(
        "/api/files/not-a-uuid/versions/1/download?objectType=EXPENSE&objectId=7"
    )
    assert invalid.status_code == 404
    assert invalid.get_json()["error"]["code"] == "FILE_NOT_FOUND"
    assert service.audit_service.events[-1]["result"] == "FAILURE"
    assert service.audit_service.events[-1]["properties"]["operation"] == "DOWNLOAD"


def test_legacy_attachment_resolution_rejects_escape_and_symlink(tmp_path):
    from app.routes.documents import _resolve_legacy_attachment

    root = tmp_path / "app"
    uploads = root / "uploads"
    uploads.mkdir(parents=True)
    stored = uploads / "invoice.pdf"
    stored.write_bytes(b"%PDF-safe")
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"%PDF-outside")
    linked = uploads / "linked.pdf"
    linked.symlink_to(outside)

    assert _resolve_legacy_attachment(root, "uploads/invoice.pdf") == stored
    assert _resolve_legacy_attachment(root, str(stored)) == stored
    assert _resolve_legacy_attachment(root, "../outside.pdf") is None
    assert _resolve_legacy_attachment(root, str(outside)) is None
    assert _resolve_legacy_attachment(root, "uploads/linked.pdf") is None


def test_postgres_file_version_and_audit_contract(tmp_path):
    if "TEST_DATABASE_URL" not in os.environ:
        pytest.skip("requires the isolated PostgreSQL contract runner")
    from app.repositories.audit import AuditRepository
    from app.services.audit import AuditService

    pg_engine = sa.create_engine(os.environ["TEST_DATABASE_URL"])
    metadata = sa.MetaData()
    users = sa.Table("users", metadata, autoload_with=pg_engine)
    expenses = sa.Table("expense_reimbursement", metadata, autoload_with=pg_engine)
    audit_events = sa.Table("audit_events", metadata, autoload_with=pg_engine)
    suffix = uuid.uuid4().hex[:10]
    with pg_engine.begin() as connection:
        user_id = connection.execute(
            users.insert().values(
                username=f"file_{suffix}", password="unused", role="BUSINESS_USER",
                name="File contract", status="active", must_change_password=False, version=1,
            ).returning(users.c.id)
        ).scalar_one()
        expense_id = connection.execute(
            expenses.insert().values(reimbursement_no=f"FILE-{suffix}").returning(expenses.c.id)
        ).scalar_one()

    pg_service = FileService(
        FilesRepository(pg_engine),
        AuditService(AuditRepository(pg_engine), app_version="test-v1"),
        storage_root=tmp_path / "pg-files", max_bytes=1024, preview_max_bytes=512,
    )
    first = pg_service.upload(
        _pdf(b"pg-v1"), original_name="report.pdf", object_type="EXPENSE",
        object_id=str(expense_id), actor_user_id=user_id, request_id=f"req-{suffix}-1",
    )
    second = pg_service.add_version(
        first["fileId"], _pdf(b"pg-v2"), original_name="report.pdf",
        object_type="EXPENSE", object_id=str(expense_id), expected_version=1,
        actor_user_id=user_id, request_id=f"req-{suffix}-2",
    )

    assert second["versionNo"] == 2
    assert pg_service.open_version(
        first["fileId"], 1, object_type="EXPENSE", object_id=str(expense_id)
    )["path"].read_bytes().endswith(b"pg-v1")
    with pg_engine.connect() as connection:
        events = connection.execute(
            sa.select(audit_events.c.metadata)
            .where(audit_events.c.actor_user_id == user_id)
            .order_by(audit_events.c.created_at, audit_events.c.id)
        ).scalars().all()
    assert [event["operation"] for event in events] == ["UPLOAD", "ADD_VERSION"]
    assert all(set(event).isdisjoint({"path", "original_name", "content"}) for event in events)


def _postgres_entities(suffix):
    pg_engine = sa.create_engine(os.environ["TEST_DATABASE_URL"])
    metadata = sa.MetaData()
    users = sa.Table("users", metadata, autoload_with=pg_engine)
    expenses = sa.Table("expense_reimbursement", metadata, autoload_with=pg_engine)
    with pg_engine.begin() as connection:
        user_id = connection.execute(
            users.insert().values(
                username=f"file_{suffix}", password="unused", role="BUSINESS_USER",
                name="File contract", status="active", must_change_password=False, version=1,
            ).returning(users.c.id)
        ).scalar_one()
        expense_id = connection.execute(
            expenses.insert().values(reimbursement_no=f"FILE-{suffix}").returning(expenses.c.id)
        ).scalar_one()
    return pg_engine, user_id, expense_id


def test_postgres_concurrent_version_writers_serialize(tmp_path):
    if "TEST_DATABASE_URL" not in os.environ:
        pytest.skip("requires the isolated PostgreSQL contract runner")
    from app.repositories.audit import AuditRepository
    from app.services.audit import AuditService

    suffix = uuid.uuid4().hex[:10]
    pg_engine, user_id, expense_id = _postgres_entities(suffix)
    pg_service = FileService(
        FilesRepository(pg_engine),
        AuditService(AuditRepository(pg_engine), app_version="test-v1"),
        storage_root=tmp_path / "concurrent-files", max_bytes=1024, preview_max_bytes=512,
    )
    first = pg_service.upload(
        _pdf(b"v1"), original_name="report.pdf", object_type="EXPENSE",
        object_id=str(expense_id), actor_user_id=user_id, request_id=f"req-{suffix}-v1",
    )
    barrier = Barrier(2)

    def add(label):
        barrier.wait(timeout=5)
        try:
            return pg_service.add_version(
                first["fileId"], _pdf(label.encode()), original_name="report.pdf",
                object_type="EXPENSE", object_id=str(expense_id), expected_version=1,
                actor_user_id=user_id, request_id=f"req-{suffix}-{label}",
            )
        except FileServiceError as error:
            return error.code

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(add, ("writer-a", "writer-b")))

    assert sum(isinstance(result, dict) for result in results) == 1
    assert results.count("VERSION_CONFLICT") == 1
    versions = sa.Table("stored_file_versions", sa.MetaData(), autoload_with=pg_engine)
    with pg_engine.connect() as connection:
        version_numbers = connection.execute(
            sa.select(versions.c.version_no)
            .where(versions.c.file_id == uuid.UUID(first["fileId"]))
            .order_by(versions.c.version_no)
        ).scalars().all()
    assert version_numbers == [1, 2]
    assert len([path for path in pg_service.storage_root.rglob("*") if path.is_file()]) == 2


def test_postgres_audit_failure_rolls_back_and_removes_moved_file(tmp_path):
    if "TEST_DATABASE_URL" not in os.environ:
        pytest.skip("requires the isolated PostgreSQL contract runner")
    suffix = uuid.uuid4().hex[:10]
    pg_engine, user_id, expense_id = _postgres_entities(suffix)
    files = sa.Table("stored_files", sa.MetaData(), autoload_with=pg_engine)
    with pg_engine.connect() as connection:
        before = connection.execute(sa.select(sa.func.count()).select_from(files)).scalar_one()
    pg_service = FileService(
        FilesRepository(pg_engine), Audit(fail=True),
        storage_root=tmp_path / "rollback-files", max_bytes=1024, preview_max_bytes=512,
    )

    with pytest.raises(FileServiceError) as error:
        pg_service.upload(
            _pdf(), original_name="report.pdf", object_type="EXPENSE",
            object_id=str(expense_id), actor_user_id=user_id, request_id=f"req-{suffix}-failure",
        )
    assert error.value.code == "FILE_OPERATION_FAILED"
    with pg_engine.connect() as connection:
        after = connection.execute(sa.select(sa.func.count()).select_from(files)).scalar_one()
    assert after == before
    assert not any(path.is_file() for path in pg_service.storage_root.rglob("*"))


def test_postgres_concurrent_first_standard_upload_is_atomic_and_archived_is_read_only(tmp_path):
    if "TEST_DATABASE_URL" not in os.environ:
        pytest.skip("requires the isolated PostgreSQL contract runner")
    from app.repositories.audit import AuditRepository
    from app.services.audit import AuditService

    suffix = uuid.uuid4().hex[:10]
    pg_engine, user_id, _expense_id = _postgres_entities(suffix)
    metadata = sa.MetaData()
    standards = sa.Table("standards", metadata, autoload_with=pg_engine)
    files = sa.Table("stored_files", metadata, autoload_with=pg_engine)
    versions = sa.Table("stored_file_versions", metadata, autoload_with=pg_engine)
    links = sa.Table("object_files", metadata, autoload_with=pg_engine)
    audit_events = sa.Table("audit_events", metadata, autoload_with=pg_engine)
    object_id = f"STD-RACE-{suffix}"
    storage_root = tmp_path / "standard-race-files"
    pg_service = FileService(
        FilesRepository(pg_engine),
        AuditService(AuditRepository(pg_engine), app_version="test-v1"),
        storage_root=storage_root, max_bytes=1024, preview_max_bytes=512,
    )
    barrier = Barrier(2)

    def create_standard(writer):
        writer.execute(
            standards.insert().values(
                doc_id=object_id, name=f"并发标准 {suffix}", status="ACTIVE",
            )
        )

    def upload(label):
        barrier.wait(timeout=5)
        try:
            return pg_service.upload_new_object(
                _pdf(label.encode()), original_name="并发标准.pdf", object_type="STANDARD",
                object_id=object_id, create_metadata=create_standard,
                actor_user_id=user_id, request_id=f"req-standard-race-{label}",
            )
        except FileServiceError as error:
            return error.code

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(upload, ("writer-a", "writer-b")))

    successful = [result for result in results if isinstance(result, dict)]
    assert len(successful) == 1
    assert results.count("FILE_OPERATION_FAILED") == 1
    first = successful[0]
    with pg_engine.connect() as connection:
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(standards).where(standards.c.doc_id == object_id)
        ) == 1
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(links).where(
                links.c.object_type == "STANDARD", links.c.object_id == object_id
            )
        ) == 1
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(files).where(
                files.c.id == uuid.UUID(first["fileId"])
            )
        ) == 1
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(versions).where(
                versions.c.file_id == uuid.UUID(first["fileId"])
            )
        ) == 1
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(audit_events).where(
                audit_events.c.actor_user_id == user_id,
                audit_events.c.metadata["operation"].as_string() == "UPLOAD",
            )
        ) == 1
    assert len([path for path in storage_root.rglob("*") if path.is_file()]) == 1

    with pg_engine.begin() as connection:
        connection.execute(
            standards.update().where(standards.c.doc_id == object_id).values(status="ARCHIVED")
        )
    with pytest.raises(FileServiceError) as version_error:
        pg_service.add_version(
            first["fileId"], _pdf(b"v2"), original_name="并发标准.pdf",
            object_type="STANDARD", object_id=object_id, expected_version=1,
            actor_user_id=user_id, request_id=f"req-standard-race-version-{suffix}",
        )
    with pytest.raises(FileServiceError) as archive_error:
        pg_service.archive(
            first["fileId"], object_type="STANDARD", object_id=object_id,
            actor_user_id=user_id, request_id=f"req-standard-race-archive-{suffix}",
        )
    assert version_error.value.code == "OBJECT_READ_ONLY"
    assert archive_error.value.code == "OBJECT_READ_ONLY"
    with pg_engine.connect() as connection:
        assert connection.scalar(
            sa.select(sa.func.count()).select_from(versions).where(
                versions.c.file_id == uuid.UUID(first["fileId"])
            )
        ) == 1
    assert len([path for path in storage_root.rglob("*") if path.is_file()]) == 1


@pytest.mark.parametrize("operation", ["upload", "add_version", "archive"])
def test_postgres_archived_standard_rechecks_after_stale_preflight_lock(
    tmp_path, monkeypatch, operation
):
    if "TEST_DATABASE_URL" not in os.environ:
        pytest.skip("requires the isolated PostgreSQL contract runner")
    from app.repositories.audit import AuditRepository
    from app.services.audit import AuditService

    suffix = uuid.uuid4().hex[:10]
    pg_engine, user_id, _expense_id = _postgres_entities(suffix)
    metadata = sa.MetaData()
    standards = sa.Table("standards", metadata, autoload_with=pg_engine)
    files = sa.Table("stored_files", metadata, autoload_with=pg_engine)
    versions = sa.Table("stored_file_versions", metadata, autoload_with=pg_engine)
    links = sa.Table("object_files", metadata, autoload_with=pg_engine)
    audit_events = sa.Table("audit_events", metadata, autoload_with=pg_engine)
    object_id = f"STD-LOCK-{suffix}-{operation}"
    storage_root = tmp_path / f"standard-lock-{operation}"
    with pg_engine.begin() as connection:
        connection.execute(
            standards.insert().values(doc_id=object_id, name="锁定标准", status="ACTIVE")
        )
    pg_service = FileService(
        FilesRepository(pg_engine),
        AuditService(AuditRepository(pg_engine), app_version="test-v1"),
        storage_root=storage_root, max_bytes=1024, preview_max_bytes=512,
    )
    first = pg_service.upload(
        _pdf(b"v1"), original_name="锁定标准.pdf", object_type="STANDARD",
        object_id=object_id, actor_user_id=user_id, request_id=f"req-standard-lock-first-{suffix}",
    )
    prechecked = Event()
    worker_pids = Queue()
    original_preflight = pg_service._validate_object_write
    original_lock = pg_service._lock_object_write

    def snapshot():
        with pg_engine.connect() as connection:
            linked_files = tuple(connection.execute(
                sa.select(files.c.id, files.c.status)
                .join(links, links.c.file_id == files.c.id)
                .where(links.c.object_type == "STANDARD", links.c.object_id == object_id)
                .order_by(files.c.id)
            ).all())
            return {
                "stored_files": connection.scalar(sa.select(sa.func.count()).select_from(files)),
                "stored_file_versions": connection.scalar(
                    sa.select(sa.func.count()).select_from(versions)
                ),
                "object_files": connection.scalar(sa.select(sa.func.count()).select_from(links)),
                "actor_audit_operations": tuple(connection.scalars(
                    sa.select(audit_events.c.metadata["operation"].as_string())
                    .where(audit_events.c.actor_user_id == user_id)
                    .order_by(audit_events.c.created_at, audit_events.c.id)
                )),
                "linked_files": linked_files,
            }

    def disk_snapshot():
        return tuple(sorted(
            (path.relative_to(storage_root).as_posix(), path.read_bytes())
            for path in storage_root.rglob("*") if path.is_file()
        ))

    def observe_lock_wait(pid):
        deadline = time.monotonic() + 3
        with pg_engine.connect() as observer:
            while time.monotonic() < deadline:
                waiting = observer.execute(
                    sa.text(
                        "SELECT wait_event_type FROM pg_stat_activity WHERE pid = :pid"
                    ),
                    {"pid": pid},
                ).scalar_one_or_none()
                if waiting == "Lock":
                    return
                time.sleep(0.02)
        raise AssertionError("mutation did not enter a PostgreSQL Lock wait")

    def record_preflight(*args, **kwargs):
        original_preflight(*args, **kwargs)
        prechecked.set()

    def record_lock(*args, **kwargs):
        connection = args[0]
        worker_pids.put(connection.scalar(sa.text("SELECT pg_backend_pid()")))
        return original_lock(*args, **kwargs)

    monkeypatch.setattr(pg_service, "_validate_object_write", record_preflight)
    monkeypatch.setattr(pg_service, "_lock_object_write", record_lock)

    def mutate():
        try:
            if operation == "upload":
                pg_service.upload(
                    _pdf(b"new"), original_name="锁定标准.pdf", object_type="STANDARD",
                    object_id=object_id, actor_user_id=user_id,
                    request_id=f"req-standard-lock-upload-{suffix}",
                )
            elif operation == "add_version":
                pg_service.add_version(
                    first["fileId"], _pdf(b"v2"), original_name="锁定标准.pdf",
                    object_type="STANDARD", object_id=object_id, expected_version=1,
                    actor_user_id=user_id, request_id=f"req-standard-lock-version-{suffix}",
                )
            else:
                pg_service.archive(
                    first["fileId"], object_type="STANDARD", object_id=object_id,
                    actor_user_id=user_id, request_id=f"req-standard-lock-archive-{suffix}",
                )
        except FileServiceError as error:
            return error.code

    before_database = snapshot()
    before_disk = disk_snapshot()
    executor = ThreadPoolExecutor(max_workers=1)
    try:
        with pg_engine.begin() as connection:
            connection.execute(
                sa.select(standards.c.doc_id)
                .where(standards.c.doc_id == object_id)
                .with_for_update(of=standards)
            )
            connection.execute(
                standards.update().where(standards.c.doc_id == object_id).values(status="ARCHIVED")
            )
            future = executor.submit(mutate)
            assert prechecked.wait(timeout=3), "mutation did not complete the stale ACTIVE precheck"
            observe_lock_wait(worker_pids.get(timeout=3))
        result = future.result(timeout=5)
    finally:
        executor.shutdown(wait=True)

    assert result == "OBJECT_READ_ONLY"
    assert snapshot() == before_database
    assert disk_snapshot() == before_disk
    assert before_database["linked_files"] == ((uuid.UUID(first["fileId"]), "ACTIVE"),)
    assert before_database["actor_audit_operations"] == ("UPLOAD",)


def test_ooxml_zip_bomb_shape_is_rejected_without_residue(service, engine):
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 300_000)
    payload.seek(0)
    with pytest.raises(FileServiceError) as error:
        service.upload(
            payload, original_name="bomb.docx", object_type="EXPENSE", object_id="7",
            actor_user_id=1, request_id="req-bomb",
        )
    assert error.value.code == "UNSUPPORTED_MEDIA_TYPE"
    assert _rows(engine, "stored_files") == []


@pytest.mark.parametrize(
    ("extension", "required_member", "parser_module", "parser_attribute"),
    [
        ("docx", "word/document.xml", "docx", "Document"),
        ("xlsx", "xl/workbook.xml", "openpyxl", "load_workbook"),
    ],
)
def test_office_preview_rejects_excessive_xml_nodes_before_library_parse(
    web_app, service, monkeypatch, extension, required_member, parser_module, parser_attribute
):
    service.max_bytes = 10 * 1024 * 1024
    service.preview_max_bytes = 10 * 1024 * 1024
    payload = io.BytesIO()
    nodes = "<root>" + ("<n/>" * 50_001) + "</root>"
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr(required_member, nodes)
    payload.seek(0)
    client = _business_client(web_app)
    uploaded = client.post(
        "/api/files",
        data={
            "objectType": "EXPENSE", "objectId": "7",
            "file": (payload, f"oversized.{extension}"),
        },
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 201

    parser = __import__(parser_module)
    called = False

    def forbidden_parser(*_args, **_kwargs):
        nonlocal called
        called = True
        raise AssertionError("OOXML parser must not run before the preview budget gate")

    monkeypatch.setattr(parser, parser_attribute, forbidden_parser)
    file_id = uploaded.get_json()["fileId"]
    response = client.get(
        f"/api/files/{file_id}/versions/1/preview?objectType=EXPENSE&objectId=7"
    )
    assert response.status_code == 409
    assert response.get_json()["error"]["code"] == "PREVIEW_UNAVAILABLE"
    assert called is False
