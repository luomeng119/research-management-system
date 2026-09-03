from __future__ import annotations

from datetime import date, datetime, timezone
import math
from pathlib import Path
import re
import tempfile
import uuid

import openpyxl
import sqlalchemy as sa


MAX_ROWS = 5000
MAX_COLUMNS = 200
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_CELL_CHARS = 20_000
MAX_ROW_BYTES = 256_000
MAX_JSON_DEPTH = 5


class GenericTablesError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


class GenericTablesService:
    def __init__(self, repository) -> None:
        self.repository = repository

    @staticmethod
    def _id(prefix: str) -> str:
        return f"{prefix}{uuid.uuid4().hex.upper()}"

    @staticmethod
    def _now():
        return datetime.now(timezone.utc)

    @staticmethod
    def _legacy(value):
        if isinstance(value, datetime):
            return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, date):
            return value.isoformat()
        return value

    @classmethod
    def _row(cls, row):
        return {key: cls._legacy(value) for key, value in row.items()} if row else None

    def get_all(self):
        return [self._row(row) for row in self.repository.list_tables()]

    def get_by_id(self, table_id):
        return self._row(self.repository.get_table(table_id))

    def get_by_id_with_current(self, table_id):
        return self._row(self.repository.get_table_with_current(table_id))

    def get_versions(self, table_id):
        return [self._row(row) for row in self.repository.list_versions(table_id)]

    def get_version_by_id(self, version_id):
        return self._row(self.repository.get_version(version_id))

    def get_columns(self, version_id):
        return [self._row(row) for row in self.repository.list_columns(version_id)]

    def get_rows(self, version_id, keyword=None):
        return [self._row(row) for row in self.repository.list_rows(version_id, keyword, limit=MAX_ROWS)]

    def _lock(self, connection, table_id):
        table = self.repository.lock_table(connection, table_id)
        if not table:
            raise GenericTablesError("TABLE_NOT_FOUND", "表格不存在", 404)
        return table

    def _version_for_table(self, connection, table_id, version_id):
        version = self.repository.get_version(version_id, connection)
        if not version or version["table_id"] != table_id:
            raise GenericTablesError("VERSION_TABLE_MISMATCH", "版本不存在或不属于该表格", 400)
        return version

    def _lock_editable_version(self, connection, version_id):
        version = self.repository.get_version(version_id, connection)
        if not version:
            raise GenericTablesError("VERSION_NOT_FOUND", "版本不存在", 404)
        table = self._lock(connection, version["table_id"])
        version = self._version_for_table(connection, table["table_id"], version_id)
        if table["current_version_id"] != version_id or bool(version["is_locked"]):
            raise GenericTablesError("VERSION_READ_ONLY", "只能修改当前编辑版本", 409)
        return table, version

    def create(self, name, description, creator):
        name = str(name or "").strip()
        if not name:
            raise GenericTablesError("INVALID_NAME", "表格名称不能为空")
        table_id, version_id, now = self._id("GT"), self._id("GTV"), self._now()
        with self.repository.engine.begin() as connection:
            connection.execute(self.repository.tables.insert().values(
                table_id=table_id, name=name, description=str(description or ""), creator=str(creator),
                created_at=now, updated_at=now, current_version_id=None,
            ))
            connection.execute(self.repository.versions.insert().values(
                version_id=version_id, table_id=table_id, version_number=1, version_label="v1",
                create_method="create", source_version_id=None, row_count=0, page_size=20,
                creator=str(creator), created_at=now, note="", is_locked=False,
            ))
            connection.execute(self.repository.tables.update().where(
                self.repository.tables.c.table_id == table_id,
            ).values(current_version_id=version_id))
        return table_id

    def update(self, table_id, **kwargs):
        values = {key: str(value or "") for key, value in kwargs.items() if key in {"name", "description"} and value is not None}
        if "name" in values and not values["name"].strip():
            raise GenericTablesError("INVALID_NAME", "表格名称不能为空")
        with self.repository.engine.begin() as connection:
            self._lock(connection, table_id)
            if values:
                values["updated_at"] = self._now()
                connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(**values))
        return True

    def delete(self, table_id):
        with self.repository.engine.begin() as connection:
            self._lock(connection, table_id)
            vids = sa.select(self.repository.versions.c.version_id).where(self.repository.versions.c.table_id == table_id)
            connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(current_version_id=None))
            connection.execute(self.repository.data.delete().where(self.repository.data.c.version_id.in_(vids)))
            connection.execute(self.repository.columns.delete().where(self.repository.columns.c.version_id.in_(vids)))
            connection.execute(self.repository.versions.delete().where(self.repository.versions.c.table_id == table_id))
            connection.execute(self.repository.tables.delete().where(self.repository.tables.c.table_id == table_id))
        return True

    def _create_version(self, connection, table_id, method, source_version_id, creator, label, note):
        number = self.repository.next_version_number(connection, table_id)
        version_id, now = self._id("GTV"), self._now()
        source = None
        if source_version_id:
            source = self._version_for_table(connection, table_id, source_version_id)
        locked = method == "snapshot"
        connection.execute(self.repository.versions.insert().values(
            version_id=version_id, table_id=table_id, version_number=number,
            version_label=label or f"v{number}_{now.strftime('%Y%m%d_%H%M')}",
            create_method=method, source_version_id=source_version_id if source else None,
            row_count=int(source["row_count"] if source else 0), page_size=int(source.get("page_size", 20) if source else 20),
            creator=str(creator), created_at=now, note=str(note or ""), is_locked=locked,
        ))
        if source:
            self.repository.copy_columns(connection, source_version_id, version_id, now)
            self.repository.copy_rows(connection, source_version_id, version_id, now)
        connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(updated_at=now))
        return version_id

    def create_version(self, table_id, method, source_version_id, creator, label, note="", _conn=None):
        if method not in {"create", "import", "snapshot", "rollback"}:
            raise GenericTablesError("INVALID_METHOD", "版本创建方式无效")
        if _conn is not None:
            self._lock(_conn, table_id)
            return self._create_version(_conn, table_id, method, source_version_id, creator, label, note)
        with self.repository.engine.begin() as connection:
            self._lock(connection, table_id)
            return self._create_version(connection, table_id, method, source_version_id, creator, label, note)

    def set_current_version(self, table_id, version_id):
        with self.repository.engine.begin() as connection:
            self._lock(connection, table_id)
            version = self._version_for_table(connection, table_id, version_id)
            if bool(version["is_locked"]):
                raise GenericTablesError("LOCKED_VERSION", "锁定版本不能设为当前版本", 409)
            connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(
                current_version_id=version_id, updated_at=self._now(),
            ))
        return True

    def save_snapshot(self, table_id, source_version_id, label, note, creator):
        with self.repository.engine.begin() as connection:
            table = self._lock(connection, table_id)
            source = self._version_for_table(connection, table_id, source_version_id)
            if table["current_version_id"] != source_version_id or bool(source["is_locked"]):
                raise GenericTablesError("VERSION_READ_ONLY", "快照源必须是当前编辑版本", 409)
            snapshot = self._create_version(connection, table_id, "snapshot", source_version_id, creator, label, note)
            current = self._create_version(connection, table_id, "import", source_version_id, creator, None, "save_snapshot_new_working")
            connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(current_version_id=current))
            return snapshot, current

    def rollback_to(self, table_id, source_version_id, creator):
        with self.repository.engine.begin() as connection:
            self._lock(connection, table_id)
            self._version_for_table(connection, table_id, source_version_id)
            current = self._create_version(connection, table_id, "rollback", source_version_id, creator, None, f"rollback_from_{source_version_id[:20]}")
            connection.execute(self.repository.tables.update().where(self.repository.tables.c.table_id == table_id).values(current_version_id=current))
            return current

    def delete_version(self, version_id):
        with self.repository.engine.begin() as connection:
            version = self.repository.get_version(version_id, connection)
            if not version:
                return False
            table = self._lock(connection, version["table_id"])
            if table["current_version_id"] == version_id:
                raise GenericTablesError("CURRENT_VERSION_DELETE", "不得删除当前版本", 409)
            connection.execute(self.repository.data.delete().where(self.repository.data.c.version_id == version_id))
            connection.execute(self.repository.columns.delete().where(self.repository.columns.c.version_id == version_id))
            connection.execute(self.repository.versions.delete().where(self.repository.versions.c.version_id == version_id))
        return True

    @classmethod
    def _validate_json(cls, value, depth=0):
        if depth > MAX_JSON_DEPTH:
            raise GenericTablesError("JSON_TOO_DEEP", "行数据嵌套过深")
        if value is None or isinstance(value, (bool, int)):
            return value
        if isinstance(value, float):
            if not math.isfinite(value):
                raise GenericTablesError("INVALID_NUMBER", "不允许 NaN 或 Infinity")
            return value
        if isinstance(value, str):
            if len(value) > MAX_CELL_CHARS:
                raise GenericTablesError("CELL_TOO_LARGE", "单元格内容过长")
            return value
        if isinstance(value, list):
            if len(value) > 200:
                raise GenericTablesError("JSON_TOO_LARGE", "列表项过多")
            return [cls._validate_json(item, depth + 1) for item in value]
        if isinstance(value, dict):
            if len(value) > MAX_COLUMNS:
                raise GenericTablesError("ROW_TOO_WIDE", "行字段过多")
            return {str(key): cls._validate_json(item, depth + 1) for key, item in value.items()}
        raise GenericTablesError("INVALID_JSON_VALUE", "行数据包含不可序列化值")

    @classmethod
    def _validated_row(cls, value):
        if not isinstance(value, dict):
            raise GenericTablesError("ROW_NOT_OBJECT", "行数据必须是对象")
        value = cls._validate_json(value)
        import json
        if len(json.dumps(value, ensure_ascii=False).encode("utf-8")) > MAX_ROW_BYTES:
            raise GenericTablesError("ROW_TOO_LARGE", "单行数据过大")
        return value

    def upsert_column(self, version_id, col_key, col_name, col_type, col_index=None, col_width=120, col_align="left", col_summary="", col_options=None, _conn=None):
        def action(connection):
            target_index = col_index
            self._lock_editable_version(connection, version_id)
            key, name = str(col_key or "").strip(), str(col_name or "").strip()
            if not key or not name:
                raise GenericTablesError("INVALID_COLUMN", "列标识和名称不能为空")
            options = self._validate_json(col_options) if col_options is not None else None
            existing = connection.execute(sa.select(self.repository.columns.c.id).where(
                self.repository.columns.c.version_id == version_id, self.repository.columns.c.col_key == key,
            )).scalar_one_or_none()
            values = dict(col_name=name, col_type=str(col_type or "text"), col_width=int(col_width), col_align=str(col_align), col_summary=str(col_summary or ""), col_options=options)
            if existing is not None:
                connection.execute(self.repository.columns.update().where(self.repository.columns.c.id == existing).values(**values))
            else:
                count = connection.execute(sa.select(sa.func.count()).select_from(self.repository.columns).where(self.repository.columns.c.version_id == version_id)).scalar_one()
                if count >= MAX_COLUMNS:
                    raise GenericTablesError("TOO_MANY_COLUMNS", "列数超过上限")
                if target_index is None:
                    maximum = connection.execute(sa.select(sa.func.max(self.repository.columns.c.col_index)).where(self.repository.columns.c.version_id == version_id)).scalar_one_or_none()
                    insert_index = int(maximum if maximum is not None else -1) + 1
                else:
                    insert_index = int(target_index)
                connection.execute(self.repository.columns.insert().values(version_id=version_id, col_key=key, col_index=insert_index, created_at=self._now(), **values))
            return True
        if _conn is not None:
            return action(_conn)
        with self.repository.engine.begin() as connection:
            return action(connection)

    def delete_column(self, version_id, col_key):
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            connection.execute(self.repository.columns.delete().where(self.repository.columns.c.version_id == version_id, self.repository.columns.c.col_key == col_key))
            for row in self.repository.list_rows(version_id, limit=MAX_ROWS, connection=connection):
                data = dict(row["row_data"])
                if col_key in data:
                    data.pop(col_key)
                    connection.execute(self.repository.data.update().where(self.repository.data.c.id == row["id"]).values(row_data=data, updated_at=self._now()))
        return True

    def reorder_columns(self, version_id, col_keys):
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            existing = {row["col_key"] for row in self.repository.list_columns(version_id, connection)}
            if len(col_keys) != len(set(col_keys)) or set(col_keys) != existing:
                raise GenericTablesError("INVALID_COLUMN_ORDER", "列顺序必须包含全部列且不重复")
            for index, key in enumerate(col_keys):
                connection.execute(self.repository.columns.update().where(self.repository.columns.c.version_id == version_id, self.repository.columns.c.col_key == key).values(col_index=index))
        return True

    def upsert_row(self, version_id, row_key, row_data_dict):
        row_data = self._validated_row(row_data_dict)
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            existing = connection.execute(sa.select(self.repository.data).where(
                self.repository.data.c.version_id == version_id, self.repository.data.c.row_key == row_key,
            )).mappings().first()
            now = self._now()
            if existing:
                connection.execute(self.repository.data.update().where(self.repository.data.c.id == existing["id"]).values(row_data=row_data, updated_at=now))
            else:
                count = self.repository.count_rows(connection, version_id)
                if count >= MAX_ROWS:
                    raise GenericTablesError("TOO_MANY_ROWS", "行数超过上限")
                maximum = connection.execute(sa.select(sa.func.max(self.repository.data.c.row_index)).where(self.repository.data.c.version_id == version_id)).scalar_one_or_none()
                connection.execute(self.repository.data.insert().values(version_id=version_id, row_key=str(row_key), row_index=int(maximum if maximum is not None else -1) + 1, row_data=row_data, row_color="", created_at=now, updated_at=now))
            connection.execute(self.repository.versions.update().where(self.repository.versions.c.version_id == version_id).values(row_count=self.repository.count_rows(connection, version_id)))
        return True

    def delete_row(self, version_id, row_key):
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            connection.execute(self.repository.data.delete().where(self.repository.data.c.version_id == version_id, self.repository.data.c.row_key == row_key))
            connection.execute(self.repository.versions.update().where(self.repository.versions.c.version_id == version_id).values(row_count=self.repository.count_rows(connection, version_id)))
        return True

    def update_row_color(self, version_id, row_key, row_color):
        if row_color not in {"", "yellow", "green", "red"}:
            raise GenericTablesError("INVALID_COLOR", "行颜色无效")
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            result = connection.execute(self.repository.data.update().where(self.repository.data.c.version_id == version_id, self.repository.data.c.row_key == row_key).values(row_color=row_color, updated_at=self._now()))
            return result.rowcount == 1

    def update_column_width(self, version_id, col_key, col_width):
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            result = connection.execute(self.repository.columns.update().where(self.repository.columns.c.version_id == version_id, self.repository.columns.c.col_key == col_key).values(col_width=int(col_width)))
            return result.rowcount == 1

    def update_page_size(self, version_id, page_size):
        if page_size not in {-1, 10, 20, 50, 100}:
            raise GenericTablesError("INVALID_PAGE_SIZE", "分页数无效")
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            result = connection.execute(self.repository.versions.update().where(self.repository.versions.c.version_id == version_id).values(page_size=page_size))
            return result.rowcount == 1

    @staticmethod
    def _name_to_key(name):
        value = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff]", "_", str(name))
        value = re.sub(r"_+", "_", value).strip("_") or "col"
        return ("c_" + value if value[0].isdigit() else value)[:30]

    @classmethod
    def _excel_value(cls, value):
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return cls._validate_json(value)

    def import_rows_from_excel(self, version_id, file_path, creator, mode="replace"):
        path = Path(file_path)
        if path.suffix.lower() != ".xlsx":
            raise GenericTablesError("INVALID_EXCEL_TYPE", "仅支持 .xlsx 文件")
        if mode not in {"replace", "append"}:
            raise GenericTablesError("INVALID_IMPORT_MODE", "导入模式无效")
        if path.stat().st_size > MAX_FILE_BYTES:
            raise GenericTablesError("FILE_TOO_LARGE", "导入文件过大")
        # Keep formula text as data. It is escaped again on export so an
        # uploaded formula never becomes executable spreadsheet content.
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
        try:
            sheet = workbook.active
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
                raise GenericTablesError("WORKBOOK_TOO_LARGE", "工作表超过 V1 上限")
            raw = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if len(raw) < 2:
            return 0
        headers = [str(value or "").strip() for value in raw[0]]
        keys, seen = [], set()
        for header in headers:
            if not header:
                keys.append(None)
                continue
            key, suffix = self._name_to_key(header), 2
            base = key
            while key in seen:
                key = f"{base[:26]}_{suffix}"
                suffix += 1
            seen.add(key)
            keys.append(key)
        parsed = []
        for values in raw[1:]:
            row = {keys[index]: self._excel_value(value) if value is not None else "" for index, value in enumerate(values) if index < len(keys) and keys[index]}
            if row:
                parsed.append(self._validated_row(row))
        with self.repository.engine.begin() as connection:
            self._lock_editable_version(connection, version_id)
            existing_count = self.repository.count_rows(connection, version_id)
            if mode == "append" and existing_count + len(parsed) > MAX_ROWS:
                raise GenericTablesError("TOO_MANY_ROWS", "导入后行数超过上限")
            if mode == "replace":
                connection.execute(self.repository.data.delete().where(self.repository.data.c.version_id == version_id))
            for index, header in enumerate(headers):
                if index < len(keys) and keys[index]:
                    self.upsert_column(version_id, keys[index], header, "text", col_index=index, _conn=connection)
            maximum = connection.execute(sa.select(sa.func.max(self.repository.data.c.row_index)).where(self.repository.data.c.version_id == version_id)).scalar_one_or_none()
            base_index = int(maximum if maximum is not None else -1) + 1
            now = self._now()
            if parsed:
                connection.execute(self.repository.data.insert(), [dict(version_id=version_id, row_key=self._id("GTR"), row_index=base_index + index, row_data=row, row_color="", created_at=now, updated_at=now) for index, row in enumerate(parsed)])
            connection.execute(self.repository.versions.update().where(self.repository.versions.c.version_id == version_id).values(row_count=self.repository.count_rows(connection, version_id)))
        return len(parsed)

    def get_column_stats(self, version_id):
        rows, columns = self.get_rows(version_id), self.get_columns(version_id)
        stats = {}
        for column in columns:
            if column["col_type"] != "number":
                continue
            values = []
            for row in rows:
                try:
                    if row["row_data"].get(column["col_key"]) not in ("", None):
                        values.append(float(row["row_data"][column["col_key"]]))
                except (TypeError, ValueError):
                    pass
            if values:
                kind = column.get("col_summary") if column.get("col_summary") in {"sum", "count", "avg"} else "sum"
                value = len(values) if kind == "count" else round(sum(values) / len(values), 2) if kind == "avg" else round(sum(values), 2)
                stats[column["col_key"]] = {"type": kind, "value": value}
        return stats

    def compare_versions(self, version_id_a, version_id_b):
        first, second = self.get_version_by_id(version_id_a), self.get_version_by_id(version_id_b)
        if not first or not second:
            raise GenericTablesError("VERSION_NOT_FOUND", "版本不存在", 404)
        if first["table_id"] != second["table_id"]:
            raise GenericTablesError("CROSS_TABLE_COMPARE", "只能比较同一表格的版本")
        cols_a = {row["col_key"]: row for row in self.get_columns(version_id_a)}
        cols_b = {row["col_key"]: row for row in self.get_columns(version_id_b)}
        rows_a = {row["row_key"]: row for row in self.get_rows(version_id_a)}
        rows_b = {row["row_key"]: row for row in self.get_rows(version_id_b)}
        col_diff = {}
        for key in cols_a.keys() | cols_b.keys():
            if key not in cols_a: status = "added"
            elif key not in cols_b: status = "deleted"
            elif any(cols_a[key].get(field) != cols_b[key].get(field) for field in ("col_name", "col_type", "col_summary")): status = "modified"
            else: status = "unchanged"
            col_diff[key] = status
        row_diff = []
        for key in sorted(rows_a.keys() | rows_b.keys()):
            old = rows_a.get(key, {}).get("row_data")
            new = rows_b.get(key, {}).get("row_data")
            status = "added" if old is None else "deleted" if new is None else "unchanged" if old == new else "modified"
            row_diff.append({"row_key": key, "status": status, "old": old, "new": new})
        return {"col_diff": col_diff, "row_diff": row_diff, "cols_a": list(cols_a), "cols_b": list(cols_b)}

    @staticmethod
    def _safe_excel(value):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    def export_to_excel(self, version_id):
        columns, rows = self.get_columns(version_id), self.get_rows(version_id)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "数据导出"
        sheet.append([column.get("col_name", column["col_key"]) for column in columns])
        for row in rows:
            sheet.append([self._safe_excel(row["row_data"].get(column["col_key"], "")) for column in columns])
        handle, path = tempfile.mkstemp(suffix=".xlsx")
        import os
        os.close(handle)
        try:
            workbook.save(path)
        except Exception:
            Path(path).unlink(missing_ok=True)
            raise
        finally:
            workbook.close()
        return path
